"""Generate a fake, deliberately messy KPI workbook for "Northwind Software".

Output: data/northwind.xlsx (all money figures in $K).

How it works:
1. TRUE_DATA holds the clean, correct numbers for 8 quarters.
2. check_true_data() confirms those numbers tie out (ARR and cash roll forward).
3. The workbook is written with mess added on purpose: odd header names,
   some numbers stored as text ("$14.3M", "120K", "5,090"), one blank
   quarter, a budget-only row for next quarter, and a junk Notes tab.

The numbers are chosen to tell the Northwind story in CLAUDE.md:
strong ARR growth, NRR sliding 108% -> 97%, rising pipeline,
burn ~20% over budget and ~11 months of runway.
"""

from pathlib import Path

from openpyxl import Workbook

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

# Where the workbook is saved, relative to this script's folder.
OUTPUT_PATH = Path(__file__).parent / "data" / "northwind.xlsx"

# The 8 actual quarters, oldest first.
QUARTERS = [
    "Q3 2024", "Q4 2024", "Q1 2025", "Q2 2025",
    "Q3 2025", "Q4 2025", "Q1 2026", "Q2 2026",
]

# This quarter's row is written with its label only - every value left blank.
BLANK_QUARTER = "Q1 2025"

# Label for the extra forecast row after the actuals (budget columns only).
BUDGET_ONLY_LABEL = "Q3 2026 (Budget)"

# ---------------------------------------------------------------------------
# The clean, correct data ($K unless noted). One list per column,
# one value per quarter, in the same order as QUARTERS.
# Q1 2025 has real values here so the chains tie out, but they never reach the file.
# ---------------------------------------------------------------------------

TRUE_DATA = {
    "starting_arr":    [12000, 13580, 15320, 17190, 19230, 21460, 23790, 25810],
    "new_arr":         [1400, 1500, 1600, 1700, 1800, 1900, 1900, 1850],
    "expansion_arr":   [360, 440, 540, 640, 770, 860, 710, 580],
    "contraction_arr": [60, 70, 90, 100, 110, 140, 200, 260],
    "churned_arr":     [120, 130, 180, 200, 230, 290, 390, 510],
    "revenue":         [3200, 3610, 4060, 4550, 5090, 5660, 6200, 6660],
    "gross_profit":    [2300, 2600, 2960, 3320, 3770, 4190, 4650, 5000],
    "net_burn":        [2400, 2600, 2800, 3000, 3200, 3400, 3650, 3900],
    "ending_cash":     [36850, 34250, 31450, 28450, 25250, 21850, 18200, 14300],
    "sm_spend":        [1700, 1800, 1900, 2000, 2100, 2250, 2350, 2400],
    "new_customers":   [38, 41, 44, 47, 50, 53, 52, 50],            # count, not $K
    "headcount":       [110, 118, 126, 134, 142, 150, 158, 165],    # count, not $K
    "pipeline":        [6000, 6800, 7500, 8300, 9200, 10100, 11200, 12500],
    "budget_new_arr":  [1400, 1500, 1650, 1750, 1850, 1950, 2000, 2000],
    "budget_arr":      [13500, 15150, 16900, 18750, 20700, 22700, 24750, 26800],
    "budget_net_burn": [2400, 2520, 2670, 2780, 2910, 2980, 3120, 3250],
}

# Budget for the quarter after the latest actual quarter (the budget-only row).
NEXT_QUARTER_BUDGET = {
    "budget_new_arr": 2050,
    "budget_arr": 28900,
    "budget_net_burn": 3300,
}

# Columns that are counts, not money - skipped by the multiple-of-10 check.
COUNT_COLUMNS = {"new_customers", "headcount"}

# ---------------------------------------------------------------------------
# The mess
# ---------------------------------------------------------------------------

# Clean column name -> the inconsistent header written in the workbook.
HEADER_NAMES = {
    "starting_arr":    "Beginning ARR",
    "new_arr":         "New ARR ($K)",
    "expansion_arr":   "expansion arr",
    "contraction_arr": "Contraction_ARR",
    "churned_arr":     " Churned ARR ",      # stray spaces on purpose
    "revenue":         "Revenue",
    "gross_profit":    "GROSS PROFIT",
    "net_burn":        "Net Burn",
    "ending_cash":     "Cash - End of Qtr",
    "sm_spend":        "S&M Spend",
    "new_customers":   "New Logos",
    "headcount":       "Headcount (FTE)",
    "pipeline":        "Qualified Pipeline",
    "budget_new_arr":  "New ARR - Budget",
    "budget_arr":      "Budget ARR",
    "budget_net_burn": "Net Burn (Bud.)",
}

# (column, quarter) -> text style. These cells are written as text instead of numbers.
#   "M"     -> "$14.3M"  (millions with a dollar sign)
#   "K"     -> "120K"    (thousands with a K)
#   "comma" -> "5,090"   (number with a thousands comma)
TEXT_CELLS = {
    ("pipeline", "Q3 2024"): "M",
    ("pipeline", "Q4 2024"): "M",
    ("pipeline", "Q2 2025"): "M",
    ("pipeline", "Q3 2025"): "M",
    ("pipeline", "Q4 2025"): "M",
    ("pipeline", "Q1 2026"): "M",
    ("pipeline", "Q2 2026"): "M",
    ("ending_cash", "Q4 2025"): "M",
    ("ending_cash", "Q1 2026"): "M",
    ("ending_cash", "Q2 2026"): "M",
    ("churned_arr", "Q3 2024"): "K",
    ("churned_arr", "Q4 2024"): "K",
    ("contraction_arr", "Q2 2026"): "K",
    ("revenue", "Q3 2025"): "comma",
}


def check_true_data():
    """Stop with an error if the clean data doesn't tie out."""
    # Every column must have exactly one value per quarter.
    for column, values in TRUE_DATA.items():
        assert len(values) == len(QUARTERS), f"{column} needs {len(QUARTERS)} values"

    # ARR roll-forward: this quarter's ending ARR must equal next quarter's starting ARR.
    d = TRUE_DATA
    for i in range(len(QUARTERS) - 1):
        ending_arr = (d["starting_arr"][i] + d["new_arr"][i] + d["expansion_arr"][i]
                      - d["contraction_arr"][i] - d["churned_arr"][i])
        assert ending_arr == d["starting_arr"][i + 1], f"ARR doesn't roll forward after {QUARTERS[i]}"

    # Cash roll-forward: this quarter's cash = last quarter's cash - this quarter's burn.
    for i in range(1, len(QUARTERS)):
        expected_cash = d["ending_cash"][i - 1] - d["net_burn"][i]
        assert d["ending_cash"][i] == expected_cash, f"Cash doesn't roll forward into {QUARTERS[i]}"

    # Money values must be multiples of 10 so "$21.85M"-style text is exact, not rounded.
    for column, values in TRUE_DATA.items():
        if column not in COUNT_COLUMNS:
            assert all(v % 10 == 0 for v in values), f"{column} has a value not divisible by 10"

    # Text cells must point at a real column and a quarter that actually gets written.
    for column, quarter in TEXT_CELLS:
        assert column in TRUE_DATA, f"Unknown column in TEXT_CELLS: {column}"
        assert quarter in QUARTERS and quarter != BLANK_QUARTER, f"Bad quarter in TEXT_CELLS: {quarter}"


def to_text(value, style):
    """Turn a $K number into messy text, e.g. 14300 -> "$14.3M"."""
    if style == "M":
        # 14300 $K / 1000 = 14.3 $M; ":g" drops trailing zeros (6.0 -> 6).
        return f"${value / 1000:g}M"
    if style == "K":
        return f"{value}K"
    if style == "comma":
        return f"{value:,}"
    raise ValueError(f"Unknown text style: {style}")


def write_kpi_sheet(ws):
    """Write the main KPI table: header row, 8 quarters, then the budget-only row."""
    columns = list(TRUE_DATA)

    # Header row: "Quarter" followed by the messy header names.
    ws.append(["Quarter"] + [HEADER_NAMES[c] for c in columns])

    # One row per actual quarter.
    for i, quarter in enumerate(QUARTERS):
        if quarter == BLANK_QUARTER:
            ws.append([quarter])  # label only, every value left empty
            continue
        row = [quarter]
        for column in columns:
            value = TRUE_DATA[column][i]
            style = TEXT_CELLS.get((column, quarter))  # None if this cell stays a number
            row.append(to_text(value, style) if style else value)
        ws.append(row)

    # Budget-only row: blank for actuals, filled for the 3 budget columns.
    ws.append([BUDGET_ONLY_LABEL] + [NEXT_QUARTER_BUDGET.get(c) for c in columns])

    # Widen columns so the sheet is readable when opened in Excel.
    for column_cells in ws.columns:
        ws.column_dimensions[column_cells[0].column_letter].width = 18


def write_notes_sheet(ws):
    """Write a junk notes tab, the kind of thing real workbooks carry around."""
    ws["A1"] = "Northwind Software - misc notes"
    ws["A3"] = "Q1 2025 actuals still with finance, do not use"
    ws["A4"] = "Pipeline from CRM export, pulled first week of quarter"
    ws["A5"] = "Board meeting moved - check calendar"
    ws["A7"] = "scratch"
    ws["B7"] = 3900
    ws["C7"] = 3250
    ws["D7"] = "20% over??"
    ws["A9"] = "TODO: update headcount plan"


def main():
    check_true_data()

    # A new workbook starts with one sheet; rename it and add the notes tab.
    wb = Workbook()
    kpi_sheet = wb.active
    kpi_sheet.title = "KPI Tracker"
    write_kpi_sheet(kpi_sheet)
    write_notes_sheet(wb.create_sheet("Notes"))

    OUTPUT_PATH.parent.mkdir(exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH.name}: {len(QUARTERS)} quarters (blank: {BLANK_QUARTER}), "
          f"1 budget-only row, {len(TEXT_CELLS)} text cells, sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
