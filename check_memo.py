"""Automated proof for the board memo (final Task 1): every number in it is in the metrics workbook.

For each of the three companies, build the memo (with output/<company>_analysis.json if it exists,
otherwise with "AI commentary unavailable"), open both saved files, and check:
1. Both files exist; the PDF is 1 or 2 pages.
2. Every number in the memo (Word body and PDF body) appears in the saved
   output/<company>_metrics.xlsx, as Excel displays it: Metrics sheet cells and quarter labels,
   Flags sheet names and thresholds, the combo rule's wording, the runway-at-budget line, and the
   Flags sheet's rows counted by status ("6 of 9 flags tripped"). The footer is checked on its own (7),
   and the What changed since the last run section (last run's values) by check_diff.py.
3. Key metrics table: each row's latest and prior cells equal that metric's Excel cells; each
   flag's threshold and status equal the Flags sheet; every flag has a row.
4. The flag count matches the company's story (check_companies.py); every tripped flag is listed;
   every metric with data missing is under Data gaps (or it says None).
5. AI text: the JSON's headline and 3 questions under "AI-drafted from computed metrics - review
   before use", and none of its wins or risks; or "AI commentary unavailable" and no AI-drafted line.
6. The PDF carries every piece of text the Word file does; neither has an em dash.
7. Footer (Word footer, and on every PDF page): fictional-data note, source file, today's date,
   the git commit, the analysis's model, and the review status from the manifest: "reviewed by" only
   if a still-valid approval lists the memo in its documents (an older approval covers the deck only).
Plus, in a temporary folder:
- An analysis with an invented number, and one quoting a raw input that analyze.py and the deck
  accept but the metrics workbook doesn't show (Northwind's ending cash), both give
  "AI commentary unavailable" while every computed number still appears.
- The number check is proven by planting a number in a copy of a saved memo: the check must fail.
- A memo with a What changed since the last run section (Northwind's Q1 2026 run, then today's)
  passes, and that section is shown to hold numbers today's workbook doesn't (so leaving it out is needed).

Expected words are typed out here, not imported from memo.py, so a wrong constant there can't
pass its own check. No API calls.
Run: python check_memo.py  -> prints "All checks passed" or stops at the first failure.
"""

import contextlib
import copy
import datetime
import io
import json
import re
import tempfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from build_deck import analysis_path, save_deck
from check_companies import COMPANIES
from check_deck import allowed_numbers, excel_display, expected_flag_count, number_tokens, read_metrics_workbook
from excel_output import save_metrics_workbook
from mapping import mapping_sha256
from memo import memo_paths, save_memo
from metrics import COMBO_FLAG_NAME, CONFIG_PATH, TRIP, load_config
from provenance import approval_status, file_sha256, git_commit, manifest_path, read_manifest

PROJECT_DIR = Path(__file__).parent
LATEST, PRIOR = "Q2 2026", "Q1 2026"
UNAVAILABLE = "AI commentary unavailable"
AI_LINE = "AI-drafted from computed metrics - review before use"
RUNWAY_CONTEXT_ROW = "Runway at next quarter's budgeted burn"
FLAG_STATUSES = ("Tripped", "Passed", "Cannot evaluate")
EM_DASH = chr(0x2014)   # the em dash, by its Unicode number
CHANGES_HEADING = "What changed since the last run"   # Task 10: its numbers are checked by check_diff.py
KEY_METRICS_HEADING = "Key metrics"


# ---------------------------------------------------------------------------
# Reading the metrics workbook: the numbers a memo may show
# ---------------------------------------------------------------------------

def flag_counts(flags):
    """The Flags sheet's rows counted by status, as text: tripped, passed, cannot evaluate, and all of them."""
    statuses = [flag["Status"].split(" ")[0] for flag in flags.values()]   # "Cannot evaluate ..." -> "Cannot"
    counts = [statuses.count(status.split(" ")[0]) for status in FLAG_STATUSES]
    return [str(count) for count in counts + [len(statuses)]]


def runway_context(path):
    """The runway-at-budget cell below the flag table, as Excel displays it."""
    for row in load_workbook(path)["Flags"].iter_rows():
        if isinstance(row[0].value, str) and row[0].value.startswith(RUNWAY_CONTEXT_ROW):
            return excel_display(row[2])
    raise AssertionError(f"{path.name}: no '{RUNWAY_CONTEXT_ROW}' row on the Flags sheet")


def workbook_allowed(path):
    """Every number the metrics workbook shows, as number tokens ('97.1%', '11.0 mo', '2026', ...)."""
    table, flags = read_metrics_workbook(path)
    allowed = allowed_numbers(table, flags)
    extra = flag_counts(flags) + [runway_context(path)]
    return table, flags, allowed.union(*(number_tokens(text) for text in extra))


# ---------------------------------------------------------------------------
# Reading the memo
# ---------------------------------------------------------------------------

def docx_body(path):
    """The Word file's paragraphs (top to bottom), and its tables as lists of cell texts."""
    document = Document(path)
    paragraphs = [item.text for item in document.paragraphs]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in document.tables]
    return paragraphs, tables


def docx_footer(path):
    return "\n".join(item.text for item in Document(path).sections[0].footer.paragraphs)


def flat(text):
    """Text with every run of spaces and line breaks as one space (a PDF wraps lines where it likes)."""
    return " ".join(text.split())


def pdf_pages(path):
    return [page.extract_text() for page in PdfReader(path).pages]


def without_footer(page, footer):
    """A PDF page's text less its footer, which is checked on its own."""
    flat_page = flat(page)
    assert flat(footer) in flat_page, f"PDF page lacks the footer: {footer!r}"
    return flat_page.replace(flat(footer), " ")


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def sentence_tokens(text):
    """check_deck.number_tokens for sentences: a comma after a number that no digit follows ends the
    sentence clause, not the number ("Q2 2026, compared with" is 2026, not "2026,")."""
    return number_tokens(re.sub(r"(\d),(?!\d)", r"\1 ", text))


def check_numbers(texts, allowed, where):
    """Every number in these texts is one the metrics workbook shows."""
    shown = set().union(*(sentence_tokens(text) for text in texts))
    invented = shown - allowed
    assert not invented, f"{where}: numbers not in the metrics workbook: {sorted(invented)}"
    return len(shown)


def check_kpi_table(tables, table, flags, name):
    """Row by row: latest and prior cells against the Metrics sheet; threshold and status against the Flags sheet."""
    assert len(tables) == 1, f"{name}: {len(tables)} tables in the memo, expected 1"
    rows = tables[0]
    assert rows[0] == ["Metric", LATEST, PRIOR, "Budget or threshold", "Status"], f"{name}: table header {rows[0]}"
    for label, latest, prior, budget, status in rows[1:]:
        if (LATEST, label) in table:  # a metric row (flag names are the metric labels)
            assert [latest, prior] == [table[(LATEST, label)], table[(PRIOR, label)]], \
                f"{name}: row {label!r} shows {latest!r}, {prior!r}; metrics workbook has " \
                f"{table[(LATEST, label)]!r}, {table[(PRIOR, label)]!r}"
        if label in flags:
            flag = flags[label]
            assert status == flag["Status"], \
                f"{name}: {label} status {status!r}, Flags sheet {flag['Status']!r}"
            if label != COMBO_FLAG_NAME:
                assert budget.endswith(flag["Threshold"]), f"{name}: {label} threshold {budget!r} vs {flag['Threshold']!r}"
    shown = [cells[0] for cells in rows[1:] if cells[0] in flags]
    assert sorted(shown) == sorted(flags), f"{name}: table shows {len(shown)} of {len(flags)} flags"


def check_flags_and_gaps(paragraphs, company, gap_labels):
    """The flag count from the story, each tripped flag, and every data gap (or None)."""
    name, text = company["name"], "\n".join(paragraphs)
    count = expected_flag_count(company)
    assert f"Flags: {count}" in paragraphs, f"{name}: flag heading should read 'Flags: {count}'"
    for flag, status in company["expected_flags"].items():
        if status == TRIP:
            assert f"{flag}: " in text, f"{name}: tripped flag {flag!r} missing from the memo"
    gaps = text.split("Data gaps (data missing)", 1)[1].split("Questions for management", 1)[0]
    for label in gap_labels:
        assert label in gaps, f"{name}: data gap {label!r} missing from the memo"
    if not gap_labels:
        assert "None: every metric and flag has the data it needs" in gaps, f"{name}: data gaps should say None"
    return count


def check_ai_text(paragraphs, summary, name):
    """The JSON's headline and questions under the AI-drafted line; or the unavailable text. Never wins or risks."""
    text = "\n".join(paragraphs)
    if summary is None:
        assert text.count(UNAVAILABLE) == 2, f"{name}: 'AI commentary unavailable' should stand for headline and questions"
        assert AI_LINE not in text, f"{name}: AI-drafted line in a memo with no AI text"
        return
    assert UNAVAILABLE not in text, f"{name}: the unavailable text is shown beside AI text"
    headline_at = paragraphs.index(summary["headline"])
    assert paragraphs[headline_at - 1] == AI_LINE, f"{name}: the headline isn't under the AI-drafted line"
    questions_at = paragraphs.index("Questions for management")
    assert paragraphs[questions_at + 1] == AI_LINE, f"{name}: the questions aren't under the AI-drafted line"
    assert paragraphs[questions_at + 2:questions_at + 2 + len(summary["questions"])] == summary["questions"], \
        f"{name}: the questions differ from the JSON"
    for point in summary["wins"] + summary["risks"]:
        assert point["detail"] not in text, f"{name}: a win or risk is in the memo: {point['title']!r}"


def expected_memo_review(workbook, output_dir):
    """How the memo's footer must end: "reviewed by" only if an approval still valid today lists the memo.

    An approval from before memos existed covers the deck only (approve.py's "documents").
    """
    manifest = read_manifest(manifest_path(workbook, output_dir))
    approval, _ = approval_status(manifest, file_sha256(workbook), file_sha256(CONFIG_PATH), mapping_sha256(workbook))
    if approval is None or "memo" not in approval.get("documents", []):
        return "AI-drafted | not reviewed"
    return f"AI-drafted | reviewed by {approval['reviewer']} on {approval['approved_at'][:10]}"


def expected_footer(workbook, output_dir, model):
    """What the footer must say, worked out here from git, the analysis and the manifest (not from memo.py)."""
    commit = git_commit()
    code = commit["commit"] + ("*" if commit["uncommitted_changes"] else "")
    return " | ".join(["Fictional data", workbook.name, datetime.date.today().isoformat(), code,
                       model or "no AI text", expected_memo_review(workbook, output_dir)])


def check_footer(docx_path, pages, footer, name):
    assert docx_footer(docx_path) == footer, f"{name}: Word footer {docx_footer(docx_path)!r}, expected {footer!r}"
    for number, page in enumerate(pages, start=1):
        assert flat(footer) in flat(page), f"{name}: PDF page {number} has no footer {footer!r}"


def check_same_text_and_no_em_dash(paragraphs, tables, pages, name):
    """Every piece of the Word file's text is in the PDF; neither file has an em dash."""
    pdf_text = flat(" ".join(pages))
    pieces = [text for text in paragraphs if text] + [cell for table in tables for row in table for cell in row]
    missing = [piece for piece in pieces if flat(piece) not in pdf_text]
    assert not missing, f"{name}: in the Word file but not the PDF: {missing[:3]}"
    for where, text in (("Word file", "\n".join(pieces)), ("PDF", pdf_text)):
        assert EM_DASH not in text, f"{name}: em dash in the {where}"
    return len(pieces)


# ---------------------------------------------------------------------------
# One company
# ---------------------------------------------------------------------------

def without_changes(text):
    """The text less the What changed section (Task 10), which holds last run's values and each move's size.

    Those numbers are in no metrics workbook built today; check_diff.py checks the section line by
    line against each company's answer key instead. Everything else must still be in the workbook.
    """
    if CHANGES_HEADING not in text:
        return text
    before, after = text.split(CHANGES_HEADING, 1)
    return before + " " + after.split(KEY_METRICS_HEADING, 1)[1]


def memo_numbers_check(docx_path, pdf_path, footer, allowed, name):
    """Check 2: every number in the Word body and the PDF body (footers left out) is in the metrics workbook."""
    paragraphs, tables = docx_body(docx_path)
    cells = [cell for table in tables for row in table for cell in row]
    count = check_numbers([without_changes("\n".join(paragraphs))] + cells, allowed, f"{name} Word file")
    pdf_body = " ".join(without_footer(page, footer) for page in pdf_pages(pdf_path))
    check_numbers([without_changes(pdf_body)], allowed, f"{name} PDF")
    return count


def check_company(company, config, output_dir):
    name, workbook = company["name"], Path(company["answer_key"].OUTPUT_PATH)
    json_path = analysis_path(workbook)
    saved = json.loads(json_path.read_text()) if json_path.exists() else {}
    summary = saved.get("summary")

    result = save_memo(workbook, config, json_path, output_dir=output_dir)
    assert (summary is None) == (result["why_unavailable"] is not None), \
        f"{name}: analysis use unexpected: {result['why_unavailable']}"
    docx_path, pdf_path = result["docx"], result["pdf"]
    assert docx_path.exists() and pdf_path.exists(), f"{name}: memo files missing"
    pages = pdf_pages(pdf_path)
    assert 1 <= len(pages) <= 2, f"{name}: the PDF is {len(pages)} pages, expected 1 or 2"

    table, flags, allowed = workbook_allowed(save_metrics_workbook(workbook, config))
    model = (saved.get("run_info") or {}).get("model") if summary else None
    footer = expected_footer(workbook, output_dir, model)
    numbers = memo_numbers_check(docx_path, pdf_path, footer, allowed, name)

    paragraphs, tables = docx_body(docx_path)
    assert paragraphs[0] == f"{name} board update: {LATEST}", f"{name}: title {paragraphs[0]!r}"
    check_kpi_table(tables, table, flags, name)
    gap_labels = sorted({label for (_, label), text in table.items() if text == "data missing"})
    count = check_flags_and_gaps(paragraphs, company, gap_labels)
    check_ai_text(paragraphs, summary, name)
    pieces = check_same_text_and_no_em_dash(paragraphs, tables, pages, name)
    check_footer(docx_path, pages, footer, name)

    ai = "AI headline and questions from the JSON" if summary else f"unavailable ({result['why_unavailable']})"
    print(f"✓ {name}: {len(pages)} page PDF + Word file, {ai}, {count}; {numbers} distinct numbers, all in "
          f"the metrics workbook; the PDF has all {pieces} pieces of the Word text; footer ends "
          f"'{expected_memo_review(workbook, output_dir)}'")
    return result


# ---------------------------------------------------------------------------
# Proofs in a temporary folder
# ---------------------------------------------------------------------------

def tampered_analysis(folder, change):
    """A copy of Northwind's saved analysis with one change applied."""
    saved = copy.deepcopy(json.loads(analysis_path(COMPANIES[0]["answer_key"].OUTPUT_PATH).read_text()))
    change(saved)
    path = Path(folder) / "tampered_analysis.json"
    path.write_text(json.dumps(saved))
    return path


def check_bad_analyses_are_unavailable(config, folder):
    """An invented number, and a raw input the workbook doesn't show, both keep the AI text out of the memo."""
    def invent_number(saved):
        saved["summary"]["headline"] = saved["summary"]["headline"].replace("11.0 mo", "11.5 mo")

    def quote_ending_cash(saved):   # 14,300 is Northwind's latest ending cash: in Claude's payload, not the workbook
        saved["summary"]["questions"][0] = "How long will the ending cash of $14,300K last at the current burn?"

    company, workbook = COMPANIES[0], Path(COMPANIES[0]["answer_key"].OUTPUT_PATH)
    table, flags, allowed = workbook_allowed(save_metrics_workbook(workbook, config, folder))
    for change, reason, deck_accepts in ((invent_number, "11.5", False), (quote_ending_cash, "14,300", True)):
        path = tampered_analysis(folder, change)
        _, deck_why = save_deck(workbook, config, path, output_dir=folder)
        assert (deck_why is None) == deck_accepts, f"{change.__name__}: the deck's verdict changed: {deck_why}"
        result = save_memo(workbook, config, path, output_dir=folder)
        why = result["why_unavailable"]
        assert why and reason in why, f"Memo accepted a tampered analysis ({change.__name__}): {why}"
        paragraphs, tables = docx_body(result["docx"])
        check_ai_text(paragraphs, None, f"Northwind ({change.__name__})")
        check_kpi_table(tables, table, flags, f"Northwind ({change.__name__})")
        deck = "the deck accepts it, the memo doesn't" if deck_accepts else "the deck rejects it too"
        print(f"✓ Northwind with a tampered analysis ({change.__name__}): AI commentary unavailable, every "
              f"computed number still there ({deck}): {why[:60]}...")


def check_number_check_catches_a_planted_number(memo, config, folder):
    """Plant a number the workbook doesn't show into a copy of a saved memo: the number check must fail."""
    workbook = Path(COMPANIES[0]["answer_key"].OUTPUT_PATH)
    _, _, allowed = workbook_allowed(save_metrics_workbook(workbook, config, folder))
    document = Document(memo["docx"])
    cell = document.tables[0].rows[5].cells[1]              # NRR's latest value
    cell.paragraphs[0].runs[0].text = re.sub(r"\d", "4", cell.text)   # 97.1% -> 44.4%
    planted = Path(folder) / "planted.docx"
    document.save(planted)
    paragraphs, tables = docx_body(planted)
    try:
        check_numbers(paragraphs + [c for table in tables for row in table for c in row], allowed, "planted")
    except AssertionError as error:
        return str(error)
    raise AssertionError("The number check didn't catch a number planted in the memo")


def check_a_memo_with_changes(config, folder):
    """A memo with a What changed section (Task 10) passes the number check, which leaves only that section out.

    Northwind's Q1 2026 workbook is run, then today's, so the memo compares the two. Its section
    holds last quarter's values and each move's size: numbers today's workbook doesn't show (proven
    here, so the exception is needed), checked line by line by check_diff.py instead.
    """
    from check_diff import last_quarter_workbook   # here, not at the top: check_diff imports main.py
    from main import run_company
    workbook, output_dir = Path(COMPANIES[0]["answer_key"].OUTPUT_PATH), Path(folder) / "changes"
    for path in (last_quarter_workbook(COMPANIES[0]["answer_key"], folder), workbook):
        with contextlib.redirect_stdout(io.StringIO()):
            run_company(path, config, True, output_dir=output_dir)
    docx_path, pdf_path = memo_paths(workbook, output_dir)
    _, _, allowed = workbook_allowed(output_dir / f"{workbook.stem}_metrics.xlsx")
    paragraphs, _ = docx_body(docx_path)
    assert CHANGES_HEADING in paragraphs, "the memo has no What changed section to check"
    section = "\n".join(paragraphs).split(CHANGES_HEADING, 1)[1].split(KEY_METRICS_HEADING, 1)[0]
    outside = sentence_tokens(section) - allowed
    assert outside, "every number in the What changed section is in the workbook: the exception isn't needed"
    memo_numbers_check(docx_path, pdf_path, expected_footer(workbook, output_dir, None), allowed, "Northwind")
    check_a_number_after_the_section_is_still_caught(docx_path, paragraphs, allowed, folder)
    return sorted(outside)


def check_a_number_after_the_section_is_still_caught(docx_path, paragraphs, allowed, folder):
    """Leaving the section out must not leave out what follows it: plant a number after it, and the check must fail."""
    document = Document(docx_path)
    start = paragraphs.index(KEY_METRICS_HEADING)
    target = next(item for item in document.paragraphs[start:] if re.search(r"\d", item.text))
    target.runs[0].text = re.sub(r"\d", "4", target.runs[0].text)   # e.g. 13.0 mo -> 44.4 mo
    planted = Path(folder) / "planted_after_changes.docx"
    document.save(planted)
    try:
        check_numbers([without_changes("\n".join(docx_body(planted)[0]))], allowed, "planted")
    except AssertionError:
        return
    raise AssertionError("A number planted after the What changed section wasn't caught")


def main():
    config = load_config()
    memos = [check_company(company, config, PROJECT_DIR / "output") for company in COMPANIES]
    with tempfile.TemporaryDirectory() as folder:
        check_bad_analyses_are_unavailable(config, folder)
        caught = check_number_check_catches_a_planted_number(memos[0], config, folder)
        print(f"✓ The number check fails on a memo with a planted number: {caught}")
        outside = check_a_memo_with_changes(config, folder)
        print(f"✓ A memo with What changed since the last run passes; that section alone holds numbers the "
              f"workbook doesn't show (checked by check_diff.py): {', '.join(outside[:4])}...")
    print("All checks passed")


if __name__ == "__main__":
    main()
