"""Build the 5-slide board deck: output/<company>_board_pack.pptx (build step 4).

Slides (every one on templates/base.pptx's "Title and Content" layout, with a footer):
1. Summary:          company and quarter, AI headline, 3 wins and 3 risks, "6 of 9 flags tripped"
2. Key metrics:      latest quarter, prior quarter, budget or threshold, status (red / green / gray)
3. ARR and cash:     two matplotlib charts (charts.py); a blank quarter shows as a gap
4. Risks and flags:  each tripped flag with value vs threshold, the combo rule; the Data gaps line beside them
5. Questions:        the AI's 3 questions for management

Inputs: metrics, flags and data gaps (metrics.py) and the AI analysis JSON (analyze.py).
The analysis is checked again here against today's numbers; if it's missing or fails,
slides 1 and 5 say "AI summary unavailable" and everything else is built as normal
(CLAUDE.md step 4 decisions K and L).

Rules:
- No math and no hand-typed numbers: every number comes from metrics.py, config.yaml or the
  validated analysis, and is formatted by metrics.format_value (% formatting happens only at output).
- Text must fit: text_fit.py shrinks it to a 12 pt floor, then stops with an error.
- Footer on every slide: fictional-data note, source file, run date, code commit, model.
- Every slide is stamped "DRAFT - NOT REVIEWED" until a person approves the deck with approve.py
  (provenance.py decides; a changed workbook or config.yaml brings the stamp back).

Run: python build_deck.py data/northwind.xlsx                  (uses output/northwind_analysis.json if it exists)
     python build_deck.py data/northwind.xlsx --no-analysis    (placeholder on slides 1 and 5)
"""

import argparse
import datetime
import json
import math
from functools import lru_cache
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pydantic import ValidationError

from analyze import BoardSummary, build_payload, payload_to_text, validate_summary
from charts import arr_chart, cash_chart, save_chart
from clean import clean_workbook
from excel_output import INFINITE_LABELS, NO_GAPS_LABEL, STATUS_COLORS, gap_label, status_label
from make_template import CONTENT_LAYOUT, DARK_GRAY, LIGHT_GRAY, MID_GRAY, NAVY, TEMPLATE_PATH, WHITE
from metrics import (CANNOT_EVALUATE, CONFIG_PATH, FLAG_RULES, METRIC_LABELS, NO_PRIOR_PERIOD, REASON_DISPLAY,
                     TRIP, compute_metrics, data_gaps, evaluate_flags, format_value, load_config, metric_reasons,
                     reason_text, runway_at_next_budget, runway_context_label)
from provenance import (NOT_REVIEWED, approval_status, deck_status, file_sha256, git_commit, manifest_path,
                        read_manifest, save_manifest)
from text_fit import MIN_FONT_PT, TextDoesNotFitError, fit_table, paragraph, shrink_to_fit

OUTPUT_DIR = Path(__file__).parent / "output"
CHART_FOLDER = "charts"   # PNGs go in output/charts/

PLACEHOLDER_TEXT = "AI summary unavailable"
PLACEHOLDER_NOTE = ("The headline, wins, risks and questions are written by Claude, and no validated version "
                    "is available for this deck. Every number on the other slides is computed in Python "
                    "and is unaffected.")
FICTIONAL_NOTE = "Fictional data, generated for demonstration"
NO_AI_MODEL = "no AI text"            # the footer's model when the deck shows the placeholder
LOCAL_CHANGES = "*"                   # after the commit: the code had uncommitted edits when this ran
                                      # (spelled out in the manifest; the footer has room for one line only)
UNKNOWN_PROMPT_VERSION = "unknown (saved before prompt versions)"
NOT_APPLICABLE = "—"
COMBO_TABLE_TEXT = "rule on Risks and flags slide"

# ---------------------------------------------------------------------------
# Look: font sizes (pt), spacing (pt) and layout sizes. Positions come from the template.
# ---------------------------------------------------------------------------

TITLE_SIZE = 28
HEADLINE_SIZE = 22
FLAG_COUNT_SIZE = 18
HEADING_SIZE = 18
BODY_SIZE = 14        # wins and risks: two columns of AI text, so a little smaller
LIST_SIZE = 16        # risks-and-flags and questions: one wide column
TABLE_SIZE = 14
FOOTER_SIZE = MIN_FONT_PT

HEADING_SPACE = 6     # after a heading
POINT_TITLE_SPACE = 2  # between a point's title and its detail
ITEM_SPACE = 10       # after each list item
SECTION_SPACE = 14    # after the last item of a section
QUESTION_SPACE = 18

# The DRAFT watermark, until a person approves the deck (provenance.py, approve.py).
WATERMARK_SIZE = 48
WATERMARK_ROTATION = 315          # degrees: bottom-left to top-right across the slide
WATERMARK_ALPHA_PERCENT = 25      # see-through, so the numbers underneath stay readable
WATERMARK_WIDTH = Inches(9.5)
WATERMARK_HEIGHT = Inches(1.2)

GAP = Inches(0.15)            # vertical gap between boxes
COLUMN_GAP = Inches(0.35)     # between the two columns (wins/risks, the two charts)
HEADLINE_HEIGHT = Inches(1.05)
FLAG_COUNT_HEIGHT = Inches(0.5)
TEXT_MARGIN_X = Inches(0.1)
TEXT_MARGIN_Y = Inches(0.05)
CELL_MARGIN_X = Inches(0.08)
CELL_MARGIN_Y = Inches(0.04)

# Key metrics table: share of the table width per column (they add up to 1).
KPI_COLUMN_SHARES = [0.27, 0.14, 0.14, 0.21, 0.24]
# Rows that aren't flags, shown first: (metric, the "vs budget" metric for its budget column or None).
CONTEXT_ROWS = [("ending_arr", "arr_vs_budget"), ("arr_yoy", None), ("gross_margin", None)]

KIND_WORDS = {"min": "trips below", "max": "trips above"}              # exactly at the threshold passes
FLAG_KINDS = {column: kind for _, column, _, kind in FLAG_RULES}       # metric -> "min" or "max"
BODY_TYPES = {PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY}


# ---------------------------------------------------------------------------
# 1. Data: everything the slides show, computed once by metrics.py
# ---------------------------------------------------------------------------

def collect_deck_data(company, source_name, actuals, next_budget, config):
    """Compute metrics, reasons, flags (latest quarter) and data gaps; return them in one dict."""
    metrics = compute_metrics(actuals)
    reasons = metric_reasons(actuals, metrics)
    flags = evaluate_flags(metrics, reasons, config)
    quarters = list(metrics.index)
    return {
        "company": company, "source_name": source_name, "config": config,
        "actuals": actuals, "next_budget": next_budget, "metrics": metrics, "reasons": reasons,
        "flags": flags, "gaps": data_gaps(actuals, metrics, flags),
        "latest": quarters[-1], "prior": quarters[-2] if len(quarters) > 1 else None,
        "runway_at_budget": runway_at_next_budget(actuals, next_budget),
    }


# ---------------------------------------------------------------------------
# 2. Text pieces (formatting only - no math)
# ---------------------------------------------------------------------------

def value_text(data, metric, quarter):
    """One metric value as slide text: the number, "∞ (...)", or why there's no number."""
    if quarter is None:
        return REASON_DISPLAY[NO_PRIOR_PERIOD]  # a one-quarter workbook has no prior quarter
    reason = data["reasons"].loc[quarter, metric]
    if isinstance(reason, str):
        return reason_text(data["actuals"], metric, quarter, reason)
    value = data["metrics"].loc[quarter, metric]
    if math.isinf(value):
        return INFINITE_LABELS[metric]
    return format_value(metric, value)


def threshold_text(flag):
    """'trips below 100.0%' or 'trips above 2.00x'."""
    return f"{KIND_WORDS[FLAG_KINDS[flag['metric']]]} {format_value(flag['metric'], flag['threshold'])}"


def flag_count_text(flags):
    """'6 of 9 flags tripped', plus ', 1 cannot evaluate' when some flags lack what they need."""
    tripped = sum(flag["status"] == TRIP for flag in flags)
    unevaluated = sum(flag["status"] == CANNOT_EVALUATE for flag in flags)
    text = f"{tripped} of {len(flags)} flags tripped"
    if unevaluated:
        text += f", {unevaluated} cannot evaluate"
    return text


def gaps_lines(gaps):
    """Every metric and flag with data missing, one line per set of quarters they miss.

    ['Q1 2025: Ending ARR ($K), NRR (annualized)', 'Q1 2025 + Q2 2025: ARR growth QoQ']
    """
    if not gaps:
        return [NO_GAPS_LABEL]
    groups = {}  # quarters -> labels, in the order first seen
    for name, quarters in gaps.items():
        groups.setdefault(" + ".join(quarters), []).append(gap_label(name))
    return [f"{quarters}: {', '.join(labels)}" for quarters, labels in groups.items()]


def gaps_text(gaps):
    """The Data gaps line as one piece of text: the groups joined with '; '."""
    return "; ".join(gaps_lines(gaps))


def points_text(ratio):
    """A drop in NRR as points: 0.01 -> '1.0 pts'."""
    return format_value("nrr", ratio).removesuffix("%") + " pts"


def combo_text(flag, config):
    """The combo rule: its status and what it tests, with the settings from config.yaml."""
    rule = (f"trips when NRR falls by at least {points_text(config['combo_min_nrr_drop'])} and pipeline rises "
            f"at every step over the last {config['combo_lookback_quarters']} quarters")
    return f"{flag['flag']}: {status_label(flag)} ({rule})"


def runway_lines(data):
    """Runway at current burn (the flag) and at next quarter's budgeted burn (context), for the cash chart."""
    current = value_text(data, "runway_months", data["latest"])
    runway = data["runway_at_budget"]
    at_budget = runway_context_label(runway, data["next_budget"] is not None) or format_value("runway_months", runway)
    return f"{METRIC_LABELS['runway_months']}: {current}\nAt next quarter's budgeted burn: {at_budget}"


# ---------------------------------------------------------------------------
# 3. The AI analysis: used only if it's valid for today's numbers
# ---------------------------------------------------------------------------

def load_analysis(path, payload):
    """Read a saved analysis and check it again. Returns (BoardSummary, None) or (None, why not).

    Checks: the file exists and is JSON; the analysis passed when it was made; it is for this
    company and this latest quarter; it has the required shape; and analyze.validate_summary
    still passes against today's payload - so every number in it is in today's data, and the
    text still fits slides 1 and 5 (ai_text_problems).
    """
    path = Path(path)
    if not path.exists():
        return None, f"no analysis file ({path.name})"
    try:
        saved = json.loads(path.read_text())
    except json.JSONDecodeError:
        return None, f"{path.name} is not valid JSON"
    if not isinstance(saved, dict) or saved.get("summary") is None:
        return None, "the analysis failed validation when it was made (no summary saved)"
    saved_payload = saved.get("payload") or {}
    for key in ("company", "latest_quarter"):
        if saved_payload.get(key) != payload[key]:
            return None, f"the analysis is for {key} {saved_payload.get(key)!r}, not {payload[key]!r}"
    try:
        summary = BoardSummary.model_validate(saved["summary"])
    except ValidationError as error:
        return None, f"the analysis doesn't have the required shape ({error.error_count()} errors)"
    problems = validate_summary(summary, payload_to_text(payload))
    if problems:  # today's numbers, and the room the slides have for the text
        return None, "the analysis fails validation for this deck: " + "; ".join(problems)
    return summary, None


def analysis_details(path):
    """How a saved analysis was made: {"model": ..., "prompt_version": ...}, or None if unreadable.

    The footer shows the model, and main.py records both in the run manifest. An analysis saved
    before prompt versions existed says so rather than guessing a version.
    """
    try:
        saved = json.loads(Path(path).read_text())
    except (OSError, json.JSONDecodeError):
        return None
    run_info = saved.get("run_info") or {}
    return {"model": run_info.get("model"),
            "prompt_version": saved.get("prompt_version") or UNKNOWN_PROMPT_VERSION}


# ---------------------------------------------------------------------------
# 4. Slide helpers: boxes, text, fitting
# ---------------------------------------------------------------------------

def layout_box(layout, placeholder_types):
    """(left, top, width, height) of the layout's first placeholder of one of these types."""
    for placeholder in layout.placeholders:
        if placeholder.placeholder_format.type in placeholder_types:
            return placeholder.left, placeholder.top, placeholder.width, placeholder.height
    raise ValueError(f"Template layout '{layout.name}' is missing a placeholder - run python make_template.py")


def find_content_layout(presentation):
    """The template's "Title and Content" layout, or a clear error if the template wasn't built."""
    layout = presentation.slide_layouts.get_by_name(CONTENT_LAYOUT)
    if layout is None:
        raise ValueError(f"{TEMPLATE_PATH.name} has no '{CONTENT_LAYOUT}' layout - run python make_template.py")
    return layout


@lru_cache(maxsize=1)  # read once: measuring text shouldn't open the template file every time
def content_area():
    """(left, top, width, height) of the area each slide fills, taken from the template."""
    return layout_box(find_content_layout(Presentation(TEMPLATE_PATH)), BODY_TYPES)


def column_box(box, position):
    """One of two side-by-side columns inside `box` (EMU): position 0 is the left one."""
    left, top, width, height = box
    column_width = (width - COLUMN_GAP) // 2
    return left + position * (column_width + COLUMN_GAP), top, column_width, height


def points(length):
    """EMU (PowerPoint's unit: 914,400 per inch) -> points."""
    return Emu(length).pt


def write_paragraphs(frame, paragraphs):
    """Write sized paragraphs into a text frame. Autofit is off: our sizes are the sizes."""
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = TEXT_MARGIN_X
    frame.margin_top = frame.margin_bottom = TEXT_MARGIN_Y
    for index, item in enumerate(paragraphs):
        target = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        target.space_after = Pt(item["space_after"])
        run = target.add_run()
        run.text = item["text"]
        run.font.size = Pt(item["size"])
        run.font.bold = item["bold"]
        run.font.color.rgb = RGBColor.from_string(item["color"] or DARK_GRAY)


def fitted(paragraphs, width, height, where):
    """Shrink paragraphs to fit a box of this size (EMU), leaving room for the text margins."""
    return shrink_to_fit(paragraphs, points(width - 2 * TEXT_MARGIN_X), points(height - 2 * TEXT_MARGIN_Y), where)


def add_text_box(slide, name, box, paragraphs, deck):
    """Add a named text box and fill it with the paragraphs, shrunk to fit (or stop with an error)."""
    left, top, width, height = box
    sized = fitted(paragraphs, width, height, f"{deck['where']}, {name}")
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.name = name
    write_paragraphs(shape.text_frame, sized)
    return shape


def add_columns(slide, columns, top, height, deck):
    """Two text boxes side by side across the content area, at matching font sizes.

    columns = [(name, paragraphs), (name, paragraphs)]. Each column is fitted on its own first;
    then both use the bigger shrink, so the two sides never show different sizes.
    """
    left, _, width, _ = deck["area"]
    boxes = [column_box((left, top, width, height), position) for position in range(len(columns))]
    shrink = 0
    for (name, paragraphs), box in zip(columns, boxes):
        sized = fitted(paragraphs, box[2], box[3], f"{deck['where']}, {name}")
        shrink = max(shrink, paragraphs[0]["size"] - sized[0]["size"])
    for (name, paragraphs), box in zip(columns, boxes):
        same_size = [{**item, "size": item["size"] - shrink} for item in paragraphs]
        add_text_box(slide, name, box, same_size, deck)


def set_title(slide, text, deck):
    """Write the slide title into the layout's title placeholder, shrunk to fit."""
    title = slide.shapes.title
    sized = fitted([paragraph(text, TITLE_SIZE, bold=True, color=NAVY)], title.width, title.height,
                   f"{deck['where']}, Title")
    write_paragraphs(title.text_frame, sized)


def new_slide(presentation, layout):
    """Add a slide and remove its body placeholder: each slide places its own boxes in that area."""
    slide = presentation.slides.add_slide(layout)
    for placeholder in list(slide.placeholders):
        if placeholder.placeholder_format.type in BODY_TYPES:
            placeholder.element.getparent().remove(placeholder.element)
    return slide


def commit_text(commit=None):
    """The code this deck was built from: "678d5e7", or "678d5e7*" when it had uncommitted edits."""
    commit = commit or git_commit()
    return commit["commit"] + (LOCAL_CHANGES if commit["uncommitted_changes"] else "")


def add_footer(slide, deck, run_date):
    """Footer on every slide: fictional-data note, source file, run date, code commit, model.

    The labels ("Source:", "Run date:") are left off on purpose: with them the line is 814 pt wide
    and wraps, without them it is 620 pt and fits on one line at 12 pt.
    """
    text = " | ".join([FICTIONAL_NOTE, deck["data"]["source_name"], run_date.isoformat(),
                       deck["commit"], deck["model"]])
    add_text_box(slide, "Footer", deck["footer_box"], [paragraph(text, FOOTER_SIZE, color=MID_GRAY)], deck)


def set_alpha(run, percent):
    """Make a run's text see-through. python-pptx sets a colour but not its opacity.

    The colour is stored as <a:solidFill><a:srgbClr val="1F2A44"/>, and DrawingML puts opacity inside
    that colour as <a:alpha val="25000"/> - a percentage in thousandths.
    """
    colour = run._r.get_or_add_rPr().find(qn("a:solidFill")).find(qn("a:srgbClr"))
    colour.append(colour.makeelement(qn("a:alpha"), {"val": str(percent * 1000)}))


def add_watermark(slide, deck):
    """Stamp DRAFT - NOT REVIEWED diagonally across the slide, on top of everything else.

    On top, not behind: slides 2 and 3 are covered by an opaque table and two chart images, and a
    watermark underneath them would be invisible on exactly the slides carrying the numbers. It is
    see-through, so those numbers stay readable through it.
    """
    slide_width, _ = deck["slide_size"]
    _, area_top, _, area_height = deck["area"]
    box = ((slide_width - WATERMARK_WIDTH) // 2, area_top + (area_height - WATERMARK_HEIGHT) // 2,
           WATERMARK_WIDTH, WATERMARK_HEIGHT)
    shape = add_text_box(slide, "Watermark", box,
                         [paragraph(NOT_REVIEWED, WATERMARK_SIZE, bold=True, color=NAVY)], deck)
    shape.rotation = WATERMARK_ROTATION
    written = shape.text_frame.paragraphs[0]
    written.alignment = PP_ALIGN.CENTER
    set_alpha(written.runs[0], WATERMARK_ALPHA_PERCENT)
    return shape


# ---------------------------------------------------------------------------
# 5. Slide 1: Summary
# ---------------------------------------------------------------------------

def points_paragraphs(heading, items):
    """A column heading, then each point's title (bold) and detail."""
    result = [paragraph(heading, HEADING_SIZE, bold=True, color=NAVY, space_after=HEADING_SPACE)]
    for item in items:
        result.append(paragraph(item.title, BODY_SIZE, bold=True, space_after=POINT_TITLE_SPACE))
        result.append(paragraph(item.detail, BODY_SIZE, space_after=ITEM_SPACE))
    return result


def summary_boxes(area):
    """Slide 1's three boxes, top to bottom: the headline, the flag count, the wins/risks columns.

    ai_text_problems measures the same boxes, so what is measured can't drift from what is drawn.
    """
    left, top, width, height = area
    count_top = top + HEADLINE_HEIGHT + GAP
    columns_top = count_top + FLAG_COUNT_HEIGHT + GAP
    return ((left, top, width, HEADLINE_HEIGHT),
            (left, count_top, width, FLAG_COUNT_HEIGHT),
            (left, columns_top, width, top + height - columns_top))


def headline_paragraph(text, from_ai):
    """The headline: navy for Claude's, gray for the "AI summary unavailable" placeholder."""
    return paragraph(text, HEADLINE_SIZE, bold=True, color=NAVY if from_ai else MID_GRAY)


def summary_columns(summary):
    """Slide 1's two columns of AI text, as [(box name, paragraphs), ...]."""
    return [("Wins", points_paragraphs("Wins", summary.wins)),
            ("Risks", points_paragraphs("Risks", summary.risks))]


def summary_slide(slide, deck):
    """Headline, flag count, then wins and risks side by side (or the placeholder note)."""
    data, summary = deck["data"], deck["summary"]
    set_title(slide, f"{data['company']}: {data['latest']} board update", deck)
    headline_box, count_box, columns_box = summary_boxes(deck["area"])

    headline = summary.headline if summary else PLACEHOLDER_TEXT
    add_text_box(slide, "Headline", headline_box, [headline_paragraph(headline, summary is not None)], deck)
    add_text_box(slide, "Flag count", count_box,
                 [paragraph(flag_count_text(data["flags"]), FLAG_COUNT_SIZE, bold=True)], deck)

    _, columns_top, _, columns_height = columns_box
    if summary is None:
        add_text_box(slide, "AI note", columns_box, [paragraph(PLACEHOLDER_NOTE, BODY_SIZE, color=MID_GRAY)], deck)
        return
    add_columns(slide, summary_columns(summary), columns_top, columns_height, deck)


# ---------------------------------------------------------------------------
# 6. Slide 2: Key metrics table
# ---------------------------------------------------------------------------

def kpi_rows(data):
    """Table body: [(cells, flag status or None)]. Context rows first, then every flag."""
    latest, prior = data["latest"], data["prior"]
    rows = []
    for metric, budget_metric in CONTEXT_ROWS:
        budget = f"vs budget: {value_text(data, budget_metric, latest)}" if budget_metric else NOT_APPLICABLE
        cells = [METRIC_LABELS[metric], value_text(data, metric, latest), value_text(data, metric, prior),
                 budget, NOT_APPLICABLE]
        rows.append((cells, None))
    for flag in data["flags"]:
        if flag["metric"] is None:  # the combo rule has no single value
            cells = [flag["flag"], NOT_APPLICABLE, NOT_APPLICABLE, COMBO_TABLE_TEXT, status_label(flag)]
        else:
            cells = [flag["flag"], value_text(data, flag["metric"], latest), value_text(data, flag["metric"], prior),
                     threshold_text(flag), status_label(flag)]
        rows.append((cells, flag["status"]))
    return rows


def kpi_header(data):
    return ["Metric", data["latest"], data["prior"] or "Prior quarter", "Budget or threshold", "Status"]


def column_widths(total_width):
    """Split the table width by KPI_COLUMN_SHARES; the last column takes any rounding remainder."""
    widths = [int(total_width * share) for share in KPI_COLUMN_SHARES[:-1]]
    return widths + [total_width - sum(widths)]


def write_cell(cell, text, size, fill_hex, text_hex, bold=False):
    """One table cell: fill, margins, and a single run of text."""
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill_hex)
    cell.margin_left = cell.margin_right = CELL_MARGIN_X
    cell.margin_top = cell.margin_bottom = CELL_MARGIN_Y
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = True
    run = cell.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size, run.font.bold = Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(text_hex)


def fill_table(table, header, rows, size):
    """Header in navy; rows striped white / light gray; each flag's status cell in its status color."""
    for column, text in enumerate(header):
        write_cell(table.cell(0, column), text, size, NAVY, WHITE, bold=True)
    for row_number, (cells, status) in enumerate(rows, start=1):
        stripe = LIGHT_GRAY if row_number % 2 == 0 else WHITE
        for column, text in enumerate(cells):
            fill_hex, text_hex = stripe, DARK_GRAY
            if status is not None and column == len(cells) - 1:
                fill_hex, text_hex = STATUS_COLORS[status]
            write_cell(table.cell(row_number, column), text, size, fill_hex, text_hex)


def kpi_slide(slide, deck):
    """Latest quarter, prior quarter, budget or threshold, status for each key metric and flag."""
    data = deck["data"]
    prior_words = f" vs {data['prior']}" if data["prior"] else ""
    set_title(slide, f"Key metrics — {data['latest']}{prior_words}", deck)
    left, top, width, height = deck["area"]
    header, rows = kpi_header(data), kpi_rows(data)
    widths = column_widths(width)

    all_text = [header] + [cells for cells, _ in rows]
    text_widths = [points(column_width - 2 * CELL_MARGIN_X) for column_width in widths]
    size, heights = fit_table(all_text, text_widths, points(height), f"{deck['where']}, KPI table",
                              TABLE_SIZE, cell_padding_pt=points(2 * CELL_MARGIN_Y))

    frame = slide.shapes.add_table(len(all_text), len(header), left, top, width, Pt(sum(heights)))
    frame.name = "KPI table"
    table = frame.table
    table.horz_banding = False  # our own stripes, not the template's
    for column, column_width in enumerate(widths):
        table.columns[column].width = column_width
    for row_number, row_height in enumerate(heights):
        table.rows[row_number].height = Pt(row_height)
    fill_table(table, header, rows, size)


# ---------------------------------------------------------------------------
# 7. Slide 3: Charts
# ---------------------------------------------------------------------------

def charts_slide(slide, deck):
    """ARR (with net new ARR) on the left, ending cash (with runway) on the right."""
    data = deck["data"]
    quarters = list(data["metrics"].index)
    set_title(slide, f"ARR and cash — {quarters[0]} to {quarters[-1]}", deck)
    left, top, width, height = deck["area"]
    chart_width = (width - COLUMN_GAP) // 2
    size_inches = (Emu(chart_width).inches, Emu(height).inches)  # drawn at slide size: 12 pt stays 12 pt

    stem = Path(data["source_name"]).stem
    charts = [
        ("ARR chart", arr_chart(quarters, data["metrics"]["ending_arr"], data["metrics"]["net_new_arr"], size_inches),
         f"{stem}_arr_chart.png"),
        ("Cash chart", cash_chart(quarters, data["actuals"]["ending_cash"], runway_lines(data), size_inches),
         f"{stem}_cash_chart.png"),
    ]
    for position, (name, figure, file_name) in enumerate(charts):
        path = save_chart(figure, deck["chart_dir"] / file_name)
        picture = slide.shapes.add_picture(str(path), left + position * (chart_width + COLUMN_GAP), top,
                                           chart_width, height)
        picture.name = name


# ---------------------------------------------------------------------------
# 8. Slide 4: Risks and flags
# ---------------------------------------------------------------------------

def section(heading, lines):
    """A bold heading and its bullet lines; extra space after the last line."""
    result = [paragraph(heading, HEADING_SIZE, bold=True, color=NAVY, space_after=HEADING_SPACE)]
    for number, line in enumerate(lines, start=1):
        space = SECTION_SPACE if number == len(lines) else ITEM_SPACE
        result.append(paragraph(f"• {line}", LIST_SIZE, space_after=space))
    return result


def flags_paragraphs(data):
    """Tripped flags (value vs threshold), flags that can't be evaluated, and the combo rule."""
    metric_flags = [flag for flag in data["flags"] if flag["metric"] is not None]
    combo_flags = [flag for flag in data["flags"] if flag["metric"] is None]
    tripped = [f"{flag['flag']}: {value_text(data, flag['metric'], flag['quarter'])} ({threshold_text(flag)})"
               for flag in metric_flags if flag["status"] == TRIP]
    unevaluated = [f"{flag['flag']}: {status_label(flag)}" for flag in metric_flags if flag["status"] == CANNOT_EVALUATE]

    result = section(f"Tripped flags ({flag_count_text(data['flags'])})", tripped or ["None"])
    if unevaluated:
        result += section("Cannot evaluate", unevaluated)
    if combo_flags:
        result += section("Combo rule", [combo_text(flag, data["config"]) for flag in combo_flags])
    return result


def risks_slide(slide, deck):
    """Flags on the left; the Data gaps line (one bullet per set of quarters) on the right."""
    data = deck["data"]
    set_title(slide, f"Risks and flags — {data['latest']}", deck)
    _, top, _, height = deck["area"]
    add_columns(slide, [("Risks and flags", flags_paragraphs(data)),
                        ("Data gaps", section("Data gaps (data missing)", gaps_lines(data["gaps"])))], top, height, deck)


# ---------------------------------------------------------------------------
# 9. Slide 5: Questions for management
# ---------------------------------------------------------------------------

def questions_paragraphs(summary):
    """The AI's questions, numbered. ai_text_problems measures these same paragraphs."""
    return [paragraph(f"{number}. {question}", LIST_SIZE, space_after=QUESTION_SPACE)
            for number, question in enumerate(summary.questions, start=1)]


def questions_slide(slide, deck):
    data, summary = deck["data"], deck["summary"]
    set_title(slide, f"Questions for management — {data['latest']}", deck)
    if summary is None:
        paragraphs = [paragraph(PLACEHOLDER_TEXT, HEADLINE_SIZE, bold=True, color=MID_GRAY, space_after=SECTION_SPACE),
                      paragraph(PLACEHOLDER_NOTE, BODY_SIZE, color=MID_GRAY)]
    else:
        paragraphs = questions_paragraphs(summary)
    add_text_box(slide, "Questions", deck["area"], paragraphs, deck)


# ---------------------------------------------------------------------------
# 10. Putting it together
# ---------------------------------------------------------------------------

SLIDE_BUILDERS = [summary_slide, kpi_slide, charts_slide, risks_slide, questions_slide]


def slide_number(build_slide):
    """Which slide a builder makes, counting from 1, so a message names the slide the reader sees."""
    return SLIDE_BUILDERS.index(build_slide) + 1


def ai_text_problems(summary):
    """Problems if Claude's text is too long for its boxes on slides 1 and 5, even at the 12 pt floor.

    analyze.validate_summary calls this, so an over-long answer is caught with the other validation
    problems and gets the one retry. If it still doesn't fit, load_analysis rejects it and the deck
    is built with the placeholder - a company is never left without a deck (decision K).
    """
    headline_box, _, columns_box = summary_boxes(content_area())
    summary_slide_number = slide_number(summary_slide)
    boxes = [(f"slide {summary_slide_number} (Headline)", headline_box,
              [headline_paragraph(summary.headline, from_ai=True)], "shorten the headline")]
    for position, (name, paragraphs) in enumerate(summary_columns(summary)):
        advice = f"shorten each {name.lower()[:-1]} detail"   # "Wins" -> "shorten each win detail"
        boxes.append((f"slide {summary_slide_number} ({name})", column_box(columns_box, position),
                      paragraphs, advice))
    boxes.append((f"slide {slide_number(questions_slide)} (Questions)", content_area(),
                  questions_paragraphs(summary), "shorten the questions"))

    problems = []
    for where, box, paragraphs, advice in boxes:
        _, _, width, height = box
        try:
            fitted(paragraphs, width, height, where)
        except TextDoesNotFitError:
            problems.append(f"AI text does not fit {where}: {advice}")
    return problems


def build_presentation(data, summary, run_date, chart_dir, approval=None, model=None):
    """Build all 5 slides on the template.

    summary = a validated BoardSummary, or None for the placeholder. approval = the record of a
    person having reviewed this deck (provenance.approval_status); without one, every slide is
    watermarked. model = the Claude model whose text is on the deck, for the footer.
    """
    presentation = Presentation(TEMPLATE_PATH)
    layout = find_content_layout(presentation)
    deck = {"data": data, "summary": summary, "chart_dir": Path(chart_dir),
            "area": layout_box(layout, BODY_TYPES), "footer_box": layout_box(layout, {PP_PLACEHOLDER.FOOTER}),
            "slide_size": (presentation.slide_width, presentation.slide_height),
            "commit": commit_text(), "model": model or NO_AI_MODEL}
    for number, build_slide in enumerate(SLIDE_BUILDERS, start=1):
        slide = new_slide(presentation, layout)
        deck["where"] = f"Slide {number}"  # names the slide in any "doesn't fit" error
        build_slide(slide, deck)
        add_footer(slide, deck, run_date)
        if approval is None:  # nobody has signed this off, so it goes out stamped as a draft
            add_watermark(slide, deck)
    return presentation


def record_deck_status(workbook_path, output_dir, approval, ai_text):
    """Update the manifest to describe the deck just written: reviewed or not, AI text or not.

    Only if a manifest is already there. main.py writes the full one (hashes, tokens, cost); this
    keeps it true when the deck is rebuilt on its own - otherwise the file could say "approved by"
    beside a deck that went back to DRAFT because the workbook changed.
    """
    path = manifest_path(workbook_path, output_dir)
    manifest = read_manifest(path)
    if manifest is None:
        return None
    manifest["deck"] = {**manifest.get("deck", {}), "ai_text": ai_text, "status": deck_status(approval)}
    return save_manifest(path, manifest)


def deck_path(workbook_path, output_dir=OUTPUT_DIR):
    """data/northwind.xlsx -> output/northwind_board_pack.pptx"""
    return Path(output_dir) / f"{Path(workbook_path).stem}_board_pack.pptx"


def analysis_path(workbook_path, output_dir=OUTPUT_DIR):
    """data/northwind.xlsx -> output/northwind_analysis.json (where analyze.py saves it)"""
    return Path(output_dir) / f"{Path(workbook_path).stem}_analysis.json"


def save_deck(workbook_path, config, analysis_file=None, run_date=None, output_dir=OUTPUT_DIR):
    """Clean one workbook, build its deck and save it. Returns (deck path, why the AI summary is unavailable or None).

    analysis_file=None means no analysis (the placeholder). An old deck is deleted first, so a
    failed build never leaves last run's deck looking current.
    """
    workbook_path, output_dir = Path(workbook_path), Path(output_dir)
    company = workbook_path.stem.title()   # same rule as analyze.py and main.py
    path = deck_path(workbook_path, output_dir)
    path.unlink(missing_ok=True)

    actuals, next_budget = clean_workbook(workbook_path)
    data = collect_deck_data(company, workbook_path.name, actuals, next_budget, config)
    if analysis_file is None:
        summary, why_unavailable = None, "no analysis requested"
    else:
        summary, why_unavailable = load_analysis(analysis_file, build_payload(company, actuals, next_budget, config))

    # Watermark unless a person approved this deck, for these exact inputs (provenance.py).
    approval, _ = approval_status(read_manifest(manifest_path(workbook_path, output_dir)),
                                  file_sha256(workbook_path), file_sha256(CONFIG_PATH))
    details = analysis_details(analysis_file) if summary else None

    chart_dir = output_dir / CHART_FOLDER
    chart_dir.mkdir(parents=True, exist_ok=True)
    presentation = build_presentation(data, summary, run_date or datetime.date.today(), chart_dir,
                                      approval=approval, model=details["model"] if details else None)
    presentation.save(path)
    record_deck_status(workbook_path, output_dir, approval, ai_text=summary is not None)
    return path, why_unavailable


def main(argv=None):
    parser = argparse.ArgumentParser(description="Build the board deck for one KPI workbook.")
    parser.add_argument("workbook", help="path to a KPI workbook, e.g. data/northwind.xlsx")
    parser.add_argument("--analysis", help="analysis JSON (default: output/<company>_analysis.json)")
    parser.add_argument("--no-analysis", action="store_true", help="build with the 'AI summary unavailable' placeholder")
    args = parser.parse_args(argv)

    analysis_file = None if args.no_analysis else (args.analysis or analysis_path(args.workbook))
    path, why_unavailable = save_deck(args.workbook, load_config(), analysis_file)
    print(f"Saved {path.relative_to(Path(__file__).parent)}")
    if why_unavailable:
        print(f"{PLACEHOLDER_TEXT}: {why_unavailable}")


if __name__ == "__main__":
    main()
