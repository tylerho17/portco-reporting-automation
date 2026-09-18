"""Save the computed metrics, flags and data gaps as an Excel workbook (build step 4b).

Output: output/<company>_metrics.xlsx with three sheets:
- Metrics:   one row per quarter, one column per metric. Cells hold real numbers
             (ratios as decimals, e.g. 0.971); Excel number formats display them as 97.1%.
             A cell with no number says why: "data missing" (gray), "n/a (no prior period)",
             or "n/m ..." (not meaningful). A cell whose flag trips in that quarter is red.
- Flags:     every flag for the latest quarter: value, threshold, status (with the reason when
             it can't be evaluated). Row color: red = tripped, green = passed, gray = cannot evaluate.
- Data gaps: every metric and flag that can't be shown because an input it uses is blank.

No math happens here. metrics.py computes everything; this file only writes and formats.

Run: python excel_output.py data/northwind.xlsx
"""

import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from clean import clean_workbook
from metrics import (CANNOT_EVALUATE, DOLLAR_COLUMNS, FLAG_RULES, METRIC_LABELS, MISSING_INPUT, MONTH_COLUMNS,
                     PASS, TRIP, compute_metrics, data_gaps, evaluate_flags, load_config, metric_reasons,
                     reason_text, runway_at_next_budget, runway_context_label)
from theme import EXCEL_STATUS_COLORS

OUTPUT_DIR = Path(__file__).parent / "output"
SHEET_NAMES = ["Metrics", "Flags", "Data gaps"]

# Excel can't store "no value" or infinity as a number, so those cells get words that say why.
# The words for the three "no number" reasons come from metrics.reason_text, shared with every output.
INFINITE_LABELS = {                            # the CLAUDE.md edge cases
    "burn_multiple": "∞ (ARR shrank)",
    "runway_months": "∞ (not burning)",
    "cac_payback_months": "∞ (never pays back)",
}

STATUS_LABELS = {TRIP: "Tripped", PASS: "Passed"}  # a flag that can't be evaluated shows its reason instead
KIND_LABELS = {"min": "below threshold", "max": "above threshold"}  # when a flag trips
FLAG_KINDS = {name: kind for name, _, _, kind in FLAG_RULES}      # flag name -> "min" or "max"

# Status -> (fill color, text color) as hex RGB. Same light red/green Excel uses for "Bad"/"Good".
# Kept in theme.py with the rest of the palette; the workbook keeps Excel's own fills (Task 3).
STATUS_COLORS = EXCEL_STATUS_COLORS

RUNWAY_CONTEXT_LABEL = "Runway at next quarter's budgeted burn (context, not a flag)"
NO_GAPS_LABEL = "None — every metric and flag has the data it needs"


# ---------------------------------------------------------------------------
# Cell values and formats
# ---------------------------------------------------------------------------

def number_format(column):
    """Excel display format for a metric column. Same rules as metrics.format_value."""
    if column in DOLLAR_COLUMNS:
        return "#,##0"                  # $K with thousands commas: 27,470
    if column in MONTH_COLUMNS:
        return '0.0" mo"'               # 11.0 mo
    if column == "burn_multiple":
        return '0.00"x"'                # 2.35x
    return "0.0%"                       # decimal 0.971 shows as 97.1%


def cell_value(actuals, metrics, reasons, column, quarter):
    """The number itself, or a text label when there is no usable number.

    No number -> the words for its reason ("data missing", "n/a (no prior period)", "n/m ...").
    Infinity -> "∞" plus the reason (e.g. burn multiple when ARR shrank).
    """
    reason = reasons.loc[quarter, column]
    if isinstance(reason, str):
        return reason_text(actuals, column, quarter, reason)
    value = metrics.loc[quarter, column]
    if math.isinf(value):
        return INFINITE_LABELS[column]  # compute_metrics only lets these three metrics be infinite
    return float(value)  # plain Python float; openpyxl writes it exactly


def runway_context_value(runway, has_budget_row):
    """Runway at next quarter's budgeted burn: a number, or a label saying why there isn't one."""
    return runway_context_label(runway, has_budget_row) or float(runway)


def status_label(flag):
    """'Tripped', 'Passed', or 'Cannot evaluate — <reason>'."""
    if flag["status"] == CANNOT_EVALUATE:
        return f"Cannot evaluate — {flag['reason']}"
    return STATUS_LABELS[flag["status"]]


# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------

def color_cell(cell, status):
    """Fill one cell with its status color (red / green / gray) and matching text color."""
    fill_color, text_color = STATUS_COLORS[status]
    cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
    cell.font = Font(color=text_color)


def write_header(sheet, headers):
    """Write bold headers in row 1 and freeze them so they stay visible while scrolling."""
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "B2"  # freezes row 1 and column A


def set_column_widths(sheet, widths):
    """Set each column's width, in characters, left to right."""
    for position, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = width


# ---------------------------------------------------------------------------
# Sheet 1: Metrics
# ---------------------------------------------------------------------------

def tripped_cells(metrics, reasons, config):
    """Every (quarter, metric column) whose flag trips in that quarter.

    The combo rule has no single metric cell, so it only shows on the Flags sheet.
    """
    tripped = set()
    for quarter in metrics.index:
        for flag in evaluate_flags(metrics, reasons, config, quarter):
            if flag["status"] == TRIP and flag["metric"] is not None:
                tripped.add((quarter, flag["metric"]))
    return tripped


def style_metric_cell(cell, column, is_gap, is_tripped):
    """Number format, right alignment, and gray (data missing) or red (tripped) fill.

    Only missing data is gray: "no prior period" and "not meaningful" are not problems to chase.
    """
    cell.number_format = number_format(column)
    cell.alignment = Alignment(horizontal="right")  # text labels line up with the numbers
    if is_gap:
        color_cell(cell, CANNOT_EVALUATE)
    elif is_tripped:
        color_cell(cell, TRIP)


def write_metrics_sheet(sheet, actuals, metrics, reasons, config):
    """One row per quarter, one column per metric, with formats and highlights."""
    columns = list(metrics.columns)
    write_header(sheet, ["Quarter"] + [METRIC_LABELS[column] for column in columns])
    tripped = tripped_cells(metrics, reasons, config)

    for row_number, quarter in enumerate(metrics.index, start=2):  # row 1 is the header
        sheet.cell(row=row_number, column=1, value=quarter)
        for column_number, column in enumerate(columns, start=2):  # column A is the quarter
            value = cell_value(actuals, metrics, reasons, column, quarter)
            cell = sheet.cell(row=row_number, column=column_number, value=value)
            is_gap = reasons.loc[quarter, column] == MISSING_INPUT
            style_metric_cell(cell, column, is_gap, (quarter, column) in tripped)

    set_column_widths(sheet, [10] + [max(len(METRIC_LABELS[c]), 12) + 2 for c in columns])


# ---------------------------------------------------------------------------
# Sheet 2: Flags
# ---------------------------------------------------------------------------

FLAG_HEADERS = ["Flag", "Quarter", "Value", "Threshold", "Trips when", "Status"]


def flag_row(flag, actuals, metrics, reasons, config):
    """One flag as a list of cell values, in FLAG_HEADERS order."""
    status = status_label(flag)
    if flag["metric"] is None:  # the combo rule is a trend test with no single value
        size = config["combo_lookback_quarters"]
        drop = config["combo_min_nrr_drop"]
        return [flag["flag"], flag["quarter"], "see NRR and Pipeline on Metrics sheet",
                f"last {size} quarters", f"NRR falls at least {drop * 100:g} pt and pipeline rises at every step",
                status]
    value = cell_value(actuals, metrics, reasons, flag["metric"], flag["quarter"])
    return [flag["flag"], flag["quarter"], value, flag["threshold"],
            KIND_LABELS[FLAG_KINDS[flag["flag"]]], status]


def style_flag_row(sheet, row_number, flag):
    """Color the whole row by status; show value and threshold in the metric's format."""
    for cell in sheet[row_number]:
        color_cell(cell, flag["status"])
    if flag["metric"] is not None:
        for column_number in (3, 4):  # Value, Threshold
            cell = sheet.cell(row=row_number, column=column_number)
            cell.number_format = number_format(flag["metric"])
            cell.alignment = Alignment(horizontal="right")


def write_runway_context(sheet, runway, has_budget_row, quarter):
    """Below the flag table, after one empty row: runway at next quarter's budgeted burn."""
    row_number = sheet.max_row + 2
    sheet.cell(row=row_number, column=1, value=RUNWAY_CONTEXT_LABEL)
    sheet.cell(row=row_number, column=2, value=quarter)
    cell = sheet.cell(row=row_number, column=3, value=runway_context_value(runway, has_budget_row))
    cell.number_format = number_format("runway_months")
    cell.alignment = Alignment(horizontal="right")


def write_flags_sheet(sheet, flags, actuals, metrics, reasons, config, runway_at_budget, has_budget_row):
    """Every flag for one quarter (the latest), colored by status, plus runway context."""
    write_header(sheet, FLAG_HEADERS)
    for flag in flags:
        sheet.append(flag_row(flag, actuals, metrics, reasons, config))
        style_flag_row(sheet, sheet.max_row, flag)
    write_runway_context(sheet, runway_at_budget, has_budget_row, flags[0]["quarter"])
    set_column_widths(sheet, [56, 10, 36, 18, 42, 32])


# ---------------------------------------------------------------------------
# Sheet 3: Data gaps
# ---------------------------------------------------------------------------

def gap_label(name):
    """'nrr' -> 'NRR (annualized)'; 'flag: Rule of 40' -> 'Flag: Rule of 40'."""
    if name.startswith("flag: "):
        return "Flag: " + name.removeprefix("flag: ")
    return METRIC_LABELS[name]


def write_gaps_sheet(sheet, gaps):
    """One row per affected metric or flag, with the quarters it can't be shown for."""
    write_header(sheet, ["Metric or flag", "Quarters with data missing"])
    if not gaps:
        sheet.append([NO_GAPS_LABEL, ""])
    for name, quarters in gaps.items():
        sheet.append([gap_label(name), ", ".join(quarters)])
    set_column_widths(sheet, [52, 30])


# ---------------------------------------------------------------------------
# Putting it together
# ---------------------------------------------------------------------------

def build_workbook(actuals, next_budget, config):
    """Compute metrics, flags and gaps, then write all three sheets into a new workbook."""
    metrics = compute_metrics(actuals)
    reasons = metric_reasons(actuals, metrics)
    flags = evaluate_flags(metrics, reasons, config)  # latest quarter
    gaps = data_gaps(actuals, metrics, flags)

    book = Workbook()
    metrics_sheet = book.active  # a new workbook starts with one empty sheet; reuse it
    metrics_sheet.title = SHEET_NAMES[0]
    write_metrics_sheet(metrics_sheet, actuals, metrics, reasons, config)
    write_flags_sheet(book.create_sheet(SHEET_NAMES[1]), flags, actuals, metrics, reasons, config,
                      runway_at_next_budget(actuals, next_budget), next_budget is not None)
    write_gaps_sheet(book.create_sheet(SHEET_NAMES[2]), gaps)
    return book


def output_path(workbook_path, output_dir=OUTPUT_DIR):
    """data/northwind.xlsx -> output/northwind_metrics.xlsx"""
    return Path(output_dir) / f"{Path(workbook_path).stem}_metrics.xlsx"


def save_metrics_workbook(workbook_path, config, output_dir=OUTPUT_DIR):
    """Clean one input workbook and save its metrics workbook. Returns the saved path."""
    actuals, next_budget = clean_workbook(workbook_path)
    path = output_path(workbook_path, output_dir)
    path.parent.mkdir(exist_ok=True)
    build_workbook(actuals, next_budget, config).save(path)
    return path


if __name__ == "__main__":
    input_path = sys.argv[1] if len(sys.argv) > 1 else "data/northwind.xlsx"
    print(f"Saved {save_metrics_workbook(input_path, load_config())}")
