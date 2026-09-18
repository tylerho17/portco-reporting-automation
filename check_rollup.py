"""Automated proof for Task 9: the portfolio rollup matches every company's answer key.

Builds output/portfolio_rollup.pptx and output/portfolio_rollup.xlsx from data/ with rollup.py,
reads both back from disk, and checks:
1. The deck has 3 slides with the titles typed here; the workbook has sheets Ranking, By status, Runway.
2. Ranking: companies in order of flags tripped, counted here from each company's story
   (check_companies.py), not from rollup.py: Fernhollow 7, Northwind 6, Alderpeak 0. Rank, quarter,
   flag count, worst flag, runway and review status agree between the slide and the workbook.
3. Worst flag: the one typed here per company, and it is one the story says trips. Its value and
   runway equal the hand formulas in check_companies.py (the workbook holds real numbers; the slide
   shows them as Excel does). The worst-flag cell is red for a company with a tripped flag, green if
   every flag passed.
4. Every number on the deck (footer aside) is a number in the rollup workbook: nothing typed.
5. Companies by status: the counts are worked out here from the stories (flag status) and from each
   company's manifest and file hashes (review status), and match both the slide and the workbook.
6. The runway chart: the figure rollup.py draws has one bar per company, shortest first, each the
   hand-formula runway, red exactly where the story says the runway flag trips, and a dashed line
   at config.yaml's runway threshold.
7. Nothing overflows (check_deck.py's re-measuring), every font is at least 12 pt, and every
   footer says fictional data, the workbook count, today's date and "no AI text", on one line.
8. A portfolio of 12 companies with 40-character names splits the ranking over 2 slides and still
   fits; an unreadable workbook is listed last, unranked, with its reason. (Temporary folders.)
9. No API: creating an Anthropic client stops the check. The rollup never asks Claude.

Run: python check_rollup.py  -> prints "All checks passed" or stops at the first failure.
"""

import datetime
import math
import shutil
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Emu

import analyze
import rollup
from check_companies import COMPANIES, LATEST, same_number
from check_deck import EXCEL_FORMATS, check_no_overflow, number_tokens, shape
from mapping import mapping_sha256
from metrics import CANNOT_EVALUATE, CONFIG_PATH, PASS, TRIP, load_config
from provenance import approval_status, file_sha256, manifest_path, read_manifest
from text_fit import MIN_FONT_PT, text_width_pt
from theme import EXCEL_STATUS_COLORS, RED, STATUS_COLORS

PROJECT_DIR = Path(__file__).parent
OUTPUT_DIR = PROJECT_DIR / "output"
PICTURE = 13   # MSO_SHAPE_TYPE.PICTURE

# Typed by hand from CLAUDE.md's stories, not imported from rollup.py.
TITLES = [f"Portfolio ranked by flags tripped, {LATEST}", f"Companies by status, {LATEST}",
          f"Runway at current burn by company, {LATEST}"]
WORST = {"Fernhollow": "Runway at current burn", "Northwind": "Runway at current burn", "Alderpeak": None}
NONE_TRIPPED = "None tripped"
RUNWAY_FLAG = "Runway at current burn"
FLAG_STATUS_ROWS = ["Flags tripped", "None tripped, some cannot evaluate", "Every flag passed",
                    "Workbook can't be read"]
REVIEW_ROWS = ["Approved", "Not reviewed", "Out of date", "Not generated"]
FOOTER_END = " | computed metrics only, no AI text"


def refuse(*args, **kwargs):
    raise AssertionError("the rollup tried to create an Anthropic client")


# ---------------------------------------------------------------------------
# What each company should show, worked out here
# ---------------------------------------------------------------------------

def story(company):
    """{"name", "tripped", "unevaluated", "checked", "worst", "runway", "runway_trips", "status"} from the answer key."""
    statuses = list(company["expected_flags"].values())
    tripped = statuses.count(TRIP)
    status = ("Flags tripped" if tripped else
              "None tripped, some cannot evaluate" if CANNOT_EVALUATE in statuses else "Every flag passed")
    return {"name": company["name"], "tripped": tripped, "unevaluated": statuses.count(CANNOT_EVALUATE),
            "checked": len(statuses), "worst": WORST[company["name"]],
            "runway": company["expected_latest"]["runway_months"],
            "runway_trips": company["expected_flags"][RUNWAY_FLAG] == TRIP, "status": status}


def expected_ranking():
    """The three stories, most flags tripped first (no two share a count, so no tie rule is needed)."""
    return sorted((story(company) for company in COMPANIES), key=lambda item: -item["tripped"])


def expected_review(workbook, output_dir):
    """The review status from the manifest and file hashes, worked out here (not with rollup.py or portfolio.py)."""
    manifest = read_manifest(manifest_path(workbook, output_dir))
    if manifest is None:
        return "Not generated"
    hashes = [(manifest.get("input") or {}).get("sha256"), (manifest.get("config") or {}).get("sha256"),
              (manifest.get("mapping") or {}).get("sha256")]
    if hashes != [file_sha256(workbook), file_sha256(CONFIG_PATH), mapping_sha256(workbook)]:
        return "Out of date"
    approval, _ = approval_status(manifest, *hashes)
    return "Approved" if approval else "Not reviewed"


# ---------------------------------------------------------------------------
# Reading the files back
# ---------------------------------------------------------------------------

def display(cell):
    """What Excel shows in a cell: text as is, whole numbers as digits, others in the cell's number format."""
    if cell.value is None or isinstance(cell.value, str):
        return cell.value
    if isinstance(cell.value, int) and cell.number_format == "General":   # counts and ranks
        return str(cell.value)
    assert cell.number_format in EXCEL_FORMATS, f"Unexpected number format {cell.number_format!r} in {cell.coordinate}"
    return EXCEL_FORMATS[cell.number_format].format(cell.value)


def sheet_dicts(sheet):
    """Every row after the header as {header: cell}."""
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2)]


def table_rows(slide, name):
    """A slide table's cell texts (header first) and the table itself."""
    table = shape(slide, name).table
    return [[cell.text for cell in row.cells] for row in table.rows], table


def fill(cell):
    """A cell's fill color (hex, no alpha) from a slide table or an Excel sheet."""
    if hasattr(cell, "fill") and hasattr(cell.fill, "fore_color"):
        return str(cell.fill.fore_color.rgb)
    return cell.fill.fgColor.rgb[-6:]


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def check_files(paths):
    presentation, book = Presentation(paths["deck"]), load_workbook(paths["excel"])
    titles = [slide.shapes.title.text for slide in presentation.slides]
    assert titles == TITLES, f"Rollup titles\nExpected: {TITLES}\nGot:      {titles}"
    assert book.sheetnames == ["Ranking", "By status", "Runway"], f"Rollup sheets: {book.sheetnames}"
    return presentation, book


def check_ranking_sheet(book, stories, reviews):
    """The Ranking sheet: order, counts, worst flag and runway as real numbers equal to the hand formulas."""
    rows = sheet_dicts(book["Ranking"])
    assert [row["Company"].value for row in rows] == [item["name"] for item in stories], "Ranking order wrong"
    for rank, (row, item) in enumerate(zip(rows, stories), start=1):
        name = item["name"]
        got = [row[key].value for key in ("Rank", "Latest quarter", "Flags tripped", "Flags that cannot evaluate",
                                          "Flags checked", "Flag status", "Review status")]
        expected = [rank, LATEST, item["tripped"], item["unevaluated"], item["checked"], item["status"], reviews[name]]
        assert got == expected, f"{name}: Ranking row\nExpected: {expected}\nGot:      {got}"
        assert row["Worst flag"].value == (item["worst"] or NONE_TRIPPED), f"{name}: worst flag {row['Worst flag'].value!r}"
        runway = row["Runway at current burn"].value
        assert isinstance(runway, (int, float)) and same_number(runway, item["runway"]), f"{name}: runway {runway!r}"
        if item["worst"] == RUNWAY_FLAG:
            assert same_number(row["Worst flag value"].value, item["runway"]), f"{name}: worst flag value"
        expected_fill = EXCEL_STATUS_COLORS[TRIP if item["tripped"] else PASS][0]
        assert fill(row["Flag status"]) == expected_fill, f"{name}: Flag status cell fill {fill(row['Flag status'])}"
    return rows


def check_ranking_slide(presentation, sheet_rows, stories):
    """Slide 1 row by row: the same words and numbers as the Ranking sheet; worst-flag cell colored by status."""
    rows, table = table_rows(presentation.slides[0], "Ranking table")
    assert rows[0] == ["Rank", "Company", "Quarter", "Flags tripped", "Worst flag", "Runway at current burn",
                       "Review"], f"Ranking header {rows[0]}"
    for number, (cells, row, item) in enumerate(zip(rows[1:], sheet_rows, stories), start=1):
        name = item["name"]
        count = f"{item['tripped']} of {item['checked']} flags tripped"
        count += f", {item['unevaluated']} cannot evaluate" if item["unevaluated"] else ""
        runway = display(row["Runway at current burn"])
        worst = NONE_TRIPPED
        if item["worst"]:
            worst = f"{item['worst']}: {display(row['Worst flag value'])} (trips below {display(row['Threshold'])})"
        expected = [str(number), name, LATEST, count, worst, runway, row["Review status"].value]
        assert cells == expected, f"{name}: slide 1 row\nExpected: {expected}\nGot:      {cells}"
        expected_fill = STATUS_COLORS[TRIP if item["tripped"] else PASS][0]
        got_fill = fill(table.cell(number, 4))
        assert got_fill == expected_fill, f"{name}: worst-flag cell {got_fill}, expected {expected_fill}"
    assert len(rows) - 1 == len(stories), f"Slide 1 shows {len(rows) - 1} companies, expected {len(stories)}"


def expected_counts(stories, reviews):
    """{(group, status): [names A to Z]}, worked out from the stories and the manifests."""
    counts = {("Flag status", label): [] for label in FLAG_STATUS_ROWS}
    counts.update({("Review status", label): [] for label in REVIEW_ROWS})
    for item in stories:
        counts[("Flag status", item["status"])].append(item["name"])
        counts[("Review status", reviews[item["name"]])].append(item["name"])
    return {key: sorted(names) for key, names in counts.items()}


def check_status(presentation, book, stories, reviews):
    """Slide 2's two tables and the By status sheet: the same counts and names, worked out here."""
    counts = expected_counts(stories, reviews)
    sheet = [(row["Group"].value, row["Status"].value, row["Companies"].value, row["Which"].value)
             for row in sheet_dicts(book["By status"])]
    expected_sheet = [(group, label, len(names), ", ".join(names) or "-") for (group, label), names in counts.items()]
    assert sheet == expected_sheet, f"By status sheet\nExpected: {expected_sheet}\nGot:      {sheet}"
    for name, group in (("Flag status table", "Flag status"), ("Review status table", "Review status")):
        rows, _ = table_rows(presentation.slides[1], name)
        expected = [[label, str(len(names)), ", ".join(names) or "-"]
                    for (kind, label), names in counts.items() if kind == group]
        assert rows[1:] == expected, f"{name}\nExpected: {expected}\nGot:      {rows[1:]}"


def check_numbers_are_in_the_workbook(presentation, book):
    """Every number on the slides (footers aside) is shown somewhere in the rollup workbook."""
    shown_in_book = set()
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            shown_in_book |= set().union(*(number_tokens(str(display(cell))) for cell in row if cell.value is not None))
    on_slides = set()
    for slide in presentation.slides:
        for item in slide.shapes:
            if item.name == "Footer":
                continue
            if item.has_table:
                on_slides |= set().union(*(number_tokens(cell.text) for row in item.table.rows for cell in row.cells))
            elif item.has_text_frame:
                on_slides |= number_tokens(item.text_frame.text)
    invented = on_slides - shown_in_book
    assert not invented, f"Numbers on the rollup deck that aren't in the rollup workbook: {sorted(invented)}"
    return len(on_slides)


def check_runway_sheet(book, stories):
    """The Runway sheet: shortest runway first, the hand-formula months, the flag result from the story."""
    ordered = sorted(stories, key=lambda item: item["runway"])
    rows = sheet_dicts(book["Runway"])
    assert [row["Company"].value for row in rows] == [item["name"] for item in ordered], "Runway sheet order"
    for row, item in zip(rows, ordered):
        assert same_number(row["Runway at current burn"].value, item["runway"]), f"{item['name']}: Runway sheet"
        expected = "Tripped" if item["runway_trips"] else "Passed"
        assert row["Runway flag"].value == expected, f"{item['name']}: runway flag {row['Runway flag'].value!r}"


def check_runway_chart(presentation, figure, stories, config):
    """The picture is on slide 3; the figure rollup.py drew has the right bars, colors and threshold line."""
    assert shape(presentation.slides[2], "Runway chart").shape_type == PICTURE, "Slide 3 has no runway picture"
    ordered = sorted(stories, key=lambda item: item["runway"])
    axis = figure.axes[0]
    names = [label.get_text() for label in axis.get_yticklabels()]
    assert names == [item["name"] for item in ordered], f"Chart order {names}"
    bars = axis.patches
    assert len(bars) == len(ordered), f"{len(bars)} bars for {len(ordered)} companies"
    red = tuple(int(RED[i:i + 2], 16) / 255 for i in (0, 2, 4))
    for bar, item in zip(bars, ordered):
        assert same_number(bar.get_width(), item["runway"]), f"{item['name']}: bar {bar.get_width()}"
        is_red = all(math.isclose(a, b) for a, b in zip(bar.get_facecolor()[:3], red))
        assert is_red == item["runway_trips"], f"{item['name']}: bar red={is_red}, runway flag trips={item['runway_trips']}"
    threshold = config["runway_min_months"]
    assert any(list(line.get_xdata()) == [threshold, threshold] for line in axis.lines), "No threshold line"


def check_footers(presentation, count):
    """Every slide: the fictional-data note, the workbook count, today's date, no AI text; one line."""
    start = f"Fictional data | Portfolio rollup of {count} workbooks | {datetime.date.today().isoformat()} | "
    for number, slide in enumerate(presentation.slides, start=1):
        box = shape(slide, "Footer")
        footer = box.text_frame.text
        assert footer.startswith(start) and footer.endswith(FOOTER_END), f"Slide {number} footer {footer!r}"
        room = Emu(box.width - box.text_frame.margin_left - box.text_frame.margin_right).pt
        assert text_width_pt(footer, MIN_FONT_PT) <= room, f"Slide {number} footer is more than one line"


def build_with_captured_chart(config, data_dir, output_dir):
    """Run rollup.save_rollup, keeping the runway figure it draws so its bars can be checked."""
    figures = []
    original = rollup.save_chart

    def keep(figure, path):
        figures.append(figure)
        return original(figure, path)
    rollup.save_chart = keep
    try:
        paths = rollup.save_rollup(config, data_dir, output_dir)
    finally:
        rollup.save_chart = original
    assert len(figures) == 1, f"Expected one chart, rollup drew {len(figures)}"
    return paths, figures[0]


def check_big_portfolio(config):
    """12 companies with 40-character names: 2 ranking slides, ranks 1 to 12, nothing overflows."""
    with tempfile.TemporaryDirectory() as folder:
        data_dir = Path(folder) / "data"
        data_dir.mkdir()
        for number in range(12):
            company = COMPANIES[number % 3]["answer_key"].OUTPUT_PATH
            shutil.copy(company, data_dir / f"Company {number:02d} {'x' * 29}.xlsx")
        presentation = Presentation(rollup.save_rollup(config, data_dir, Path(folder) / "output")["deck"])
        titles = [slide.shapes.title.text for slide in presentation.slides]
        assert titles[:2] == [f"{TITLES[0]} (1 of 2)", f"{TITLES[0]} (2 of 2)"], f"Big portfolio titles {titles}"
        ranks = [row[0] for slide in list(presentation.slides)[:2] for row in table_rows(slide, "Ranking table")[0][1:]]
        assert ranks == [str(number) for number in range(1, 13)], f"Big portfolio ranks {ranks}"
        check_no_overflow(presentation, "12-company rollup")
    print("✓ 12 companies with 40-character names: ranking over 2 slides, ranks 1 to 12, nothing overflows")


def check_unreadable_workbook(config):
    """A workbook clean.py can't read is listed last, unranked, with its reason; the others rank as before."""
    with tempfile.TemporaryDirectory() as folder:
        data_dir = Path(folder) / "data"
        data_dir.mkdir()
        for company in COMPANIES:
            shutil.copy(company["answer_key"].OUTPUT_PATH, data_dir)
        book = Workbook()
        book.active.append(["Quarter", "Starting ARR"])
        book.active.append(["Q1 2025", 100])
        book.save(data_dir / "broken.xlsx")
        paths = rollup.save_rollup(config, data_dir, Path(folder) / "output")
        rows = sheet_dicts(load_workbook(paths["excel"])["Ranking"])
        names = [row["Company"].value for row in rows]
        assert names == [item["name"] for item in expected_ranking()] + ["Broken"], f"With a broken workbook: {names}"
        broken = rows[-1]
        assert broken["Rank"].value is None and broken["Flag status"].value == "Workbook can't be read"
        assert broken["Note"].value and "Traceback" not in broken["Note"].value, "No plain reason for the broken workbook"
        check_no_overflow(Presentation(paths["deck"]), "rollup with a broken workbook")
    print("✓ An unreadable workbook is listed last, unranked, with clean.py's reason; the rest rank as before")


def main():
    analyze.anthropic.Anthropic = refuse   # the rollup has no AI step: creating a client is a bug
    config = load_config()
    stories = expected_ranking()
    workbooks = {company["name"]: Path(company["answer_key"].OUTPUT_PATH) for company in COMPANIES}
    reviews = {name: expected_review(path, OUTPUT_DIR) for name, path in workbooks.items()}

    paths, figure = build_with_captured_chart(config, PROJECT_DIR / "data", OUTPUT_DIR)
    presentation, book = check_files(paths)
    print(f"✓ {paths['deck'].name}: 3 slides with the expected titles; {paths['excel'].name}: 3 sheets")
    sheet_rows = check_ranking_sheet(book, stories, reviews)
    check_ranking_slide(presentation, sheet_rows, stories)
    order = ", ".join(f"{item['name']} {item['tripped']}" for item in stories)
    print(f"✓ Ranked by flags tripped ({order}); worst flag, runway and review match the answer keys on both files")
    numbers = check_numbers_are_in_the_workbook(presentation, book)
    print(f"✓ {numbers} distinct numbers on the deck, every one in the rollup workbook")
    check_status(presentation, book, stories, reviews)
    print(f"✓ Companies by status match the stories and the manifests ({', '.join(sorted(set(reviews.values())))})")
    check_runway_sheet(book, stories)
    check_runway_chart(presentation, figure, stories, config)
    print("✓ Runway chart and sheet: shortest first, hand-formula months, red exactly where the flag trips, threshold line")
    check_no_overflow(presentation, "rollup")
    check_footers(presentation, len(COMPANIES))
    print("✓ Nothing overflows, no font below 12 pt, every footer on one line and saying no AI text")
    check_big_portfolio(config)
    check_unreadable_workbook(config)
    print("All checks passed")


if __name__ == "__main__":
    main()
