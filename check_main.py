"""Automated proof for build step 5: main.py runs the batch and survives a bad company.

Checks:
1. `python main.py --all --skip-ai` (a real command, run as a separate process) exits 0,
   finds exactly the three company workbooks, saves a fresh Excel file and a fresh deck for
   each, skips AI with a clear message, and its summary table shows all three "OK (AI skipped)"
   with the flag and gap counts from each company's story (check_companies.py), not from main.py.
2. One company failing doesn't stop the batch: a broken workbook, a missing file and a
   simulated code bug sit between good companies, and every good company still succeeds.
3. A failed run exits with code 1; wrong arguments exit with code 2 (argparse's usage error).
4. --skip-ai never calls Claude and never needs an API key: with analyze(), the key check and
   the Anthropic client all replaced by functions that stop the check if called, every deck
   still gets the "AI summary unavailable" placeholder, and no analysis JSON is written
   (an earlier one is left untouched).
5. --all ignores Excel's "~$" lock files and anything that isn't .xlsx.
6. The batch also writes output/batch_summary.csv with the same counts, and prints no
   quarter warning when every company ends on the same quarter.

No API calls: every run uses --skip-ai. The main.py process also runs with no API key and the
API address pointed at a dead local port, so even a bug that called Claude couldn't reach it.
The AI-on paths ("OK", "OK (AI failed)" after the retry, an API error, a missing key) are proven
in tests/test_main.py with a fake Claude client. (check_deck.py checks what is inside the decks.)

Run: python check_main.py  -> prints "All checks passed" or stops at the first failure.
"""

import csv
import io
import os
import subprocess
import sys
import tempfile
import time
from contextlib import redirect_stderr, redirect_stdout
from pathlib import Path

from openpyxl import Workbook
from pptx import Presentation

import analyze
import main
from build_deck import PLACEHOLDER_TEXT, analysis_path, deck_path
from check_companies import COMPANIES, expected_gaps
from clean import clean_workbook
from excel_output import output_path
from metrics import CANNOT_EVALUATE, TRIP, compute_metrics, load_config

PROJECT_DIR = Path(__file__).parent
EXPECTED_OK = "OK (AI skipped)"
DEAD_API_URL = "http://127.0.0.1:9"  # nothing listens on port 9, so a call can't reach any API


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def run_main(*args):
    """Run main.py as its own process, like typing it in the terminal. Returns the finished process.

    The process gets no API key and a dead API address, so it can never reach Claude.
    """
    env = {**os.environ, "ANTHROPIC_BASE_URL": DEAD_API_URL}
    env.pop("ANTHROPIC_API_KEY", None)
    return subprocess.run([sys.executable, "main.py", *args], cwd=PROJECT_DIR,
                          capture_output=True, text=True, env=env)


def summary_table(stdout):
    """Rows of the summary table as {company: [company, flags, gaps, result]}.

    Columns are separated by at least two spaces; the header and dashes lines are skipped.
    """
    lines = stdout.split("=== Summary ===")[1].strip().splitlines()
    rows = {}
    for line in lines[2:]:           # skip the header and the dashes
        if not line.strip():
            break                    # a blank line ends the table
        cells = [cell for cell in line.split("  ") if cell.strip()]
        rows[cells[0].strip()] = [cell.strip() for cell in cells]
    return rows


def story_flags_text(company):
    """Expected 'Flags tripped' cell, counted from the company's story in check_companies.py."""
    statuses = list(company["expected_flags"].values())
    text = f"{statuses.count(TRIP)} of {len(statuses)}"
    if statuses.count(CANNOT_EVALUATE):
        text += f", {statuses.count(CANNOT_EVALUATE)} cannot evaluate"
    return text


def story_gaps_text(company):
    """Expected 'Data gaps' cell, from the gaps the CLAUDE.md rules predict for this company."""
    actuals, _ = clean_workbook(company["answer_key"].OUTPUT_PATH)
    count = len(expected_gaps(company, compute_metrics(actuals).columns))
    if count == 0:
        return "none"
    return f"{count} metrics/flags (blank: {company['answer_key'].BLANK_QUARTER})"


def write_broken_workbook(folder):
    """A workbook with a Quarter header but most required columns missing."""
    path = Path(folder) / "broken.xlsx"
    book = Workbook()
    book.active.append(["Quarter", "Starting ARR"])
    book.active.append(["Q2 2026", 1000])
    book.save(path)
    return path


def run_quietly(function, *args, **kwargs):
    """Call a main.py function without printing; returns (result, stdout text, stderr text)."""
    stdout, stderr = io.StringIO(), io.StringIO()
    with redirect_stdout(stdout), redirect_stderr(stderr):
        result = function(*args, **kwargs)
    return result, stdout.getvalue(), stderr.getvalue()


def analysis_file_state(workbook):
    """(size, modified time) of a company's analysis JSON, or None if there isn't one."""
    path = analysis_path(workbook)
    return (path.stat().st_size, path.stat().st_mtime) if path.exists() else None


def headline_on_deck(path):
    """The text in slide 1's Headline box."""
    slide = Presentation(path).slides[0]
    return next(shape.text_frame.text for shape in slide.shapes if shape.name == "Headline")


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def check_finds_the_three_companies():
    """--all picks up exactly the three company workbooks in data/."""
    found = {path.name for path in main.find_workbooks()}
    expected = {Path(company["answer_key"].OUTPUT_PATH).name for company in COMPANIES}
    assert found == expected, f"--all finds {sorted(found)}, expected {sorted(expected)}"


def check_batch_run():
    """main.py --all --skip-ai: exit 0, fresh Excel files and decks, a clear AI skip message, a correct summary."""
    started = time.time()
    workbooks = [company["answer_key"].OUTPUT_PATH for company in COMPANIES]
    analyses_before = [analysis_file_state(workbook) for workbook in workbooks]
    process = run_main("--all", "--skip-ai")
    assert process.returncode == 0, f"Exit code {process.returncode}\n{process.stdout}\n{process.stderr}"
    assert process.stderr == "", f"Unexpected error output:\n{process.stderr}"
    assert process.stdout.count("AI commentary: skipped (--skip-ai)") == len(COMPANIES), "AI skip message missing"
    assert process.stdout.count("✓ Deck: output/") == len(COMPANIES), "Deck line missing"
    assert process.stdout.count("AI summary unavailable on slides 1 and 5") == len(COMPANIES), \
        "--skip-ai decks should say they carry the AI placeholder"
    assert f"{len(COMPANIES)} of {len(COMPANIES)} companies succeeded" in process.stdout, "Success count wrong"
    assert "AI summary unavailable for" not in process.stdout, "--skip-ai isn't an AI failure"
    assert [analysis_file_state(workbook) for workbook in workbooks] == analyses_before, \
        "--skip-ai must not write, replace or delete any analysis JSON"

    rows = summary_table(process.stdout)
    assert set(rows) == {company["name"] for company in COMPANIES}, f"Summary rows: {sorted(rows)}"
    for company in COMPANIES:
        expected = [company["name"], story_flags_text(company), story_gaps_text(company), EXPECTED_OK]
        assert rows[company["name"]] == expected, \
            f"Summary row wrong.\nExpected: {expected}\nGot:      {rows[company['name']]}"
        for saved in (output_path(company["answer_key"].OUTPUT_PATH), deck_path(company["answer_key"].OUTPUT_PATH)):
            assert saved.exists() and saved.stat().st_mtime >= started - 1, f"{saved} wasn't saved by this run"
    check_summary_csv(started)
    assert "⚠" not in process.stdout, "All three companies end on Q2 2026, so there should be no quarter warning"
    return rows


def check_summary_csv(started):
    """output/batch_summary.csv was written by this run: one row per company, counts from each story."""
    path = main.SUMMARY_CSV_PATH
    assert path.exists() and path.stat().st_mtime >= started - 1, f"{path} wasn't saved by this run"
    with open(path, newline="") as file:
        rows = {row["Company"]: row for row in csv.DictReader(file)}
    assert set(rows) == {company["name"] for company in COMPANIES}, f"CSV companies: {sorted(rows)}"
    for company in COMPANIES:
        statuses = list(company["expected_flags"].values())
        row = rows[company["name"]]
        expected = {"Latest quarter": "Q2 2026", "Flags tripped": str(statuses.count(TRIP)),
                    "Flags total": str(len(statuses)), "Cannot evaluate": str(statuses.count(CANNOT_EVALUATE)),
                    "Result": EXPECTED_OK}
        got = {key: row[key] for key in expected}
        assert got == expected, f"CSV row for {company['name']}.\nExpected: {expected}\nGot:      {got}"


def check_failure_does_not_stop_batch(folder, config):
    """Bad companies in the middle of the batch are recorded; the good ones after them still run."""
    northwind, alderpeak = (COMPANIES[0]["answer_key"].OUTPUT_PATH, COMPANIES[1]["answer_key"].OUTPUT_PATH)
    paths = [northwind, write_broken_workbook(folder), Path(folder) / "missing.xlsx", alderpeak]
    results, stdout, stderr = run_quietly(main.run_batch, paths, config, True)

    assert [r["company"] for r in results] == ["Northwind", "Broken", "Missing", "Alderpeak"], \
        f"Batch order wrong: {[r['company'] for r in results]}"
    assert results[0]["error"] is None and results[3]["error"] is None, "A good company failed"
    assert "missing columns" in results[1]["error"], f"Broken workbook error unclear: {results[1]['error']}"
    assert results[2]["error"].startswith("FileNotFoundError"), f"Missing file error: {results[2]['error']}"
    assert stdout.count("✗ FAILED") == 2, "Each failure should print one FAILED line"
    assert stderr == "", f"Input errors shouldn't print a traceback:\n{stderr}"
    return results


def check_code_bug_does_not_stop_batch(config):
    """An unexpected error (a bug, not bad input) is recorded with a traceback; the batch continues."""
    real_compute = main.compute_metrics
    calls = []

    def buggy_first_call(actuals):  # fails for the first company only
        calls.append(1)
        if len(calls) == 1:
            raise KeyError("simulated bug")
        return real_compute(actuals)

    main.compute_metrics = buggy_first_call
    try:
        paths = [COMPANIES[0]["answer_key"].OUTPUT_PATH, COMPANIES[2]["answer_key"].OUTPUT_PATH]
        results, _, stderr = run_quietly(main.run_batch, paths, config, True)
    finally:
        main.compute_metrics = real_compute  # always put the real function back
    assert results[0]["error"] == "KeyError: 'simulated bug'", f"Bug not recorded: {results[0]['error']}"
    assert results[1]["error"] is None, "The company after the bug should still succeed"
    assert "Traceback" in stderr, "An unexpected error should print its traceback"


def check_exit_codes(broken_path):
    """1 when a company fails; 2 for no arguments or for both a file and --all."""
    process = run_main(str(broken_path), "--skip-ai")
    assert process.returncode == 1, f"Failed company: exit code {process.returncode}, expected 1"
    assert "0 of 1 companies succeeded" in process.stdout, "Single-file failure summary missing"
    assert summary_table(process.stdout)["Broken"][3].startswith("FAILED: ValueError"), "Result should say FAILED"
    for args in [(), ("data/northwind.xlsx", "--all")]:
        process = run_main(*args, "--skip-ai")
        assert process.returncode == 2, f"main.py {args}: exit code {process.returncode}, expected 2"
        assert "give one workbook path, or --all" in process.stderr, f"main.py {args}: unclear usage error"


def refuse(what):
    """A stand-in function that stops the check if it is ever called."""
    def called(*args, **kwargs):
        raise AssertionError(f"--skip-ai {what}")
    return called


def check_skip_ai_never_calls_claude(folder, config):
    """--skip-ai with every way to Claude blocked: each company still gets a placeholder deck, and no analysis JSON."""
    blocked = [(main, "analyze", refuse("called analyze()")),
               (main, "api_key_problem", refuse("looked for an API key")),
               (analyze.anthropic, "Anthropic", refuse("created an Anthropic client"))]
    originals = [(module, name, getattr(module, name)) for module, name, _ in blocked]
    for module, name, stand_in in blocked:
        setattr(module, name, stand_in)
    try:
        workbooks = [company["answer_key"].OUTPUT_PATH for company in COMPANIES]
        results, stdout, _ = run_quietly(main.run_batch, workbooks, config, True, output_dir=Path(folder))
    finally:
        for module, name, original in originals:  # always put the real functions back
            setattr(module, name, original)
    for company, result, workbook in zip(COMPANIES, results, workbooks):
        assert main.result_text(result) == EXPECTED_OK, f"{company['name']}: {main.result_text(result)}"
        assert headline_on_deck(deck_path(workbook, folder)) == PLACEHOLDER_TEXT, f"{company['name']}: no placeholder"
        assert not analysis_path(workbook, folder).exists(), f"{company['name']}: --skip-ai saved an analysis JSON"
    placeholder_lines = stdout.count(f"({PLACEHOLDER_TEXT} on slides 1 and 5)")
    assert placeholder_lines == len(COMPANIES), "Each deck line should say it carries the placeholder"


def check_ignores_lock_and_other_files(folder):
    """--all skips '~$' lock files and non-.xlsx files, and sorts by name."""
    for name in ["b.xlsx", "a.xlsx", "~$a.xlsx", "notes.txt", "old.xls"]:
        (Path(folder) / name).write_text("")
    found = [path.name for path in main.find_workbooks(folder)]
    assert found == ["a.xlsx", "b.xlsx"], f"find_workbooks found {found}"


def main_check():
    config = load_config()
    check_finds_the_three_companies()
    rows = check_batch_run()
    for company in COMPANIES:
        print(f"✓ main.py --all --skip-ai: {' | '.join(rows[company['name']])}")

    with tempfile.TemporaryDirectory() as folder:
        results = check_failure_does_not_stop_batch(folder, config)
        print(f"✓ Batch continues past a broken workbook and a missing file ({results[1]['error'][:60]}...)")
        check_code_bug_does_not_stop_batch(config)
        print("✓ Batch continues past an unexpected code error, and prints its traceback")
        check_exit_codes(Path(folder) / "broken.xlsx")
        print("✓ Exit codes: 0 all OK, 1 any company failed, 2 wrong arguments")

    with tempfile.TemporaryDirectory() as folder:
        check_skip_ai_never_calls_claude(folder, config)
        print("✓ --skip-ai never calls Claude or needs a key: placeholder on every deck, no analysis JSON saved")

    with tempfile.TemporaryDirectory() as folder:
        check_ignores_lock_and_other_files(folder)
        print("✓ --all ignores Excel lock files (~$) and non-.xlsx files")
    print("All checks passed")


if __name__ == "__main__":
    main_check()
