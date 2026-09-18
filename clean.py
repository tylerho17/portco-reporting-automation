"""Turn a messy KPI workbook into a clean table of numbers.

Main entry point: clean_workbook(path) -> (actuals, next_budget)
- actuals: one row per actual quarter, the 16 standard columns, all in $K.
  Missing values are NaN ("not a number", pandas' marker for "no value").
- next_budget: the 3 budget numbers from the budget-only forecast row, or None.

Nothing here guesses or fills in numbers. If something can't be read with
certainty, the script stops with an error that says what and where.

A header that is neither a standard column nor in HEADER_ALIASES is looked up in the company's
confirmed mapping (mappings/<company>.yaml, written by mapping.py once a person confirmed it). If it
isn't there either, the script stops with mapping.py's proposal for it (UnconfirmedMappingError).

The result is cached (cache.py, Task 17) by the hash of the workbook and of its mapping file, so
the deck, memo and Excel steps of one run don't read the same file again.

Run directly to print the clean table:  python clean.py data/northwind.xlsx
"""

import numbers
import re
import sys
import zipfile
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from cache import ResultCache
from provenance import file_sha256

# ---------------------------------------------------------------------------
# Standard column names (from CLAUDE.md)
# ---------------------------------------------------------------------------

ACTUAL_COLUMNS = [
    "starting_arr", "new_arr", "expansion_arr", "contraction_arr", "churned_arr",
    "revenue", "gross_profit", "net_burn", "ending_cash", "sm_spend",
    "new_customers", "headcount", "pipeline",
]
BUDGET_COLUMNS = ["budget_new_arr", "budget_arr", "budget_net_burn"]
STANDARD_COLUMNS = ACTUAL_COLUMNS + BUDGET_COLUMNS

# Normalized messy header -> standard name. Only needed when they differ:
# a header that already normalizes to a standard name (e.g. "expansion arr") needs no entry.
HEADER_ALIASES = {
    "beginning_arr": "starting_arr",
    "new_arr_k": "new_arr",
    "cash_end_of_qtr": "ending_cash",
    "s_m_spend": "sm_spend",
    "new_logos": "new_customers",
    "headcount_fte": "headcount",
    "qualified_pipeline": "pipeline",
    "new_arr_budget": "budget_new_arr",
    "net_burn_bud": "budget_net_burn",
}

# A row label containing any of these words is the budget-only forecast row.
BUDGET_LABEL_WORDS = ("budget", "bud", "plan")

# Quarter labels look like "Q2 2026": Q, a digit 1-4, a space, a 4-digit year.
QUARTER_PATTERN = re.compile(r"^Q([1-4])\s+(\d{4})$")

# The same quarter anywhere inside a longer label, e.g. "Q3 2026 (Budget)" or "Q3 2026 - Bud".
QUARTER_IN_LABEL = re.compile(r"\bQ([1-4])\s+(\d{4})\b")

# Number text, once "$" and spaces are removed and letters are uppercased: an optional minus,
# digits (commas only between groups of 3, like "1,250,000"), optional decimals, optional K or M.
# Examples that match: "12.5M", "-1.2M", "1,250K", "5,090", "850".
NUMBER_TEXT = re.compile(r"^-?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?[KM]?$")

# Added to "Can't read ... as a number" errors so the fix is obvious.
NUMBER_HINT = " (expected e.g. 1250, '$1.2M' or '850K'; leave the cell empty if there's no data)"

# The header that marks the quarter-label column.
LABEL_HEADER = "quarter"

EXCEL_WORKBOOK_PART = "xl/workbook.xml"   # inside every .xlsx (which is a zip file)


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

class UnknownHeadersError(ValueError):
    """Headers that are neither a standard column, an alias nor a confirmed mapping.

    Carries what mapping.py needs to propose a column for each: `unknown` (their positions) and
    `known` ({position: name} for the headers that were recognised).
    """
    def __init__(self, message, unknown, known):
        super().__init__(message)
        self.unknown, self.known = unknown, known


class UnconfirmedMappingError(ValueError):
    """The stop for an unknown header: names each one and mapping.py's proposal. `proposals` is for the web page."""
    def __init__(self, message, proposals):
        super().__init__(message)
        self.proposals = proposals


def is_blank(value):
    """True for an empty cell: missing (None/NaN) or text that is only spaces."""
    if isinstance(value, str):
        return value.strip() == ""
    return pd.isna(value)


def normalize_header(text):
    """'Net Burn (Bud.)' -> 'net_burn_bud'. Lowercase; each run of symbols/spaces becomes one '_'."""
    return re.sub(r"[^a-z0-9]+", "_", str(text).lower()).strip("_")


def standard_column(header, confirmed=None):
    """Map a messy header to its standard column name, or stop if it's unknown.

    confirmed = {normalized header: column} from the company's mapping file (mapping.confirmed_aliases).
    A standard name or HEADER_ALIASES entry always wins, so a saved line can't redefine a known header.
    """
    name = normalize_header(header)
    name = HEADER_ALIASES.get(name, name)  # use the alias if there is one, else keep as is
    if name not in STANDARD_COLUMNS:
        name = (confirmed or {}).get(name, name)
    if name not in STANDARD_COLUMNS:
        raise ValueError(f"Unknown column header {header!r}: if it means one of the {len(STANDARD_COLUMNS)} "
                         f"input columns, confirm a mapping with mapping.py; otherwise delete the column")
    return name


def parse_number(value):
    """Turn one cell into a number in $K.

    '$12.5M' -> 12500.0   '120K' -> 120.0   '5,090' -> 5090.0   blank -> NaN
    """
    if is_blank(value):
        return float("nan")
    if isinstance(value, bool):  # a TRUE/FALSE cell; Python would otherwise count True as 1
        raise ValueError(f"Can't read {value!r} as a number - the cell holds TRUE/FALSE{NUMBER_HINT}")
    if isinstance(value, numbers.Real):  # already a number (int, float, numpy number)
        return float(value)

    # Strip the decoration, then check the text has a shape we know before reading it.
    text = str(value).replace("$", "").replace(" ", "").upper()
    if not NUMBER_TEXT.match(text):  # stops "n/a", "(120)", "12.5%", "1.250,5", "inf", "1e3"
        raise ValueError(f"Can't read {value!r} as a number{NUMBER_HINT}")

    # Read the unit letter at the end.
    multiplier = 1
    if text.endswith("M"):
        multiplier = 1000  # millions -> thousands
        text = text[:-1]
    elif text.endswith("K"):
        text = text[:-1]   # already thousands

    # Decimal does exact decimal math, so '21.85' * 1000 is exactly 21850 (no binary rounding).
    return float(Decimal(text.replace(",", "")) * multiplier)


def parse_quarter(label):
    """'Q2 2026' -> (2026, 2). Stops if the label isn't in that format."""
    match = QUARTER_PATTERN.match(label)
    if not match:
        raise ValueError(f"Can't read quarter label {label!r} (expected e.g. 'Q2 2026') - write the label like "
                         f"that, or if the row is a note, move it to another tab")
    return int(match.group(2)), int(match.group(1))


def next_quarter(year, q):
    """(2025, 4) -> (2026, 1): the quarter after this one. Q4 rolls into next year's Q1."""
    return (year, q + 1) if q < 4 else (year + 1, 1)


def quarter_name(year, q):
    """(2026, 1) -> 'Q1 2026'."""
    return f"Q{q} {year}"


def at_row(rows, i):
    """'row 7: ' for the i-th label when Excel row numbers are known, else '' (for error messages)."""
    return f"row {rows[i]}: " if rows else ""


def check_quarters_in_order(labels, rows=None):
    """Stop unless quarters run oldest -> newest with none skipped or repeated.

    QoQ looks back 1 row and YoY looks back 4 rows, so a missing row would
    silently compare the wrong quarters. (A blank row is fine; a missing row is not.)
    `rows` (optional) holds each label's Excel row number, so errors can say where to look.
    """
    quarters = []
    for i, label in enumerate(labels):
        try:
            quarters.append(parse_quarter(label))
        except ValueError as error:
            raise ValueError(f"{at_row(rows, i)}{error}") from None

    for i in range(1, len(quarters)):
        expected = next_quarter(*quarters[i - 1])
        if quarters[i] != expected:
            raise ValueError(f"{at_row(rows, i)}After {labels[i - 1]} expected {quarter_name(*expected)}, "
                             f"found {labels[i]} - quarters must run oldest to newest with none skipped "
                             f"or repeated (a quarter with no data still needs its own row)")


# ---------------------------------------------------------------------------
# Reading the workbook
# ---------------------------------------------------------------------------

# Error messages use Excel's own addresses (row 7, column F, cell F7) so a person can go
# straight to the problem. pandas counts rows and columns from 0; Excel counts rows from 1
# and names columns A, B, C... (checked: pandas keeps empty rows and columns at the top and left).

def excel_row(row_index):
    """pandas row index -> Excel row number (index 6 is row 7)."""
    return row_index + 1


def excel_column(position):
    """pandas column position -> Excel column letter (position 5 is column F)."""
    return get_column_letter(position + 1)


def header_row_of(sheet):
    """The row index (within the first 10 rows) holding a 'Quarter' header, or None."""
    for row_index in range(min(10, len(sheet))):
        cells = [normalize_header(v) for v in sheet.iloc[row_index] if not is_blank(v)]
        if LABEL_HEADER in cells:
            return row_index
    return None


def is_excel_workbook(path):
    """True if the file is a zip holding an Excel workbook part (every .xlsx has xl/workbook.xml).

    Being a zip isn't enough: a .pptx or .docx renamed .xlsx is a zip too, and pandas' error on
    one is cryptic.
    """
    if not zipfile.is_zipfile(path):
        return False
    with zipfile.ZipFile(path) as archive:
        return EXCEL_WORKBOOK_PART in archive.namelist()


def check_is_workbook(path):
    """Stop, in plain words, on a file that isn't there or isn't an Excel workbook, before pandas reads it.

    pandas' own errors ("Excel file format cannot be determined, you must specify an engine
    manually", "[Errno 2] No such file or directory") don't say what to do.
    """
    if not Path(path).exists():
        raise FileNotFoundError(f"Can't find {path}: check the file name and folder, then run again")
    if not is_excel_workbook(path):
        raise ValueError(f"{Path(path).name} isn't a readable Excel workbook: open it in Excel, save it as an "
                         f"Excel Workbook (.xlsx), then run again")


def find_kpi_sheet(path):
    """Return (sheet_name, sheet, header_row) for the one tab whose header row contains 'Quarter'.

    Checks the first 10 rows of each tab, so title rows above the table are fine.
    Tabs without a 'Quarter' header (like Notes) are ignored. Two or more KPI-looking tabs stop:
    a stale copy could be picked by mistake, so we don't guess.
    """
    # Every tab, raw, no header guessing. Only a truly empty cell counts as missing: by default pandas
    # would also quietly blank out text like "n/a", "NA" or "NULL", and those must stop in parse_number.
    sheets = pd.read_excel(path, sheet_name=None, header=None, keep_default_na=False, na_values=[""])
    found = [(name, sheet, header_row_of(sheet)) for name, sheet in sheets.items() if header_row_of(sheet) is not None]
    if not found:
        raise ValueError(f"No tab in {Path(path).name} has a 'Quarter' header in its first 10 rows (tabs checked: "
                         f"{', '.join(repr(name) for name in sheets)}): add a 'Quarter' header above the column of "
                         f"quarter labels, like 'Q2 2026'")
    if len(found) > 1:
        names = [f"'{name}'" for name, _, _ in found]
        each = "both" if len(found) == 2 else "all"
        raise ValueError(f"Tabs {', '.join(names[:-1])} and {names[-1]} {each} have a 'Quarter' header - keep one "
                         f"KPI tab and delete or rename the others")
    return found[0]


def check_no_error_cells(path, sheet_name):
    """Stop on Excel error values like #DIV/0! or #REF! in the KPI tab.

    pandas reads them as empty cells, which would turn a broken formula into "data missing"
    instead of a problem to fix. openpyxl still sees them (cell type "e" = error).
    """
    workbook = load_workbook(path, read_only=True, data_only=True)  # data_only: saved results, not formulas
    try:
        for row in workbook[sheet_name].iter_rows():
            for cell in row:
                if cell.data_type == "e":
                    raise ValueError(f"cell {cell.coordinate}: Excel error value {cell.value!r} - fix the formula "
                                     f"in the workbook, or clear the cell if there's no data")
    finally:
        workbook.close()  # read-only mode keeps the file open until closed


def header_name(header, confirmed=None):
    """'Quarter' -> 'quarter' (the label column); any other header -> its standard column name."""
    if normalize_header(header) == LABEL_HEADER:
        return LABEL_HEADER
    return standard_column(header, confirmed)


def name_headers(headers, header_row, confirmed=None):
    """Return {position: name} for every filled header. Stops on a repeated meaning, then on unknown headers.

    Every unknown header is gathered before stopping (UnknownHeadersError), so one message can list
    them all, each with mapping.py's proposal.
    """
    names, unknown, first_error = {}, [], None
    for position, header in headers.items():
        if is_blank(header):
            continue
        try:
            name = header_name(header, confirmed)
        except ValueError as error:
            unknown.append(position)
            first_error = first_error or f"cell {excel_column(position)}{excel_row(header_row)} (header): {error}"
            continue
        if name in names.values():  # e.g. "Beginning ARR" and "Starting ARR": which one is right?
            first = next(p for p, n in names.items() if n == name)
            raise ValueError(f"row {excel_row(header_row)} (header): columns {excel_column(first)} "
                             f"({headers[first]!r}) and {excel_column(position)} ({header!r}) both mean "
                             f"{name!r} - keep one and delete the other")
        names[position] = name
    if unknown:
        raise UnknownHeadersError(first_error, unknown, names)
    return names


def map_columns(headers, header_row, confirmed=None):
    """Return (label_position, {position: standard name}) from the header row. Stops if a column is missing."""
    names = name_headers(headers, header_row, confirmed)
    missing = [c for c in STANDARD_COLUMNS if c not in names.values()]
    if missing:
        raise ValueError(f"row {excel_row(header_row)} (header) is missing columns: {', '.join(missing)} - add a "
                         f"column headed with each name (its cells can stay empty if there's no data)")
    label_position = next(p for p, name in names.items() if name == LABEL_HEADER)
    column_map = {p: name for p, name in names.items() if name != LABEL_HEADER}
    return label_position, column_map


def check_no_headerless_values(sheet, header_row, known_positions):
    """Stop if a column without a header has values under the header row. They would be silently ignored."""
    below_header = sheet.iloc[header_row + 1:]
    for position in range(sheet.shape[1]):
        if position in known_positions:
            continue
        filled = [i for i, value in below_header.iloc[:, position].items() if not is_blank(value)]
        if filled:
            column = excel_column(position)
            raise ValueError(f"column {column} has values (first in cell {column}{excel_row(filled[0])}) but no "
                             f"header in row {excel_row(header_row)} - add a header or delete the values")


def check_unlabelled_row_is_empty(row_index, raw_row, column_map, label_position):
    """Stop if a row has numbers but no quarter label. Skipping it would silently lose them."""
    if any(not is_blank(raw_row[position]) for position in column_map):
        raise ValueError(f"row {excel_row(row_index)} has values but no quarter label in column "
                         f"{excel_column(label_position)} - add the label or delete the row")


def parse_row(label, row_index, raw_row, column_map):
    """Parse every cell in one row. Errors name the cell, quarter and column, e.g. 'cell F7 (Q2 2025, revenue)'."""
    row = {}
    for position, name in column_map.items():
        try:
            row[name] = parse_number(raw_row[position])
        except ValueError as error:
            cell = f"{excel_column(position)}{excel_row(row_index)}"
            raise ValueError(f"cell {cell} ({label}, {name}): {error}") from None
    return row


def is_budget_label(label):
    """True for a label naming the budget-only forecast row, e.g. 'Q3 2026 (Budget)' or 'Q3 2026 Plan'."""
    return any(word in label.lower() for word in BUDGET_LABEL_WORDS)


def is_budget_only_row(label, row_index, row, column_map):
    """True for the forecast row, e.g. 'Q3 2026 (Budget)'. Stops if that row also has actuals."""
    if not is_budget_label(label):
        return False
    filled_actuals = [f"{excel_column(position)}{excel_row(row_index)} ({name})"
                      for position, name in column_map.items()
                      if name in ACTUAL_COLUMNS and not pd.isna(row[name])]
    if filled_actuals:
        raise ValueError(f"row {excel_row(row_index)} ({label!r}) is labelled as a budget-only row but also has "
                         f"actual values in {', '.join(filled_actuals)} - a budget-only row may fill only "
                         f"{', '.join(BUDGET_COLUMNS)}: move the actuals to their own quarter row, or if this is "
                         f"an actual quarter, take {', '.join(repr(w) for w in BUDGET_LABEL_WORDS[:-1])} or "
                         f"{BUDGET_LABEL_WORDS[-1]!r} out of its label")
    return True


def check_budget_row_is_next_quarter(label, budget_row_number, last_actual_label):
    """Stop unless the budget-only row is for the quarter right after the last actual quarter.

    'Runway at next quarter's budgeted burn' would otherwise quietly use some other quarter's budget.
    """
    expected = quarter_name(*next_quarter(*parse_quarter(last_actual_label)))
    match = QUARTER_IN_LABEL.search(label)
    if not match:
        raise ValueError(f"row {budget_row_number} ({label!r}): a budget-only row label must name its quarter, "
                         f"e.g. '{expected} (Budget)'")
    found = quarter_name(int(match.group(2)), int(match.group(1)))
    if found != expected:
        raise ValueError(f"row {budget_row_number} ({label!r}): this budget-only row is for {found}, but the "
                         f"quarter after the last actual quarter ({last_actual_label}) is {expected} - fix the "
                         f"label, or keep only next quarter's budget")


def clean_sheet(sheet, header_row, confirmed=None):
    """Turn the raw KPI tab into (actuals, next_budget). Errors say which row, column or cell to fix.

    confirmed = the company's confirmed header mappings (see standard_column).
    """
    label_position, column_map = map_columns(sheet.iloc[header_row], header_row, confirmed)
    check_no_headerless_values(sheet, header_row, [label_position, *column_map])

    actual_rows, row_numbers = {}, {}  # quarter label -> parsed row, quarter label -> Excel row number
    next_budget, budget_row_number = None, None
    for row_index, raw_row in sheet.iloc[header_row + 1:].iterrows():  # every row below the header
        if is_blank(raw_row[label_position]):
            check_unlabelled_row_is_empty(row_index, raw_row, column_map, label_position)
            continue  # skip empty rows under the table
        label = str(raw_row[label_position]).strip()
        row = parse_row(label, row_index, raw_row, column_map)

        if is_budget_only_row(label, row_index, row, column_map):
            if next_budget is not None:
                raise ValueError(f"row {excel_row(row_index)} ({label!r}) is a second budget-only row (the first "
                                 f"is {next_budget.name!r}) - keep only next quarter's budget")
            next_budget = pd.Series({c: row[c] for c in BUDGET_COLUMNS}, name=label)
            budget_row_number = excel_row(row_index)
        elif label in actual_rows:  # a repeated label would overwrite the earlier row without a word
            raise ValueError(f"row {excel_row(row_index)}: quarter {label!r} appears twice (also in row "
                             f"{row_numbers[label]}) - delete one of the rows")
        else:
            actual_rows[label] = row  # blank quarters land here too, as all-NaN rows
            row_numbers[label] = excel_row(row_index)

    if not actual_rows:
        raise ValueError(f"row {excel_row(header_row)} (header) has no quarter rows under it: add one row per "
                         f"quarter under the header, labelled like 'Q2 2026'")
    actuals = pd.DataFrame.from_dict(actual_rows, orient="index", columns=STANDARD_COLUMNS)
    actuals.index.name = "quarter"
    check_quarters_in_order(list(actual_rows), list(row_numbers.values()))
    if next_budget is not None:
        check_budget_row_is_next_quarter(next_budget.name, budget_row_number, actuals.index[-1])
    return actuals, next_budget


_CLEANED = ResultCache()   # {(workbook hash, mapping file hash): (actuals, next_budget)}


def clean_workbook(path, mappings_dir=None):
    """Read a messy KPI workbook and return (actuals, next_budget).

    Every problem in the KPI tab stops with a message that starts with the sheet name, e.g.
    "Sheet 'KPI Tracker', cell F7 (Q2 2025, revenue): Can't read 'n/a' as a number ..."
    Headers are also matched against the company's confirmed mapping (mappings/<company>.yaml; a
    test passes its own mappings_dir). An unknown header stops with mapping.py's proposal for it.
    The same workbook with the same mapping is read once: later calls get a copy of the first answer.
    """
    import mapping   # here, not at the top of the file: mapping.py imports this file

    check_is_workbook(path)   # first, so a missing file stops in plain words, not in the hashing
    key = (file_sha256(path), mapping.mapping_sha256(path, mappings_dir))
    return _CLEANED.get(key, lambda: read_workbook(path, mappings_dir))


def clear_cache():
    """Forget every cleaned workbook (benchmark.py and the tests start from nothing)."""
    _CLEANED.clear()


def read_workbook(path, mappings_dir=None):
    """clean_workbook's work, without the cache: open the workbook, find the KPI tab and clean it."""
    import mapping

    sheet_name, sheet, header_row = find_kpi_sheet(path)
    try:
        confirmed = mapping.confirmed_aliases(path, mappings_dir)
        check_no_error_cells(path, sheet_name)
        return clean_sheet(sheet, header_row, confirmed)
    except UnknownHeadersError as error:
        proposals = mapping.propose(sheet, header_row, error.unknown, error.known)
        raise UnconfirmedMappingError(f"Sheet {sheet_name!r}, {mapping.stop_message(proposals, path)}",
                                      proposals) from None
    except ValueError as error:
        raise ValueError(f"Sheet {sheet_name!r}, {error}") from None


if __name__ == "__main__":
    workbook_path = sys.argv[1] if len(sys.argv) > 1 else "data/northwind.xlsx"
    actuals, next_budget = clean_workbook(workbook_path)
    print(actuals.T.to_string())  # .T flips the table: quarters across, columns down
    print(f"\nBudget-only row: {next_budget.name if next_budget is not None else 'none'}")
    if next_budget is not None:
        print(next_budget.to_string())
