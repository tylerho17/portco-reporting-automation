"""Automated proof for Task 11: the exports carry the same values as the metrics workbook, and the email is Outlook-safe.

For each company, in a temporary folder (output/ is never touched):
1. excel_output.py saves the metrics workbook and `python export.py <workbook>` (export.main) saves the four
   exports. Everything is read back from disk the way another tool would: the workbook with openpyxl,
   the CSVs with the csv module, the JSON with json (which is told to refuse NaN and Infinity).
2. Metrics CSV and JSON: exactly one row per workbook cell, none missing, none extra. Each number
   equals the cell's number exactly (not rounded); each text equals the cell as Excel shows it (worked
   out here from the cell's own number format); each unit matches that format; the status matches
   the words in a cell with no number; flag_tripped is true exactly where the cell is red, and
   "missing input" exactly where it is gray.
3. Flags CSV and JSON against the Flags sheet: same flags in the same order, value, threshold,
   "trips when", status words, and a status that matches the row's color.
4. JSON: the quarters, runway at next quarter's budgeted burn, and the data gaps equal the workbook's.
5. The latest quarter's exported values equal the hand formulas in check_companies.py, and the flag
   statuses equal each company's story.
6. The email: follows every Outlook rule (outlook_problems); its table shows the workbook's values for
   the latest and prior quarters, thresholds and statuses, each status cell in its status color;
   the flag count, runway context and every data gap agree with the workbook.
7. The web page's download (export.export_bytes) is byte for byte the file the command line wrote.
8. No API: creating an Anthropic client stops the check. Nothing here asks Claude.

Run: python check_export.py  -> prints "All checks passed" or stops at the first failure.
"""

import csv
import io
import json
import math
import re
import tempfile
from contextlib import redirect_stdout
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

import analyze
import export
from check_companies import COMPANIES, same_number
from excel_output import save_metrics_workbook
from metrics import CANNOT_EVALUATE, METRIC_LABELS, PASS, TRIP, load_config
from portfolio import load_company
from theme import EXCEL_STATUS_COLORS, STATUS_COLORS

PROJECT_DIR = Path(__file__).parent
DATA_DIR = PROJECT_DIR / "data"
EM_DASH = chr(0x2014)

# Typed here from the workbook's own formats and words, not imported from export.py.
UNIT_OF_FORMAT = {"#,##0": "$K", '0.0" mo"': "months", '0.00"x"': "multiple", "0.0%": "ratio"}
SHOWN_AS = {"#,##0": "{:,.0f}", '0.0" mo"': "{:.1f} mo", '0.00"x"': "{:.2f}x", "0.0%": "{:.1%}"}
STATUS_OF_WORDS = [("data missing", "missing input"), ("n/a (no prior period)", "no prior period"),
                   ("n/m", "not meaningful"), ("∞", "infinite")]   # (the words start with, status)
FILL_OF = {status: colors[0] for status, colors in EXCEL_STATUS_COLORS.items()}
STATUS_OF_FILL = {fill: status for status, fill in FILL_OF.items()}
LABEL_TO_METRIC = {label: metric for metric, label in METRIC_LABELS.items()}
RUNWAY_ROW_START = "Runway at next quarter's budgeted burn"

# The Outlook rules (Outlook draws email with Word's HTML engine).
FORBIDDEN_TAGS = {"style", "script", "link", "img", "iframe", "svg", "video", "audio", "form", "input",
                  "button", "object", "embed"}   # dropped, blocked, or shown as a broken picture
TEXT_TAGS = {"td", "th", "p"}                     # each needs its own font, or Outlook shows Times New Roman
ALLOWED_CSS = {"font-family", "font-size", "font-weight", "color", "background-color", "margin", "padding",
               "width", "border", "border-bottom", "border-collapse", "text-align", "vertical-align", "line-height"}
UNSUPPORTED_VALUES = ("rgb(", "rgba(", "hsl(", "var(", "calc(", "url(")
HEX_COLOR = re.compile(r"#[0-9A-Fa-f]{6}")
MAX_WIDTH = 640
FONT_START = "Arial"


def refuse(*args, **kwargs):
    raise AssertionError("the export tried to create an Anthropic client")


# ---------------------------------------------------------------------------
# Reading the files back, as another tool would
# ---------------------------------------------------------------------------

def fill_of(cell):
    """A cell's solid fill as 6-digit hex ('FFC7CE'), or None."""
    if cell.fill is None or cell.fill.fill_type != "solid":
        return None
    return str(cell.fill.fgColor.rgb)[-6:].upper()


def sheet_rows(sheet):
    """[{header: {"value", "format", "fill"}}] for every row under the header, blank rows as None."""
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for row in sheet.iter_rows(min_row=2):
        if all(cell.value is None for cell in row):
            rows.append(None)
            continue
        rows.append({header: {"value": cell.value, "format": cell.number_format, "fill": fill_of(cell)}
                     for header, cell in zip(headers, row)})
    return rows


def read_workbook(path):
    """The metrics workbook as {"quarters", "metrics": {(quarter, label): cell}, "flags", "runway", "gaps"}."""
    book = load_workbook(path)
    metric_rows = sheet_rows(book["Metrics"])
    metrics = {(row["Quarter"]["value"], label): cell for row in metric_rows for label, cell in row.items()
               if label != "Quarter"}
    flag_rows = sheet_rows(book["Flags"])
    flags = [row for row in flag_rows if row and not str(row["Flag"]["value"]).startswith(RUNWAY_ROW_START)]
    runway = next(row for row in flag_rows if row and str(row["Flag"]["value"]).startswith(RUNWAY_ROW_START))
    gaps = [(row["Metric or flag"]["value"], row["Quarters with data missing"]["value"] or "")
            for row in sheet_rows(book["Data gaps"]) if row]
    return {"quarters": [row["Quarter"]["value"] for row in metric_rows], "metrics": metrics, "flags": flags,
            "runway": runway, "gaps": [gap for gap in gaps if not gap[0].startswith("None:")]}


def refuse_constant(name):
    raise ValueError(f"{name} in the JSON: strict JSON readers reject it")


def read_exports(paths):
    """{"metrics_csv": [rows], "flags_csv": [rows], "json": dict, "email": text}, read from disk."""
    def csv_rows(path):
        with open(path, newline="", encoding="utf-8") as file:
            return list(csv.DictReader(file))
    return {"metrics_csv": csv_rows(paths["metrics_csv"]), "flags_csv": csv_rows(paths["flags_csv"]),
            "json": json.loads(paths["json"].read_text(encoding="utf-8"), parse_constant=refuse_constant),
            "email": paths["email"].read_text(encoding="utf-8")}


# ---------------------------------------------------------------------------
# Comparing an export with the workbook (a CSV row holds text; a JSON row holds numbers, true/false, null)
# ---------------------------------------------------------------------------

def number(value):
    """A CSV or JSON value as a float, or None when it's empty."""
    return None if value in (None, "") else float(value)


def words(value):
    """A CSV or JSON text, with empty and null both as ''."""
    return "" if value is None else str(value)


def is_true(value):
    return value is True or value == "true"


def shown(cell):
    """A workbook cell as Excel shows it: its number in its own number format, or its words."""
    if isinstance(cell["value"], (int, float)):
        return SHOWN_AS[cell["format"]].format(cell["value"])
    return cell["value"]


def status_of(cell_words):
    """The status the words in a cell with no number stand for."""
    return next((status for start, status in STATUS_OF_WORDS if cell_words.startswith(start)), None)


def expected_status(cell):
    """The status a cell stands for: a number, or what the words in a cell with no number mean."""
    return export.NUMBER if isinstance(cell["value"], (int, float)) else status_of(cell["value"])


def value_problems(where, row, cell):
    """The row's value and text against one workbook cell."""
    problems = []
    if isinstance(cell["value"], (int, float)):
        if number(row["value"]) != cell["value"]:
            problems.append(f"{where}: value {row['value']} but the workbook has {cell['value']!r}")
    elif number(row["value"]) is not None:
        problems.append(f"{where}: value {row['value']} but the workbook has no number ({cell['value']})")
    if words(row["text"]) != shown(cell):
        problems.append(f"{where}: text {row['text']!r} but the workbook shows {shown(cell)!r}")
    return problems


def cell_problems(row, cell):
    """One metric row against its Metrics sheet cell: value, text, status, label, unit and color."""
    where = f"{row['quarter']} {row['label']}"
    problems = value_problems(where, row, cell)
    if row["status"] != expected_status(cell):
        problems.append(f"{where}: status {row['status']!r}, expected {expected_status(cell)!r}")
    if METRIC_LABELS.get(row["metric"]) != row["label"]:
        problems.append(f"{where}: metric {row['metric']!r} doesn't have this label")
    if row["unit"] != UNIT_OF_FORMAT[cell["format"]]:
        problems.append(f"{where}: unit {row['unit']!r} but the cell's format is {cell['format']!r}")
    if is_true(row["flag_tripped"]) != (cell["fill"] == FILL_OF[TRIP]):
        problems.append(f"{where}: flag_tripped {row['flag_tripped']} but the cell's fill is {cell['fill']}")
    if (row["status"] == "missing input") != (cell["fill"] == FILL_OF[CANNOT_EVALUATE]):
        problems.append(f"{where}: status {row['status']!r} but the cell's fill is {cell['fill']}")
    return problems


def metric_problems(rows, excel):
    """Every metric row against the Metrics sheet: one row per cell, each cell's value, text and color."""
    problems, seen = [], set()
    for row in rows:
        key = (row["quarter"], row["label"])
        if key not in excel["metrics"]:
            problems.append(f"{key[0]} {key[1]}: exported but not in the workbook")
        elif key in seen:
            problems.append(f"{key[0]} {key[1]}: exported twice")
        else:
            seen.add(key)
            problems += cell_problems(row, excel["metrics"][key])
    problems += [f"{quarter} {label}: in the workbook but not exported"
                 for quarter, label in excel["metrics"] if (quarter, label) not in seen]
    return problems


def flag_status_problems(where, row, sheet_row):
    """Status words, the status the row's color stands for, and the reason."""
    problems = []
    status_words = sheet_row["Status"]["value"]
    if row["status_text"] != status_words:
        problems.append(f"{where}: status {row['status_text']!r} but the workbook says {status_words!r}")
    if row["status"] != STATUS_OF_FILL.get(sheet_row["Status"]["fill"]):
        problems.append(f"{where}: status {row['status']!r} but the row's fill is {sheet_row['Status']['fill']}")
    reason = status_words.split(": ", 1)[1] if ": " in status_words else ""
    if words(row["reason"]) != reason:
        problems.append(f"{where}: reason {row['reason']!r}, expected {reason!r}")
    return problems


def flag_problems(rows, excel):
    """Every flag row against the Flags sheet, in the same order."""
    names, expected = [row["flag"] for row in rows], [row["Flag"]["value"] for row in excel["flags"]]
    if names != expected:
        return [f"flags {names} but the workbook has {expected}"]
    problems = []
    for row, sheet_row in zip(rows, excel["flags"]):
        where = f"Flag {row['flag']}"
        if row["quarter"] != sheet_row["Quarter"]["value"]:
            problems.append(f"{where}: quarter {row['quarter']} but the workbook has {sheet_row['Quarter']['value']}")
        problems += flag_status_problems(where, row, sheet_row)
        if words(row["metric"]) == "":   # the combo rule: no value or threshold, only its words
            trips_when = f"{sheet_row['Trips when']['value']}, {sheet_row['Threshold']['value']}"
            if number(row["value"]) is not None or number(row["threshold"]) is not None:
                problems.append(f"{where}: the combo rule has no single value or threshold")
        else:
            trips_when = sheet_row["Trips when"]["value"]
            problems += value_problems(where, row, sheet_row["Value"])
            if number(row["threshold"]) != sheet_row["Threshold"]["value"]:
                problems.append(f"{where}: threshold {row['threshold']} but the workbook has "
                                f"{sheet_row['Threshold']['value']}")
        if row["trips_when"] != trips_when:
            problems.append(f"{where}: trips when {row['trips_when']!r}, expected {trips_when!r}")
    return problems


def summary_problems(record, excel):
    """The JSON's flag counts against the Flags sheet's colors."""
    statuses = [STATUS_OF_FILL.get(row["Status"]["fill"]) for row in excel["flags"]]
    expected = {"checked": len(statuses), "tripped": statuses.count(TRIP), "passed": statuses.count(PASS),
                "cannot_evaluate": statuses.count(CANNOT_EVALUATE)}
    got = {key: record["flag_summary"][key] for key in expected}
    return [] if got == expected else [f"flag_summary {got} but the workbook's colors give {expected}"]


def record_problems(record, excel):
    """The JSON's quarters, flag counts, runway at next quarter's budget and data gaps against the workbook."""
    problems = []
    if record["quarters"] != excel["quarters"] or record["latest_quarter"] != excel["quarters"][-1]:
        problems.append(f"quarters {record['quarters']} but the workbook has {excel['quarters']}")
    problems += summary_problems(record, excel)
    runway, cell = record["runway_at_next_budget"], excel["runway"]["Value"]
    if record["runway_at_next_budget"]["quarter"] != excel["runway"]["Quarter"]["value"]:
        problems.append(f"runway at budget: quarter {runway['quarter']}")
    problems += value_problems("Runway at next quarter's budgeted burn", runway, cell)
    gaps = [(gap["label"], ", ".join(gap["quarters"])) for gap in record["data_gaps"]]
    if gaps != excel["gaps"]:
        problems.append(f"data gaps {gaps} but the workbook has {excel['gaps']}")
    return problems


# ---------------------------------------------------------------------------
# The email: reading its HTML, the Outlook rules, and its values
# ---------------------------------------------------------------------------

class EmailParser(HTMLParser):
    """Collects every element (tag and attributes), each table's cells, and each paragraph's text."""

    def __init__(self):
        super().__init__(convert_charrefs=True)   # "&amp;" arrives as "&"
        self.elements, self.tables, self.paragraphs = [], [], []
        self.cell = self.paragraph = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        self.elements.append((tag, attrs))
        if tag == "table":
            self.tables.append([])
        elif tag == "tr" and self.tables:
            self.tables[-1].append([])
        elif tag in ("td", "th") and self.tables and self.tables[-1]:
            self.cell = {"text": "", "attrs": attrs}
            self.tables[-1][-1].append(self.cell)
        elif tag == "p":
            self.paragraph = []

    def handle_endtag(self, tag):
        if tag in ("td", "th"):
            self.cell = None
        elif tag == "p" and self.paragraph is not None:
            self.paragraphs.append("".join(self.paragraph))
            self.paragraph = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell["text"] += data
        if self.paragraph is not None:
            self.paragraph.append(data)


def parse_email(html):
    parser = EmailParser()
    parser.feed(html)
    parser.close()
    return parser


def email_tables(html):
    """Each table as rows of cell texts."""
    return [[[cell["text"] for cell in row] for row in table] for table in parse_email(html).tables]


def declarations(style):
    """'color: #334155; font-size: 14px' -> {"color": "#334155", "font-size": "14px"}."""
    result = {}
    for part in (style or "").split(";"):
        name, _, value = part.partition(":")
        if name.strip():
            result[name.strip().lower()] = value.strip()
    return result


def style_problems(where, styles, attrs):
    """One element's inline CSS: known properties, supported values, hex colors, a fill also as bgcolor."""
    problems = [f"{where}: Outlook ignores {name}" for name in styles if name not in ALLOWED_CSS]
    for name, value in styles.items():
        if any(bad in value for bad in UNSUPPORTED_VALUES):
            problems.append(f"{where}: {name}: {value} (Outlook needs plain values)")
        elif name.endswith("color") and not HEX_COLOR.fullmatch(value):
            problems.append(f"{where}: {name}: {value} is not a 6-digit hex color")
    if "background-color" in styles and (attrs.get("bgcolor") or "").upper() != styles["background-color"].upper():
        problems.append(f"{where}: a fill without the same bgcolor (Outlook may drop the fill)")
    return problems


def table_problems(where, attrs):
    """A table needs its size and spacing as attributes, and must fit the email's width."""
    problems = [f"{where}: no {name} attribute" for name in ("width", "cellpadding", "cellspacing", "border")
                if name not in attrs]
    width = attrs.get("width") or ""
    if width and (not width.isdigit() or int(width) > MAX_WIDTH):
        problems.append(f"{where}: width {width} (at most {MAX_WIDTH} pixels)")
    return problems


def outlook_problems(html):
    """Everything in the email that Outlook would drop, block or draw differently; [] if none."""
    problems = [] if '<meta charset="utf-8">' in html else ["no <meta charset=\"utf-8\">: ∞ would show as junk"]
    if EM_DASH in html:
        problems.append("an em dash")
    for number, (tag, attrs) in enumerate(parse_email(html).elements):
        where = f"<{tag}> (element {number + 1})"
        if tag in FORBIDDEN_TAGS:
            problems.append(f"{where}: Outlook drops or blocks it")
        if "class" in attrs:
            problems.append(f"{where}: a class (Outlook drops style sheets, so a class does nothing)")
        styles = declarations(attrs.get("style"))
        problems += style_problems(where, styles, attrs)
        if tag == "table":
            problems += table_problems(where, attrs)
        if tag in TEXT_TAGS and not styles.get("font-family", "").startswith(FONT_START):
            problems.append(f"{where}: no font-family starting with {FONT_START} (Outlook shows Times New Roman)")
    return problems


def expected_threshold(sheet_row, metric):
    """'trips below 100.0%': the Flags sheet's threshold in its own format, and which side trips."""
    side = sheet_row["Trips when"]["value"].split()[0]   # "below threshold" -> "below"
    return f"trips {side} {shown(sheet_row['Threshold'])}"


def flag_cell_problems(row_cells, sheet_row, excel):
    """A flag's threshold and status cells in the email against the Flags sheet."""
    label = sheet_row["Flag"]["value"]
    metric = LABEL_TO_METRIC.get(label)
    threshold = (expected_threshold(sheet_row, metric) if metric else
                 f"{sheet_row['Trips when']['value']}, {sheet_row['Threshold']['value']}")
    problems = []
    if row_cells[3]["text"] != threshold:
        problems.append(f"email {label}: threshold {row_cells[3]['text']!r}, expected {threshold!r}")
    if row_cells[4]["text"] != sheet_row["Status"]["value"]:
        problems.append(f"email {label}: status {row_cells[4]['text']!r}, expected {sheet_row['Status']['value']!r}")
    fill = STATUS_COLORS[STATUS_OF_FILL[sheet_row["Status"]["fill"]]][0]
    if (row_cells[4]["attrs"].get("bgcolor") or "").upper() != f"#{fill}":
        problems.append(f"email {label}: status cell filled {row_cells[4]['attrs'].get('bgcolor')}, expected #{fill}")
    return problems


def email_row_problems(row_cells, latest, prior, excel):
    """One row of the email's table: its two quarters' values, and a flag's threshold and status."""
    label = row_cells[0]["text"]
    problems = []
    if label in LABEL_TO_METRIC:
        for column, quarter in ((1, latest), (2, prior)):
            expected = shown(excel["metrics"][(quarter, label)])
            if row_cells[column]["text"] != expected:
                problems.append(f"email {label} {quarter}: {row_cells[column]['text']!r}, the workbook shows {expected!r}")
    sheet_row = next((row for row in excel["flags"] if row["Flag"]["value"] == label), None)
    if sheet_row is not None:
        problems += flag_cell_problems(row_cells, sheet_row, excel)
    return problems


def flag_count_words(excel):
    """'6 of 9 flags tripped' (', 1 cannot evaluate'), counted here from the Flags sheet's colors."""
    statuses = [STATUS_OF_FILL.get(row["Status"]["fill"]) for row in excel["flags"]]
    text = f"{statuses.count(TRIP)} of {len(statuses)} flags tripped"
    unevaluated = statuses.count(CANNOT_EVALUATE)
    return text + (f", {unevaluated} cannot evaluate" if unevaluated else "")


def email_problems(html, excel):
    """The email's table, flag count, runway context and data gaps against the workbook."""
    parser = parse_email(html)
    rows = parser.tables[0]
    latest, prior = rows[0][1]["text"], rows[0][2]["text"]
    problems = [] if latest == excel["quarters"][-1] else [f"email: latest quarter {latest}"]
    labels = [row[0]["text"] for row in rows[1:]]
    missing = [row["Flag"]["value"] for row in excel["flags"] if row["Flag"]["value"] not in labels]
    problems += [f"email: flag {name} not in the table" for name in missing]
    for row_cells in rows[1:]:
        problems += email_row_problems(row_cells, latest, prior, excel)
    ending_arr_budget = f"vs budget: {shown(excel['metrics'][(latest, METRIC_LABELS['arr_vs_budget'])])}"
    if not any(row[0]["text"] == METRIC_LABELS["ending_arr"] and row[3]["text"] == ending_arr_budget for row in rows):
        problems.append(f"email: ending ARR's budget cell isn't {ending_arr_budget!r}")
    text = "\n".join(parser.paragraphs)
    expected_lines = [flag_count_words(excel),
                      f"{RUNWAY_ROW_START}: {shown(excel['runway']['Value'])} (context, not a flag)"]
    problems += [f"email: no line {line!r}" for line in expected_lines if line not in parser.paragraphs]
    problems += [f"email: data gap {label} not listed" for label, _ in excel["gaps"] if label not in text]
    return problems


# ---------------------------------------------------------------------------
# The answer keys, and one company end to end
# ---------------------------------------------------------------------------

def exported_value(row):
    """A metric row as a float for same_number: the number, inf for infinite, NaN for no number."""
    if row["status"] == export.INFINITE:
        return math.inf
    value = number(row["value"])
    return math.nan if value is None else value


def answer_key_problems(company, exports):
    """The latest quarter's values against check_companies.py's hand formulas, and the flags against the story."""
    rows = {row["metric"]: row for row in exports["metrics_csv"] if row["quarter"] == exports["json"]["latest_quarter"]}
    problems = [f"{company['name']} {metric}: exported {rows[metric]['value']!r} ({rows[metric]['status']}), "
                f"hand formula {expected}" for metric, expected in company["expected_latest"].items()
                if not same_number(exported_value(rows[metric]), expected)]
    statuses = {row["flag"]: row["status"] for row in exports["flags_csv"]}
    if statuses != company["expected_flags"]:
        problems.append(f"{company['name']} flags {statuses}, the story says {company['expected_flags']}")
    return problems


def check_company(company, config, folder):
    """Save the workbook and the exports, read both back, and compare everything."""
    workbook = DATA_DIR / f"{company['name'].lower()}.xlsx"
    excel = read_workbook(save_metrics_workbook(workbook, config, folder))
    with redirect_stdout(io.StringIO()):   # the command line's "Saved ..." lines would only repeat the folder
        assert export.main([str(workbook), "--output-dir", str(folder)]) == 0, "python export.py failed"
    paths = export.export_paths(workbook, folder)
    exports = read_exports(paths)
    checks = [("metrics CSV", metric_problems(exports["metrics_csv"], excel)),
              ("JSON metrics", metric_problems(exports["json"]["metrics"], excel)),
              ("flags CSV", flag_problems(exports["flags_csv"], excel)),
              ("JSON flags", flag_problems(exports["json"]["flags"], excel)),
              ("JSON runway and gaps", record_problems(exports["json"], excel)),
              ("answer key", answer_key_problems(company, exports)),
              ("email values", email_problems(exports["email"], excel)),
              ("email Outlook rules", outlook_problems(exports["email"]))]
    for what, problems in checks:
        assert problems == [], f"{company['name']} {what}:\n" + "\n".join(problems)
    data, _ = load_company(workbook, config)
    for kind, path in paths.items():
        assert export.export_bytes(kind, data, workbook) == path.read_bytes(), f"{company['name']}: page's {kind} differs"
    print(f"✓ {company['name']}: {len(exports['metrics_csv'])} metric values, {len(exports['flags_csv'])} flags, "
          f"runway at budget and {len(excel['gaps'])} data gaps equal the workbook; email Outlook-safe")


def main():
    analyze.anthropic.Anthropic = refuse   # any attempt to reach Claude stops the check
    config = load_config()
    with tempfile.TemporaryDirectory() as folder:
        for company in COMPANIES:
            check_company(company, config, Path(folder))
    print("All checks passed")


if __name__ == "__main__":
    main()
