"""Automated proof for build step 4b: the Excel output matches the metrics table.

For all three companies this builds output/<company>_metrics.xlsx, reads it back
from disk with openpyxl, and checks:
1. The file is where it should be, with sheets Metrics, Flags, Data gaps in that order.
2. Metrics sheet: every cell equals the metrics table to 15 significant digits (a real
   number, not rounded text), ratios are decimals with a % format, and a value with no
   number has the exact label for its reason ("data missing", "n/a (no prior period)",
   "n/m ...") and every infinity has its exact "∞ (...)" label.
   Latest-quarter cells also match the hand formulas in check_companies.py.
3. Metrics highlights: red where a flag trips in that quarter, gray only where an input is
   missing, no fill anywhere else.
4. Flags sheet: name, value, threshold, status and row color match the company's story.
5. Data gaps sheet lists exactly the gaps metrics.data_gaps found, no more and no fewer.

Expected labels and formats are written out here on purpose instead of imported from
excel_output.py, so a wrong constant there can't make its own check pass. Colors come from
theme.py, the only file that types one (tests/test_theme.py types them by hand).

Run: python check_excel_output.py  -> prints "All checks passed" or stops at the first failure.
"""

import math
from pathlib import Path

from openpyxl import load_workbook

from check_companies import ALL_FLAG_NAMES, COMPANIES, LATEST, same_number
from clean import clean_workbook
from excel_output import save_metrics_workbook
from metrics import (CANNOT_EVALUATE, DOLLAR_COLUMNS, FLAG_RULES, METRIC_LABELS, MISSING_INPUT, MONTH_COLUMNS,
                     NO_PRIOR_PERIOD, NOT_MEANINGFUL, PASS, TRIP, compute_metrics, data_gaps, evaluate_flags,
                     load_config, metric_reasons, runway_at_next_budget)
from theme import EXCEL_STATUS_COLORS

OUTPUT_DIR = Path(__file__).parent / "output"

# Typed out by hand (not imported from excel_output.py) so a wrong constant there can't pass its own check.
STATUS_TEXT = {TRIP: "Tripped", PASS: "Passed"}
REASON_WORDS = {MISSING_INPUT: "data missing", NO_PRIOR_PERIOD: "n/a (no prior period)"}
INFINITE_TEXT = {
    "burn_multiple": "∞ (ARR shrank)",
    "runway_months": "∞ (not burning)",
    "cac_payback_months": "∞ (never pays back)",
}
# Fills from theme.py, the only file that types a color (tests/test_theme.py types these values by hand).
STATUS_FILL = {status: fill for status, (fill, _) in EXCEL_STATUS_COLORS.items()}
RED, GRAY = STATUS_FILL[TRIP], STATUS_FILL[CANNOT_EVALUATE]


def status_text(status, reason):
    """'Tripped', 'Passed', or 'Cannot evaluate: <reason>'."""
    return f"Cannot evaluate: {reason}" if status == CANNOT_EVALUATE else STATUS_TEXT[status]


# ---------------------------------------------------------------------------
# Reading cells back
# ---------------------------------------------------------------------------

def fill_color(cell):
    """The cell's fill as 6-digit hex (FFC7CE is Excel's red), or None if it has no fill."""
    if cell.fill.fill_type is None:
        return None
    return cell.fill.fgColor.rgb[-6:]  # openpyxl adds 2 alpha digits in front: '00FFC7CE'


def is_number(value):
    """True for a real number cell (Excel reads 27470.0 back as the int 27470)."""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def same_as_saved(got, expected):
    """True if a number read back from Excel is the metric value, to 15 significant digits.

    openpyxl saves numbers with 16 significant digits, but a Python float can need 17,
    so 1.0706921944035346 comes back as 1.070692194403535. Excel itself only keeps 15
    digits, so that is the most precise "equal" can mean. It still catches any real
    difference, e.g. 97.1 saved instead of 0.971, or a value rounded to 0.971.
    """
    return math.isclose(got, expected, rel_tol=1e-15, abs_tol=1e-15)


def excel_number(cell):
    """Turn a cell back into a metric value: number -> float, a reason label -> NaN, an exact ∞ label -> inf."""
    if is_number(cell.value):
        return float(cell.value)
    if cell.value in REASON_WORDS.values() or str(cell.value).startswith("n/m"):
        return math.nan
    if cell.value in INFINITE_TEXT.values():
        return math.inf
    raise AssertionError(f"{cell.coordinate}: can't read {cell.value!r} as a metric value")


def expected_format(column):
    """The Excel number format each kind of metric must have."""
    if column in DOLLAR_COLUMNS:
        return "#,##0"
    if column in MONTH_COLUMNS:
        return '0.0" mo"'
    if column == "burn_multiple":
        return '0.00"x"'
    return "0.0%"  # every other metric is a ratio stored as a decimal


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def check_file(path, answer_key):
    """Saved as output/<input file name>_metrics.xlsx, with the three sheets in order."""
    expected_path = OUTPUT_DIR / f"{Path(answer_key.OUTPUT_PATH).stem}_metrics.xlsx"
    assert Path(path) == expected_path, f"Saved to {path}, expected {expected_path}"
    book = load_workbook(path)  # read back from disk, not the in-memory workbook
    assert book.sheetnames == ["Metrics", "Flags", "Data gaps"], f"Sheets wrong: {book.sheetnames}"
    return book


def check_value_cell(cell, column, value, reason):
    """One cell holds exactly the metric value, or the exact label for why there's no number."""
    where = f"{cell.parent.title}!{cell.coordinate} ({column})"
    if reason == NOT_MEANINGFUL:
        assert isinstance(cell.value, str) and cell.value.startswith("n/m"), f"{where}: expected n/m, got {cell.value!r}"
    elif isinstance(reason, str):
        expected = REASON_WORDS[reason]
        assert cell.value == expected, f"{where}: expected {expected!r}, got {cell.value!r}"
    elif math.isinf(value):
        expected = INFINITE_TEXT[column]
        assert cell.value == expected, f"{where}: expected {expected!r}, got {cell.value!r}"
    else:
        assert is_number(cell.value), f"{where}: expected a number, got {cell.value!r}"
        assert same_as_saved(cell.value, value), f"{where}: expected {float(value)!r}, got {cell.value!r}"
    assert cell.number_format == expected_format(column), \
        f"{where}: format {cell.number_format!r}, expected {expected_format(column)!r}"


def tripped_by_quarter(metrics, reasons, config):
    """{(quarter, metric column)} for every metric flag that trips, in every quarter."""
    return {(quarter, flag["metric"])
            for quarter in metrics.index
            for flag in evaluate_flags(metrics, reasons, config, quarter)
            if flag["status"] == TRIP and flag["metric"] is not None}


def check_metrics_sheet(sheet, metrics, reasons, gaps, config):
    """Headers, quarters, every value and format, and the red/gray highlights."""
    columns = list(metrics.columns)
    header = [cell.value for cell in sheet[1]]
    assert header == ["Quarter"] + [METRIC_LABELS[c] for c in columns], f"Metrics header wrong: {header}"
    assert sheet.max_row == len(metrics) + 1, f"Metrics has {sheet.max_row - 1} rows, expected {len(metrics)}"
    tripped = tripped_by_quarter(metrics, reasons, config)

    for row, quarter in zip(sheet.iter_rows(min_row=2), metrics.index):
        assert row[0].value == quarter, f"Metrics {row[0].coordinate}: expected {quarter}, got {row[0].value}"
        for cell, column in zip(row[1:], columns):
            is_gap = quarter in gaps.get(column, [])  # data_gaps lists missing input only
            check_value_cell(cell, column, metrics.loc[quarter, column], reasons.loc[quarter, column])
            expected_fill = GRAY if is_gap else RED if (quarter, column) in tripped else None
            assert fill_color(cell) == expected_fill, \
                f"Metrics {cell.coordinate} ({quarter}, {column}): fill {fill_color(cell)}, expected {expected_fill}"


def check_latest_against_hand_formulas(sheet, company, metrics):
    """Latest-quarter cells equal the hand formulas typed from the answer key."""
    latest_row = sheet[sheet.max_row]
    assert latest_row[0].value == LATEST, f"Last Metrics row is {latest_row[0].value}, expected {LATEST}"
    cells = dict(zip(metrics.columns, latest_row[1:]))  # metric column -> its cell
    for column, expected in company["expected_latest"].items():
        got = excel_number(cells[column])
        assert same_number(got, expected), f"Excel {LATEST} {column}: expected {expected}, got {got}"


def check_story_highlights(sheet, company, metrics):
    """Latest-quarter red cells are exactly the metrics whose flags the story says trip."""
    story_trips = {column for name, column, _, _ in FLAG_RULES if company["expected_flags"][name] == TRIP}
    cells = dict(zip(metrics.columns, sheet[sheet.max_row][1:]))
    red = {column for column, cell in cells.items() if fill_color(cell) == RED}
    assert red == story_trips, f"{LATEST} red cells {sorted(red)}, story says {sorted(story_trips)}"


def check_flag_row(row, name, column, config_key, company, metrics, reasons, config):
    """One metric flag row: name, quarter, value, threshold, status, and the row's color."""
    flag_cell, quarter_cell, value_cell, threshold_cell, _, status_cell = row
    status = company["expected_flags"][name]
    assert flag_cell.value == name, f"Flags {flag_cell.coordinate}: expected {name!r}, got {flag_cell.value!r}"
    assert quarter_cell.value == LATEST, f"Flags {quarter_cell.coordinate}: got {quarter_cell.value!r}"
    check_value_cell(value_cell, column, metrics.loc[LATEST, column], reasons.loc[LATEST, column])
    assert threshold_cell.value == config[config_key], \
        f"Flags {threshold_cell.coordinate}: threshold {threshold_cell.value}, expected {config[config_key]}"
    assert threshold_cell.number_format == expected_format(column), f"Flags {threshold_cell.coordinate}: wrong format"
    check_status_and_color(row, name, status)


def check_status_and_color(row, name, status):
    """Status text matches (in these stories, can't-evaluate is always missing input), and the row's fill."""
    status_cell = row[-1]
    expected = status_text(status, MISSING_INPUT)
    assert status_cell.value == expected, f"{name}: status {status_cell.value!r}, expected {expected!r}"
    fills = {fill_color(cell) for cell in row}
    assert fills == {STATUS_FILL[status]}, f"{name}: row fills {fills}, expected {STATUS_FILL[status]}"


def check_flags_sheet(sheet, company, metrics, reasons, config, runway_budget):
    """Every flag row, the combo row, and the runway-at-budget context line."""
    header = [cell.value for cell in sheet[1]]
    assert header == ["Flag", "Quarter", "Value", "Threshold", "Trips when", "Status"], f"Flags header wrong: {header}"
    rows = list(sheet.iter_rows(min_row=2, max_row=len(ALL_FLAG_NAMES) + 1))
    assert [row[0].value for row in rows] == ALL_FLAG_NAMES, f"Flag names wrong: {[r[0].value for r in rows]}"

    for row, (name, column, config_key, _) in zip(rows, FLAG_RULES):
        check_flag_row(row, name, column, config_key, company, metrics, reasons, config)
    combo_name = ALL_FLAG_NAMES[-1]
    check_status_and_color(rows[-1], combo_name, company["expected_flags"][combo_name])

    # One empty row, then the context line. Nothing else on the sheet.
    context_row = len(ALL_FLAG_NAMES) + 3
    assert sheet.max_row == context_row, f"Flags sheet has {sheet.max_row} rows, expected {context_row}"
    assert all(cell.value is None for cell in sheet[context_row - 1]), "Row above runway context should be empty"
    label_cell, _, runway_cell = sheet[context_row][:3]
    assert label_cell.value.startswith("Runway at next quarter's budgeted burn"), f"Got {label_cell.value!r}"
    assert is_number(runway_cell.value) and same_as_saved(runway_cell.value, runway_budget), \
        f"Runway at budget {runway_cell.value!r}, expected {runway_budget}"
    assert same_number(excel_number(runway_cell), company["expected_runway_at_budget"]), "Runway at budget vs hand formula"
    assert fill_color(label_cell) is None, "Runway context is not a flag, so it must not be colored"


def check_gaps_sheet(sheet, gaps):
    """One row per gap from data_gaps, in the same order, with the same quarters."""
    header = [cell.value for cell in sheet[1]]
    assert header == ["Metric or flag", "Quarters with data missing"], f"Data gaps header wrong: {header}"
    got = [(a.value, b.value) for a, b in sheet.iter_rows(min_row=2, max_col=2)]

    if not gaps:
        assert len(got) == 1 and got[0][0].startswith("None"), f"No gaps expected, sheet says {got}"
        return
    expected = []
    for name, quarters in gaps.items():
        label = "Flag: " + name.removeprefix("flag: ") if name.startswith("flag: ") else METRIC_LABELS[name]
        expected.append((label, ", ".join(quarters)))
    assert got == expected, f"Data gaps wrong.\nExpected: {expected}\nGot:      {got}"


def check_company(company, config):
    """Build the Excel file for one company, read it back, and run every check."""
    name, answer_key = company["name"], company["answer_key"]
    actuals, next_budget = clean_workbook(answer_key.OUTPUT_PATH)
    metrics = compute_metrics(actuals)
    reasons = metric_reasons(actuals, metrics)
    gaps = data_gaps(actuals, metrics, evaluate_flags(metrics, reasons, config))

    book = check_file(save_metrics_workbook(answer_key.OUTPUT_PATH, config), answer_key)
    check_metrics_sheet(book["Metrics"], metrics, reasons, gaps, config)
    check_latest_against_hand_formulas(book["Metrics"], company, metrics)
    check_story_highlights(book["Metrics"], company, metrics)
    print(f"✓ {name}: Metrics sheet matches the metrics table ({len(metrics)} quarters: values, formats, highlights)")
    check_flags_sheet(book["Flags"], company, metrics, reasons, config, runway_at_next_budget(actuals, next_budget))
    print(f"✓ {name}: Flags sheet matches the story (values, thresholds, statuses, colors)")
    check_gaps_sheet(book["Data gaps"], gaps)
    print(f"✓ {name}: Data gaps sheet lists exactly the {len(gaps)} gaps")


def main():
    config = load_config()
    for company in COMPANIES:
        check_company(company, config)
    print("All checks passed")


if __name__ == "__main__":
    main()
