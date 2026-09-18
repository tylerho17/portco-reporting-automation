"""Build the board memo: output/<company>_board_memo.docx and .pdf, beside the deck (final Task 1).

A 1 to 2 page written update for board members who read rather than present:
1. Title and quarter:          "<Company> board update: Q2 2026", compared with the prior quarter
2. Headline:                   the AI headline under "AI-drafted from computed metrics - review before use"
3. Key metrics:                the same table as slide 1 (latest, prior, budget or threshold, status),
                               then runway at next quarter's budgeted burn
4. Flags:                      "6 of 9 flags tripped", each tripped flag with its value and threshold,
                               flags that can't be evaluated, and the combo rule
5. Data gaps:                  every metric and flag with data missing
6. Questions for management:   the AI's 3 questions
Footer on every page: fictional-data note, source file, run date, commit, model, review status.

The content is built once as a list of "blocks" (heading, text, bullets, table), then written twice:
write_docx (python-docx) and write_pdf (reportlab). Both files carry the same words.

Rules (the same as build_deck.py):
- No math and no hand-typed numbers: every number comes from metrics.py, config.yaml or the
  validated analysis, through the deck's own text helpers. tests/test_build_deck.py's
  no-digit test covers this file too.
- The AI text is used only if it passes everything the deck checks (build_deck.load_analysis),
  AND every number in the headline and questions is one the metrics workbook shows. Otherwise the
  memo says "AI commentary unavailable" and every computed number still appears.
- No em dashes: two labels shared with the deck and Excel have one (the "Cannot evaluate" status
  and the "None" data gaps line); the memo shows a colon instead.

Run: python memo.py data/northwind.xlsx                  (uses output/northwind_analysis.json if it exists)
     python memo.py data/northwind.xlsx --no-analysis    (AI commentary unavailable)
"""

import argparse
import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.shared import Inches as DocxInches
from matplotlib import font_manager
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from analyze import build_payload, numbers_in
from build_deck import (AI_DRAFTED, FICTIONAL_NOTE, NO_AI_MODEL, OUTPUT_DIR, QUESTIONS_HEADING, analysis_details,
                        analysis_path, collect_deck_data, commit_text, flag_count_text, gaps_lines, kpi_header,
                        load_analysis, review_text, threshold_text, value_text)
from clean import clean_workbook
from excel_output import STATUS_COLORS, flag_row, runway_context_value, status_label
from make_template import DARK_GRAY, FONT, LIGHT_GRAY, MID_GRAY, NAVY, WHITE
from metrics import CANNOT_EVALUATE, CONFIG_PATH, METRIC_LABELS, TRIP, format_value, load_config
from provenance import approval_status, file_sha256, git_commit, manifest_path, read_manifest, save_manifest

MEMO_UNAVAILABLE = "AI commentary unavailable"
MEMO_UNAVAILABLE_NOTE = ("The headline and questions are written by Claude, and no validated version is "
                         "available for this memo. Every number in it is computed in Python and is unaffected.")
AI_DRAFTED_LINE = "AI-drafted from computed metrics - review before use"
NOT_A_FLAG = "-"                        # the budget or status cell of a row that isn't a flag
COMBO_CELL = "see Flags below"          # the combo rule has no single value; its rule is in the Flags section
NO_TRIPPED_FLAGS = "No flag tripped"
RUNWAY_CONTEXT = "Runway at next quarter's budgeted burn (context, not a flag)"
MEMO_SUFFIX = "_board_memo"
EM_DASH = chr(0x2014)                   # the em dash, by its Unicode number, so this file never shows one

# Rows of the key metrics table that aren't flags, shown first: (metric, its "vs budget" metric or None).
CONTEXT_ROWS = [("ending_arr", "arr_vs_budget"), ("net_new_arr", None), ("arr_yoy", None), ("gross_margin", None)]

# ---------------------------------------------------------------------------
# Look: font sizes (pt), page margins (inches) and table column shares
# ---------------------------------------------------------------------------

TITLE_SIZE = 18
HEADING_SIZE = 12
HEADLINE_SIZE = 11
BODY_SIZE = 9.5
TABLE_SIZE = 8.5
FOOTER_SIZE = 7
SPACE_BEFORE_HEADING = 8
SPACE_AFTER_TEXT = 3
PAGE_MARGIN = 0.7                   # inches, every side, in both files
KPI_COLUMN_SHARES = [0.30, 0.13, 0.13, 0.24, 0.20]

# The PDF embeds DejaVu Sans (it ships with matplotlib, and text_fit.py measures with it): the
# PDF's built-in fonts have no "∞" and would print a box.
PDF_FONT, PDF_BOLD_FONT = "DejaVuSans", "DejaVuSans-Bold"

# A text block's style -> (font size, bold, color). Both files use this, so they look alike.
TEXT_STYLES = {
    "body": (BODY_SIZE, False, DARK_GRAY),
    "note": (BODY_SIZE, False, MID_GRAY),
    "headline": (HEADLINE_SIZE, True, NAVY),
    "unavailable": (HEADLINE_SIZE, True, MID_GRAY),
}


# ---------------------------------------------------------------------------
# 1. The AI analysis: used only if valid today, and every number it shows is in the metrics workbook
# ---------------------------------------------------------------------------

def workbook_texts(data):
    """Every text the metrics workbook shows for this company (excel_output.py), as the memo writes them.

    Metric values for every quarter and the quarter labels (Metrics sheet), flag names and
    thresholds, the combo rule's settings and the runway context (Flags sheet), and the flag
    count, which is the Flags sheet's rows counted by status.
    """
    texts = [flag_count_text(data["flags"])]
    for quarter in data["metrics"].index:
        texts.append(quarter)
        texts += [value_text(data, metric, quarter) for metric in data["metrics"].columns]
    for flag in data["flags"]:
        if flag["metric"] is None:  # the combo rule: its row is words, with the settings from config.yaml
            texts += [cell for cell in flag_cells(data, flag)]
        else:  # the name, and the threshold as Excel displays it (0.15 shows as 15.0%)
            texts += [flag["flag"], threshold_text(flag)]
    texts.append(runway_context_text(data))
    return texts


def workbook_numbers(data):
    """The numbers in workbook_texts, as analyze.numbers_in reads them (sign-aware, units dropped)."""
    return numbers_in("\n".join(workbook_texts(data)))


def unlisted_numbers(summary, data):
    """Numbers in the AI headline and questions (the parts the memo shows) that the metrics workbook doesn't show."""
    used = numbers_in(" ".join([summary.headline] + summary.questions))
    return sorted(used - workbook_numbers(data))


def memo_analysis(analysis_file, payload, data):
    """(BoardSummary, None) if the memo may show the AI text, else (None, why not).

    First every check the deck makes (build_deck.load_analysis), so the memo and the deck show the
    same AI text or neither does. Then one more: analyze.py allows any number in Claude's payload,
    which includes raw inputs (net burn, ending cash) the metrics workbook doesn't show, and every
    number in the memo must be one the workbook shows.
    """
    if analysis_file is None:
        return None, "no analysis requested"
    summary, why = load_analysis(analysis_file, payload)
    if summary is None:
        return None, why
    unlisted = unlisted_numbers(summary, data)
    if unlisted:
        shown = ", ".join(f"{number:,g}" for number in unlisted)
        return None, f"the AI text quotes numbers the metrics workbook doesn't show: {shown}"
    return summary, None


# ---------------------------------------------------------------------------
# 2. Text pieces (formatting only - no math)
# ---------------------------------------------------------------------------

def no_em_dash(text):
    """The memo's words for a label with an em dash (EM_DASH): a colon, e.g. 'Cannot evaluate: missing input'."""
    return text.replace(f" {EM_DASH} ", ": ").replace(EM_DASH, "-")


def flag_cells(data, flag):
    """One flag's row on the metrics workbook's Flags sheet (excel_output.flag_row): the same words."""
    return flag_row(flag, data["actuals"], data["metrics"], data["reasons"], data["config"])


def runway_context_text(data):
    """Runway at next quarter's budgeted burn: '13.0 mo', or the words for why there's no number."""
    value = runway_context_value(data["runway_at_budget"], data["next_budget"] is not None)
    return value if isinstance(value, str) else format_value("runway_months", value)


def combo_line(data, flag):
    """'NRR falling while pipeline rising: Passed (trips when NRR falls at least 1 pt ..., over the last 3 quarters)'.

    The rule's words are the Flags sheet's own ("Trips when" and "Threshold" cells), so every
    number in them is one the metrics workbook shows.
    """
    _, _, _, window, trips_when, status = flag_cells(data, flag)
    return f"{flag['flag']}: {status} (trips when {trips_when}, over the {window})"


def kpi_rows(data):
    """Key metrics table: [(cells, flag status or None)]. Context rows first, then every flag."""
    latest, prior = data["latest"], data["prior"]
    rows = []
    for metric, budget_metric in CONTEXT_ROWS:
        budget = f"vs budget: {value_text(data, budget_metric, latest)}" if budget_metric else NOT_A_FLAG
        cells = [METRIC_LABELS[metric], value_text(data, metric, latest), value_text(data, metric, prior),
                 budget, NOT_A_FLAG]
        rows.append(([no_em_dash(cell) for cell in cells], None))
    for flag in data["flags"]:
        if flag["metric"] is None:  # the combo rule has no single value
            cells = [flag["flag"], NOT_A_FLAG, NOT_A_FLAG, COMBO_CELL, status_label(flag)]
        else:
            cells = [flag["flag"], value_text(data, flag["metric"], latest), value_text(data, flag["metric"], prior),
                     threshold_text(flag), status_label(flag)]
        rows.append(([no_em_dash(cell) for cell in cells], flag["status"]))
    return rows


def flag_lines(data):
    """Tripped flags (value and threshold), then flags that can't be evaluated, then the combo rule."""
    metric_flags = [flag for flag in data["flags"] if flag["metric"] is not None]
    tripped = [f"{flag['flag']}: {value_text(data, flag['metric'], flag['quarter'])} ({threshold_text(flag)})"
               for flag in metric_flags if flag["status"] == TRIP]
    unevaluated = [f"{flag['flag']}: {status_label(flag)}" for flag in metric_flags
                   if flag["status"] == CANNOT_EVALUATE]
    combos = [combo_line(data, flag) for flag in data["flags"] if flag["metric"] is None]
    return [no_em_dash(line) for line in (tripped or [NO_TRIPPED_FLAGS]) + unevaluated + combos]


# ---------------------------------------------------------------------------
# 3. The memo as blocks: what both files show, in order
# ---------------------------------------------------------------------------

def heading(text):
    return {"kind": "heading", "text": text}


def text(words, style="body", ai=False):
    """A paragraph. style: "body", "note" (gray), "headline" (bold navy), "unavailable" (bold gray).

    ai=True marks the AI's slots (headline, questions): Claude's words, or "AI commentary unavailable"
    in their place. Everything else is computed, and is the same with or without the AI text.
    """
    return {"kind": "text", "text": words, "style": style, "ai": ai}


def bullets(items, ai=False):
    return {"kind": "bullets", "items": items, "ai": ai}


def ai_blocks(summary, part):
    """The AI headline or questions under the AI-drafted line, or the 'unavailable' text in their place."""
    if summary is None:
        return [text(MEMO_UNAVAILABLE, "unavailable", ai=True), text(MEMO_UNAVAILABLE_NOTE, "note", ai=True)]
    body = text(summary.headline, "headline", ai=True) if part == "headline" else bullets(summary.questions, ai=True)
    return [text(AI_DRAFTED_LINE, "note", ai=True), body]


def memo_blocks(data, summary):
    """Every block of the memo, top to bottom. summary = a validated BoardSummary, or None."""
    prior = f", compared with {data['prior']}" if data["prior"] else ""
    return [
        {"kind": "title", "text": f"{data['company']} board update: {data['latest']}"},
        text(f"Quarterly KPI update for {data['latest']}{prior}. Numbers are computed in Python from the "
             f"company's KPI workbook; the headline and questions are drafted by AI.", "note"),
        heading("Headline"),
        *ai_blocks(summary, "headline"),
        heading("Key metrics"),
        {"kind": "table", "header": kpi_header(data), "rows": kpi_rows(data)},
        text(f"{RUNWAY_CONTEXT}: {runway_context_text(data)}", "note"),
        heading(f"Flags: {flag_count_text(data['flags'])}"),
        bullets(flag_lines(data)),
        heading("Data gaps (data missing)"),
        bullets([no_em_dash(line) for line in gaps_lines(data["gaps"])]),
        heading(QUESTIONS_HEADING),
        *ai_blocks(summary, "questions"),
    ]


def block_texts(blocks):
    """Every piece of text in the blocks, in order: what a reader of either file sees (except the footer)."""
    texts = []
    for block in blocks:
        if block["kind"] == "bullets":
            texts += block["items"]
        elif block["kind"] == "table":
            texts += block["header"] + [cell for cells, _ in block["rows"] for cell in cells]
        else:
            texts.append(block["text"])
    return texts


def memo_footer(data, run_date, model, approval, commit=None):
    """'Fictional data | northwind.xlsx | 2026-09-17 | 2a215a9 | claude-sonnet-5 | AI-drafted | not reviewed'.

    The same parts, in the same order, as the deck's footer. model=None means no AI text.
    """
    return " | ".join([FICTIONAL_NOTE, data["source_name"], run_date.isoformat(), commit_text(commit),
                       model or NO_AI_MODEL, AI_DRAFTED, review_text(approval)])


# ---------------------------------------------------------------------------
# 4. Writing the Word file (python-docx)
# ---------------------------------------------------------------------------

def docx_run(paragraph, words, size, bold=False, color=DARK_GRAY):
    """Add one run of text in the memo's font."""
    run = paragraph.add_run(words)
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def docx_paragraph(document, words, size, bold=False, color=DARK_GRAY, space_before=0, style=None):
    """Add a paragraph with one run and the memo's spacing."""
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(SPACE_AFTER_TEXT)
    docx_run(paragraph, words, size, bold, color)
    return paragraph


def keep_with_next(paragraph):
    """Word moves this paragraph to the next page with the one after it, so a heading never ends a page."""
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def shade_cell(cell, fill_hex):
    """Fill a Word table cell with a color. python-docx has no setting for this, so it's the XML: <w:shd w:fill=...>."""
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill_hex)
    cell._tc.get_or_add_tcPr().append(shading)


def docx_cell(cell, words, fill_hex, text_hex, bold=False):
    """One table cell: fill color and a single run of text."""
    shade_cell(cell, fill_hex)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    docx_run(paragraph, words, TABLE_SIZE, bold, text_hex)


def docx_table(document, block):
    """The key metrics table: navy header, striped rows, each flag's status cell in its status color."""
    table = document.add_table(rows=len(block["rows"]) + 1, cols=len(block["header"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    width = LETTER[0] / inch - 2 * PAGE_MARGIN
    for column, share in enumerate(KPI_COLUMN_SHARES):
        for cell in table.columns[column].cells:
            cell.width = DocxInches(width * share)
    for column, words in enumerate(block["header"]):
        docx_cell(table.cell(0, column), words, NAVY, WHITE, bold=True)
    for row_number, (cells, status) in enumerate(block["rows"], start=1):
        stripe = LIGHT_GRAY if row_number % 2 == 0 else WHITE
        for column, words in enumerate(cells):
            fill_hex, text_hex = stripe, DARK_GRAY
            if status is not None and column == len(cells) - 1:
                fill_hex, text_hex = STATUS_COLORS[status]
            docx_cell(table.cell(row_number, column), words, fill_hex, text_hex)


def docx_block(document, block):
    """Write one block into the Word document."""
    kind = block["kind"]
    if kind == "title":
        docx_paragraph(document, block["text"], TITLE_SIZE, bold=True, color=NAVY)
    elif kind == "heading":
        keep_with_next(docx_paragraph(document, block["text"], HEADING_SIZE, bold=True, color=NAVY,
                                      space_before=SPACE_BEFORE_HEADING))
    elif kind == "table":
        docx_table(document, block)
    elif kind == "bullets":  # a list stays on one page: every item but the last keeps with the next
        for number, item in enumerate(block["items"], start=1):
            paragraph = docx_paragraph(document, item, BODY_SIZE, style="List Bullet")
            if number < len(block["items"]):
                keep_with_next(paragraph)
    else:
        size, bold, color = TEXT_STYLES[block["style"]]
        paragraph = docx_paragraph(document, block["text"], size, bold=bold, color=color)
        if block["style"] == "note" and block["ai"]:  # "AI-drafted ..." stays with the AI text under it
            keep_with_next(paragraph)


def write_docx(blocks, footer, path):
    """Save the memo as a Word file: US Letter, the memo's margins, the footer on every page."""
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = DocxInches(LETTER[0] / inch), DocxInches(LETTER[1] / inch)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, DocxInches(PAGE_MARGIN))
    for block in blocks:
        docx_block(document, block)
    footer_paragraph = section.footer.paragraphs[0]
    docx_run(footer_paragraph, footer, FOOTER_SIZE, color=MID_GRAY)
    document.save(path)
    return Path(path)


# ---------------------------------------------------------------------------
# 5. Writing the PDF (reportlab)
# ---------------------------------------------------------------------------

def register_pdf_fonts():
    """Tell reportlab where DejaVu Sans is (regular and bold). Safe to call more than once."""
    for name, weight in ((PDF_FONT, "normal"), (PDF_BOLD_FONT, "bold")):
        if name not in pdfmetrics.getRegisteredFontNames():
            path = font_manager.findfont(font_manager.FontProperties(family="DejaVu Sans", weight=weight))
            pdfmetrics.registerFont(TTFont(name, path))


def pdf_color(hex_text):
    return colors.HexColor("#" + hex_text)


def pdf_style(size, bold=False, color=DARK_GRAY, space_before=0):
    """A reportlab paragraph style in the memo's font."""
    return ParagraphStyle("memo", fontName=PDF_BOLD_FONT if bold else PDF_FONT, fontSize=size,
                          leading=size * 1.25, textColor=pdf_color(color), spaceBefore=space_before,
                          spaceAfter=SPACE_AFTER_TEXT)


def pdf_paragraph(words, style):
    """reportlab reads a little HTML inside a paragraph, so &, < and > are escaped first."""
    return Paragraph(escape(words), style)


def pdf_table(block):
    """The key metrics table, in the same colors as the Word file and the deck."""
    width = LETTER[0] - 2 * PAGE_MARGIN * inch
    header_style = pdf_style(TABLE_SIZE, bold=True, color=WHITE)
    rows = [[pdf_paragraph(words, header_style) for words in block["header"]]]
    commands = [("BACKGROUND", (0, 0), (-1, 0), pdf_color(NAVY)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
    for row_number, (cells, status) in enumerate(block["rows"], start=1):
        stripe = LIGHT_GRAY if row_number % 2 == 0 else WHITE
        commands.append(("BACKGROUND", (0, row_number), (-1, row_number), pdf_color(stripe)))
        text_colors = [DARK_GRAY] * len(cells)
        if status is not None:
            fill_hex, text_colors[-1] = STATUS_COLORS[status]
            commands.append(("BACKGROUND", (-1, row_number), (-1, row_number), pdf_color(fill_hex)))
        rows.append([pdf_paragraph(words, pdf_style(TABLE_SIZE, color=colour))
                     for words, colour in zip(cells, text_colors)])
    table = Table(rows, colWidths=[width * share for share in KPI_COLUMN_SHARES])
    table.setStyle(TableStyle(commands))
    return table


def pdf_flowables(block):
    """One block as reportlab 'flowables' (the pieces it lays out down the page)."""
    kind = block["kind"]
    if kind == "title":
        return [pdf_paragraph(block["text"], pdf_style(TITLE_SIZE, bold=True, color=NAVY))]
    if kind == "heading":
        return [pdf_paragraph(block["text"], pdf_style(HEADING_SIZE, bold=True, color=NAVY,
                                                       space_before=SPACE_BEFORE_HEADING))]
    if kind == "table":
        return [pdf_table(block), Spacer(0, SPACE_AFTER_TEXT)]
    if kind == "bullets":
        items = [ListItem(pdf_paragraph(item, pdf_style(BODY_SIZE))) for item in block["items"]]
        return [ListFlowable(items, bulletType="bullet", bulletFontName=PDF_FONT, bulletFontSize=BODY_SIZE)]
    size, bold, color = TEXT_STYLES[block["style"]]
    return [pdf_paragraph(block["text"], pdf_style(size, bold=bold, color=color))]


def write_pdf(blocks, footer, path):
    """Save the memo as a PDF: US Letter, the memo's margins, the footer drawn at the foot of every page."""
    register_pdf_fonts()
    margin = PAGE_MARGIN * inch

    def draw_footer(canvas, document):
        """Called by reportlab on every page. A paragraph, so a long file name wraps instead of running off the page."""
        line = pdf_paragraph(footer, pdf_style(FOOTER_SIZE, color=MID_GRAY))
        _, height = line.wrap(document.width, margin)
        line.drawOn(canvas, margin, (margin - height) / 2)

    document = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=margin, rightMargin=margin,
                                 topMargin=margin, bottomMargin=margin, title=blocks[0]["text"])
    story = [KeepTogether(section) for section in pdf_sections(blocks)]
    document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
    return Path(path)


def pdf_sections(blocks):
    """The memo's flowables grouped by section (a heading and everything up to the next one).

    Each group is kept on one page when it fits, so a heading is never left at the foot of a page
    with its text on the next.
    """
    sections = []
    for block in blocks:
        if block["kind"] in ("title", "heading") or not sections:
            sections.append([])
        sections[-1] += pdf_flowables(block)
    return sections


# ---------------------------------------------------------------------------
# 6. Putting it together
# ---------------------------------------------------------------------------

def memo_paths(workbook_path, output_dir=OUTPUT_DIR):
    """data/northwind.xlsx -> (output/northwind_board_memo.docx, output/northwind_board_memo.pdf)"""
    stem = Path(output_dir) / f"{Path(workbook_path).stem}{MEMO_SUFFIX}"
    return stem.with_suffix(".docx"), stem.with_suffix(".pdf")


def save_memo(workbook_path, config, analysis_file=None, run_date=None, output_dir=OUTPUT_DIR):
    """Clean one workbook, build its memo and save both files.

    Returns {"docx": path, "pdf": path, "why_unavailable": why the AI text isn't in it, or None}.
    Old memo files are deleted first, so a failed build never leaves last run's memo looking current.
    The review status comes from the manifest, exactly as for the deck (provenance.approval_status).
    """
    workbook_path, output_dir = Path(workbook_path), Path(output_dir)
    docx_path, pdf_path = memo_paths(workbook_path, output_dir)
    for path in (docx_path, pdf_path):
        path.unlink(missing_ok=True)

    company = workbook_path.stem.title()   # same rule as analyze.py, build_deck.py and main.py
    actuals, next_budget = clean_workbook(workbook_path)
    data = collect_deck_data(company, workbook_path.name, actuals, next_budget, config)
    summary, why_unavailable = memo_analysis(analysis_file, build_payload(company, actuals, next_budget, config), data)

    approval, _ = approval_status(read_manifest(manifest_path(workbook_path, output_dir)),
                                  file_sha256(workbook_path), file_sha256(CONFIG_PATH))
    details = analysis_details(analysis_file) if summary else None
    footer = memo_footer(data, run_date or datetime.date.today(), details["model"] if details else None,
                         memo_approval(approval), git_commit())

    output_dir.mkdir(parents=True, exist_ok=True)
    blocks = memo_blocks(data, summary)
    write_docx(blocks, footer, docx_path)
    write_pdf(blocks, footer, pdf_path)
    record_memo_status(workbook_path, output_dir, [docx_path.name, pdf_path.name], ai_text=summary is not None)
    return {"docx": docx_path, "pdf": pdf_path, "why_unavailable": why_unavailable}


def memo_approval(approval):
    """The approval, if it covers the memo (approve.py lists "memo" in its documents), else None.

    provenance.approval_status decides whether an approval still counts for these inputs; this adds
    one more question: did the reviewer have a memo in front of them? An approval recorded before
    memos existed approved a deck only, so the memo's footer says "not reviewed".
    """
    return approval if approval and "memo" in approval.get("documents", []) else None


def record_memo_status(workbook_path, output_dir, files, ai_text):
    """Update the manifest to describe the memo just written: its files, and AI text or not.

    Only if a manifest is already there (main.py writes the full one), so a memo rebuilt on its own
    never leaves the manifest describing the last one. The same rule as build_deck.record_deck_status.
    """
    path = manifest_path(workbook_path, output_dir)
    manifest = read_manifest(path)
    if manifest is None:
        return None
    manifest["memo"] = {"files": files, "ai_text": ai_text}
    return save_manifest(path, manifest)


def main(argv=None, output_dir=OUTPUT_DIR):
    parser = argparse.ArgumentParser(description="Build the board memo (Word and PDF) for one KPI workbook.")
    parser.add_argument("workbook", help="path to a KPI workbook, e.g. data/northwind.xlsx")
    parser.add_argument("--analysis", help="analysis JSON (default: output/<company>_analysis.json)")
    parser.add_argument("--no-analysis", action="store_true", help="build with 'AI commentary unavailable'")
    args = parser.parse_args(argv)

    analysis_file = None if args.no_analysis else (args.analysis or analysis_path(args.workbook, output_dir))
    result = save_memo(args.workbook, load_config(), analysis_file, output_dir=output_dir)
    project = Path(__file__).parent
    docx = result["docx"]
    print(f"Saved {docx.relative_to(project) if docx.is_relative_to(project) else docx} and {result['pdf'].name}")
    if result["why_unavailable"]:
        print(f"{MEMO_UNAVAILABLE}: {result['why_unavailable']}")


if __name__ == "__main__":
    main()
