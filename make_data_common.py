"""Shared helpers for the make_data scripts (one script per fictional company).

Each company script (make_data.py, make_data_alderpeak.py, make_data_fernhollow.py)
holds only its own numbers and mess settings, then calls save_workbook().

How a workbook is built:
1. check_true_data() confirms the clean numbers tie out (ARR and cash roll forward).
2. The KPI tab is written with mess added on purpose: odd header names, some
   numbers stored as text ("$14.3M", "120K", "5,090"), an optional blank quarter,
   and a budget-only row for next quarter.
3. A junk Notes tab is added, the kind of thing real workbooks carry around.

All money figures are $K.
"""

from openpyxl import Workbook

# Columns that are counts, not money - skipped by the multiple-of-10 check.
COUNT_COLUMNS = {"new_customers", "headcount"}


# ---------------------------------------------------------------------------
# Checking the answer key
# ---------------------------------------------------------------------------

def ending_arr_at(true_data, i):
    """Ending ARR for quarter number i = starting + new + expansion - contraction - churn."""
    d = true_data
    return (d["starting_arr"][i] + d["new_arr"][i] + d["expansion_arr"][i]
            - d["contraction_arr"][i] - d["churned_arr"][i])


def check_true_data(company):
    """Stop with an error if the clean data doesn't tie out or the mess settings are wrong."""
    quarters, true_data = company["quarters"], company["true_data"]
    blank_quarter, text_cells = company["blank_quarter"], company["text_cells"]
    header_names = company["header_names"]

    # Every column must have exactly one value per quarter, and a header to write.
    for column, values in true_data.items():
        assert len(values) == len(quarters), f"{column} needs {len(quarters)} values"
    assert set(header_names) == set(true_data), "HEADER_NAMES must cover exactly the TRUE_DATA columns"

    # ARR roll-forward: this quarter's ending ARR must equal next quarter's starting ARR.
    for i in range(len(quarters) - 1):
        assert ending_arr_at(true_data, i) == true_data["starting_arr"][i + 1], \
            f"ARR doesn't roll forward after {quarters[i]}"

    # Cash roll-forward: this quarter's cash = last quarter's cash - this quarter's burn.
    cash, burn = true_data["ending_cash"], true_data["net_burn"]
    for i in range(1, len(quarters)):
        assert cash[i] == cash[i - 1] - burn[i], f"Cash doesn't roll forward into {quarters[i]}"

    # Money values must be multiples of 10 so "$21.85M"-style text is exact, not rounded.
    for column, values in true_data.items():
        if column not in COUNT_COLUMNS:
            assert all(v % 10 == 0 for v in values), f"{column} has a value not divisible by 10"

    # The blank quarter (if any) must be a real quarter.
    assert blank_quarter is None or blank_quarter in quarters, f"Unknown blank quarter: {blank_quarter}"

    # Text cells must point at a real column and a quarter that actually gets written.
    for column, quarter in text_cells:
        assert column in true_data, f"Unknown column in TEXT_CELLS: {column}"
        assert quarter in quarters and quarter != blank_quarter, f"Bad quarter in TEXT_CELLS: {quarter}"


# ---------------------------------------------------------------------------
# Writing the workbook
# ---------------------------------------------------------------------------

def to_text(value, style):
    """Turn a $K number into messy text, e.g. 14300 -> "$14.3M"."""
    if style == "M":
        # 14300 $K / 1000 = 14.3 $M; ":g" drops trailing zeros (6.0 -> 6).
        return f"${value / 1000:g}M"
    if style == "K":
        return f"{value}K"
    if style == "comma":
        return f"{value:,}"
    raise ValueError(f"Unknown text style: {style}")


def quarter_row(i, quarter, columns, true_data, text_cells):
    """One actual quarter's row: the label, then each value as a number or messy text."""
    row = [quarter]
    for column in columns:
        value = true_data[column][i]
        style = text_cells.get((column, quarter))  # None if this cell stays a number
        row.append(to_text(value, style) if style else value)
    return row


def write_kpi_sheet(ws, company):
    """Write the KPI table: optional title, header row, the quarters, then the budget-only row."""
    # Column order follows HEADER_NAMES, so a company can list its columns in its own order.
    columns = list(company["header_names"])

    if company.get("title"):
        ws.append([company["title"]])  # a title line above the table
        ws.append([])                  # and an empty row under it

    ws.append(["Quarter"] + [company["header_names"][c] for c in columns])

    for i, quarter in enumerate(company["quarters"]):
        if quarter == company["blank_quarter"]:
            ws.append([quarter])  # label only, every value left empty
        else:
            ws.append(quarter_row(i, quarter, columns, company["true_data"], company["text_cells"]))

    # Budget-only row: blank for actuals, filled for the 3 budget columns.
    next_budget = company["next_quarter_budget"]
    ws.append([company["budget_only_label"]] + [next_budget.get(c) for c in columns])

    # Widen columns so the sheet is readable when opened in Excel.
    for column_cells in ws.columns:
        ws.column_dimensions[column_cells[0].column_letter].width = 18


def write_notes_sheet(ws, notes):
    """Write the junk notes tab. `notes` maps a cell address to its content, e.g. {"A1": "scratch"}."""
    for cell, value in notes.items():
        ws[cell] = value


def save_workbook(company):
    """Check the answer key, then write the messy workbook to company["output_path"].

    `company` is a dictionary built by each make_data script (see make_data.py for the keys).
    """
    check_true_data(company)

    # A new workbook starts with one sheet; rename it and add the notes tab.
    wb = Workbook()
    kpi_sheet = wb.active
    kpi_sheet.title = "KPI Tracker"
    write_kpi_sheet(kpi_sheet, company)
    write_notes_sheet(wb.create_sheet("Notes"), company["notes"])
    if company.get("notes_first"):
        wb.move_sheet("Notes", offset=-1)  # put the junk tab in front of the KPI tab

    output_path = company["output_path"]
    output_path.parent.mkdir(exist_ok=True)
    wb.save(output_path)
    blank = company["blank_quarter"] or "none"
    print(f"Wrote {output_path.name}: {len(company['quarters'])} quarters (blank: {blank}), "
          f"1 budget-only row, {len(company['text_cells'])} text cells, sheets: {wb.sheetnames}")
