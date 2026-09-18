"""Golden files (final Task 8): approved text dumps of every company's deck, memo and metrics workbook.

A golden file is a copy of an output that a person has read and approved. The test rebuilds the
output, turns it into text the same way, and fails if anything differs, showing the changed lines.

Why text, not the files themselves: a .pptx, .docx or .xlsx is a zip whose bytes change on every
save (timestamps inside), so comparing bytes would fail every time and say nothing useful. The
dump keeps what a reader sees: every word, its size, weight and color, where each box sits, table
fills, Excel number formats, the memo's page-break rules and the PDF's text page by page.

Two things change from run to run and would make every golden fail for no reason, so they're fixed
here: the run date (RUN_DATE) and the git commit in the footer (COMMIT). The AI text comes from
the saved analyses in tests/golden/analysis/, so nothing calls the API.

    python golden.py            rebuild and compare; prints any differences, exit 1 if there are any
    python golden.py --update   rebuild and overwrite the goldens: only after you've read the change
"""

import argparse
import datetime
import difflib
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pypdf import PdfReader

import build_deck
import memo
from build_deck import save_deck
from excel_output import save_metrics_workbook
from memo import save_memo
from metrics import load_config

PROJECT_DIR = Path(__file__).parent
DATA_DIR = PROJECT_DIR / "data"
GOLDEN_DIR = PROJECT_DIR / "tests" / "golden"
FIXTURE_DIR = GOLDEN_DIR / "analysis"      # copies of output/<company>_analysis.json from the live run
KINDS = ("deck", "memo", "metrics")

RUN_DATE = datetime.date(2026, 7, 15)      # a fixed footer date: a real one would change the goldens every day
COMMIT = {"commit": "0000000", "uncommitted_changes": False}   # a fixed commit, for the same reason
UPDATE_COMMAND = "python golden.py --update"
EMU_PER_INCH = 914400
PICTURE = 13                               # MSO_SHAPE_TYPE.PICTURE


# ---------------------------------------------------------------------------
# Where things are
# ---------------------------------------------------------------------------

def companies():
    """Every company with a workbook in data/, by file name: ["alderpeak", "fernhollow", "northwind"]."""
    return sorted(path.stem for path in DATA_DIR.glob("*.xlsx") if not path.name.startswith("~$"))


def golden_path(company, kind, folder=GOLDEN_DIR):
    """("northwind", "deck") -> tests/golden/northwind_deck.txt"""
    return Path(folder) / f"{company}_{kind}.txt"


def fixture_path(company):
    """The saved analysis the goldens are built with: tests/golden/analysis/northwind_analysis.json"""
    return FIXTURE_DIR / f"{company}_analysis.json"


# ---------------------------------------------------------------------------
# Deck -> text
# ---------------------------------------------------------------------------

def inches(emu):
    """PowerPoint's unit (EMU) -> inches, 2 decimals: '0.50'."""
    return f"{emu / EMU_PER_INCH:.2f}"


def run_style(run):
    """'15 pt bold 1F2A44': a run's size, weight and color, leaving out whatever it inherits."""
    parts = [f"{run.font.size.pt:g} pt"] if run.font.size else []
    if run.font.bold:
        parts.append("bold")
    if run.font.color.type is not None:
        parts.append(str(run.font.color.rgb))
    return " ".join(parts)


def text_lines(frame, indent="    "):
    """Each paragraph of a text frame on its own line, with the style of its first run."""
    lines = []
    for paragraph in frame.paragraphs:
        style = run_style(paragraph.runs[0]) if paragraph.runs else ""
        lines.append(f"{indent}{paragraph.text!r}  {style}".rstrip())
    return lines


def cell_fill(cell):
    """A table cell's solid fill as hex, or '' when it has none."""
    return f"fill {cell.fill.fore_color.rgb}" if cell.fill.type == MSO_FILL_TYPE.SOLID else ""


def table_lines(table):
    """Each table row: every cell's text, style and fill."""
    lines = []
    for number, row in enumerate(table.rows, start=1):
        cells = [f"{cell.text!r} {run_style(cell.text_frame.paragraphs[0].runs[0])} {cell_fill(cell)}".rstrip()
                 if cell.text_frame.paragraphs[0].runs else f"'' {cell_fill(cell)}".rstrip() for cell in row.cells]
        lines.append(f"    row {number}: " + " | ".join(cells))
    return lines


def shape_lines(shape):
    """A shape's name, kind, position and size, then its text or table."""
    kind = "picture" if shape.shape_type == PICTURE else "table" if shape.has_table else "text"
    lines = [f"  [{shape.name}] {kind} at {inches(shape.left)}, {inches(shape.top)} in, "
             f"size {inches(shape.width)} x {inches(shape.height)} in"]
    if shape.has_table:
        lines += table_lines(shape.table)
    elif shape.has_text_frame:
        lines += text_lines(shape.text_frame)
    return lines


def dump_deck(path):
    """Every slide of a deck as text, shape by shape in the order they're drawn."""
    slides = Presentation(path).slides
    lines = []
    for number, slide in enumerate(slides, start=1):
        lines.append(f"== Slide {number} of {len(slides)} (layout: {slide.slide_layout.name}) ==")
        for shape in slide.shapes:
            lines += shape_lines(shape)
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Memo -> text (the Word file, its footer, and the PDF's text)
# ---------------------------------------------------------------------------

def docx_style(paragraph):
    """'List Bullet, 11 pt, keep with next': what decides how a Word paragraph looks and where pages break."""
    parts = [] if paragraph.style.name == "Normal" else [paragraph.style.name]
    if paragraph.runs:
        font = paragraph.runs[0].font
        parts += [f"{font.size.pt:g} pt"] if font.size else []
        parts += ["bold"] if font.bold else []
        parts += [str(font.color.rgb)] if font.color.type is not None else []
    if paragraph.paragraph_format.keep_with_next:
        parts.append("keep with next")
    return ", ".join(parts)


def docx_cell_fill(cell):
    """A Word table cell's fill (memo.shade_cell writes it as <w:shd w:fill=...>), or ''."""
    properties = cell._tc.tcPr
    shading = properties.find(qn("w:shd")) if properties is not None else None
    return f" fill {shading.get(qn('w:fill'))}" if shading is not None else ""


def docx_lines(document):
    """The Word file's body in order: paragraphs with their style, tables row by row."""
    lines = []
    for item in document.element.body.iterchildren():
        if item.tag == qn("w:p"):
            paragraph = next(p for p in document.paragraphs if p._p is item)
            lines.append(f"  {paragraph.text!r}  {docx_style(paragraph)}".rstrip())
        elif item.tag == qn("w:tbl"):
            table = next(t for t in document.tables if t._tbl is item)
            lines.append(f"  table, aligned {table.alignment.name.lower() if table.alignment is not None else 'left'}")
            for number, row in enumerate(table.rows, start=1):
                lines.append(f"  row {number}: " + " | ".join(f"{cell.text!r}{docx_cell_fill(cell)}"
                                                             for cell in row.cells))
    return lines


def dump_memo(docx_path, pdf_path):
    """The memo as text: the Word file's page setup, body and footer, then each PDF page's text."""
    document = Document(docx_path)
    section = document.sections[0]
    lines = ["== Word file ==",
             f"  page {section.page_width.inches:.2f} x {section.page_height.inches:.2f} in, margins "
             f"{section.left_margin.inches:.2f} in"]
    lines += docx_lines(document)
    lines.append("-- Footer --")
    lines += [f"  {paragraph.text!r}  {docx_style(paragraph)}" for paragraph in section.footer.paragraphs]
    for number, page in enumerate(PdfReader(pdf_path).pages, start=1):
        lines.append(f"== PDF page {number} ==")
        lines += [f"  {line.rstrip()}" for line in page.extract_text().splitlines()]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Metrics workbook -> text
# ---------------------------------------------------------------------------

def cell_value_text(value):
    """A cell's value: text in quotes, a decimal to 12 significant digits (the last bits of a float are noise)."""
    if isinstance(value, float):
        return f"{value:.12g}"
    return repr(value) if isinstance(value, str) else str(value)


def excel_cell_line(cell):
    """'B2  0.971  [0.0%]  fill FFC7CE  bold  right': value, number format, fill, weight, font color, alignment."""
    parts = [cell.coordinate, cell_value_text(cell.value)]
    if cell.number_format != "General":
        parts.append(f"[{cell.number_format}]")
    if cell.fill.fill_type == "solid":
        parts.append(f"fill {cell.fill.fgColor.rgb[-6:]}")
    if cell.font.b:
        parts.append("bold")
    if cell.font.color is not None and cell.font.color.type == "rgb":
        parts.append(f"text {cell.font.color.rgb[-6:]}")
    if cell.alignment.horizontal:
        parts.append(cell.alignment.horizontal)
    return "  ".join(parts)


def dump_workbook(path):
    """Every sheet: column widths, frozen panes, and every filled cell."""
    lines = []
    for sheet in load_workbook(path).worksheets:
        widths = ", ".join(f"{letter} {dimension.width:g}" for letter, dimension in sorted(sheet.column_dimensions.items())
                           if dimension.width)
        lines += [f"== Sheet {sheet.title} ==", f"  column widths: {widths or 'default'}",
                  f"  frozen at: {sheet.freeze_panes or 'none'}"]
        lines += [f"  {excel_cell_line(cell)}" for row in sheet.iter_rows() for cell in row if cell.value is not None]
    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Building, comparing, updating
# ---------------------------------------------------------------------------

def build_dumps(company, folder):
    """Build one company's deck, memo and metrics workbook into `folder` and dump each: {kind: text}.

    The commit is patched only while building: build_deck and memo each look up git_commit when
    they write the footer, and a real one changes with every commit.
    """
    workbook, analysis, config = DATA_DIR / f"{company}.xlsx", fixture_path(company), load_config()
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    with patch.object(build_deck, "git_commit", lambda: COMMIT), patch.object(memo, "git_commit", lambda: COMMIT):
        deck, _ = save_deck(workbook, config, analysis, run_date=RUN_DATE, output_dir=folder)
        memo_files = save_memo(workbook, config, analysis, run_date=RUN_DATE, output_dir=folder)
    metrics = save_metrics_workbook(workbook, config, output_dir=folder)
    return {"deck": dump_deck(deck), "memo": dump_memo(memo_files["docx"], memo_files["pdf"]),
            "metrics": dump_workbook(metrics)}


def build_all(folder):
    """Every company's dumps, each built in its own subfolder: {company: {kind: text}}."""
    return {company: build_dumps(company, Path(folder) / company) for company in companies()}


def compare(actual, path):
    """None if the text matches the golden file, else a message with the changed lines and how to accept them."""
    path = Path(path)
    if not path.exists():
        return f"{path.name}: no golden file yet. To create it: {UPDATE_COMMAND}"
    expected = path.read_text()
    if actual == expected:
        return None
    diff = difflib.unified_diff(expected.splitlines(), actual.splitlines(), lineterm="",
                                fromfile=f"{path.name} (approved)", tofile=f"{path.name} (rebuilt now)")
    return (f"{path.name} differs from its approved golden ('-' approved, '+' now):\n" + "\n".join(diff) +
            f"\nIf this change is intended: {UPDATE_COMMAND}, read `git diff tests/golden`, and commit the goldens.")


def write_goldens(dumps, folder=GOLDEN_DIR):
    """Write every dump whose text changed, delete goldens for outputs that no longer exist; return those names."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    changed = []
    for company, kinds in dumps.items():
        for kind, text in kinds.items():
            path = golden_path(company, kind, folder)
            if not path.exists() or path.read_text() != text:
                path.write_text(text)
                changed.append(path.name)
    expected = {golden_path(company, kind, folder).name for company, kinds in dumps.items() for kind in kinds}
    for path in folder.glob("*.txt"):
        if path.name not in expected:
            path.unlink()
            changed.append(path.name)
    return sorted(changed)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Compare every output with its approved golden text dump.")
    parser.add_argument("--update", action="store_true", help="overwrite the goldens with what the code makes now")
    args = parser.parse_args(argv)

    with tempfile.TemporaryDirectory() as folder:
        dumps = build_all(folder)
    if args.update:
        changed = write_goldens(dumps)
        print(f"Updated: {', '.join(changed)}" if changed else "Nothing changed.")
        print("Read `git diff tests/golden` before committing: every changed line is now the approved output.")
        return 0
    problems = [problem for company, kinds in dumps.items() for kind, text in kinds.items()
                if (problem := compare(text, golden_path(company, kind)))]
    total = sum(len(kinds) for kinds in dumps.values())
    print("\n\n".join(problems + [f"{total - len(problems)} of {total} match their goldens"]))
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
