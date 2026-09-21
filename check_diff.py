"""Automated proof for Task 10: what changed since the last run, against each company's answer key.

For each company, "last quarter's workbook" is written from its make_data script without its
latest quarter (Q1 2026 is then the latest, and the budget-only row is Q2 2026's budget). That
workbook is run first, then today's, into a temporary output folder, as a quarter apart would be.
For each company this checks:
1. The first run has nothing to compare with: the manifest says so and the memo has no section.
2. Today's run is compared with the Q1 2026 run: the flags that flipped, the metrics that moved
   (with both values and the size of the move) and the new and resolved data gaps equal the
   lists typed below, worked out by hand from the answer key (the formulas are in the comments).
3. The memo (Word and PDF) has the section after the headline, with exactly those lines.
4. The web page's comparison (portfolio.run_changes) is the memo's, line for line.
5. Approve, then rebuild from the same numbers: the comparison is still with Q1 2026, not "nothing
   changed" against the run ten seconds earlier; the memo's footer now says reviewed.
Plus: a first run of today's workbook alone shows no section, and python diff_runs.py prints the
same lines as the memo.

Expected words are typed out here, not imported from diff_runs.py. No API calls: every run skips
the AI step. The real output/ folder is never touched.
Run: python check_diff.py  -> prints "All checks passed" or stops at the first failure.
"""

import contextlib
import io
import tempfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader

import make_data as northwind
import make_data_alderpeak as alderpeak
import make_data_fernhollow as fernhollow
from approve import approve
from diff_runs import main as diff_runs_main
from main import load_config, run_company
from make_data_common import save_workbook
from portfolio import load_company, run_changes
from provenance import manifest_path, read_manifest

PROJECT_DIR = Path(__file__).parent
HEADING = "What changed since the last run"
COMPARED = "Compared with the run of {when}, whose latest quarter was Q1 2026 (now Q2 2026)."
MOVED_TITLE = "Metrics that moved more than 5.0 pts (percentages) or 10.0% (other metrics)"
REVIEWER = "Check Reviewer"
SHORT_TITLES = {"Flags that flipped": "flags flipped", MOVED_TITLE: "metrics moved",
                "New data gaps": "new gap line(s)", "Resolved data gaps": "resolved gap line(s)"}

# ---------------------------------------------------------------------------
# The expected changes, Q1 2026 run -> Q2 2026 run, by hand from each answer key.
# Moves: percentages by points (more than 5.0), everything else by % of the old value (more than 10%).
# ---------------------------------------------------------------------------

NORTHWIND = {
    "Flags that flipped": [
        "NRR (annualized): Tripped (was Passed)",              # 1 + 4*(710-200-390)/23790 = 102.0% -> 97.1%
        "Burn multiple: Tripped (was Passed)",                  # 3650/2020 = 1.81x -> 3900/1660 = 2.35x
        "Runway at current burn: Tripped (was Passed)",         # 18200/(3650/3) = 15.0 -> 14300/(3900/3) = 11.0 mo
        "Rule of 40: Tripped (was Cannot evaluate: missing input)",   # Q1 2026's YoY needs blank Q1 2025
        # NRR steps Q3 2025 -> Q1 2026: 108.9% -> 108.0% (-0.9 pts, under 1 point) -> 102.0%: not falling;
        # Q4 2025 -> Q2 2026: -6.0, -4.9 pts with pipeline up both steps: trips.
        "NRR falling while pipeline rising: Tripped (was Passed)",
    ],
    MOVED_TITLE: [
        "Net new ARR ($K): 2,020 to 1,660 (down 17.8%)",       # 1900+710-200-390 -> 1850+580-260-510; 360/2020
        "ARR growth YoY: data missing to 42.8%",                # 27470/19230 - 1; Q1 2026's needs blank Q1 2025
        "Revenue growth YoY: data missing to 46.4%",            # 6660/4550 - 1
        "Pipeline ($K): 11,200 to 12,500 (up 11.6%)",           # 1300/11200
        "Rule of 40: data missing to -12.2%",                   # 46.4% - 3900/6660
        "Burn multiple: 1.81x to 2.35x (up 30.0%)",             # 2.349/1.807 - 1
        "Net new ARR vs budget: -1.5% to -19.0% (down 17.6 pts)",   # 2020/2050-1 -> 1660/2050-1: 0.1756
        "Runway at current burn: 15.0 mo to 11.0 mo (down 26.5%)",  # 11.0/14.96 - 1
    ],
    # Not listed (under the setting): ending ARR +6.4%, net burn +6.8% (3650 -> 3900), NRR -4.9 pts, GRR -2.0 pts,
    # CAC payback +4.8%, ...
    "Resolved data gaps": ["Q1 2026: Flag: Rule of 40"],   # Q1 2026's Rule of 40 flag lacked Q1 2025 revenue
}

ALDERPEAK = {
    # No flag flips (healthy both quarters); no gap opens or closes (no blank quarter).
    MOVED_TITLE: [
        "Net burn ($K): 250 to 200 (down 20.0%)",               # 200/250 - 1; net burn is a metric row since the memo fix
        "Burn multiple: 0.16x to 0.12x (down 24.9%)",           # 250/1530 = 0.163 -> 200/1630 = 0.123
        "Runway at current burn: 88.8 mo to 108.0 mo (up 21.6%)",   # 7400/(250/3) -> 7200/(200/3)
    ],
    # Not listed: net new ARR +6.5%, ending ARR +9.7%, pipeline +7.0%, CAC payback -0.1%, FCF margin +1.8 pts, ...
}

FERNHOLLOW = {
    "Flags that flipped": ["Rule of 40: Cannot evaluate: missing input (was Tripped)"],   # Q2 2026's YoY needs blank Q2 2025
    MOVED_TITLE: [
        "Net new ARR ($K): -110 to -240 (down 118.2%)",         # 240+70-130-290 -> 180+60-150-330; 130/110
        "ARR growth YoY: 4.4% to data missing",                 # 7590/7270 - 1; Q2 2026's needs blank Q2 2025
        "Revenue growth YoY: 7.7% to data missing",             # 1820/1690 - 1
        "Pipeline ($K): 2,800 to 2,500 (down 10.7%)",           # 300/2800
        "FCF margin: -85.7% to -93.2% (down 7.5 pts)",          # -1560/1820 -> -1650/1770
        "Rule of 40: -78.0% to data missing",                   # 7.7% - 85.7%
        "Net new ARR vs budget: -122.9% to -150.0% (down 27.1 pts)",   # -110/480-1 -> -240/480-1
        "Ending ARR vs budget: -18.8% to -25.2% (down 6.4 pts)",   # 7590/9350-1 -> 7350/9830-1
        "CAC payback: 96.2 mo to 134.4 mo (up 39.7%)",          # 1120/(240*1060/1820)*12 -> 1150/(180*1010/1770)*12
        "Runway at current burn: 9.5 mo to 6.0 mo (down 37.0%)",   # 4950/(1560/3) -> 3300/(1650/3)
    ],
    # Not listed: NRR -3.9 pts, GRR -3.5 pts, net burn +5.8% (1560 -> 1650), burn multiple ∞ both runs
    # ("ARR shrank"), ending ARR -3.2%, ...
    "New data gaps": ["Q2 2026: ARR growth YoY, Revenue growth YoY, Rule of 40, Flag: Rule of 40"],
}

COMPANIES = [("Northwind", northwind, NORTHWIND), ("Alderpeak", alderpeak, ALDERPEAK),
             ("Fernhollow", fernhollow, FERNHOLLOW)]


# ---------------------------------------------------------------------------
# Last quarter's workbook
# ---------------------------------------------------------------------------

def earlier_budget_label(label, dropped):
    """The company's own budget-row wording, a quarter earlier: 'Q3 2026 (Budget)' -> 'Q2 2026 (Budget)'.

    The label starts with the quarter after the latest actual one ("Q3 2026", two words); with
    the latest quarter dropped, that's the dropped quarter.
    """
    next_quarter = " ".join(label.split(" ")[:2])
    return label.replace(next_quarter, dropped)


def last_quarter_workbook(answer_key, folder):
    """Write the company's workbook as it stood a quarter ago (its latest quarter left off) into folder.

    answer_key is the company's make_data module. The file keeps the company's name
    (folder/northwind.xlsx), so its manifest is the same file as today's workbook's. The dropped
    quarter's budget columns become the budget-only row, as they were before that quarter closed.
    The title line (Fernhollow) is kept; the Notes tab's position isn't (it doesn't change a number).
    """
    dropped = answer_key.QUARTERS[-1]
    path = Path(folder) / answer_key.OUTPUT_PATH.name
    company = {
        "output_path": path,
        "quarters": answer_key.QUARTERS[:-1],
        "blank_quarter": answer_key.BLANK_QUARTER,
        "budget_only_label": earlier_budget_label(answer_key.BUDGET_ONLY_LABEL, dropped),
        "true_data": {column: values[:-1] for column, values in answer_key.TRUE_DATA.items()},
        "next_quarter_budget": {column: answer_key.TRUE_DATA[column][-1] for column in answer_key.NEXT_QUARTER_BUDGET},
        "header_names": answer_key.HEADER_NAMES,
        "text_cells": {cell: style for cell, style in answer_key.TEXT_CELLS.items() if cell[1] != dropped},
        "notes": answer_key.NOTES,
        "title": getattr(answer_key, "TITLE", None),
    }
    with contextlib.redirect_stdout(io.StringIO()):   # save_workbook prints a line per file
        save_workbook(company)
    return path


# ---------------------------------------------------------------------------
# Reading what a run left behind
# ---------------------------------------------------------------------------

def quietly(function, *args, **kwargs):
    """Call a function with its printout captured; returns (its answer, what it printed)."""
    printed = io.StringIO()
    with contextlib.redirect_stdout(printed):
        answer = function(*args, **kwargs)
    return answer, printed.getvalue()


def build(workbook, output_dir):
    """main.run_company with no AI step, as `python main.py <workbook> --skip-ai` does. Returns what it printed."""
    return quietly(run_company, workbook, load_config(), True, output_dir=output_dir)[1]


def memo_paragraphs(stem, output_dir):
    return [item.text for item in Document(Path(output_dir) / f"{stem}_board_memo.docx").paragraphs]


def memo_section(paragraphs):
    """The What changed section's paragraphs (heading to Key metrics), or [] if the memo has none."""
    if HEADING not in paragraphs:
        return []
    start = paragraphs.index(HEADING)
    return paragraphs[start:paragraphs.index("Key metrics")]


def expected_section(expected, when):
    """The section's paragraphs, as the memo must write them: heading, which run, each title then its lines."""
    lines = [HEADING, COMPARED.format(when=when)]
    for title, items in expected.items():
        lines += [title] + items
    return lines


def pdf_words(stem, output_dir):
    pages = PdfReader(Path(output_dir) / f"{stem}_board_memo.pdf").pages
    return " ".join(" ".join(page.extract_text() for page in pages).split())


def page_lines(workbook, output_dir):
    """What the web page's card shows, as (compared-with line, [titles and lines])."""
    from diff_runs import change_sections, compared_with_text   # the page's own words, to compare with the memo's
    config = load_config()
    data, problem = load_company(workbook, config)
    assert problem is None, problem
    report, settings, problem = run_changes(workbook, data, config, output_dir)
    assert problem is None and report is not None, f"the page has nothing to compare with: {problem}"
    lines = [compared_with_text(report)]
    for title, items in change_sections(report, settings):
        lines += [title] + items
    return lines


# ---------------------------------------------------------------------------
# One company
# ---------------------------------------------------------------------------

def check_first_run(name, answer_key, folder, output_dir):
    """Last quarter's workbook on its own: nothing to compare with, and no section in the memo."""
    earlier = last_quarter_workbook(answer_key, folder)
    printed = build(earlier, output_dir)
    manifest = read_manifest(manifest_path(earlier, output_dir))
    assert manifest["results"]["quarter"] == "Q1 2026", f"{name}: first run recorded {manifest['results']['quarter']}"
    assert manifest["previous_run"] is None, f"{name}: a first run has a previous run"
    assert "Nothing to compare with yet: this is the first run on record." in printed, f"{name}: printout {printed}"
    assert memo_section(memo_paragraphs(answer_key.OUTPUT_PATH.stem, output_dir)) == [], \
        f"{name}: a first run's memo has a What changed section"
    return manifest["run_at"]


def check_quarter_later(name, answer_key, expected, output_dir, first_run_at):
    """Today's workbook a quarter later: the manifest, the memo (Word and PDF) and the page all show the expected lines."""
    workbook, stem = answer_key.OUTPUT_PATH, answer_key.OUTPUT_PATH.stem
    build(workbook, output_dir)
    manifest = read_manifest(manifest_path(workbook, output_dir))
    assert manifest["previous_run"]["run_at"] == first_run_at, f"{name}: compared with the wrong run"
    assert manifest["previous_run"]["results"]["quarter"] == "Q1 2026", f"{name}: compared with the wrong quarter"

    when = first_run_at.replace("T", " ")[:16]
    want = expected_section(expected, when)
    paragraphs = memo_paragraphs(stem, output_dir)
    got = memo_section(paragraphs)
    assert got == want, f"{name}: the memo's section differs:\n  got  {got}\n  want {want}"
    assert paragraphs.index("Headline") < paragraphs.index(HEADING), f"{name}: the section isn't after the headline"
    pdf = pdf_words(stem, output_dir)
    missing = [line for line in want if " ".join(line.split()) not in pdf]
    assert not missing, f"{name}: in the Word memo but not the PDF: {missing}"
    assert page_lines(workbook, output_dir) == want[1:], f"{name}: the page's lines differ from the memo's"
    return want


def check_rebuild_after_approval(name, answer_key, output_dir, want, first_run_at):
    """Approve, rebuild from the same numbers: still compared with the Q1 2026 run; the footer says reviewed."""
    workbook, stem = answer_key.OUTPUT_PATH, answer_key.OUTPUT_PATH.stem
    approve(stem, REVIEWER, workbook.parent, output_dir)
    build(workbook, output_dir)
    manifest = read_manifest(manifest_path(workbook, output_dir))
    assert manifest["previous_run"]["run_at"] == first_run_at, f"{name}: the rebuild lost last quarter's comparison"
    assert memo_section(memo_paragraphs(stem, output_dir)) == want, f"{name}: the rebuilt memo's section differs"
    footer = Document(Path(output_dir) / f"{stem}_board_memo.docx").sections[0].footer.paragraphs[0].text
    assert f"reviewed by {REVIEWER}" in footer, f"{name}: the rebuilt memo isn't reviewed: {footer}"


def check_command_line(answer_key, output_dir, want):
    """python diff_runs.py prints the memo's lines (titles end in a colon, lines start '  - ')."""
    status, printed = quietly(diff_runs_main, [str(answer_key.OUTPUT_PATH), "--output-dir", str(output_dir)])
    assert status == 0
    lines = printed.splitlines()
    assert lines[1] == want[1], f"command line: {lines[1]!r}"
    stripped = [line.removeprefix("  - ").removesuffix(":") for line in lines[2:]]
    assert stripped == want[2:], f"command line lines differ from the memo's: {stripped}"


def check_company(name, answer_key, expected):
    with tempfile.TemporaryDirectory() as folder:
        output_dir = Path(folder) / "output"
        first_run_at = check_first_run(name, answer_key, folder, output_dir)
        want = check_quarter_later(name, answer_key, expected, output_dir, first_run_at)
        check_command_line(answer_key, output_dir, want)
        check_rebuild_after_approval(name, answer_key, output_dir, want, first_run_at)
    counts = ", ".join(f"{len(items)} {SHORT_TITLES[title]}" for title, items in expected.items())
    print(f"✓ {name}: Q1 2026 run then Q2 2026 run: {counts}; memo (Word and PDF), page and command line "
          f"agree; still compared with Q1 2026 after approve and rebuild")


def check_today_alone():
    """Today's workbook with no earlier run: nothing to compare with, no section."""
    with tempfile.TemporaryDirectory() as folder:
        build(northwind.OUTPUT_PATH, folder)
        assert memo_section(memo_paragraphs("northwind", folder)) == [], "a first run's memo has a section"
        assert read_manifest(manifest_path(northwind.OUTPUT_PATH, folder))["previous_run"] is None
    print("✓ Northwind's workbook run alone: no earlier run, no What changed section")


def main():
    for name, answer_key, expected in COMPANIES:
        check_company(name, answer_key, expected)
    check_today_alone()
    print("All checks passed")


if __name__ == "__main__":
    main()
