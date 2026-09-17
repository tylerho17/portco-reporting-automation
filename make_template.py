"""Build templates/base.pptx: the brand template every board deck is built on.

The brand is fictional and neutral ("Example Capital"): navy and gray, Arial, 16:9.
It has exactly two layouts:
- "Title Slide":        navy background, white title (for a cover page if one is ever wanted)
- "Title and Content":  navy title at the top, a content area, a footer strip at the bottom

build_deck.py reads the positions of the content area and the footer from the
"Title and Content" layout, so moving them here moves them on every deck.

python-pptx can't create a template from nothing, so this starts from its built-in
default (4:3, eleven Office layouts), then resizes, restyles and trims it.

Run: python make_template.py  -> saves templates/base.pptx
"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.oxml.shapes.autoshape import CT_Shape
from pptx.shapes.autoshape import Shape
from pptx.util import Emu, Inches, Pt

TEMPLATE_PATH = Path(__file__).parent / "templates" / "base.pptx"

# ---------------------------------------------------------------------------
# Brand: one place for colors, font and name. build_deck.py imports these.
# ---------------------------------------------------------------------------

BRAND_NAME = "Example Capital"   # fictional
NAVY = "1F2A44"                  # titles, accents, chart bars
DARK_GRAY = "333333"             # body text
MID_GRAY = "7F7F7F"              # footer text, rules
LIGHT_GRAY = "F2F2F2"            # table stripes, panels
WHITE = "FFFFFF"
FONT = "Arial"

TITLE_LAYOUT = "Title Slide"
CONTENT_LAYOUT = "Title and Content"

# ---------------------------------------------------------------------------
# Geometry (16:9 = 13.333 x 7.5 inches)
# ---------------------------------------------------------------------------

SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)
MARGIN = Inches(0.5)
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN

TITLE_BOX = (MARGIN, Inches(0.3), CONTENT_WIDTH, Inches(0.85))        # (left, top, width, height)
BODY_BOX = (MARGIN, Inches(1.3), CONTENT_WIDTH, Inches(5.45))         # ends at 6.75 in
FOOTER_BOX = (MARGIN, Inches(6.95), Inches(9.5), Inches(0.4))         # left part of the bottom strip
BRAND_BOX = (MARGIN + Inches(9.5), Inches(6.95), CONTENT_WIDTH - Inches(9.5), Inches(0.4))
TOP_BAR = (0, 0, SLIDE_WIDTH, Inches(0.12))                           # thin navy band at the very top
FOOTER_RULE = (MARGIN, Inches(6.87), CONTENT_WIDTH, Emu(9525))        # hairline above the footer (0.75 pt)

COVER_TITLE_BOX = (MARGIN, Inches(2.4), CONTENT_WIDTH, Inches(1.5))
COVER_SUBTITLE_BOX = (MARGIN, Inches(4.0), CONTENT_WIDTH, Inches(1.0))

A_NAMESPACE = "http://schemas.openxmlformats.org/drawingml/2006/main"  # the "a:" prefix in slide XML

# Placeholders the decks don't use: the run date lives in the footer text instead.
UNUSED_PLACEHOLDERS = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}


# ---------------------------------------------------------------------------
# Theme: colors and fonts that every layout inherits
# ---------------------------------------------------------------------------

def set_theme(presentation):
    """Rewrite the theme's colors (navy/gray) and fonts (Arial) in its XML."""
    theme_part = presentation.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme_part.blob)
    colors = {"dk2": NAVY, "lt2": LIGHT_GRAY, "accent1": NAVY, "accent2": MID_GRAY}
    for slot, hex_color in colors.items():
        root.find(f".//{qn('a:' + slot)}/{qn('a:srgbClr')}").set("val", hex_color)
    for font_kind in ("a:majorFont", "a:minorFont"):  # major = headings, minor = body
        root.find(f".//{qn(font_kind)}/{qn('a:latin')}").set("typeface", FONT)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def set_text_style(level_element, size_pt, hex_color, bold=False, align=None):
    """Set size, color and boldness on one text-style level (e.g. <a:lvl1pPr>)."""
    if align:
        level_element.set("algn", align)
    run_defaults = level_element.find(qn("a:defRPr"))
    run_defaults.set("sz", str(size_pt * 100))  # PowerPoint stores sizes in 1/100 pt
    run_defaults.set("b", "1" if bold else "0")
    old_fill = run_defaults.find(qn("a:solidFill"))
    new_fill = etree.fromstring(f'<a:solidFill xmlns:a="{A_NAMESPACE}"><a:srgbClr val="{hex_color}"/></a:solidFill>')
    run_defaults.replace(old_fill, new_fill)


def set_master_text_styles(master):
    """Titles: navy, bold, 28 pt, left-aligned. Body: dark gray, 20 pt then 18 pt."""
    styles = master.element.find(qn("p:txStyles"))
    set_text_style(styles.find(f"{qn('p:titleStyle')}/{qn('a:lvl1pPr')}"), 28, NAVY, bold=True, align="l")
    body = styles.find(qn("p:bodyStyle"))
    set_text_style(body.find(qn("a:lvl1pPr")), 20, DARK_GRAY)
    set_text_style(body.find(qn("a:lvl2pPr")), 18, DARK_GRAY)


# ---------------------------------------------------------------------------
# Layouts and placeholders
# ---------------------------------------------------------------------------

def keep_only_two_layouts(presentation):
    """Delete every layout except the title layout and the title-and-content layout."""
    for layout in list(presentation.slide_layouts):
        if layout.name not in (TITLE_LAYOUT, CONTENT_LAYOUT):
            presentation.slide_layouts.remove(layout)


def remove_unused_placeholders(shapes):
    """Take the date and slide-number placeholders off a master or layout."""
    for placeholder in list(shapes.placeholders):
        if placeholder.placeholder_format.type in UNUSED_PLACEHOLDERS:
            placeholder.element.getparent().remove(placeholder.element)


def place(shape, box):
    """Move and resize a shape to box = (left, top, width, height)."""
    shape.left, shape.top, shape.width, shape.height = box


def box_for(placeholder_type, cover=False):
    """Where a placeholder of this type goes (None = leave it where it is)."""
    if placeholder_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
        return COVER_TITLE_BOX if cover else TITLE_BOX
    if placeholder_type == PP_PLACEHOLDER.SUBTITLE:
        return COVER_SUBTITLE_BOX
    if placeholder_type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
        return BODY_BOX
    if placeholder_type == PP_PLACEHOLDER.FOOTER:
        return FOOTER_BOX
    return None


def position_placeholders(shapes, cover=False):
    """Move every placeholder on a master or layout to its 16:9 position."""
    for placeholder in shapes.placeholders:
        box = box_for(placeholder.placeholder_format.type, cover)
        if box:
            place(placeholder, box)


def set_placeholder_style(placeholder, size_pt, hex_color, bold=False):
    """Give one layout placeholder its own text style (a <a:lstStyle> that slides inherit)."""
    list_style = placeholder.text_frame._txBody.find(qn("a:lstStyle"))
    level = etree.fromstring(
        f'<a:lvl1pPr xmlns:a="{A_NAMESPACE}"><a:defRPr sz="{size_pt * 100}" b="{int(bold)}">'
        f'<a:solidFill><a:srgbClr val="{hex_color}"/></a:solidFill></a:defRPr></a:lvl1pPr>')
    list_style.append(level)


def style_cover_layout(layout):
    """Title Slide: navy background, white title, light gray subtitle, no master decorations."""
    layout.element.set("showMasterSp", "0")  # hide the top bar and footer rule from the master
    layout.background.fill.solid()
    layout.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    for placeholder in layout.placeholders:
        kind = placeholder.placeholder_format.type
        if kind == PP_PLACEHOLDER.CENTER_TITLE:
            set_placeholder_style(placeholder, 40, WHITE, bold=True)
        elif kind == PP_PLACEHOLDER.SUBTITLE:
            set_placeholder_style(placeholder, 20, LIGHT_GRAY)


# ---------------------------------------------------------------------------
# Master decorations: top bar, footer rule, brand name
# ---------------------------------------------------------------------------

def add_master_shape(master, name, box, textbox=False):
    """Add a rectangle (or text box) to the slide master, behind the placeholders.

    python-pptx only has "add shape" for slides, so this builds the shape's XML directly
    and puts it right after the two header elements of the shape tree (= drawn first, at the back).
    """
    shape_id = max(int(element.get("id")) for element in master.element.iter(qn("p:cNvPr"))) + 1
    left, top, width, height = box
    if textbox:
        element = CT_Shape.new_textbox_sp(shape_id, name, left, top, width, height)
    else:
        element = CT_Shape.new_autoshape_sp(shape_id, name, "rect", left, top, width, height)
    master.shapes._spTree.insert(2, element)
    return Shape(element, master.shapes)


def fill_solid(shape, hex_color):
    """Solid fill, no outline."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(hex_color)
    shape.line.fill.background()


def add_brand_name(master):
    """Brand name at the bottom right of every content slide: navy, bold, 12 pt."""
    shape = add_master_shape(master, "Brand name", BRAND_BOX, textbox=True)
    frame = shape.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT
    run = paragraph.add_run()
    run.text = BRAND_NAME
    run.font.size, run.font.bold = Pt(12), True
    run.font.color.rgb = RGBColor.from_string(NAVY)


def decorate_master(master):
    """Thin navy bar across the top, gray hairline above the footer, brand name."""
    fill_solid(add_master_shape(master, "Top bar", TOP_BAR), NAVY)
    fill_solid(add_master_shape(master, "Footer rule", FOOTER_RULE), MID_GRAY)
    add_brand_name(master)


# ---------------------------------------------------------------------------
# Putting it together
# ---------------------------------------------------------------------------

def build_template():
    """Return the brand template as an in-memory Presentation with no slides."""
    presentation = Presentation()  # python-pptx's built-in default template
    presentation.slide_width, presentation.slide_height = SLIDE_WIDTH, SLIDE_HEIGHT
    set_theme(presentation)
    keep_only_two_layouts(presentation)

    master = presentation.slide_master
    set_master_text_styles(master)
    remove_unused_placeholders(master)
    position_placeholders(master)
    decorate_master(master)

    for layout in presentation.slide_layouts:
        remove_unused_placeholders(layout)
        position_placeholders(layout, cover=layout.name == TITLE_LAYOUT)
    style_cover_layout(presentation.slide_layouts.get_by_name(TITLE_LAYOUT))
    return presentation


def save_template(path=TEMPLATE_PATH):
    """Build the template and save it. Returns the path."""
    path = Path(path)
    path.parent.mkdir(exist_ok=True)
    build_template().save(path)
    return path


if __name__ == "__main__":
    print(f"Saved {save_template()}")
