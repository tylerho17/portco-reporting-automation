"""The portfolio rollup (Task 9): one summary deck and one workbook across every company in data/.

Output: output/portfolio_rollup.pptx and output/portfolio_rollup.xlsx.

The deck (templates/base.pptx, like every company deck, with the same footer rules):
1. Portfolio ranked by flags tripped: rank, company, quarter, "7 of 9 flags tripped", the worst
   tripped flag with its value and threshold, runway at current burn, review status. More than
   ROWS_PER_SLIDE companies continue on a second ranking slide, so the table always fits.
2. Companies by status: how many companies have a tripped flag / only flags that can't be
   evaluated / every flag passed / a workbook that can't be read; and how many decks are
   approved / not reviewed / out of date / not generated (the portfolio page's Deck status).
3. Runway at current burn by company: a bar chart (charts.py), red where the runway flag trips,
   with the flag's threshold from config.yaml as a dashed line.
The workbook has the same three things as sheets: Ranking, By status, Runway.

Which flag is "worst": the tripped flag highest in WORST_FIRST, a fixed order an investor would
read them in (reasons beside it). Ranking: most flags tripped first; a tie goes to the company
whose worst flag is worse, then by name. A workbook that can't be read is listed last, unranked.

Rules:
- No new math and no typed numbers: every number comes from metrics.py (through portfolio.load_company,
  the same numbers the web page shows), config.yaml, or is a count of companies or flags.
  Formatting is metrics.format_value's and excel_output's.
- No AI: the rollup carries computed metrics only, and says so in the footer. It never calls the API.
- Text must fit (text_fit.py, 12 pt floor); no em dash anywhere.
- Numbers are worked out from each workbook now; only the review status reads output/ (the manifests).

Run: python rollup.py      (writes output/portfolio_rollup.pptx and .xlsx)
"""

import argparse
import datetime
import math
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu, Pt

from build_deck import (BODY_TYPES, CELL_MARGIN_X, CELL_MARGIN_Y, FICTIONAL_NOTE, FOOTER_SIZE, TABLE_SIZE,
                        add_text_box, column_box, commit_text, find_content_layout, flag_count_text, layout_box,
                        new_slide, points, set_title, threshold_text, value_text, write_cell)
from charts import runway_chart, save_chart
from excel_output import cell_value, color_cell, number_format, set_column_widths, status_label, write_header
from main import DATA_DIR, OUTPUT_DIR, company_name, find_workbooks
from make_template import TEMPLATE_PATH
from metrics import CANNOT_EVALUATE, COMBO_FLAG_NAME, FLAG_RULES, METRIC_LABELS, PASS, TRIP, load_config
from portfolio import NOT_GENERATED, OUT_OF_DATE, load_company, run_state
from provenance import NOT_REVIEWED
from text_fit import fit_table, paragraph
from theme import MID_GRAY, NAVY, SLATE, STATUS_COLORS, SURFACE, WHITE

ROLLUP_STEM = "portfolio_rollup"
CHART_FOLDER = "charts"
NO_VALUE = "-"
NONE_TRIPPED = "None tripped"
NO_AI_TEXT = "computed metrics only, no AI text"    # the footer's last part: nothing here is AI-drafted
NO_COMPANIES = "No KPI workbooks in data/: there is nothing to roll up."
RUNWAY = "runway_months"
RUNWAY_THRESHOLD_KEY = next(key for _, column, key, _ in FLAG_RULES if column == RUNWAY)   # config.yaml

# The worst flag first. A fixed order, because the flags' units can't be compared (months, %, x):
WORST_FIRST = [METRIC_LABELS[column] for column in (
    "runway_months",           # cash runs out: the company has months to raise or cut, nothing else matters first
    "nrr",                     # the customer base shrinks without new sales
    "grr",                     # customers are leaving: a product or churn problem
    "burn_multiple",           # growth costs too much cash (infinite when ARR shrank)
    "burn_vs_budget",          # the budget the board approved is no longer credible
    "net_new_arr_vs_budget",   # the growth case is off track
    "cac_payback_months",      # each new customer ties up too much cash
    "rule_of_40",              # a summary test: its causes usually trip one of the flags above
)] + [COMBO_FLAG_NAME]         # explains why NRR falls; NRR itself is already higher up

# A company's flag status: its flags' worst status, or a workbook that can't be read.
UNREADABLE = "unreadable"
STATUS_ORDER = [TRIP, CANNOT_EVALUATE, PASS, UNREADABLE]
STATUS_LABELS = {TRIP: "Flags tripped", CANNOT_EVALUATE: "None tripped, some cannot evaluate",
                 PASS: "Every flag passed", UNREADABLE: "Workbook can't be read"}
UNREADABLE_TEXT = "Workbook can't be read: see the portfolio page"
COLOR_OF = {TRIP: TRIP, CANNOT_EVALUATE: CANNOT_EVALUATE, PASS: PASS, UNREADABLE: CANNOT_EVALUATE}

# Review status: the portfolio page's Deck status in a word or two.
REVIEW_ORDER = ["Approved", "Not reviewed", "Out of date", "Not generated"]
OUT_OF_DATE_START = OUT_OF_DATE.split(":")[0]   # "Out of date"
APPROVED_START = "approved by"                  # provenance.deck_status's wording

# The deck's tables.
ROWS_PER_SLIDE = 7     # companies per ranking slide: 7 rows of 3 lines each fit at 12 pt (40-character names)
MAX_NAMES = 3          # company names listed per status on the slide (40-character names fit); the workbook lists them all
RANKING_HEADER = ["Rank", "Company", "Quarter", "Flags tripped", "Worst flag", "Runway at current burn", "Review"]
RANKING_SHARES = [0.06, 0.21, 0.09, 0.17, 0.25, 0.11, 0.11]   # share of the table width per column
COUNT_SHARES = [0.36, 0.16, 0.48]
WORST_COLUMN = RANKING_HEADER.index("Worst flag")

# The workbook's sheets.
SHEET_NAMES = ["Ranking", "By status", "Runway"]
RANKING_HEADERS = ["Rank", "Company", "Latest quarter", "Flags tripped", "Flags that cannot evaluate",
                   "Flags checked", "Worst flag", "Worst flag value", "Threshold", "Runway at current burn",
                   "Flag status", "Review status", "Note"]
STATUS_HEADERS = ["Group", "Status", "Companies", "Which"]
RUNWAY_HEADERS = ["Company", "Latest quarter", "Runway at current burn", "Threshold", "Runway flag"]


# ---------------------------------------------------------------------------
# 1. One company: its flags summed up, worst flag, status
# ---------------------------------------------------------------------------

def severity(name):
    """A flag's place in WORST_FIRST: 0 is the worst."""
    return WORST_FIRST.index(name)


def worst_flag(flags):
    """The tripped flag highest in WORST_FIRST, or None if no flag tripped."""
    tripped = [flag for flag in flags if flag["status"] == TRIP]
    return min(tripped, key=lambda flag: severity(flag["flag"]), default=None)


def company_status(flags):
    """TRIP if any flag tripped; else CANNOT_EVALUATE if any couldn't be evaluated; else PASS."""
    statuses = {flag["status"] for flag in flags}
    if TRIP in statuses:
        return TRIP
    return CANNOT_EVALUATE if CANNOT_EVALUATE in statuses else PASS


def review_label(page_status):
    """The portfolio page's Deck status ('approved by Tyler Ho on ...') as 'Approved', 'Not reviewed', ..."""
    if page_status == NOT_GENERATED:
        return "Not generated"
    if page_status == NOT_REVIEWED:
        return "Not reviewed"
    if page_status.startswith(OUT_OF_DATE_START):
        return "Out of date"
    if page_status.startswith(APPROVED_START):
        return "Approved"
    return page_status   # a wording this file doesn't know is shown as it is, not guessed


def company_entry(workbook_path, config, output_dir):
    """Everything the rollup shows about one company. A workbook that can't be read keeps its message."""
    data, problem = load_company(workbook_path, config)
    flags = data["flags"] if data else []
    return {"company": company_name(workbook_path), "data": data, "problem": problem,
            "tripped": sum(flag["status"] == TRIP for flag in flags),
            "unevaluated": sum(flag["status"] == CANNOT_EVALUATE for flag in flags),
            "checked": len(flags), "worst": worst_flag(flags),
            "status": company_status(flags) if data else UNREADABLE,
            "review": review_label(run_state(workbook_path, output_dir)["status"])}


# ---------------------------------------------------------------------------
# 2. The portfolio: ranking and counts
# ---------------------------------------------------------------------------

def rank_key(entry):
    """Most flags tripped first; then the worse worst flag; then the company name."""
    worst = severity(entry["worst"]["flag"]) if entry["worst"] else len(WORST_FIRST)
    return -entry["tripped"], worst, entry["company"]


def ranked(entries):
    """The entries in rank order, each with its "rank" (1 = most attention). Unreadable ones last, rank None."""
    readable = sorted((entry for entry in entries if entry["data"] is not None), key=rank_key)
    unreadable = sorted((entry for entry in entries if entry["data"] is None), key=lambda entry: entry["company"])
    return ([{**entry, "rank": number} for number, entry in enumerate(readable, start=1)]
            + [{**entry, "rank": None} for entry in unreadable])


def collect_rollup(config, data_dir=DATA_DIR, output_dir=OUTPUT_DIR):
    """Every workbook in data/, ranked. Reads output/ only for each company's review status."""
    return ranked([company_entry(path, config, output_dir) for path in find_workbooks(data_dir)])


def counts(entries, key, order, labels=None):
    """[(label, how many companies, their names A to Z)] for every value in `order`, zeros included."""
    return [((labels or {}).get(value, value), sum(entry[key] == value for entry in entries),
             sorted(entry["company"] for entry in entries if entry[key] == value)) for value in order]


def status_counts(entries):
    """Companies by flag status: Flags tripped, None tripped but some cannot evaluate, Every flag passed, unreadable."""
    return counts(entries, "status", STATUS_ORDER, STATUS_LABELS)


def review_counts(entries):
    """Companies by review status: Approved, Not reviewed, Out of date, Not generated."""
    return counts(entries, "review", REVIEW_ORDER)


def names_text(names):
    """'A, B, C', or the first MAX_NAMES and 'and 7 more'; '-' for none."""
    if not names:
        return NO_VALUE
    shown = ", ".join(names[:MAX_NAMES])
    return shown if len(names) <= MAX_NAMES else f"{shown} and {len(names) - MAX_NAMES} more"


def quarter_words(entries):
    """The one latest quarter every company shares ('Q2 2026'), or "each company's latest quarter"."""
    quarters = {entry["data"]["latest"] for entry in entries if entry["data"] is not None}
    if len(quarters) == 1:
        return quarters.pop()
    return "each company's latest quarter" if quarters else NO_VALUE


# ---------------------------------------------------------------------------
# 3. Text for one company (formatting only: build_deck's own words and numbers)
# ---------------------------------------------------------------------------

def worst_flag_text(entry):
    """'Runway at current burn: 6.0 mo (trips below 12.0 mo)', 'None tripped', or why the workbook can't be read."""
    if entry["data"] is None:
        return UNREADABLE_TEXT
    flag = entry["worst"]
    if flag is None:
        return NONE_TRIPPED
    if flag["metric"] is None:   # the combo rule has no single value
        return f"{flag['flag']}: {status_label(flag)}"
    return f"{flag['flag']}: {value_text(entry['data'], flag['metric'], flag['quarter'])} ({threshold_text(flag)})"


def runway_flag(entry):
    """The company's runway flag (its value, threshold and status)."""
    return next(flag for flag in entry["data"]["flags"] if flag["metric"] == RUNWAY)


def runway_text(entry):
    """'6.0 mo', '∞ (not burning)', 'data missing' (build_deck.value_text), or '-' for an unreadable workbook."""
    if entry["data"] is None:
        return NO_VALUE
    return value_text(entry["data"], RUNWAY, entry["data"]["latest"])


def runway_value(entry):
    """Months of runway (inf = not burning), or NaN when there's no number."""
    if entry["data"] is None:
        return math.nan
    return entry["data"]["metrics"].loc[entry["data"]["latest"], RUNWAY]


def runway_order(entries):
    """The chart's order: shortest runway first, then not burning (∞), then no number, then unreadable."""
    def key(entry):
        value = runway_value(entry)
        group = 0 if math.isfinite(value) else 1 if math.isinf(value) else 2 if entry["data"] else 3
        return group, value if group == 0 else 0, entry["company"]
    return sorted(entries, key=key)


# ---------------------------------------------------------------------------
# 4. The deck
# ---------------------------------------------------------------------------

def widths_of(total_width, shares):
    """Split a width by shares; the last column takes any rounding remainder."""
    widths = [int(total_width * share) for share in shares[:-1]]
    return widths + [total_width - sum(widths)]


def draw_table(slide, name, box, header, rows, shares, deck):
    """A table in the box: navy header, striped rows, fitted with text_fit (or stop naming the slide).

    rows = [(cells, {column number: flag status whose colors that cell takes})].
    """
    left, top, width, height = box
    widths = widths_of(width, shares)
    all_text = [header] + [cells for cells, _ in rows]
    text_widths = [points(column_width - 2 * CELL_MARGIN_X) for column_width in widths]
    size, heights = fit_table(all_text, text_widths, points(height), f"{deck['where']}, {name}", TABLE_SIZE,
                              cell_padding_pt=points(2 * CELL_MARGIN_Y))
    frame = slide.shapes.add_table(len(all_text), len(header), left, top, width, Pt(sum(heights)))
    frame.name = name
    table = frame.table
    table.horz_banding = False   # our own stripes, not the template's
    for column, column_width in enumerate(widths):
        table.columns[column].width = column_width
    for row_number, row_height in enumerate(heights):
        table.rows[row_number].height = Pt(row_height)
    for column, text in enumerate(header):
        write_cell(table.cell(0, column), text, size, NAVY, WHITE, bold=True)
    for row_number, (cells, colored) in enumerate(rows, start=1):
        stripe = SURFACE if row_number % 2 == 0 else WHITE
        for column, text in enumerate(cells):
            fill_hex, text_hex = STATUS_COLORS[colored[column]] if column in colored else (stripe, SLATE)
            write_cell(table.cell(row_number, column), text, size, fill_hex, text_hex)


def ranking_row(entry):
    """One company's cells on the ranking slide, with the worst-flag cell in the company's status colors."""
    data = entry["data"]
    cells = [str(entry["rank"]) if entry["rank"] else NO_VALUE, entry["company"],
             data["latest"] if data else NO_VALUE, flag_count_text(data["flags"]) if data else NO_VALUE,
             worst_flag_text(entry), runway_text(entry), entry["review"]]
    return cells, {WORST_COLUMN: COLOR_OF[entry["status"]]}


def pages(entries):
    """The ranking split into slides of at most ROWS_PER_SLIDE companies (always at least one slide)."""
    return [entries[start:start + ROWS_PER_SLIDE] for start in range(0, max(len(entries), 1), ROWS_PER_SLIDE)]


def ranking_slide(slide, deck, page, number, total):
    """One ranking slide: '(1 of 2)' in the title when the ranking needs more than one."""
    part = f" ({number} of {total})" if total > 1 else ""
    set_title(slide, f"Portfolio ranked by flags tripped, {deck['quarters']}{part}", deck)
    draw_table(slide, "Ranking table", deck["area"], RANKING_HEADER, [ranking_row(entry) for entry in page],
               RANKING_SHARES, deck)


def count_rows(rows, colored_by=None):
    """(label, count, names) rows as table cells; colored_by maps a label to its flag status colors."""
    return [([label, str(count), names_text(names)], {0: colored_by[label]} if colored_by else {})
            for label, count, names in rows]


def status_slide(slide, deck):
    """Two small tables side by side: companies by flag status (colored) and by review status (not:
    red and green mean a flag's status only)."""
    set_title(slide, f"Companies by status, {deck['quarters']}", deck)
    colors = {STATUS_LABELS[status]: COLOR_OF[status] for status in STATUS_ORDER}
    tables = [("Flag status table", "Flag status", count_rows(status_counts(deck["entries"]), colors)),
              ("Review status table", "Review status", count_rows(review_counts(deck["entries"])))]
    for position, (name, heading, rows) in enumerate(tables):
        draw_table(slide, name, column_box(deck["area"], position), [heading, "Companies", "Which"], rows,
                   COUNT_SHARES, deck)


def runway_threshold(config):
    """(months, 'trips below 12.0 mo'): the runway flag's threshold from config.yaml, in the flags' words."""
    months = config[RUNWAY_THRESHOLD_KEY]
    return months, threshold_text({"metric": RUNWAY, "threshold": months})


def runway_slide(slide, deck):
    """The runway chart, drawn at the size it takes on the slide (so 13 pt in the chart is 13 pt here)."""
    set_title(slide, f"Runway at current burn by company, {deck['quarters']}", deck)
    left, top, width, height = deck["area"]
    order = runway_order(deck["entries"])
    months, words = runway_threshold(deck["config"])
    figure = runway_chart([entry["company"] for entry in order], [runway_value(entry) for entry in order],
                          [runway_text(entry) if entry["data"] else STATUS_LABELS[UNREADABLE] for entry in order],
                          [entry["data"] is not None and runway_flag(entry)["status"] == TRIP for entry in order],
                          months, words, (Emu(width).inches, Emu(height).inches))
    path = save_chart(figure, deck["chart_dir"] / f"{ROLLUP_STEM}_runway_chart.png")
    picture = slide.shapes.add_picture(str(path), left, top, width, height)
    picture.name = "Runway chart"


def footer_text(deck, run_date):
    """'Fictional data | Portfolio rollup of 3 workbooks | 2026-09-18 | 2cf0d17 | computed metrics only, no AI text'."""
    count = len(deck["entries"])
    workbooks = f"Portfolio rollup of {count} workbook{'s' if count != 1 else ''}"
    return " | ".join([FICTIONAL_NOTE, workbooks, run_date.isoformat(), deck["commit"], NO_AI_TEXT])


def build_rollup_presentation(entries, config, run_date, chart_dir):
    """The ranking slide(s), the status slide and the runway slide, each with the footer."""
    presentation = Presentation(TEMPLATE_PATH)
    layout = find_content_layout(presentation)
    deck = {"entries": entries, "config": config, "chart_dir": Path(chart_dir), "quarters": quarter_words(entries),
            "area": layout_box(layout, BODY_TYPES), "footer_box": layout_box(layout, {PP_PLACEHOLDER.FOOTER}),
            "commit": commit_text()}
    ranking = pages(entries)
    builders = [lambda slide, page=page, number=number: ranking_slide(slide, deck, page, number, len(ranking))
                for number, page in enumerate(ranking, start=1)]
    builders += [lambda slide: status_slide(slide, deck), lambda slide: runway_slide(slide, deck)]
    for number, build_slide in enumerate(builders, start=1):
        slide = new_slide(presentation, layout)
        deck["where"] = f"Rollup slide {number}"   # names the slide in any "doesn't fit" error
        build_slide(slide)
        add_text_box(slide, "Footer", deck["footer_box"],
                     [paragraph(footer_text(deck, run_date), FOOTER_SIZE, color=MID_GRAY)], deck)
    return presentation


# ---------------------------------------------------------------------------
# 5. The workbook
# ---------------------------------------------------------------------------

def number_cell(sheet, row, column, metric):
    """Give one cell its metric's number format (right-aligned, so words line up with numbers)."""
    cell = sheet.cell(row=row, column=column)
    cell.number_format = number_format(metric)
    cell.alignment = Alignment(horizontal="right")


def metric_cell(entry, metric):
    """A metric's latest value as excel_output writes it: the number, or the words for why there's none."""
    data = entry["data"]
    return cell_value(data["actuals"], data["metrics"], data["reasons"], metric, data["latest"])


def ranking_values(entry):
    """One company's row on the Ranking sheet, in RANKING_HEADERS order."""
    if entry["data"] is None:
        return [None, entry["company"], None, None, None, None, None, None, None, None,
                STATUS_LABELS[UNREADABLE], entry["review"], entry["problem"]]
    worst = entry["worst"]
    if worst is None:
        name, value, threshold = NONE_TRIPPED, None, None
    elif worst["metric"] is None:   # the combo rule has no single value or threshold
        name, value, threshold = worst["flag"], status_label(worst), None
    else:
        name, value, threshold = worst["flag"], metric_cell(entry, worst["metric"]), worst["threshold"]
    return [entry["rank"], entry["company"], entry["data"]["latest"], entry["tripped"], entry["unevaluated"],
            entry["checked"], name, value, threshold, metric_cell(entry, RUNWAY), STATUS_LABELS[entry["status"]],
            entry["review"], None]


def write_ranking_sheet(sheet, entries):
    """One row per company in rank order; worst flag red, flag status in its color, numbers in their formats."""
    write_header(sheet, RANKING_HEADERS)
    column = {header: number for number, header in enumerate(RANKING_HEADERS, start=1)}
    for entry in entries:
        sheet.append(ranking_values(entry))
        row = sheet.max_row
        color_cell(sheet.cell(row=row, column=column["Flag status"]), COLOR_OF[entry["status"]])
        if entry["worst"] is not None:
            color_cell(sheet.cell(row=row, column=column["Worst flag"]), TRIP)
            if entry["worst"]["metric"] is not None:
                for header in ("Worst flag value", "Threshold"):
                    number_cell(sheet, row, column[header], entry["worst"]["metric"])
        if entry["data"] is not None:
            number_cell(sheet, row, column["Runway at current burn"], RUNWAY)
    set_column_widths(sheet, [7, 44, 15, 14, 26, 14, 34, 18, 12, 24, 34, 16, 60])


def write_status_sheet(sheet, entries):
    """Companies by flag status (colored) and by review status, with every name."""
    write_header(sheet, STATUS_HEADERS)
    for status, (label, count, names) in zip(STATUS_ORDER, status_counts(entries)):
        sheet.append(["Flag status", label, count, ", ".join(names)])
        color_cell(sheet.cell(row=sheet.max_row, column=2), COLOR_OF[status])
    for label, count, names in review_counts(entries):
        sheet.append(["Review status", label, count, ", ".join(names)])
    set_column_widths(sheet, [16, 36, 12, 80])


def write_runway_sheet(sheet, entries, config):
    """The chart's numbers, in the chart's order, with the runway flag's result."""
    write_header(sheet, RUNWAY_HEADERS)
    months, _ = runway_threshold(config)
    for entry in runway_order(entries):
        if entry["data"] is None:
            sheet.append([entry["company"], None, None, months, STATUS_LABELS[UNREADABLE]])
            color_cell(sheet.cell(row=sheet.max_row, column=5), CANNOT_EVALUATE)
        else:
            flag = runway_flag(entry)
            sheet.append([entry["company"], entry["data"]["latest"], metric_cell(entry, RUNWAY), months,
                          status_label(flag)])
            color_cell(sheet.cell(row=sheet.max_row, column=5), flag["status"])
            number_cell(sheet, sheet.max_row, 3, RUNWAY)
        number_cell(sheet, sheet.max_row, 4, RUNWAY)
    set_column_widths(sheet, [44, 15, 24, 12, 34])


def build_rollup_workbook(entries, config):
    """The three sheets: Ranking, By status, Runway."""
    book = Workbook()
    book.active.title = SHEET_NAMES[0]   # a new workbook starts with one empty sheet; reuse it
    write_ranking_sheet(book.active, entries)
    write_status_sheet(book.create_sheet(SHEET_NAMES[1]), entries)
    write_runway_sheet(book.create_sheet(SHEET_NAMES[2]), entries, config)
    return book


# ---------------------------------------------------------------------------
# 6. Saving, the web page's download, and the command line
# ---------------------------------------------------------------------------

def rollup_paths(output_dir=OUTPUT_DIR):
    """{"deck": output/portfolio_rollup.pptx, "excel": output/portfolio_rollup.xlsx}"""
    return {"deck": Path(output_dir) / f"{ROLLUP_STEM}.pptx", "excel": Path(output_dir) / f"{ROLLUP_STEM}.xlsx"}


def write_rollup(entries, config, folder, run_date=None):
    """Write the deck and the workbook into `folder`. Old ones are deleted first, so a failed build
    never leaves last run's rollup looking current. Stops if there is no company at all."""
    if not entries:
        raise ValueError(NO_COMPANIES)
    paths = rollup_paths(folder)
    for path in paths.values():
        path.unlink(missing_ok=True)
    chart_dir = Path(folder) / CHART_FOLDER
    chart_dir.mkdir(parents=True, exist_ok=True)
    build_rollup_presentation(entries, config, run_date or datetime.date.today(), chart_dir).save(paths["deck"])
    build_rollup_workbook(entries, config).save(paths["excel"])
    return paths


def save_rollup(config, data_dir=DATA_DIR, output_dir=OUTPUT_DIR, run_date=None):
    """Roll up every company in data/ into output/. Returns {"deck": path, "excel": path}."""
    return write_rollup(collect_rollup(config, data_dir, output_dir), config, output_dir, run_date)


def rollup_download(kind, config, data_dir=DATA_DIR, output_dir=OUTPUT_DIR):
    """(file name, bytes) of the rollup deck ("deck") or workbook ("excel") for the web page's button.

    Built in a temporary folder: a download never writes to output/ (it only reads the manifests there).
    """
    entries = collect_rollup(config, data_dir, output_dir)
    with tempfile.TemporaryDirectory() as folder:
        path = write_rollup(entries, config, folder)[kind]
        return path.name, path.read_bytes()


def main(argv=None):
    parser = argparse.ArgumentParser(description="Roll every company in data/ up into one deck and one workbook.")
    parser.parse_args(argv)
    config = load_config()
    entries = collect_rollup(config, DATA_DIR, OUTPUT_DIR)
    paths = write_rollup(entries, config, OUTPUT_DIR)
    for entry in entries:
        rank = entry["rank"] or NO_VALUE
        print(f"{rank:>2}  {entry['company']:<20} {STATUS_LABELS[entry['status']]:<24} {worst_flag_text(entry)}")
    project = Path(__file__).parent
    for path in paths.values():
        print(f"Saved {path.relative_to(project) if path.is_relative_to(project) else path}")


if __name__ == "__main__":
    main()
