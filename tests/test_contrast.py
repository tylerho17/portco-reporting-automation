"""Contrast tests (final Task 18): every text color on every background it sits on passes WCAG AA.

WCAG 2 AA asks for a contrast ratio of at least 4.5 : 1 between normal-size text and what's behind
it, and 3 : 1 for the parts of a control you need to see (a checkbox, a chart bar). The ratio is
worked out from each color's relative luminance, in theme.py (contrast_ratio).

Three kinds of test, so a future color change can't slip past:
- the math itself, against values worked out by hand;
- theme.TEXT_PAIRS, the list of every pair and where it's used (charts, template, memo included);
- the pairs read back from what the code really makes: the status colors, the web page's style
  sheet, Streamlit's theme, and every colored run in four real decks.
Run from the project folder:  pytest
"""

import re

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

import build_deck
import golden
import theme
from build_deck import save_deck
from metrics import load_config

AA_TEXT = 4.5         # typed from WCAG 2, not copied from theme.py
AA_NON_TEXT = 3.0


# ---------------------------------------------------------------------------
# The math, against values worked out by hand
# ---------------------------------------------------------------------------

def test_black_on_white_is_21_to_1():
    # Luminance 0 and 1: (1 + 0.05) / (0 + 0.05) = 21.
    assert theme.relative_luminance("000000") == 0
    assert theme.relative_luminance("FFFFFF") == 1
    assert theme.contrast_ratio("000000", "FFFFFF") == pytest.approx(21)


def test_a_color_on_itself_is_1_to_1_and_order_does_not_matter():
    assert theme.contrast_ratio("0B2545", "0B2545") == pytest.approx(1)
    assert theme.contrast_ratio("0B2545", "FFFFFF") == theme.contrast_ratio("FFFFFF", "0B2545")


def test_the_lightest_gray_that_passes_on_white():
    # 76 hex = 118; 118 / 255 = 0.4627; ((0.4627 + 0.055) / 1.055) ** 2.4 = 0.1812 for each of R, G, B,
    # so luminance 0.1812 and the ratio is 1.05 / (0.1812 + 0.05) = 4.54: passes.
    assert theme.contrast_ratio("767676", "FFFFFF") == pytest.approx(4.54, abs=0.01)
    # One step lighter: 77 hex = 119 -> luminance 0.1843 -> 1.05 / 0.2343 = 4.48: fails.
    assert theme.contrast_ratio("777777", "FFFFFF") == pytest.approx(4.48, abs=0.01)


def test_dark_channels_use_the_straight_line_part_of_the_curve():
    # 0A hex = 10; 10 / 255 = 0.0392 is under 0.04045, so it's divided by 12.92: 0.003035.
    assert theme.relative_luminance("0A0A0A") == pytest.approx(0.003035, abs=0.000001)


def test_the_aa_threshold_is_wcag_s():
    assert theme.AA_TEXT_RATIO == AA_TEXT and theme.AA_NON_TEXT_RATIO == AA_NON_TEXT


# ---------------------------------------------------------------------------
# The declared pairs
# ---------------------------------------------------------------------------

def pair_id(pair):
    return pair[2]


@pytest.mark.parametrize("text, background, where", theme.TEXT_PAIRS, ids=[pair_id(p) for p in theme.TEXT_PAIRS])
def test_every_declared_text_pair_passes_aa(text, background, where):
    assert theme.contrast_ratio(text, background) >= AA_TEXT, where


@pytest.mark.parametrize("mark, background, where", theme.NON_TEXT_PAIRS,
                         ids=[pair_id(p) for p in theme.NON_TEXT_PAIRS])
def test_every_declared_control_and_chart_mark_passes_3_to_1(mark, background, where):
    assert theme.contrast_ratio(mark, background) >= AA_NON_TEXT, where


def test_the_declared_pairs_name_only_palette_colors():
    palette = {theme.NAVY, theme.NAVY_DARK, theme.SLATE, theme.MID_GRAY, theme.LINE, theme.SURFACE, theme.WHITE,
               theme.RED, theme.RED_FILL, theme.GREEN, theme.GREEN_FILL, theme.GRAY_FILL,
               theme.AMBER}   # Task 19: the charts' shrinking-quarter color
    for first, second, where in theme.TEXT_PAIRS + theme.NON_TEXT_PAIRS:
        assert {first, second} <= palette, where


# ---------------------------------------------------------------------------
# Pairs read back from what the code makes
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("colors", [theme.STATUS_COLORS, theme.EXCEL_STATUS_COLORS], ids=["deck and web", "excel"])
def test_every_status_text_passes_on_its_own_fill(colors):
    for status, (fill, text) in colors.items():
        assert theme.contrast_ratio(text, fill) >= AA_TEXT, status


CSS_RULE = re.compile(r"^(?P<selector>.+?) \{ (?P<body>.*) \}$")
DECLARATION = re.compile(r"([a-z-]+): ([^;]+);")
PAGE_BACKGROUNDS = (theme.SURFACE, theme.WHITE)   # text on the page itself, or in a white card on it


def css_rules():
    """streamlit_css() as [(selector parts, {property: value})], one per rule."""
    rules = []
    for line in theme.streamlit_css().splitlines():
        match = CSS_RULE.match(line)
        rules.append(([part.strip() for part in match["selector"].split(",")], dict(DECLARATION.findall(match["body"]))))
    return rules


def hex_of(css_value):
    """'#0B2545' -> '0B2545'; None for anything that isn't one color ('1px solid #E2E8F0')."""
    return css_value[1:] if re.fullmatch(r"#[0-9A-Fa-f]{6}", css_value) else None


def backgrounds_behind(part, declarations, rules):
    """The background a rule's text sits on: its own, else its parent rule's, else the page's two."""
    if "background-color" in declarations:
        return [hex_of(declarations["background-color"])]
    for parent_parts, parent in rules:
        if "background-color" in parent and any(part.startswith(p + " ") for p in parent_parts):
            return [hex_of(parent["background-color"])]
    return list(PAGE_BACKGROUNDS)


def css_text_pairs():
    """Every (text, background, selector) the style sheet can put on screen."""
    rules = css_rules()
    return [(hex_of(declarations["color"]), background, part)
            for parts, declarations in rules if "color" in declarations
            for part in parts for background in backgrounds_behind(part, declarations, rules)]


def test_the_style_sheet_reader_finds_every_colored_rule():
    # The check below is only as good as this reader: buttons, headings, captions, tables, the portfolio header.
    selectors = {selector for _, _, selector in css_text_pairs()}
    assert '[class*="st-key-portfolio-header"] p' in selectors and ".bp-table th" in selectors
    assert ".stApp h1" in selectors and 'button[data-testid="stBaseButton-primary"]:hover' in selectors
    header_text = [bg for text, bg, sel in css_text_pairs() if sel == '[class*="st-key-portfolio-header"] p']
    assert header_text == [theme.NAVY]    # white text found on the navy header, not on the page


def test_every_text_color_in_the_style_sheet_passes_aa():
    failures = [(sel, text, bg, round(theme.contrast_ratio(text, bg), 2)) for text, bg, sel in css_text_pairs()
                if theme.contrast_ratio(text, bg) < AA_TEXT]
    assert failures == []


def test_streamlit_s_own_text_and_links_pass_aa_and_its_accent_passes_3_to_1():
    streamlit = {name: value.lstrip("#") for name, value in theme.streamlit_theme().items()
                 if isinstance(value, str) and value.startswith("#")}
    for background in (streamlit["backgroundColor"], streamlit["secondaryBackgroundColor"]):
        assert theme.contrast_ratio(streamlit["textColor"], background) >= AA_TEXT
        assert theme.contrast_ratio(streamlit["linkColor"], background) >= AA_TEXT
        assert theme.contrast_ratio(streamlit["primaryColor"], background) >= AA_NON_TEXT   # checkbox, progress bar


# ---------------------------------------------------------------------------
# The decks
# ---------------------------------------------------------------------------

WATERMARK = "Watermark"   # see-through on purpose (see test_the_watermark_is_the_only_exception)


@pytest.fixture(scope="module")
def decks(tmp_path_factory):
    """Each company's deck from its saved analysis, as a draft (so the watermark is there too),
    plus Northwind with no analysis (the gray placeholder text). No API call."""
    folder, config, paths = tmp_path_factory.mktemp("decks"), load_config(), []
    for company in golden.companies():
        path, _ = save_deck(golden.DATA_DIR / f"{company}.xlsx", config, golden.fixture_path(company),
                            output_dir=folder / company, draft=True)
        paths.append(path)
    path, _ = save_deck(golden.DATA_DIR / "northwind.xlsx", config, None, output_dir=folder / "no-ai")
    return [Presentation(path) for path in paths + [path]]


def solid_fill(fill_owner):
    """A shape's or table cell's solid fill as hex, or None when it has none."""
    return str(fill_owner.fill.fore_color.rgb) if fill_owner.fill.type == MSO_FILL_TYPE.SOLID else None


def colored_runs(text_frame):
    """Every run in a text frame whose color is set on the run itself."""
    return [run for paragraph in text_frame.paragraphs for run in paragraph.runs if run.font.color.type is not None]


def theme_background(master):
    """The template theme's light background color ("lt1"), which a master's <p:bgRef> points to."""
    light = etree.fromstring(master.part.part_related_by(RT.THEME).blob).find(f".//{qn('a:lt1')}")[0]
    return light.get("lastClr") or light.get("val")   # <a:sysClr lastClr="FFFFFF"/> or <a:srgbClr val="..."/>


def slide_background(slide):
    """The slide's own solid background, else its layout's, else the master's, which is the theme's."""
    for part in (slide, slide.slide_layout):
        if part._element.cSld.find(qn("p:bg")) is not None:
            return solid_fill(part.background)
    return theme_background(slide.slide_layout.slide_master)


def deck_text_pairs(deck):
    """Every (text, background, where) in a deck: table cells on their fill, text boxes on the slide."""
    pairs = []
    for number, slide in enumerate(deck.slides, start=1):
        behind = slide_background(slide)
        for shape in slide.shapes:
            where = f"slide {number} {shape.name}"
            if shape.has_table:
                pairs += [(str(run.font.color.rgb), solid_fill(cell) or behind, where)
                          for row in shape.table.rows for cell in row.cells for run in colored_runs(cell.text_frame)]
            elif shape.has_text_frame and shape.name != WATERMARK:
                pairs += [(str(run.font.color.rgb), solid_fill(shape) or behind, where)
                          for run in colored_runs(shape.text_frame)]
    return pairs


def test_the_deck_reader_finds_the_status_cells_and_the_footer(decks):
    pairs = [pair for deck in decks for pair in deck_text_pairs(deck)]
    assert (theme.GREEN, theme.GREEN_FILL) in {(text, bg) for text, bg, _ in pairs}   # Alderpeak's passed flags
    assert (theme.RED, theme.RED_FILL) in {(text, bg) for text, bg, _ in pairs}       # Fernhollow's tripped flags
    assert any(where.endswith("Footer") for _, _, where in pairs)


def test_every_colored_text_on_every_slide_passes_aa(decks):
    failures = {(text, bg, where, round(theme.contrast_ratio(text, bg), 2))
                for deck in decks for text, bg, where in deck_text_pairs(deck)
                if theme.contrast_ratio(text, bg) < AA_TEXT}
    assert failures == set()


def test_the_watermark_is_the_only_exception(decks):
    # "DRAFT - NOT REVIEWED" is 25% see-through navy so the numbers under it stay readable, which
    # can't pass 4.5 : 1 by design. The same words are in the footer, which does pass (checked above).
    for deck in decks[:-1]:
        for slide in deck.slides:
            names = [shape.name for shape in slide.shapes]
            footer = next(shape for shape in slide.shapes if shape.name == "Footer").text_frame.text
            assert WATERMARK in names and "not reviewed" in footer
            assert build_deck.WATERMARK_ALPHA_PERCENT < 100


def test_the_template_s_text_styles_pass_aa():
    # Text that takes its color from the template (titles, body, the cover page) has no color of its own.
    template = Presentation(golden.PROJECT_DIR / "templates" / "base.pptx")
    styles = template.slide_master._element.find(qn("p:txStyles"))
    for style in ("p:titleStyle", "p:bodyStyle"):
        color = styles.find(f"{qn(style)}/{qn('a:lvl1pPr')}//{qn('a:srgbClr')}").get("val")
        assert theme.contrast_ratio(color, theme.WHITE) >= AA_TEXT, style
    assert theme_background(template.slide_master) == theme.WHITE   # what the content slides sit on
    # The navy cover: a placeholder with no color of its own would take the master's dark gray.
    cover = next(layout for layout in template.slide_layouts if layout.name == "Title Slide")
    for placeholder in cover.placeholders:
        own = placeholder._element.find(f".//{qn('a:lstStyle')}//{qn('a:srgbClr')}")
        assert own is not None, f"{placeholder.name} has no color of its own on the navy cover"
        assert theme.contrast_ratio(own.get("val"), solid_fill(cover.background)) >= AA_TEXT, placeholder.name
