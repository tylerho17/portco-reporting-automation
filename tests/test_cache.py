"""Tests for the clean and metrics caches (Task 17): each workbook is read once per run, never out of date.

What is proved:
- once: a whole company run (Excel, AI reuse check, deck, memo, manifest) opens and parses its
  workbook once, and works out its metric table once;
- never out of date: an edited workbook, a deleted mapping file or a changed number in the
  inputs is worked out again, never answered from the cache;
- safe to change: changing a table the cache handed out doesn't change what it hands out next;
- bounded: the cache keeps at most cache.MAX_ENTRIES results, so the web page can't fill memory.
The goldens (tests/test_golden.py) prove the outputs themselves didn't change.

Run from the project folder:  python -m pytest -q
"""

import shutil
from pathlib import Path

import numpy as np
import pytest

import cache
import clean
import main
import mapping
import metrics
from clean import UnconfirmedMappingError, clean_workbook
from golden import fixture_path
from metrics import compute_metrics, load_config, metric_reasons

PROJECT_DIR = Path(__file__).parent.parent
DATA_DIR = PROJECT_DIR / "data"


@pytest.fixture(autouse=True)
def empty_caches(tmp_path, monkeypatch):
    """Every test starts with nothing cached, and a temporary mappings/ folder."""
    clean.clear_cache()
    metrics.clear_cache()
    monkeypatch.setattr(mapping, "MAPPINGS_DIR", tmp_path / "mappings")


def counting(monkeypatch, module, name):
    """Wrap module.name so each call is counted. Returns the list the calls are added to."""
    calls, original = [], getattr(module, name)

    def counted(*args, **kwargs):
        calls.append(args)
        return original(*args, **kwargs)

    monkeypatch.setattr(module, name, counted)
    return calls


# ---------------------------------------------------------------------------
# Once per run
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("company", ["northwind", "alderpeak", "fernhollow"])
def test_a_company_run_reads_its_workbook_once(company, tmp_path, monkeypatch):
    reads = counting(monkeypatch, clean, "find_kpi_sheet")          # opens and parses the file
    tables = counting(monkeypatch, metrics, "metrics_table")        # works out every metric
    shutil.copy(fixture_path(company), tmp_path)   # the saved analysis: reused, so no API call
    main.run_company(DATA_DIR / f"{company}.xlsx", load_config(), skip_ai=True, output_dir=tmp_path,
                     reuse_saved=True)
    assert len(reads) == 1
    assert len(tables) == 1


def test_a_second_clean_of_the_same_workbook_is_not_read_again(monkeypatch):
    reads = counting(monkeypatch, clean, "find_kpi_sheet")
    first, _ = clean_workbook(DATA_DIR / "northwind.xlsx")
    second, _ = clean_workbook(DATA_DIR / "northwind.xlsx")
    assert len(reads) == 1
    assert first.equals(second)


# ---------------------------------------------------------------------------
# Never out of date
# ---------------------------------------------------------------------------

def test_an_edited_workbook_is_read_again(tmp_path):
    # Same file name, new contents: the cache is keyed on the contents (their hash), not the name.
    path = tmp_path / "acme.xlsx"
    shutil.copy(DATA_DIR / "northwind.xlsx", path)
    clean_workbook(path)
    shutil.copy(DATA_DIR / "alderpeak.xlsx", path)
    actuals, next_budget = clean_workbook(path)
    expected_actuals, expected_budget = clean_workbook(DATA_DIR / "alderpeak.xlsx")
    assert actuals.equals(expected_actuals)
    assert next_budget.equals(expected_budget)


def test_a_deleted_mapping_file_is_noticed(tmp_path):
    # The confirmed mapping is an input too: without it, the renamed header must stop again.
    path = tmp_path / "northwind.xlsx"
    shutil.copy(DATA_DIR / "northwind.xlsx", path)
    rename_header(path, "Qualified Pipeline", "Sales Funnel")
    saved = mapping.save_mapping(path, {"Sales Funnel": "pipeline"})
    clean_workbook(path)
    saved.unlink()
    with pytest.raises(UnconfirmedMappingError):
        clean_workbook(path)


def test_a_problem_is_not_remembered(tmp_path):
    # A stop isn't cached: once the mapping is confirmed, the same workbook cleans.
    path = tmp_path / "northwind.xlsx"
    shutil.copy(DATA_DIR / "northwind.xlsx", path)
    rename_header(path, "Qualified Pipeline", "Sales Funnel")
    with pytest.raises(UnconfirmedMappingError):
        clean_workbook(path)
    mapping.save_mapping(path, {"Sales Funnel": "pipeline"})
    actuals, _ = clean_workbook(path)
    assert actuals.equals(clean_workbook(DATA_DIR / "northwind.xlsx")[0])


def test_changed_numbers_get_new_metrics():
    actuals, _ = clean_workbook(DATA_DIR / "northwind.xlsx")
    before = compute_metrics(actuals)
    latest = actuals.index[-1]
    actuals.loc[latest, "new_arr"] += 100
    after = compute_metrics(actuals)
    assert after.loc[latest, "ending_arr"] == before.loc[latest, "ending_arr"] + 100
    assert after.equals(metrics.metrics_table(actuals))


def test_a_blanked_input_gets_new_reasons():
    actuals, _ = clean_workbook(DATA_DIR / "alderpeak.xlsx")   # no blank quarter
    latest = actuals.index[-1]
    assert metric_reasons(actuals, compute_metrics(actuals)).loc[latest, "nrr"] is None
    actuals.loc[latest, "churned_arr"] = np.nan
    assert metric_reasons(actuals, compute_metrics(actuals)).loc[latest, "nrr"] == metrics.MISSING_INPUT


def rename_header(path, old, new):
    """Rename one header cell of the KPI tab, in place."""
    from openpyxl import load_workbook
    book = load_workbook(path)
    for row in book["KPI Tracker"].iter_rows(max_row=10):
        for cell in row:
            if cell.value == old:
                cell.value = new
    book.save(path)


# ---------------------------------------------------------------------------
# Safe to change, and bounded
# ---------------------------------------------------------------------------

def test_changing_a_cleaned_table_does_not_change_the_next_one():
    # Three times: the first answer is worked out, the next two come from the cache. Each is changed
    # after it's checked, so the third shows whether changing a cached answer reached the cache.
    for _ in range(3):
        actuals, next_budget = clean_workbook(DATA_DIR / "northwind.xlsx")
        assert actuals.iloc[0, 0] != -1.0
        assert not (next_budget == -1.0).any()
        actuals.iloc[0, 0] = -1.0
        next_budget.iloc[:] = -1.0


def test_changing_a_metric_table_does_not_change_the_next_one():
    actuals, _ = clean_workbook(DATA_DIR / "northwind.xlsx")
    for _ in range(3):   # worked out, then twice from the cache (see the test above)
        metric_table = compute_metrics(actuals)
        reasons = metric_reasons(actuals, compute_metrics(actuals))
        assert metric_table.iloc[-1, 0] != -1.0
        assert reasons.iloc[-1, 0] != "changed"
        metric_table.iloc[-1, 0] = -1.0
        reasons.iloc[-1, 0] = "changed"


def test_the_cache_keeps_at_most_max_entries():
    results = cache.ResultCache(max_entries=2)
    for key in ("a", "b", "c"):
        results.get(key, lambda: key.upper())
    assert len(results) == 2
    assert results.get("a", lambda: "worked out again") == "worked out again"   # the oldest was dropped
    assert results.get("c", lambda: "worked out again") == "C"                  # the newest was kept
