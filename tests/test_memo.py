"""Unit tests for memo.py: the board memo's content, the AI gate, and the saved Word and PDF files.

Small made-up tables stand in for a workbook (the same ones tests/test_build_deck.py uses).
Expected text is written out by hand. No API calls: the analyses are small JSON files written here.
Run from the project folder:  pytest
"""

import datetime
import json
import math
import re

import pandas as pd
import pytest
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

import theme
from analyze import BoardSummary, build_payload
from build_deck import collect_deck_data
from clean import STANDARD_COLUMNS
from diff_runs import HEADING, run_results
from memo import (AI_DRAFTED_LINE, MEMO_UNAVAILABLE, block_texts, memo_analysis, memo_approval, memo_blocks,
                  memo_footer, memo_paths, save_memo, unlisted_numbers, workbook_numbers, write_docx, write_pdf)
from metrics import CONFIG_PATH
from provenance import file_sha256, manifest_path, read_manifest, save_manifest

NAN = math.nan
RUN_DATE = datetime.date(2026, 9, 17)
COMMIT = {"commit": "abc1234", "uncommitted_changes": False}

# Thresholds written out here, so tuning config.yaml never breaks these tests.
TEST_CONFIG = {
    "nrr_min": 1.00, "grr_min": 0.85, "burn_multiple_max": 2.0, "burn_over_budget_max": 0.15,
    "runway_min_months": 12, "cac_payback_max_months": 24, "net_new_arr_vs_budget_min": -0.20,
    "rule_of_40_min": 0.40, "nrr_falling_pipeline_rising_flag": True, "combo_lookback_quarters": 3,
    "combo_min_nrr_drop": 0.01,
}


def three_quarters(blank=None):
    """Three complete quarters, every input 100 (ending cash 1200), with one quarter optionally blank."""
    actuals = pd.DataFrame(100.0, index=["Q4 2025", "Q1 2026", "Q2 2026"], columns=STANDARD_COLUMNS)
    actuals["ending_cash"] = 1200.0
    if blank:
        actuals.loc[blank] = NAN
    return actuals


def memo_data(blank=None):
    return collect_deck_data("Testco", "testco.xlsx", three_quarters(blank), None, TEST_CONFIG)


def summary_dict(headline="Retention is the main question for the board.", question="What drives churn?"):
    """A valid summary. With the defaults it has no numbers, so no number check has anything to reject."""
    return {"headline": headline, "wins": [{"title": "Win title", "detail": "A win detail."}] * 3,
            "risks": [{"title": "Risk title", "detail": "A risk detail."}] * 3,
            "questions": [question, "Where is pipeline coming from?", "How is hiring going?"]}


def summary_from(summary):
    return BoardSummary.model_validate(summary)


def write_analysis(tmp_path, summary, quarter="Q2 2026"):
    path = tmp_path / "testco_analysis.json"
    path.write_text(json.dumps({"summary": summary, "payload": {"company": "Testco", "latest_quarter": quarter},
                                "run_info": {"model": "claude-sonnet-5"}}))
    return path


@pytest.fixture
def payload():
    return build_payload("Testco", three_quarters(), None, TEST_CONFIG)


def all_text(blocks):
    return "\n".join(block_texts(blocks))


# ---------------------------------------------------------------------------
# Content: every computed number is there, with or without the AI text
# ---------------------------------------------------------------------------

def test_the_memo_has_title_quarter_and_every_section_in_order():
    text = all_text(memo_blocks(memo_data(), None))
    assert text.startswith("Testco board update: Q2 2026")
    order = ["Headline", "Key metrics", "Flags", "Data gaps", "Questions for management"]
    positions = [text.index(heading) for heading in order]
    assert positions == sorted(positions)


def test_without_an_analysis_the_memo_says_ai_commentary_unavailable_and_keeps_every_number():
    with_ai = memo_blocks(memo_data(), summary_from(summary_dict()))
    without_ai = memo_blocks(memo_data(), None)
    assert MEMO_UNAVAILABLE in all_text(without_ai)
    assert AI_DRAFTED_LINE not in all_text(without_ai)
    # The computed parts (tables, flags, gaps) are the same blocks either way.
    computed = [block for block in with_ai if not block.get("ai")]
    assert computed == [block for block in without_ai if not block.get("ai")]


def test_the_key_metrics_table_has_every_flag_with_value_threshold_and_status():
    table = next(block for block in memo_blocks(memo_data(), None) if block["kind"] == "table")
    assert table["header"] == ["Metric", "Q2 2026", "Q1 2026", "Budget or threshold", "Status"]
    rows = {cells[0]: cells for cells, _ in table["rows"]}
    # Every input 100: NRR = 1 + 4 x (100 - 100 - 100) / 100 = -300.0%. Runway = 1,200 / (100 / 3) = 36.0 mo.
    assert rows["NRR (annualized)"] == ["NRR (annualized)", "-300.0%", "-300.0%", "trips below 100.0%", "Tripped"]
    assert rows["Runway at current burn"][1:] == ["36.0 mo", "36.0 mo", "trips below 12.0 mo", "Passed"]
    assert rows["Rule of 40"][4] == "Cannot evaluate: no prior period"
    assert rows["NRR falling while pipeline rising"][1:] == ["-", "-", "see Flags below", "Passed"]
    assert len(rows) == 4 + 9   # 4 context rows (Ending ARR, Net new ARR, ARR growth YoY, Gross margin) + 9 flags


def test_tripped_flags_are_listed_with_their_value_and_threshold():
    text = all_text(memo_blocks(memo_data(), None))
    # Every input 100: net new ARR = 100 + 100 - 100 - 100 = 0 while burning 100, so burn multiple is ∞.
    assert "Flags: 3 of 9 flags tripped, 2 cannot evaluate" in text
    for line in ("NRR (annualized): -300.0% (trips below 100.0%)", "GRR (annualized): -700.0% (trips below 85.0%)",
                 "Burn multiple: ∞ (ARR shrank) (trips above 2.00x)", "Rule of 40: Cannot evaluate: no prior period",
                 "Net new ARR vs budget: Cannot evaluate: not meaningful"):
        assert line in text


def test_the_combo_rule_uses_the_metrics_workbook_wording():
    # The memo's words for the combo rule come from the Flags sheet, so check_memo.py finds its numbers there.
    text = all_text(memo_blocks(memo_data(), None))
    # NRR is the same every quarter, so it isn't falling: the rule passes.
    assert ("NRR falling while pipeline rising: Passed "
            "(trips when NRR falls at least 1 pt and pipeline rises at every step, over the last 3 quarters)") in text


def test_a_blank_quarter_is_listed_under_data_gaps():
    text = all_text(memo_blocks(memo_data(blank="Q1 2026"), None))
    assert "Q1 2026: Ending ARR ($K)" in text


def test_no_gaps_says_so():
    assert "None: every metric and flag has the data it needs" in all_text(memo_blocks(memo_data(), None))


def test_no_em_dash_anywhere_in_the_memo():
    # Two labels shared with the deck and Excel carry one (the "Cannot evaluate" status, the "None" gaps line).
    for blocks in (memo_blocks(memo_data(blank="Q1 2026"), None), memo_blocks(memo_data(), None),
                   memo_blocks(memo_data(), summary_from(summary_dict()))):
        assert "—" not in all_text(blocks)


def test_with_a_valid_analysis_the_memo_shows_the_headline_and_questions_but_not_wins_or_risks():
    text = all_text(memo_blocks(memo_data(), summary_from(summary_dict())))
    assert AI_DRAFTED_LINE in text
    assert "Retention is the main question for the board." in text
    assert "What drives churn?" in text and "How is hiring going?" in text
    assert "A win detail." not in text and "A risk detail." not in text


def test_questions_are_bullets_not_numbers():
    # A numbered list would put "1.", "2.", "3." in the memo: numbers that aren't in the metrics workbook.
    blocks = memo_blocks(memo_data(), summary_from(summary_dict()))
    questions = next(block for block in blocks if block["kind"] == "bullets" and block.get("ai"))
    assert questions["items"] == ["What drives churn?", "Where is pipeline coming from?", "How is hiring going?"]


# ---------------------------------------------------------------------------
# The AI gate: only an analysis valid today, whose numbers are all in the metrics workbook
# ---------------------------------------------------------------------------

def test_a_missing_analysis_file_is_unavailable(tmp_path, payload):
    summary, why = memo_analysis(tmp_path / "nothing.json", payload, memo_data())
    assert summary is None and "no analysis file" in why


def test_no_analysis_requested_is_unavailable(payload):
    summary, why = memo_analysis(None, payload, memo_data())
    assert summary is None and why == "no analysis requested"


def test_an_analysis_that_failed_when_made_is_unavailable(tmp_path, payload):
    summary, why = memo_analysis(write_analysis(tmp_path, None), payload, memo_data())
    assert summary is None and "failed validation" in why


def test_an_analysis_for_another_quarter_is_unavailable(tmp_path, payload):
    summary, why = memo_analysis(write_analysis(tmp_path, summary_dict(), quarter="Q1 2026"), payload, memo_data())
    assert summary is None and "Q1 2026" in why


def test_an_analysis_with_a_number_not_in_the_data_is_unavailable(tmp_path, payload):
    path = write_analysis(tmp_path, summary_dict(headline="ARR grew 555.5% this quarter."))
    summary, why = memo_analysis(path, payload, memo_data())
    assert summary is None and "555.5" in why


def test_a_valid_analysis_is_used(tmp_path, payload):
    summary, why = memo_analysis(write_analysis(tmp_path, summary_dict()), payload, memo_data())
    assert why is None and summary.headline == "Retention is the main question for the board."


def test_workbook_numbers_are_the_metric_values_thresholds_quarters_and_flag_counts():
    numbers = workbook_numbers(memo_data())
    assert {100.0, 36.0, 12.0, 2026.0, 9.0, 1.0} <= numbers   # NRR %, runway, a threshold, a year, flag counts
    assert 1200.0 not in numbers   # ending cash is an input: Claude sees it, the metrics workbook doesn't show it


def test_thresholds_count_as_the_workbook_displays_them():
    # The Flags sheet stores 0.15 and shows 15.0%: "over budget (threshold 15.0%)" quotes the workbook.
    # (Bug found building Fernhollow's memo: the raw 0.15 was listed, so a correct question was rejected.)
    numbers = workbook_numbers(memo_data())
    assert {15.0, 85.0, 2.0, 24.0, -20.0, 40.0} <= numbers
    assert 0.15 not in numbers


def test_an_ai_number_that_is_in_the_payload_but_not_the_metrics_workbook_is_listed():
    # Ending cash (1,200) is in Claude's payload, so analyze.py accepts it; the memo can't prove it from the workbook.
    summary = summary_from(summary_dict(question="How long will the 1,200 of cash last?"))
    assert unlisted_numbers(summary, memo_data()) == [1200.0]


def test_the_gate_rejects_an_ai_number_the_metrics_workbook_does_not_show(tmp_path, payload):
    path = write_analysis(tmp_path, summary_dict(question="How long will the 1,200 of cash last?"))
    summary, why = memo_analysis(path, payload, memo_data())
    assert summary is None and "1,200" in why and "metrics workbook" in why


def test_only_the_headline_and_questions_are_gated():
    # The wins and risks aren't in the memo, so a number there that the workbook lacks doesn't matter.
    summary = summary_dict()
    summary["risks"] = [{"title": "Cash", "detail": "Cash is 1,200."}] * 3
    assert unlisted_numbers(summary_from(summary), memo_data()) == []


# ---------------------------------------------------------------------------
# Footer
# ---------------------------------------------------------------------------

def test_footer_says_where_the_memo_came_from_and_that_nobody_reviewed_it():
    footer = memo_footer(memo_data(), RUN_DATE, "claude-sonnet-5", None, COMMIT)
    assert footer == "Fictional data | testco.xlsx | 2026-09-17 | abc1234 | claude-sonnet-5 | AI-drafted | not reviewed"


def test_footer_names_the_reviewer_and_says_when_there_is_no_ai_text():
    approval = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00"}
    footer = memo_footer(memo_data(), RUN_DATE, None, approval, {**COMMIT, "uncommitted_changes": True})
    assert footer == ("Fictional data | testco.xlsx | 2026-09-17 | abc1234* | no AI text | AI-drafted | "
                      "reviewed by Tyler Ho on 2026-09-17")


# ---------------------------------------------------------------------------
# The saved files
# ---------------------------------------------------------------------------

def docx_text(path):
    """Every paragraph and table cell in a Word file, top to bottom (body only)."""
    document = Document(path)
    texts = [item.text for item in document.paragraphs]
    texts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(texts)


def test_the_word_file_holds_the_same_text_and_the_footer(tmp_path):
    blocks = memo_blocks(memo_data(), summary_from(summary_dict()))
    path = write_docx(blocks, "the footer line", tmp_path / "memo.docx")
    text = docx_text(path)
    for piece in block_texts(blocks):
        assert piece in text
    assert Document(path).sections[0].footer.paragraphs[0].text == "the footer line"


def test_the_pdf_holds_the_same_text_on_one_or_two_pages_with_the_footer_on_each(tmp_path):
    blocks = memo_blocks(memo_data(blank="Q1 2026"), summary_from(summary_dict()))
    path = write_pdf(blocks, "the footer line", tmp_path / "memo.pdf")
    pages = PdfReader(path).pages
    assert 1 <= len(pages) <= 2
    text = " ".join(page.extract_text() for page in pages)
    flat = " ".join(text.split())
    for piece in block_texts(blocks):
        assert " ".join(piece.split()) in flat
    assert all("the footer line" in page.extract_text() for page in pages)


def test_a_heading_is_never_left_at_the_foot_of_a_pdf_page(tmp_path):
    # Found on Northwind's first memo: "Questions for management" ended page 1, its questions began page 2.
    # 44 lines leave room on page 1 for the heading but not its list: without keeping the section
    # together the heading ends page 1 (tried by hand; 52 lines pushed the heading over either way).
    filler = [{"kind": "text", "text": "Filler line to push the last section down the page.", "style": "body"}] * 44
    blocks = filler + [{"kind": "heading", "text": "Last heading"},
                       {"kind": "bullets", "items": ["First question?", "Second question?"]}]
    pages = PdfReader(write_pdf(blocks, "footer", tmp_path / "memo.pdf")).pages
    assert len(pages) == 2
    assert "Last heading" in pages[1].extract_text() and "First question?" in pages[1].extract_text()


def test_a_long_footer_wraps_inside_the_pdf_page_instead_of_running_off_it(tmp_path):
    footer = "Fictional data | " + "a very long workbook file name " * 8 + ".xlsx | 2026-09-17"
    page = PdfReader(write_pdf([{"kind": "title", "text": "Title"}], footer, tmp_path / "memo.pdf")).pages[0]
    assert " ".join(footer.split()) in " ".join(page.extract_text().split())   # every word is on the page


def test_word_keeps_headings_with_the_text_under_them(tmp_path):
    blocks = memo_blocks(memo_data(), summary_from(summary_dict()))
    paragraphs = Document(write_docx(blocks, "footer", tmp_path / "memo.docx")).paragraphs
    headings = [item for item in paragraphs if item.text in ("Headline", "Flags: 3 of 9 flags tripped, 2 cannot evaluate",
                                                             "Questions for management")]
    assert len(headings) == 3 and all(item.paragraph_format.keep_with_next for item in headings)


def test_the_pdf_draws_characters_outside_basic_fonts(tmp_path):
    # "∞ (ARR shrank)" must come out as ∞, not a missing-glyph box: the PDF embeds Arial (or DejaVu Sans).
    blocks = [{"kind": "text", "text": "Burn multiple: ∞ (ARR shrank)", "style": "body"}]
    text = PdfReader(write_pdf(blocks, "footer", tmp_path / "memo.pdf")).pages[0].extract_text()
    assert "∞ (ARR shrank)" in text


def pdf_font_names(path):
    """The fonts that draw text in a PDF, e.g. {'/AAAAAA+ArialMT', '/AAAAAA+Arial-BoldMT'}.

    Only fonts used inside a text block that draws something (Tj): reportlab's page setup names
    Helvetica once in an empty block, which draws nothing.
    """
    names = set()
    for page in PdfReader(path).pages:
        fonts = {key: str(font.get_object()["/BaseFont"]) for key, font in page["/Resources"]["/Font"].items()}
        for block in re.findall(r"BT(.*?)ET", page.get_contents().get_data().decode("latin-1"), re.DOTALL):
            if "Tj" in block:
                names |= {fonts[key] for key in re.findall(r"(/F[\w+]+) [\d.]+ Tf", block)}
    return names


def test_the_pdf_is_set_in_arial_or_the_next_installed_font(tmp_path):
    blocks = [{"kind": "title", "text": "Title"}, {"kind": "text", "text": "Body", "style": "body"}]
    names = pdf_font_names(write_pdf(blocks, "footer", tmp_path / "memo.pdf"))
    family, _ = theme.font_file()
    assert names and all(family.replace(" ", "") in name for name in names)


def cell_fill(cell):
    """A Word table cell's shading color, as memo.shade_cell wrote it."""
    shading = cell._tc.tcPr.find(qn("w:shd"))
    return shading.get(qn("w:fill"))


def test_word_table_is_navy_over_white_and_surface_with_status_fills(tmp_path):
    blocks = memo_blocks(memo_data(), summary_from(summary_dict()))
    table = Document(write_docx(blocks, "footer", tmp_path / "memo.docx")).tables[0]
    header = table.cell(0, 0)
    assert cell_fill(header) == "0B2545" and str(header.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert [cell_fill(table.cell(row, 0)) for row in (1, 2)] == ["FFFFFF", "F8FAFC"]
    statuses = {table.cell(row, 0).text: cell_fill(table.cell(row, len(table.columns) - 1))
                for row in range(1, len(table.rows))}
    assert statuses["NRR (annualized)"] == "FDE8E6"                        # tripped: the palette's red fill
    assert statuses["Runway at current burn"] == "EAF6EF"                  # passed: green fill
    body_run = Document(tmp_path / "memo.docx").paragraphs[0].runs[0]
    assert body_run.font.name == "Arial"


def test_save_memo_writes_both_files_beside_the_deck(tmp_path, monkeypatch):
    workbook = tmp_path / "testco.xlsx"
    monkeypatch.setattr("memo.clean_workbook", lambda path: (three_quarters(), None))
    workbook.write_bytes(b"not read: clean_workbook is replaced")
    result = save_memo(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert result["docx"] == tmp_path / "testco_board_memo.docx" and result["docx"].exists()
    assert result["pdf"] == tmp_path / "testco_board_memo.pdf" and result["pdf"].exists()
    assert result["why_unavailable"] == "no analysis requested"
    assert MEMO_UNAVAILABLE in docx_text(result["docx"])


def test_a_memo_rebuilt_on_its_own_updates_an_existing_manifest(tmp_path, monkeypatch):
    workbook = tmp_path / "testco.xlsx"
    workbook.write_bytes(b"not read: clean_workbook is replaced")
    monkeypatch.setattr("memo.clean_workbook", lambda path: (three_quarters(), None))
    save_memo(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert not manifest_path(workbook, tmp_path).exists()   # no manifest: none is made (main.py writes it)

    save_manifest(manifest_path(workbook, tmp_path), {"company": "Testco", "memo": {"ai_text": True}})
    save_memo(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert read_manifest(manifest_path(workbook, tmp_path))["memo"] == {
        "files": ["testco_board_memo.docx", "testco_board_memo.pdf"], "ai_text": False}


def test_the_memo_counts_as_reviewed_only_if_the_approval_covers_it():
    # An approval recorded before memos existed (no "documents") approved a deck; nobody read a memo.
    approval = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00"}
    assert memo_approval(approval) is None
    assert memo_approval({**approval, "documents": ["deck"]}) is None
    assert memo_approval({**approval, "documents": ["deck", "memo"]})["reviewer"] == "Tyler Ho"
    assert memo_approval(None) is None


def test_a_deck_only_approval_leaves_the_memo_footer_not_reviewed(tmp_path, monkeypatch):
    workbook = tmp_path / "testco.xlsx"
    workbook.write_bytes(b"not read: clean_workbook is replaced")
    monkeypatch.setattr("memo.clean_workbook", lambda path: (three_quarters(), None))
    hashes = {"input_sha256": file_sha256(workbook), "config_sha256": file_sha256(CONFIG_PATH)}
    for documents, review in ((["deck"], "not reviewed"), (["deck", "memo"], "reviewed by Tyler Ho on 2026-09-17")):
        approval = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00", "documents": documents, **hashes}
        save_manifest(manifest_path(workbook, tmp_path), {"company": "Testco", "approval": approval})
        result = save_memo(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
        assert Document(result["docx"]).sections[0].footer.paragraphs[0].text.endswith("AI-drafted | " + review)


# ---------------------------------------------------------------------------
# What changed since the last run (Task 10): only when an earlier run exists
# ---------------------------------------------------------------------------

def changes_report():
    """An earlier run's comparison, written by hand: one flag flipped, NRR down 12 points."""
    return {"since": "2026-06-18T09:05:41", "quarter_before": "Q1 2026", "quarter_now": "Q2 2026",
            "flags": [("Runway at current burn", "Passed", "Tripped")],
            "moved": [("NRR (annualized)", "102.0%", "90.0%", "down 12.0 pts")],
            "new_gaps": {}, "resolved_gaps": {"flag: Rule of 40": ["Q1 2026"]}}


def test_a_memo_with_no_earlier_run_has_no_what_changed_section():
    assert HEADING not in all_text(memo_blocks(memo_data(), None))


def test_what_changed_comes_after_the_headline_and_before_key_metrics():
    text = all_text(memo_blocks(memo_data(), None, changes=changes_report()))
    positions = [text.index(words) for words in ("Headline", HEADING, "Key metrics")]
    assert positions == sorted(positions)
    for line in ("Compared with the run of 2026-06-18 09:05, whose latest quarter was Q1 2026 (now Q2 2026).",
                 "Flags that flipped", "Runway at current burn: Tripped (was Passed)",
                 "NRR (annualized): 102.0% to 90.0% (down 12.0 pts)", "Resolved data gaps",
                 "Q1 2026: Flag: Rule of 40"):
        assert line in text


def test_what_changed_uses_the_move_settings_from_config():
    data = memo_data()
    data["config"] = {**TEST_CONFIG, "diff_min_points": 0.1}
    assert "Metrics that moved more than 10.0 pts (percentages)" in all_text(
        memo_blocks(data, None, changes=changes_report()))


def test_save_memo_compares_with_the_run_in_the_manifest(tmp_path, monkeypatch):
    workbook = tmp_path / "testco.xlsx"
    workbook.write_bytes(b"not read: clean_workbook is replaced")
    monkeypatch.setattr("memo.clean_workbook", lambda path: (three_quarters(), None))
    earlier = run_results(memo_data())
    earlier["flags"]["Runway at current burn"] = "Tripped"
    save_manifest(manifest_path(workbook, tmp_path), {"company": "Testco", "run_at": "2026-06-18T09:05:41",
                                                      "results": earlier})
    result = save_memo(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert "Runway at current burn: Passed (was Tripped)" in docx_text(result["docx"])
    pdf_words = " ".join(" ".join(page.extract_text() for page in PdfReader(result["pdf"]).pages).split())
    assert "Runway at current burn: Passed (was Tripped)" in pdf_words


def test_memo_paths():
    docx, pdf = memo_paths("data/northwind.xlsx", "output")
    assert (str(docx), str(pdf)) == ("output/northwind_board_memo.docx", "output/northwind_board_memo.pdf")


def test_an_old_memo_is_deleted_before_a_build_that_fails(tmp_path, monkeypatch):
    # A failed build must never leave last run's memo looking current.
    workbook = tmp_path / "testco.xlsx"
    old_docx, old_pdf = memo_paths(workbook, tmp_path)
    old_docx.write_text("old")
    old_pdf.write_text("old")

    def broken(path):
        raise ValueError("unreadable workbook")
    monkeypatch.setattr("memo.clean_workbook", broken)
    with pytest.raises(ValueError):
        save_memo(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert not old_docx.exists() and not old_pdf.exists()
