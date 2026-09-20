"""Unit tests for build_deck.py: slide text, the AI placeholder, status colors, footer, and loud failures.

Small made-up tables stand in for a workbook. Expected text is written out by hand.
Run from the project folder:  pytest
"""

import ast
import datetime
import json
import math
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Emu

from analyze import build_payload
from build_deck import (AI_DRAFTED_LINE, FICTIONAL_NOTE, NO_AI_MODEL, NOT_APPLICABLE, PLACEHOLDER_NOTE, PLACEHOLDER_TEXT,
                        build_presentation, collect_deck_data, deck_path, flag_count_text, gaps_text, load_analysis,
                        save_deck, shorten_middle, threshold_text)
from clean import STANDARD_COLUMNS
from metrics import CANNOT_EVALUATE, METRIC_LABELS, MISSING_INPUT, PASS, TRIP
from provenance import (NOT_REVIEWED, build_manifest, file_sha256, git_commit, manifest_path, read_manifest,
                        save_manifest)
from text_fit import MIN_FONT_PT, TextDoesNotFitError, text_width_pt

NAN = math.nan
PROJECT_DIR = Path(__file__).parent.parent
RUN_DATE = datetime.date(2026, 9, 17)

# Thresholds written out here, so tuning config.yaml never breaks these tests.
TEST_CONFIG = {
    "nrr_min": 1.00, "grr_min": 0.85, "burn_multiple_max": 2.0, "burn_over_budget_max": 0.15,
    "runway_min_months": 12, "cac_payback_max_months": 24, "net_new_arr_vs_budget_min": -0.20,
    "rule_of_40_min": 0.40, "nrr_falling_pipeline_rising_flag": True, "combo_lookback_quarters": 3,
    "combo_min_nrr_drop": 0.01,
}

# Status fills from theme.STATUS_COLORS, typed from the Task 3 brief: red, green, gray.
# (The Excel workbook keeps Excel's own fills; the deck takes the palette's.)
RED, GREEN, GRAY = "FDE8E6", "EAF6EF", "EDF0F3"


def flag(name, status, reason=None, metric="nrr", threshold=1.0):
    return {"flag": name, "metric": metric, "quarter": "Q2 2026", "value": 0.97, "threshold": threshold,
            "status": status, "reason": reason}


def three_quarters(blank=None):
    """Three complete quarters, every input 100 (ending cash 1200), with one quarter optionally blank."""
    actuals = pd.DataFrame(100.0, index=["Q4 2025", "Q1 2026", "Q2 2026"], columns=STANDARD_COLUMNS)
    actuals["ending_cash"] = 1200.0
    if blank:
        actuals.loc[blank] = NAN
    return actuals


def deck_data(company="Testco", blank=None):
    return collect_deck_data(company, "testco.xlsx", three_quarters(blank), None, TEST_CONFIG)


# 69 words, no numbers in it, so only the length and fit checks can fail it.
PLAIN_DIAGNOSIS = ("Retention is the question this quarter. " + "The installed base shrank while new sales held "
                   "up, and spend stayed where the plan put it. ") * 3


def summary_dict():
    """A valid v5 summary whose only number is $100K, which is what every Testco input is."""
    point = {"title": "Steady base", "detail": "Customers stayed."}
    question = {"theme": "Definitions and assumptions",
                "question": "Which definition gives net burn at $100K this quarter?"}
    return {"headline": "Retention is the main question for the board.",
            "diagnosis": PLAIN_DIAGNOSIS,
            "risks": [point] * 3,
            "questions": [question] * 8}


def shape(slide, name):
    matches = [s for s in slide.shapes if s.name == name]
    assert len(matches) == 1, f"expected one shape named {name!r}, found {len(matches)}"
    return matches[0]


# ---------------------------------------------------------------------------
# Text pieces
# ---------------------------------------------------------------------------

def test_flag_count_text():
    flags = [flag("a", TRIP)] * 6 + [flag("b", PASS)] * 3
    assert flag_count_text(flags) == "6 of 9 flags tripped"


def test_flag_count_text_mentions_flags_that_cannot_be_evaluated():
    flags = [flag("a", TRIP)] * 7 + [flag("b", PASS), flag("c", CANNOT_EVALUATE, MISSING_INPUT)]
    assert flag_count_text(flags) == "7 of 9 flags tripped, 1 cannot evaluate"


def test_threshold_text_says_which_way_the_flag_trips():
    assert threshold_text(flag("NRR (annualized)", TRIP, metric="nrr", threshold=1.0)) == "trips below 100.0%"
    assert threshold_text(flag("Burn multiple", TRIP, metric="burn_multiple", threshold=2.0)) == "trips above 2.00x"


def test_gaps_text_none():
    assert gaps_text({}) == "None: every metric and flag has the data it needs"


def test_gaps_text_groups_metrics_by_the_quarters_they_miss():
    gaps = {"ending_arr": ["Q1 2025"], "nrr": ["Q1 2025"], "arr_qoq": ["Q1 2025", "Q2 2025"],
            "flag: Rule of 40": ["Q2 2026"]}
    assert gaps_text(gaps) == ("Q1 2025: Ending ARR ($K), NRR (annualized); "
                               "Q1 2025 + Q2 2025: ARR growth QoQ; "
                               "Q2 2026: Flag: Rule of 40")


# ---------------------------------------------------------------------------
# Loading the AI analysis: anything not valid today shows the placeholder
# ---------------------------------------------------------------------------

@pytest.fixture
def payload():
    return build_payload("Testco", three_quarters(), None, TEST_CONFIG)


def write_analysis(tmp_path, summary, company="Testco", quarter="Q2 2026"):
    path = tmp_path / "testco_analysis.json"
    path.write_text(json.dumps({"summary": summary, "payload": {"company": company, "latest_quarter": quarter}}))
    return path


def test_valid_analysis_loads(tmp_path, payload):
    summary, reason = load_analysis(write_analysis(tmp_path, summary_dict()), payload)
    assert reason is None and summary.headline == summary_dict()["headline"]


def test_missing_analysis_file(tmp_path, payload):
    summary, reason = load_analysis(tmp_path / "nothing.json", payload)
    assert summary is None and "no analysis file" in reason


def test_analysis_that_failed_validation_when_made(tmp_path, payload):
    # analyze.py saves "summary": null when both attempts failed.
    summary, reason = load_analysis(write_analysis(tmp_path, None), payload)
    assert summary is None and "failed validation" in reason


def test_analysis_that_is_not_json(tmp_path, payload):
    path = tmp_path / "bad.json"
    path.write_text("{not json")
    summary, reason = load_analysis(path, payload)
    assert summary is None and "not valid JSON" in reason


def test_analysis_with_the_wrong_shape(tmp_path, payload):
    broken = summary_dict()
    del broken["questions"]
    summary, reason = load_analysis(write_analysis(tmp_path, broken), payload)
    assert summary is None and "shape" in reason


def test_analysis_for_another_quarter_is_stale(tmp_path, payload):
    summary, reason = load_analysis(write_analysis(tmp_path, summary_dict(), quarter="Q1 2026"), payload)
    assert summary is None and "Q1 2026" in reason


def test_analysis_for_another_company(tmp_path, payload):
    summary, reason = load_analysis(write_analysis(tmp_path, summary_dict(), company="Othercorp"), payload)
    assert summary is None and "Othercorp" in reason


def test_analysis_with_a_number_not_in_todays_data(tmp_path, payload):
    invented = summary_dict()
    invented["headline"] = "ARR grew 12345% this quarter."
    summary, reason = load_analysis(write_analysis(tmp_path, invented), payload)
    assert summary is None and "12345" in reason


def test_analysis_too_long_for_the_slides_gives_the_placeholder(tmp_path, payload):
    # Decision K: the numbers are valid, so the deck is still built - it just can't carry this text
    # (review finding 1a). Before this, the build stopped and the company had no deck at all.
    too_long = summary_dict()
    too_long["risks"] = [{"title": "Steady base", "detail": "customers " * 45}] * 3
    summary, reason = load_analysis(write_analysis(tmp_path, too_long), payload)
    assert summary is None and "does not fit slide 3" in reason


def test_analysis_with_seven_questions(tmp_path, payload):
    short = summary_dict()
    short["questions"] = short["questions"][:7]
    summary, reason = load_analysis(write_analysis(tmp_path, short), payload)
    assert summary is None and "questions must have 8 to 10 items" in reason


# ---------------------------------------------------------------------------
# The whole deck
# ---------------------------------------------------------------------------

def all_text(slide):
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def build(tmp_path, summary=None, approval=None, model=None, draft=False, **data_options):
    from analyze import BoardSummary
    parsed = BoardSummary.model_validate(summary) if summary else None
    return build_presentation(deck_data(**data_options), parsed, RUN_DATE, tmp_path, approval=approval, model=model,
                              draft=draft)


def test_four_slides_in_order(tmp_path):
    slides = build(tmp_path).slides
    assert [slide.shapes.title.text for slide in slides] == [
        "Testco: key metrics, Q2 2026 vs Q1 2026", "ARR and cash, Q4 2025 to Q2 2026",
        "Risks and flags, Q2 2026", "AI commentary, Q2 2026"]


def test_placeholder_on_slide_4_when_there_is_no_analysis(tmp_path):
    slide = build(tmp_path).slides[3]
    assert shape(slide, "Headline").text_frame.text == PLACEHOLDER_TEXT
    assert shape(slide, "AI note").text_frame.text == PLACEHOLDER_NOTE
    # Nothing on the slide is AI-drafted, so the "AI-drafted ... review before use" line isn't shown.
    assert AI_DRAFTED_LINE not in all_text(slide)


def test_analysis_text_on_slide_4(tmp_path):
    slide = build(tmp_path, summary=summary_dict()).slides[3]
    assert shape(slide, "Headline").text_frame.text == summary_dict()["headline"]
    assert shape(slide, "Diagnosis").text_frame.text == summary_dict()["diagnosis"]
    questions = all_text(slide)
    assert all(item["question"] in questions for item in summary_dict()["questions"])
    assert "Definitions and assumptions" in questions, "the theme heads its questions"
    assert PLACEHOLDER_TEXT not in questions


def test_the_ai_line_sits_under_the_title_and_above_the_headline(tmp_path):
    slide = build(tmp_path, summary=summary_dict()).slides[3]
    line = shape(slide, "AI-drafted line")
    assert line.text_frame.text == "AI-drafted from computed metrics - review before use"
    assert slide.shapes.title.top + slide.shapes.title.height <= line.top
    assert line.top + line.height <= shape(slide, "Headline").top


def test_the_risks_are_on_slide_3_under_the_data_gaps(tmp_path):
    summary = summary_dict()
    summary["risks"] = [{"title": "Retention slipping", "detail": "Customers left."}] * 3
    slides = build(tmp_path, summary=summary).slides
    column = shape(slides[2], "Data gaps and risks").text_frame.text
    assert column.index("Data gaps") < column.index("Risks (AI-drafted"), "the gaps come first"
    assert "Retention slipping" in column and "Customers left." in column
    assert "Retention slipping" not in all_text(slides[3]), "the risks are not on slide 4 as well"


def test_without_an_analysis_slide_3_says_the_risks_are_unavailable(tmp_path):
    column = shape(build(tmp_path).slides[2], "Data gaps and risks").text_frame.text
    assert "Risks (AI-drafted" in column and PLACEHOLDER_TEXT in column


def test_the_flag_count_is_on_the_risks_slide(tmp_path):
    # It was on the Summary slide, which is gone. Python counts it, so it shows with or without AI text.
    text = shape(build(tmp_path).slides[2], "Risks and flags").text_frame.text
    assert "Tripped flags (" in text and "flags tripped" in text


def test_footer_on_every_slide(tmp_path):
    for slide in build(tmp_path).slides:
        footer = shape(slide, "Footer").text_frame.text
        assert FICTIONAL_NOTE in footer and "testco.xlsx" in footer and "2026-09-17" in footer


def status_fills(table):
    """{row label: status cell fill hex} for rows that are flags."""
    fills = {}
    for row in list(table.rows)[1:]:
        cell = row.cells[len(row.cells) - 1]
        if cell.fill.type is not None and cell.text != NOT_APPLICABLE:
            fills[row.cells[0].text] = str(cell.fill.fore_color.rgb)
    return fills


def test_status_cells_are_red_green_or_gray(tmp_path):
    # Every input 100: NRR = 1 + 4 * (100 - 100 - 100) / 100 = -300% -> trips.
    # GRR = 1 - 4 * 200 / 100 = -700% -> trips. Net new ARR = 0 while burning -> burn multiple ∞ -> trips.
    # Runway = 1200 / (100 / 3) = 36 months -> passes. 3 quarters: Rule of 40 needs 4 back -> can't evaluate.
    table = shape(build(tmp_path).slides[0], "KPI table").table
    fills = status_fills(table)
    assert fills["NRR (annualized)"] == RED
    assert fills["Burn multiple"] == RED
    assert fills["Runway at current burn"] == GREEN
    assert fills["Rule of 40"] == GRAY


def status_text_colors(table):
    """{row label: status cell text color} for rows that are flags."""
    last = len(table.columns) - 1
    return {row.cells[0].text: str(row.cells[last].text_frame.paragraphs[0].runs[0].font.color.rgb)
            for row in list(table.rows)[1:] if row.cells[last].text != NOT_APPLICABLE}


def test_status_text_is_red_green_or_slate_and_the_table_is_navy_over_white_and_surface(tmp_path):
    table = shape(build(tmp_path).slides[0], "KPI table").table
    colors = status_text_colors(table)
    assert colors["NRR (annualized)"] == "C0392B" and colors["Runway at current burn"] == "1A7742"
    assert colors["Rule of 40"] == "334155"
    header = table.cell(0, 0)
    assert str(header.fill.fore_color.rgb) == "0B2545"
    assert str(header.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    stripes = [str(table.cell(row, 0).fill.fore_color.rgb) for row in (1, 2)]
    assert stripes == ["FFFFFF", "F8FAFC"]                     # white, then the surface color


def test_type_sizes_on_the_slides(tmp_path):
    # Title 28, section 20, body 15, caption 13, table 14; the footer stays at the 12 pt floor.
    deck = build(tmp_path)
    assert run_sizes(deck.slides[0].shapes.title) == {28}
    assert max(run_sizes(shape(deck.slides[2], "Data gaps and risks"))) == 20   # its heading
    assert min(run_sizes(shape(deck.slides[2], "Data gaps and risks"))) == 15   # its lines
    slide_4 = build(tmp_path, summary=summary_dict()).slides[3]
    assert run_sizes(shape(slide_4, "AI-drafted line")) == {13}
    assert run_sizes(shape(slide_4, "Headline")) == {20}
    assert run_sizes(shape(deck.slides[0], "Footer")) == {12}
    table = shape(deck.slides[0], "KPI table").table
    assert {run.font.size.pt for row in table.rows for cell in row.cells
            for run in cell.text_frame.paragraphs[0].runs} <= {14, 13, 12}   # 14, shrunk only if it must


def test_kpi_table_shows_why_a_value_is_missing(tmp_path):
    # Q1 2026 blank: the prior-quarter column says "data missing"; Rule of 40 says no prior period.
    table = shape(build(tmp_path, blank="Q1 2026").slides[0], "KPI table").table
    rows = {row.cells[0].text: [cell.text for cell in row.cells] for row in table.rows}
    assert rows["NRR (annualized)"][2] == "data missing"
    assert rows["Rule of 40"][1] == "n/a (no prior period)"
    assert rows["Burn multiple"][1] == "∞ (ARR shrank)"
    assert rows["Rule of 40"][4] == "Cannot evaluate: no prior period"


def test_risks_slide_lists_tripped_flags_combo_and_data_gaps(tmp_path):
    slide = build(tmp_path, blank="Q1 2026").slides[2]
    text = shape(slide, "Risks and flags").text_frame.text
    assert "NRR (annualized): -300.0% (trips below 100.0%)" in text
    assert "NRR falling while pipeline rising: Cannot evaluate: missing input" in text
    gaps = shape(slide, "Data gaps and risks").text_frame.text
    assert gaps.startswith("Data gaps") and "Q1 2026 + Q2 2026: ARR growth QoQ" in gaps


def run_sizes(text_shape):
    return {run.font.size.pt for item in text_shape.text_frame.paragraphs for run in item.runs}


def test_side_by_side_columns_use_the_same_font_sizes(tmp_path):
    # Slide 3's flags column is long and its gaps column short, so the gaps column must shrink with it.
    slide = build(tmp_path, blank="Q1 2026").slides[2]
    flags, gaps = shape(slide, "Risks and flags"), shape(slide, "Data gaps and risks")
    assert max(run_sizes(flags)) == max(run_sizes(gaps))   # both headings at the same size
    assert min(run_sizes(flags)) == min(run_sizes(gaps))   # and both bodies too


def test_charts_slide_has_two_pictures(tmp_path):
    slide = build(tmp_path).slides[1]
    assert shape(slide, "ARR chart").shape_type == 13   # MSO_SHAPE_TYPE.PICTURE
    assert shape(slide, "Cash chart").shape_type == 13


def test_text_too_long_for_its_box_fails_loudly(tmp_path):
    with pytest.raises(TextDoesNotFitError, match="Slide 1"):
        build(tmp_path, company="Testco " * 60)


# ---------------------------------------------------------------------------
# Provenance and review status in the footer; the DRAFT watermark only with --draft
# ---------------------------------------------------------------------------

APPROVED = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00",
            "input_sha256": "input-hash", "config_sha256": "config-hash"}


def test_the_footer_carries_the_commit_and_the_model(tmp_path):
    footer = shape(build(tmp_path, model="claude-sonnet-5").slides[0], "Footer").text_frame.text
    assert git_commit()["commit"] in footer
    assert "claude-sonnet-5" in footer
    assert FICTIONAL_NOTE in footer and "testco.xlsx" in footer and RUN_DATE.isoformat() in footer


def test_the_footer_says_when_there_is_no_ai_text(tmp_path):
    assert NO_AI_MODEL in shape(build(tmp_path).slides[0], "Footer").text_frame.text


def footer(presentation):
    return shape(presentation.slides[0], "Footer").text_frame.text


def watermarks(presentation):
    return [item.name for slide in presentation.slides for item in slide.shapes if item.name == "Watermark"]


def test_no_watermark_by_default(tmp_path):
    # The watermark is opt-in (--draft): the footer already says whether anyone has reviewed the deck.
    assert watermarks(build(tmp_path)) == []


def test_draft_watermarks_every_slide_of_an_unreviewed_deck(tmp_path):
    for slide in build(tmp_path, draft=True).slides:
        assert shape(slide, "Watermark").text_frame.text == NOT_REVIEWED


def test_the_watermark_is_drawn_over_the_content_not_under_it(tmp_path):
    # Slides 1 and 2 are covered by an opaque table and two chart images, so a watermark added
    # first would be invisible on exactly the slides that carry the numbers.
    for slide in build(tmp_path, draft=True).slides:
        assert slide.shapes[-1].name == "Watermark"


def test_draft_never_stamps_an_approved_deck(tmp_path):
    # The stamp says NOT REVIEWED; on a deck someone approved, that would be false.
    assert watermarks(build(tmp_path, approval=APPROVED, draft=True)) == []


def test_the_footer_says_an_unreviewed_deck_is_not_reviewed(tmp_path):
    assert footer(build(tmp_path, model="claude-sonnet-5")).endswith(" | claude-sonnet-5 | AI-drafted | not reviewed")


def test_the_footer_names_the_reviewer_and_the_day(tmp_path):
    # approved_at is to the second; the footer shows the day only.
    assert footer(build(tmp_path, approval=APPROVED)).endswith(" | AI-drafted | reviewed by Tyler Ho on 2026-09-17")


def test_the_footer_is_on_every_slide_with_the_same_review_status(tmp_path):
    texts = {shape(slide, "Footer").text_frame.text for slide in build(tmp_path, approval=APPROVED).slides}
    assert len(texts) == 1


def footer_room_pt(presentation):
    """Width the footer text has on one line: the footer box less its two side margins."""
    box = shape(presentation.slides[0], "Footer")
    return Emu(box.width - box.text_frame.margin_left - box.text_frame.margin_right).pt


def test_the_footer_fits_one_line_with_the_longest_file_name_and_a_reviewer(tmp_path):
    # Worst case today: the longest workbook name, the default model, and a reviewer (plus the "*"
    # for uncommitted code whenever these tests run mid-edit). Measured on width, not only height.
    presentation = build_presentation(
        collect_deck_data("Fernhollow", "fernhollow.xlsx", three_quarters(), None, TEST_CONFIG), None, RUN_DATE,
        tmp_path, approval=APPROVED, model="claude-sonnet-5")
    text = footer(presentation)
    assert "reviewed by Tyler Ho on 2026-09-17" in text
    assert text_width_pt(text, MIN_FONT_PT) <= footer_room_pt(presentation)


def test_a_reviewer_name_too_long_for_the_footer_leaves_the_name_to_the_manifest(tmp_path):
    # A name is typed by a person, so it can be any length. The footer keeps the review day on one
    # line; who reviewed stays in the manifest, which approve.py wrote.
    long_name = {**APPROVED, "reviewer": "Alexandra Richardson-Whitfield of the Portfolio Operations Team"}
    presentation = build(tmp_path, approval=long_name, model="claude-sonnet-5")
    text = footer(presentation)
    assert text.endswith(" | AI-drafted | reviewed on 2026-09-17")
    assert text_width_pt(text, MIN_FONT_PT) <= footer_room_pt(presentation)


def test_a_file_name_too_long_for_the_footer_is_shortened_not_a_failed_deck(tmp_path):
    # The web page takes whatever name the file has. Before the fix, this name stopped the whole
    # deck with "Slide 1, Footer: text doesn't fit". The footer keeps the start and the extension.
    long_file = "Northwind KPI workbook Q2 2026 final version for the board.xlsx"
    presentation = build_presentation(collect_deck_data("Testco", long_file, three_quarters(), None, TEST_CONFIG),
                                      None, RUN_DATE, tmp_path, approval=APPROVED, model="claude-sonnet-5")
    text = footer(presentation)
    assert text.startswith(f"{FICTIONAL_NOTE} | North") and "….xlsx |" in text
    assert text.endswith(" | AI-drafted | reviewed by Tyler Ho on 2026-09-17")
    assert text_width_pt(text, MIN_FONT_PT) <= footer_room_pt(presentation)


def test_a_file_name_that_fits_is_never_shortened():
    assert shorten_middle("fernhollow.xlsx", lambda text: True) == "fernhollow.xlsx"


def northwind_manifest(tmp_path, input_sha256=None, config_sha256=None):
    """A manifest for Northwind in tmp_path, approved against the hashes given (or today's)."""
    workbook, config = PROJECT_DIR / "data" / "northwind.xlsx", PROJECT_DIR / "config.yaml"
    approval = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00",
                "input_sha256": input_sha256 or file_sha256(workbook),
                "config_sha256": config_sha256 or file_sha256(config)}
    manifest = build_manifest("Northwind", workbook, config, "northwind_board_pack.pptx", {},
                              ai_text=False, approval=approval)
    save_manifest(manifest_path(workbook, tmp_path), manifest)
    return workbook


def watermarks_in(path):
    return watermarks(Presentation(path))


def test_rebuilding_an_approved_deck_drops_the_watermark(tmp_path):
    workbook = northwind_manifest(tmp_path)
    path, _ = save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path, draft=True)
    assert watermarks_in(path) == []


def test_the_footer_reads_the_review_status_from_the_manifest(tmp_path):
    workbook = northwind_manifest(tmp_path)
    path, _ = save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert footer(Presentation(path)).endswith(" | AI-drafted | reviewed by Tyler Ho on 2026-09-17")


def test_without_a_manifest_the_footer_says_not_reviewed_and_there_is_no_watermark(tmp_path):
    path, _ = save_deck(PROJECT_DIR / "data" / "northwind.xlsx", TEST_CONFIG, None, run_date=RUN_DATE,
                        output_dir=tmp_path)
    assert footer(Presentation(path)).endswith(" | AI-drafted | not reviewed")
    assert watermarks_in(path) == []


def test_a_rebuilt_deck_updates_what_the_manifest_says_about_it(tmp_path):
    # The manifest has to describe the deck on disk, or an audit trail is worse than none.
    workbook = northwind_manifest(tmp_path, input_sha256="the-hash-of-an-older-workbook")
    save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    saved = read_manifest(manifest_path(workbook, tmp_path))
    assert saved["deck"]["status"] == NOT_REVIEWED      # it was "approved by Tyler Ho ..." before
    assert saved["approval"]["reviewer"] == "Tyler Ho"  # the record of the approval is kept


def test_a_workbook_changed_since_approval_goes_back_to_draft(tmp_path):
    # Nobody has reviewed a deck built from numbers that arrived after the approval.
    workbook = northwind_manifest(tmp_path, input_sha256="the-hash-of-an-older-workbook")
    path, _ = save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path, draft=True)
    assert len(watermarks_in(path)) == 4
    assert footer(Presentation(path)).endswith(" | AI-drafted | not reviewed")


def test_changed_thresholds_send_an_approved_deck_back_to_draft(tmp_path):
    workbook = northwind_manifest(tmp_path, config_sha256="the-hash-of-older-thresholds")
    path, _ = save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path, draft=True)
    assert len(watermarks_in(path)) == 4
    assert footer(Presentation(path)).endswith(" | AI-drafted | not reviewed")


def test_command_line_draft_flag(tmp_path):
    # build_deck.py --draft turns the watermark on; without it the deck has none.
    import build_deck
    workbook = str(PROJECT_DIR / "data" / "northwind.xlsx")
    build_deck.main([workbook, "--no-analysis", "--draft"], output_dir=tmp_path)
    assert len(watermarks_in(tmp_path / "northwind_board_pack.pptx")) == 4
    build_deck.main([workbook, "--no-analysis"], output_dir=tmp_path)
    assert watermarks_in(tmp_path / "northwind_board_pack.pptx") == []


def test_deck_path():
    assert deck_path("data/northwind.xlsx", "output") == Path("output/northwind_board_pack.pptx")


# ---------------------------------------------------------------------------
# The appendix slide (--appendix only): every metric for every quarter
# ---------------------------------------------------------------------------
# Worked out by hand for three_quarters() (every input 100, ending cash 1,200):
#   Ending ARR = 100 + 100 + 100 - 100 - 100 = 100; net new ARR = 0
#   NRR = 1 + 4 x (100 - 100 - 100) / 100 = -300.0%, below the 100% threshold: tripped in every quarter
#   Burn multiple: net new ARR 0 with burn 100 -> infinite ("ARR shrank"), tripped
#   Runway = 1,200 / (100 / 3) = 36.0 mo; CAC payback = 100 / (100 x 100%) x 12 = 12.0 mo
#   Net new ARR vs budget: budgeted net new ARR is 100 - 100 = 0 -> not meaningful (n/m)
#   Every QoQ metric in Q4 2025, the first quarter, has no prior period (n/a)

def appendix(tmp_path, **data_options):
    """The appendix slide of a deck built with appendix=True."""
    presentation = build_presentation(deck_data(**data_options), None, RUN_DATE, tmp_path, appendix=True)
    return presentation.slides[-1]


def appendix_rows(slide):
    return [[cell.text for cell in row.cells] for row in shape(slide, "Appendix table").table.rows]


def appendix_row(slide, label):
    return next(row for row in appendix_rows(slide) if row[0] == label)


def cell_fill(slide, label, column):
    table = shape(slide, "Appendix table").table
    row_number = [row[0] for row in appendix_rows(slide)].index(label)
    return str(table.cell(row_number, column).fill.fore_color.rgb)


def test_no_appendix_by_default(tmp_path):
    assert len(build(tmp_path).slides) == 4


def test_the_appendix_adds_one_slide_after_the_four(tmp_path):
    presentation = build_presentation(deck_data(), None, RUN_DATE, tmp_path, appendix=True)
    titles = [slide.shapes.title.text for slide in presentation.slides]
    assert len(titles) == 5
    assert titles[:4] == [slide.shapes.title.text for slide in build(tmp_path).slides]
    assert titles[4] == "Appendix: every metric, Q4 2025 to Q2 2026"


def test_the_appendix_has_every_metric_and_every_quarter(tmp_path):
    rows = appendix_rows(appendix(tmp_path))
    assert rows[0] == ["Metric", "Q4 2025", "Q1 2026", "Q2 2026"]
    assert [row[0] for row in rows[1:]] == list(METRIC_LABELS.values())   # all 19, in metrics.py's order


def test_the_appendix_shows_numbers_in_the_deck_formats(tmp_path):
    slide = appendix(tmp_path)
    assert appendix_row(slide, "Ending ARR ($K)") == ["Ending ARR ($K)", "100", "100", "100"]
    assert appendix_row(slide, "NRR (annualized)") == ["NRR (annualized)", "-300.0%", "-300.0%", "-300.0%"]
    assert appendix_row(slide, "Runway at current burn") == ["Runway at current burn", "36.0 mo", "36.0 mo", "36.0 mo"]
    assert appendix_row(slide, "CAC payback") == ["CAC payback", "12.0 mo", "12.0 mo", "12.0 mo"]


def test_the_appendix_uses_short_marks_that_the_key_explains(tmp_path):
    # "n/a (no prior period)" is wider than a quarter column, so the cell shows the mark and the key
    # under the table says what it means. "data missing" fits, so it is never shortened.
    slide = appendix(tmp_path, blank="Q1 2026")
    assert appendix_row(slide, "ARR growth QoQ") == ["ARR growth QoQ", "n/a", "data missing", "data missing"]
    assert appendix_row(slide, "Burn multiple") == ["Burn multiple", "∞", "data missing", "∞"]
    assert appendix_row(slide, "Net new ARR vs budget")[1:] == ["n/a", "data missing", "data missing"]
    key = shape(slide, "Appendix key").text_frame.text
    assert key == ("Key: n/a = no prior period; ∞ in Burn multiple = ARR shrank; "
                   "gray = data missing; red, bold = flag tripped")


def test_the_key_lists_only_the_marks_on_the_slide(tmp_path):
    key = shape(appendix(tmp_path), "Appendix key").text_frame.text
    assert key == ("Key: n/a = no prior period; n/m = not meaningful (see the metrics workbook); "
                   "∞ in Burn multiple = ARR shrank; red, bold = flag tripped")


def test_appendix_colors_gaps_gray_and_tripped_cells_red_and_bold(tmp_path):
    slide = appendix(tmp_path, blank="Q1 2026")
    assert cell_fill(slide, "Ending ARR ($K)", 2) == GRAY          # Q1 2026 is blank
    assert cell_fill(slide, "NRR (annualized)", 1) == RED          # trips below 100%
    assert cell_fill(slide, "Runway at current burn", 1) not in (RED, GRAY, GREEN)   # passes: no color
    table = shape(slide, "Appendix table").table
    nrr_row = [row[0] for row in appendix_rows(slide)].index("NRR (annualized)")
    assert table.cell(nrr_row, 1).text_frame.paragraphs[0].runs[0].font.bold   # not color alone
    runway_row = [row[0] for row in appendix_rows(slide)].index("Runway at current burn")
    assert not table.cell(runway_row, 1).text_frame.paragraphs[0].runs[0].font.bold


def test_the_appendix_fits_at_the_twelve_point_floor_or_above(tmp_path):
    slide = appendix(tmp_path, blank="Q1 2026")
    sizes = {run.font.size.pt for row in shape(slide, "Appendix table").table.rows for cell in row.cells
             for run in cell.text_frame.paragraphs[0].runs}
    assert min(sizes) >= MIN_FONT_PT
    assert shape(slide, "Footer").text_frame.text == footer(build(tmp_path))


def ten_quarters():
    """Ten complete quarters, Q1 2024 to Q2 2026, every input 100 (ending cash 1,200)."""
    quarters = [f"Q{quarter} {year}" for year in (2024, 2025, 2026) for quarter in (1, 2, 3, 4)][:10]
    actuals = pd.DataFrame(100.0, index=quarters, columns=STANDARD_COLUMNS)
    actuals["ending_cash"] = 1200.0
    return actuals


def test_the_appendix_shows_the_last_eight_quarters_of_a_longer_workbook(tmp_path):
    # Nine or more quarter columns don't fit at 12 pt; the board reads the recent two years.
    data = collect_deck_data("Testco", "testco.xlsx", ten_quarters(), None, TEST_CONFIG)
    slide = build_presentation(data, None, RUN_DATE, tmp_path, appendix=True).slides[-1]
    assert appendix_rows(slide)[0][1:] == ["Q3 2024", "Q4 2024", "Q1 2025", "Q2 2025",
                                           "Q3 2025", "Q4 2025", "Q1 2026", "Q2 2026"]
    assert slide.shapes.title.text == "Appendix: every metric, Q3 2024 to Q2 2026"


def test_an_appendix_that_cannot_fit_stops_naming_the_slide(tmp_path, monkeypatch):
    import build_deck
    monkeypatch.setattr(build_deck, "appendix_text", lambda data, metric, quarter: "much too long " * 8)
    with pytest.raises(TextDoesNotFitError, match="Slide 5, Appendix table"):
        appendix(tmp_path)


def test_the_appendix_gets_the_watermark_too(tmp_path):
    presentation = build_presentation(deck_data(), None, RUN_DATE, tmp_path, draft=True, appendix=True)
    assert len(watermarks(presentation)) == 5


def test_command_line_appendix_flag(tmp_path):
    import build_deck
    workbook = str(PROJECT_DIR / "data" / "northwind.xlsx")
    build_deck.main([workbook, "--no-analysis", "--appendix"], output_dir=tmp_path)
    assert len(Presentation(tmp_path / "northwind_board_pack.pptx").slides) == 5
    build_deck.main([workbook, "--no-analysis"], output_dir=tmp_path)
    assert len(Presentation(tmp_path / "northwind_board_pack.pptx").slides) == 4


def test_the_manifest_records_whether_the_deck_has_the_appendix(tmp_path):
    workbook = northwind_manifest(tmp_path)
    save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path, appendix=True)
    assert read_manifest(manifest_path(workbook, tmp_path))["deck"]["appendix"] is True
    save_deck(workbook, TEST_CONFIG, None, run_date=RUN_DATE, output_dir=tmp_path)
    assert read_manifest(manifest_path(workbook, tmp_path))["deck"]["appendix"] is False


# ---------------------------------------------------------------------------
# No number typed by hand
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("file_name", ["build_deck.py", "charts.py", "memo.py"])
def test_no_digit_in_any_text_written_in_the_code(file_name):
    # Every number a reader sees comes from metrics, flags, config or the analysis - never from a
    # string typed into the code. (Layout sizes like 0.2 inches are code numbers, not text.)
    tree = ast.parse((PROJECT_DIR / file_name).read_text())
    docstrings = {id(node.body[0].value) for node in ast.walk(tree)
                  if isinstance(node, (ast.Module, ast.FunctionDef, ast.ClassDef)) and node.body
                  and isinstance(node.body[0], ast.Expr) and isinstance(node.body[0].value, ast.Constant)}
    typed = [node.value for node in ast.walk(tree)
             if isinstance(node, ast.Constant) and isinstance(node.value, str) and id(node) not in docstrings]
    with_digits = [text for text in typed if any(character.isdigit() for character in text)]
    assert with_digits == []
