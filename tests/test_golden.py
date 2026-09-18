"""Tests for the golden files (final Task 8): approved text dumps of every company's deck, memo and metrics workbook.

The main test rebuilds each company's three outputs into a temporary folder (from the committed
analysis fixtures in tests/golden/analysis/, so no API call), dumps them to text with golden.py,
and compares the text with the approved file in tests/golden/. Any difference fails with a
readable diff and the command to accept it on purpose. The other tests check the helper itself:
that a diff names the changed lines, that the dumps don't change from day to day, and that they
hold what the value checks (check_deck.py and friends) don't look at.

Run from the project folder:  pytest
"""

import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import golden

PROJECT_DIR = Path(__file__).parent.parent
COMPANIES = ["alderpeak", "fernhollow", "northwind"]


@pytest.fixture(scope="module")
def rebuilt(tmp_path_factory):
    """Every company's outputs built once for this file: {company: {"deck": text, "memo": text, "metrics": text}}."""
    folder = tmp_path_factory.mktemp("golden_build")
    return {company: golden.build_dumps(company, folder / company) for company in COMPANIES}


# ---------------------------------------------------------------------------
# The goldens themselves
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("kind", golden.KINDS)
@pytest.mark.parametrize("company", COMPANIES)
def test_the_output_matches_its_approved_golden(rebuilt, company, kind):
    problem = golden.compare(rebuilt[company][kind], golden.golden_path(company, kind))
    if problem:
        pytest.fail(problem, pytrace=False)


def test_every_company_in_data_has_three_goldens_and_an_analysis_fixture():
    assert golden.companies() == COMPANIES
    for company in COMPANIES:
        assert golden.fixture_path(company).exists()
        for kind in golden.KINDS:
            assert golden.golden_path(company, kind).exists()


def test_no_stray_golden_files():
    """A golden for a company or output that no longer exists would never be compared: it must go."""
    expected = {golden.golden_path(company, kind).name for company in COMPANIES for kind in golden.KINDS}
    assert {path.name for path in golden.GOLDEN_DIR.glob("*.txt")} == expected


@pytest.mark.parametrize("company", COMPANIES)
def test_the_saved_analysis_still_passes_the_deck_and_memo_checks(rebuilt, company):
    """If a fixture stopped matching today's numbers, the goldens would quietly show the placeholder instead."""
    assert "AI summary unavailable" not in rebuilt[company]["deck"]
    assert "AI summary unavailable" not in rebuilt[company]["memo"]
    assert "AI-drafted from computed metrics" in rebuilt[company]["deck"]


# ---------------------------------------------------------------------------
# The dumps don't change from day to day, and hold what value checks skip
# ---------------------------------------------------------------------------

def test_the_footer_uses_the_fixed_date_and_commit_not_todays(rebuilt):
    deck = rebuilt["northwind"]["deck"]
    assert golden.RUN_DATE.isoformat() in deck and golden.COMMIT["commit"] in deck
    assert datetime.date.today().isoformat() not in deck   # RUN_DATE is in the past, so today never matches it

def test_building_twice_gives_the_same_text(rebuilt, tmp_path):
    again = golden.build_dumps("northwind", tmp_path / "again")
    assert again == rebuilt["northwind"]


def test_the_deck_dump_shows_titles_positions_sizes_colors_and_fills(rebuilt):
    deck = rebuilt["northwind"]["deck"]
    assert "== Slide 1 of 4" in deck and "== Slide 4 of 4" in deck
    assert "at 0." in deck and " in, size " in deck           # where each shape sits
    assert " pt" in deck and "bold" in deck                    # font size and weight
    assert "fill " in deck                                     # table cell fills (status colors)
    assert "picture" in deck                                   # the two charts, by name and size


def test_the_memo_dump_has_the_word_file_its_footer_and_the_pdf(rebuilt):
    memo = rebuilt["northwind"]["memo"]
    assert "== Word file" in memo and "-- Footer" in memo and "== PDF page 1" in memo
    assert "keep with next" in memo                            # the page-break rules
    assert "table, aligned center" in memo


def test_the_metrics_dump_shows_every_sheet_with_formats_and_fills(rebuilt):
    metrics = rebuilt["northwind"]["metrics"]
    for sheet in ("Metrics", "Flags", "Data gaps"):
        assert f"== Sheet {sheet}" in metrics
    assert "[0.0%]" in metrics and "fill " in metrics


def test_a_workbook_cell_shows_its_value_format_fill_and_bold(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.title = "Test"
    sheet["A1"] = "Header"
    sheet["A1"].font = Font(bold=True)
    sheet["B2"] = 0.123456789012345
    sheet["B2"].number_format = "0.0%"
    sheet["B2"].fill = PatternFill("solid", fgColor="FFC7CE")
    sheet["B2"].alignment = Alignment(horizontal="right")
    path = tmp_path / "test.xlsx"
    book.save(path)
    lines = golden.dump_workbook(path).splitlines()
    assert "  A1  'Header'  bold" in lines
    assert "  B2  0.123456789012  [0.0%]  fill FFC7CE  right" in lines   # 12 significant digits: no float noise


# ---------------------------------------------------------------------------
# Comparing, and updating on purpose
# ---------------------------------------------------------------------------

def test_a_difference_shows_the_changed_lines_and_how_to_accept_it(tmp_path):
    path = tmp_path / "northwind_deck.txt"
    path.write_text("Slide 1\nNRR 97.1%\nRunway 11.0 mo\n")
    problem = golden.compare("Slide 1\nNRR 97.2%\nRunway 11.0 mo\n", path)
    assert "-NRR 97.1%" in problem and "+NRR 97.2%" in problem
    assert "northwind_deck.txt" in problem
    assert "python golden.py --update" in problem


def test_the_same_text_is_no_problem(tmp_path):
    path = tmp_path / "northwind_deck.txt"
    path.write_text("Slide 1\n")
    assert golden.compare("Slide 1\n", path) is None


def test_a_missing_golden_says_how_to_make_it(tmp_path):
    problem = golden.compare("Slide 1\n", tmp_path / "northwind_deck.txt")
    assert "no golden file" in problem and "python golden.py --update" in problem


def test_update_writes_only_what_changed_and_says_which(tmp_path):
    dumps = {"northwind": {"deck": "a\n", "memo": "b\n", "metrics": "c\n"}}
    assert golden.write_goldens(dumps, tmp_path) == ["northwind_deck.txt", "northwind_memo.txt",
                                                     "northwind_metrics.txt"]
    dumps["northwind"]["memo"] = "b2\n"
    assert golden.write_goldens(dumps, tmp_path) == ["northwind_memo.txt"]
    assert (tmp_path / "northwind_memo.txt").read_text() == "b2\n"


def test_update_removes_a_golden_for_an_output_that_no_longer_exists(tmp_path):
    (tmp_path / "oldco_deck.txt").write_text("gone\n")
    golden.write_goldens({"northwind": {"deck": "a\n", "memo": "b\n", "metrics": "c\n"}}, tmp_path)
    assert not (tmp_path / "oldco_deck.txt").exists()


def test_the_check_command_exits_1_on_a_difference_and_0_when_all_match(monkeypatch, capsys):
    dumps = {company: {kind: golden.golden_path(company, kind).read_text() for kind in golden.KINDS}
             for company in COMPANIES}
    monkeypatch.setattr(golden, "build_all", lambda folder: dumps)
    assert golden.main([]) == 0
    assert "9 of 9 match" in capsys.readouterr().out
    dumps["northwind"]["deck"] += "one more line\n"
    assert golden.main([]) == 1
    assert "+one more line" in capsys.readouterr().out
