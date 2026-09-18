"""Unit tests for theme.py: the one palette, font and type scale every output uses.

Expected values are typed from the Task 3 brief, not copied from theme.py.
Run from the project folder:  pytest
"""

import re
import tomllib
from pathlib import Path

import pytest
from matplotlib import font_manager
from pptx import Presentation

import excel_output
import text_fit
import theme
from metrics import CANNOT_EVALUATE, PASS, TRIP

PROJECT_DIR = Path(__file__).parent.parent


# ---------------------------------------------------------------------------
# The palette
# ---------------------------------------------------------------------------

def test_the_palette_is_the_brief_s_colors():
    assert (theme.NAVY, theme.NAVY_DARK, theme.SLATE, theme.MID_GRAY) == ("0B2545", "08192F", "334155", "64748B")
    assert (theme.LINE, theme.SURFACE, theme.WHITE) == ("E2E8F0", "F8FAFC", "FFFFFF")
    assert (theme.RED, theme.RED_FILL) == ("C0392B", "FDE8E6")
    assert (theme.GREEN, theme.GREEN_FILL) == ("1E8449", "EAF6EF")
    assert theme.GRAY_FILL == "EDF0F3"


def test_status_colors_are_red_green_and_gray_on_their_fills():
    assert theme.STATUS_COLORS == {TRIP: ("FDE8E6", "C0392B"), PASS: ("EAF6EF", "1E8449"),
                                   CANNOT_EVALUATE: ("EDF0F3", "334155")}


def test_the_excel_workbook_keeps_its_fills():
    # The brief: "Keep existing Excel fills" (Excel's own "Bad" / "Good" / gray).
    assert excel_output.STATUS_COLORS == {TRIP: ("FFC7CE", "9C0006"), PASS: ("C6EFCE", "006100"),
                                          CANNOT_EVALUATE: ("D9D9D9", "404040")}
    assert excel_output.STATUS_COLORS is theme.EXCEL_STATUS_COLORS


def test_css_color_adds_the_hash():
    assert theme.css_color("0B2545") == "#0B2545"


HEX_COLOR = re.compile(r"""["']#?[0-9A-Fa-f]{6}["']|#[0-9A-Fa-f]{6}\b""")


def code_files():
    """Every Python file of the project except tests (tests type expected colors on purpose)."""
    return sorted(path for path in PROJECT_DIR.glob("*.py"))


def test_no_file_but_theme_py_types_a_color():
    typed = {path.name: HEX_COLOR.findall(path.read_text()) for path in code_files() if path.name != "theme.py"}
    assert {name: found for name, found in typed.items() if found} == {}


def test_the_hex_color_finder_finds_colors_written_every_way():
    # The check above is only as good as this pattern.
    for line in ('NAVY = "1F2A44"', "fill='ffc7ce'", "color: #0B2545;", 'hex_color("#334155")'):
        assert HEX_COLOR.search(line), line
    assert not HEX_COLOR.search("Emu(12192000)") and not HEX_COLOR.search('"Q2 2026"')


# ---------------------------------------------------------------------------
# Nothing from another brand
# ---------------------------------------------------------------------------

OTHER_BRAND = "K" + "1"   # spelled in two pieces, so this file doesn't name it either


def test_no_code_config_or_template_names_the_other_brand():
    files = code_files() + [PROJECT_DIR / ".streamlit" / "config.toml", PROJECT_DIR / "run_app.command"]
    named = [path.name for path in files if re.search(rf"\b{OTHER_BRAND}\b", path.read_text())]
    template_text = " ".join(shape.text_frame.text for shape in Presentation(theme_template()).slide_master.shapes
                             if shape.has_text_frame)
    assert named == [] and OTHER_BRAND not in template_text
    assert theme.BRAND_NAME == "Example Capital"


def theme_template():
    return PROJECT_DIR / "templates" / "base.pptx"


# ---------------------------------------------------------------------------
# Font and sizes
# ---------------------------------------------------------------------------

def test_arial_first_then_helvetica_then_dejavu_sans():
    assert theme.FONT == "Arial"
    assert theme.FONT_STACK == ("Arial", "Helvetica", "DejaVu Sans")
    assert theme.css_font_stack() == "Arial, Helvetica, 'DejaVu Sans', sans-serif"


def missing(*families):
    """A stand-in for matplotlib's findfont that can't find these families."""
    real = font_manager.findfont

    def find(properties, fallback_to_default=True, **kwargs):
        if properties.get_family()[0] in families:
            raise ValueError("not installed")
        return real(properties, fallback_to_default=fallback_to_default, **kwargs)
    return find


@pytest.mark.parametrize("not_installed, expected", [
    ((), "Arial"),
    (("Arial",), "Helvetica"),
    (("Arial", "Helvetica"), "DejaVu Sans"),
])
def test_the_first_installed_font_of_the_stack_is_used(monkeypatch, not_installed, expected):
    if not theme.installed(expected):
        pytest.skip(f"{expected} isn't installed on this computer")
    monkeypatch.setattr(font_manager, "findfont", missing(*not_installed))
    assert theme.first_installed_font() == expected


def test_the_pdf_font_is_a_single_font_file_it_can_embed(monkeypatch):
    # Helvetica on a Mac is a .ttc collection (several fonts in one file); the PDF needs a .ttf,
    # so without Arial the memo falls through to DejaVu Sans.
    for bold in (False, True):
        _, path = theme.font_file(bold)
        assert Path(path).suffix.lower() == ".ttf"
    theme.font_file.cache_clear()
    monkeypatch.setattr(font_manager, "findfont", missing("Arial"))
    family, path = theme.font_file(bold=True)
    theme.font_file.cache_clear()
    assert family == "DejaVu Sans" and path.endswith("DejaVuSans-Bold.ttf")


def test_type_sizes_and_the_12_pt_floor():
    assert (theme.TITLE_PT, theme.SECTION_PT, theme.BODY_PT, theme.CAPTION_PT, theme.TABLE_PT) == (28, 20, 15, 13, 14)
    assert theme.MIN_PT == 12 and text_fit.MIN_FONT_PT == theme.MIN_PT
    assert min(theme.TITLE_PT, theme.SECTION_PT, theme.BODY_PT, theme.CAPTION_PT, theme.TABLE_PT) >= theme.MIN_PT


# ---------------------------------------------------------------------------
# The web page's look (Streamlit)
# ---------------------------------------------------------------------------

def css_rules(css):
    """{selector: declarations} for every rule in a style sheet (no nesting in ours)."""
    return {selector.strip(): body.strip() for selector, body in re.findall(r"([^{}]+)\{([^{}]*)\}", css)}


def rule(css, *words, plain=False):
    """The declarations of the one rule whose selector holds every word (plain=True: and no ':hover' or the like)."""
    found = [body for selector, body in css_rules(css).items()
             if all(word in selector for word in words) and not (plain and ":" in selector)]
    assert len(found) == 1, f"{len(found)} rules match {words}"
    return found[0]


def test_primary_buttons_are_navy_with_white_text():
    css = theme.streamlit_css()
    primary = rule(css, "primary", plain=True)
    for declaration in ("background-color: #0B2545", "color: #FFFFFF", "border-radius: 6px", "padding: 10px 18px"):
        assert declaration in primary
    assert "background-color: #08192F" in rule(css, "primary", ":hover")
    disabled = rule(css, "primary", ":disabled")
    assert "background-color: #F8FAFC" in disabled and "color: #64748B" in disabled


def test_secondary_buttons_are_white_with_a_navy_border():
    secondary = rule(theme.streamlit_css(), "secondary", plain=True)
    for declaration in ("background-color: #FFFFFF", "color: #0B2545", "border: 1px solid #0B2545",
                        "border-radius: 6px", "padding: 10px 18px"):
        assert declaration in secondary


def test_no_button_is_ever_red_or_green():
    button_rules = [body for selector, body in css_rules(theme.streamlit_css()).items() if "button" in selector.lower()]
    assert button_rules
    for body in button_rules:
        for color in (theme.RED, theme.RED_FILL, theme.GREEN, theme.GREEN_FILL):
            assert color.lower() not in body.lower()


def test_one_column_about_1100_px_with_white_cards_over_the_surface():
    css = theme.streamlit_css()
    assert "max-width: 1100px" in rule(css, "block-container")
    assert "background-color: #F8FAFC" in css_rules(css)[".stApp"]
    card = rule(css, "st-key-card")
    assert "background-color: #FFFFFF" in card and "border: 1px solid #E2E8F0" in card


def test_tables_have_a_navy_header_and_40_px_rows():
    css = theme.streamlit_css()
    header = rule(css, "bp-table", "th")
    assert "background-color: #0B2545" in header and "color: #FFFFFF" in header
    assert "height: 40px" in rule(css, "bp-table", "td")
    row_header = css_rules(css)['[class*="st-key-portfolio-header"]']
    assert "background-color: #0B2545" in row_header and "color: #FFFFFF" in row_header
    assert "min-height: 40px" in rule(css, "portfolio-row")


def test_page_text_is_arial_in_the_brief_s_sizes():
    css = theme.streamlit_css()
    assert "font-family: Arial, Helvetica, 'DejaVu Sans', sans-serif" in rule(css, "html")
    assert "font-size: 28px" in rule(css, "h1")
    assert "font-size: 20px" in rule(css, "h3")
    assert "font-size: 13px" in rule(css, "stCaption")
    assert "font-size: 14px" in css_rules(css)[".bp-table"]


def test_streamlit_s_config_carries_the_same_colors_and_font():
    # Streamlit reads its accent colors (checkbox, progress bar, links) only from config.toml, so the
    # file repeats theme.py's values; this keeps the two from drifting.
    config = tomllib.loads((PROJECT_DIR / ".streamlit" / "config.toml").read_text())
    assert config["theme"] == theme.streamlit_theme()
    assert config["theme"]["primaryColor"] == "#0B2545"
