"""Tests for rollup.py: the portfolio rollup deck and workbook across every company (Task 9).

Every test copies the three workbooks into a temporary data/ folder and builds into a temporary
output/ folder, so the real data/ and output/ are never touched. The rollup never asks Claude; a
guard makes creating a real Anthropic client fail anyway. Expected values are the company stories
(CLAUDE.md), typed here by hand: Fernhollow 7 tripped, Northwind 6, Alderpeak 0.
Run from the project folder:  pytest
"""

import datetime
import math
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from pptx import Presentation

import analyze
import mapping
import rollup
from charts import runway_chart
from check_deck import check_no_overflow
from metrics import CANNOT_EVALUATE, COMBO_FLAG_NAME, FLAG_RULES, PASS, TRIP, load_config
from provenance import NOT_REVIEWED, save_manifest
from theme import EXCEL_STATUS_COLORS, RED, STATUS_COLORS

PROJECT_DIR = Path(__file__).parent.parent
COMPANIES = ["alderpeak", "fernhollow", "northwind"]
RUN_DATE = datetime.date(2026, 7, 15)
EM_DASH = chr(0x2014)


@pytest.fixture(autouse=True)
def no_real_client(monkeypatch):
    """Creating a real Anthropic client fails the test: the rollup has no AI step at all."""
    def refuse(*args, **kwargs):
        raise AssertionError("a test tried to create a real Anthropic client")
    monkeypatch.setattr(analyze.anthropic, "Anthropic", refuse)


@pytest.fixture(autouse=True)
def mappings_dir(tmp_path, monkeypatch):
    """A temporary mappings/ folder: the project's own is never read or written."""
    folder = tmp_path / "mappings"
    monkeypatch.setattr(mapping, "MAPPINGS_DIR", folder)
    return folder


@pytest.fixture
def folders(tmp_path):
    """(data folder with copies of the three workbooks, empty output folder)."""
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    data_dir.mkdir()
    for name in COMPANIES:
        shutil.copy(PROJECT_DIR / "data" / f"{name}.xlsx", data_dir)
    return data_dir, output_dir


@pytest.fixture
def config():
    return load_config()


def flag(name, status):
    """A flag dict with only what the ranking reads."""
    return {"flag": name, "status": status, "metric": None, "reason": None}


def entries(folders, config):
    return rollup.collect_rollup(config, *folders)


# ---------------------------------------------------------------------------
# Worst flag, company status, review status
# ---------------------------------------------------------------------------

def test_every_flag_has_exactly_one_place_in_the_worst_first_order():
    # A new flag in metrics.py must be given a place, or the rollup can't say which is worst.
    every_flag = [name for name, _, _, _ in FLAG_RULES] + [COMBO_FLAG_NAME]
    assert sorted(rollup.WORST_FIRST) == sorted(every_flag)
    assert len(set(rollup.WORST_FIRST)) == len(rollup.WORST_FIRST)


def test_runway_is_the_worst_flag_and_the_combo_rule_the_least():
    # The order typed from the task's reasoning: cash running out comes first.
    assert rollup.WORST_FIRST[0] == "Runway at current burn"
    assert rollup.WORST_FIRST[-1] == COMBO_FLAG_NAME


def test_the_worst_flag_is_the_tripped_one_highest_in_the_order():
    flags = [flag("Rule of 40", TRIP), flag("Runway at current burn", PASS), flag("NRR (annualized)", TRIP),
             flag("Burn multiple", CANNOT_EVALUATE)]
    assert rollup.worst_flag(flags)["flag"] == "NRR (annualized)"


def test_no_tripped_flag_means_no_worst_flag():
    assert rollup.worst_flag([flag("Rule of 40", PASS), flag("Burn multiple", CANNOT_EVALUATE)]) is None


@pytest.mark.parametrize("statuses, expected", [
    ([TRIP, PASS, CANNOT_EVALUATE], TRIP),
    ([PASS, CANNOT_EVALUATE], CANNOT_EVALUATE),
    ([PASS, PASS], PASS),
])
def test_a_company_s_status_is_its_flags_worst_status(statuses, expected):
    assert rollup.company_status([flag("Rule of 40", status) for status in statuses]) == expected


@pytest.mark.parametrize("page_status, expected", [
    ("Not generated yet", "Not generated"),
    ("Out of date: the workbook has changed since the last run. Generate again.", "Out of date"),
    (NOT_REVIEWED, "Not reviewed"),
    ("approved by Tyler Ho on 2026-07-15T10:00:00", "Approved"),
])
def test_review_status_is_the_portfolio_page_s_status_in_one_word_or_two(page_status, expected):
    assert rollup.review_label(page_status) == expected


# ---------------------------------------------------------------------------
# Ranking
# ---------------------------------------------------------------------------

def test_the_portfolio_is_ranked_by_flags_tripped(folders, config):
    ranked = entries(folders, config)
    assert [entry["company"] for entry in ranked] == ["Fernhollow", "Northwind", "Alderpeak"]
    assert [entry["rank"] for entry in ranked] == [1, 2, 3]
    assert [entry["tripped"] for entry in ranked] == [7, 6, 0]
    assert [entry["unevaluated"] for entry in ranked] == [1, 0, 0]
    assert [entry["status"] for entry in ranked] == [TRIP, TRIP, PASS]


def test_each_company_s_worst_flag_matches_its_story(folders, config):
    worst = {entry["company"]: entry["worst"] for entry in entries(folders, config)}
    assert worst["Fernhollow"]["flag"] == "Runway at current burn"     # 6.0 mo
    assert worst["Northwind"]["flag"] == "Runway at current burn"      # 11.0 mo
    assert worst["Alderpeak"] is None


def test_the_worst_flag_text_shows_value_and_threshold(folders, config):
    texts = {entry["company"]: rollup.worst_flag_text(entry) for entry in entries(folders, config)}
    assert texts["Fernhollow"] == "Runway at current burn: 6.0 mo (trips below 12.0 mo)"
    assert texts["Northwind"] == "Runway at current burn: 11.0 mo (trips below 12.0 mo)"
    assert texts["Alderpeak"] == rollup.NONE_TRIPPED


def test_a_tie_on_flags_tripped_puts_the_worse_worst_flag_first_then_the_name():
    def entry(company, tripped, worst):
        return {"company": company, "tripped": tripped, "worst": flag(worst, TRIP) if worst else None,
                "data": object(), "problem": None}
    ranked = rollup.ranked([entry("Bravo", 2, "Rule of 40"), entry("Alpha", 2, "Rule of 40"),
                            entry("Charlie", 2, "Runway at current burn"), entry("Delta", 3, "CAC payback")])
    assert [item["company"] for item in ranked] == ["Delta", "Charlie", "Alpha", "Bravo"]
    assert [item["rank"] for item in ranked] == [1, 2, 3, 4]


def test_an_unreadable_workbook_goes_last_with_no_rank_and_the_others_still_rank(folders, config):
    book = Workbook()
    book.active.append(["Quarter", "Starting ARR"])
    book.active.append(["Q1 2025", 100])
    book.save(folders[0] / "broken.xlsx")
    ranked = entries(folders, config)
    assert [entry["company"] for entry in ranked] == ["Fernhollow", "Northwind", "Alderpeak", "Broken"]
    broken = ranked[-1]
    assert broken["rank"] is None and broken["status"] == rollup.UNREADABLE and broken["problem"]
    assert "Traceback" not in broken["problem"]


def test_status_counts_list_every_status_in_order_with_zeros(folders, config):
    counts = rollup.status_counts(entries(folders, config))
    assert [(label, count) for label, count, _ in counts] == [
        ("Flags tripped", 2), ("None tripped, some cannot evaluate", 0), ("Every flag passed", 1),
        ("Workbook can't be read", 0)]
    assert counts[0][2] == ["Fernhollow", "Northwind"]


def test_review_counts_read_each_company_s_manifest(folders, config):
    counts = dict((label, count) for label, count, _ in rollup.review_counts(entries(folders, config)))
    assert counts == {"Approved": 0, "Not reviewed": 0, "Out of date": 0, "Not generated": 3}


def test_a_list_of_names_that_is_too_long_ends_with_how_many_more():
    names = [f"Company {letter}" for letter in "ABCDEFGHIJ"]
    assert rollup.names_text(names) == "Company A, Company B, Company C and 7 more"
    assert rollup.names_text([]) == rollup.NO_VALUE


# ---------------------------------------------------------------------------
# The runway chart (charts.py)
# ---------------------------------------------------------------------------

def test_the_runway_chart_draws_a_bar_per_number_red_when_tripped_and_the_threshold():
    figure = runway_chart(["Fernhollow", "Northwind", "Alderpeak", "Blank Co", "Cash Rich"],
                          [6.0, 11.0, 108.0, math.nan, math.inf],
                          ["6.0 mo", "11.0 mo", "108.0 mo", "data missing", "∞ (not burning)"],
                          [True, True, False, False, False], 12.0, "trips below 12.0 mo", (6, 4))
    axis = figure.axes[0]
    widths = [bar.get_width() for bar in axis.patches]
    assert widths == [6.0, 11.0, 108.0]                   # no bar for data missing or infinity
    colors = [bar.get_facecolor()[:3] for bar in axis.patches]
    red = tuple(int(RED[i:i + 2], 16) / 255 for i in (0, 2, 4))
    assert [color == pytest.approx(red) for color in colors] == [True, True, False]
    assert any(list(line.get_xdata()) == [12.0, 12.0] for line in axis.lines)   # the threshold line
    labels = [text.get_text() for text in axis.texts]
    for label in ("6.0 mo", "108.0 mo", "data missing", "∞ (not burning)"):
        assert any(label in text for text in labels), label
    assert [tick.get_text() for tick in axis.get_yticklabels()] == [
        "Fernhollow", "Northwind", "Alderpeak", "Blank Co", "Cash Rich"]
    assert axis.yaxis_inverted()                           # first company at the top


def test_the_chart_order_is_shortest_runway_first_then_infinite_then_no_number(folders, config):
    order = [entry["company"] for entry in rollup.runway_order(entries(folders, config))]
    assert order == ["Fernhollow", "Northwind", "Alderpeak"]


# ---------------------------------------------------------------------------
# The deck
# ---------------------------------------------------------------------------

def deck(folders, config):
    paths = rollup.save_rollup(config, *folders, run_date=RUN_DATE)
    return paths, Presentation(paths["deck"])


def test_the_rollup_deck_has_three_slides_with_titles(folders, config):
    paths, presentation = deck(folders, config)
    assert paths["deck"] == folders[1] / "portfolio_rollup.pptx"
    assert [slide.shapes.title.text for slide in presentation.slides] == [
        "Portfolio ranked by flags tripped, Q2 2026", "Companies by status, Q2 2026",
        "Runway at current burn by company, Q2 2026"]


def table_rows(slide, name):
    frame = next(item for item in slide.shapes if item.name == name)
    return [[cell.text for cell in row.cells] for row in frame.table.rows], frame.table


def test_the_ranking_table_shows_each_company_in_rank_order(folders, config):
    _, presentation = deck(folders, config)
    rows, table = table_rows(presentation.slides[0], "Ranking table")
    assert rows[0] == ["Rank", "Company", "Quarter", "Flags tripped", "Worst flag", "Runway at current burn", "Review"]
    assert rows[1] == ["1", "Fernhollow", "Q2 2026", "7 of 9 flags tripped, 1 cannot evaluate",
                       "Runway at current burn: 6.0 mo (trips below 12.0 mo)", "6.0 mo", "Not generated"]
    assert rows[2][:4] == ["2", "Northwind", "Q2 2026", "6 of 9 flags tripped"]
    assert rows[3][:5] == ["3", "Alderpeak", "Q2 2026", "0 of 9 flags tripped", rollup.NONE_TRIPPED]
    worst_fills = [str(table.cell(row, 4).fill.fore_color.rgb) for row in (1, 2, 3)]
    assert worst_fills == [STATUS_COLORS[TRIP][0], STATUS_COLORS[TRIP][0], STATUS_COLORS[PASS][0]]


def test_the_status_slide_counts_companies_by_flag_status_and_review_status(folders, config):
    _, presentation = deck(folders, config)
    flag_rows, _ = table_rows(presentation.slides[1], "Flag status table")
    assert flag_rows[0] == ["Flag status", "Companies", "Which"]
    assert flag_rows[1] == ["Flags tripped", "2", "Fernhollow, Northwind"]
    assert flag_rows[3] == ["Every flag passed", "1", "Alderpeak"]
    review_rows, _ = table_rows(presentation.slides[1], "Review status table")
    assert review_rows[4] == ["Not generated", "3", "Alderpeak, Fernhollow, Northwind"]


def test_the_runway_slide_is_one_chart_picture(folders, config):
    _, presentation = deck(folders, config)
    pictures = [item for item in presentation.slides[2].shapes if item.name == "Runway chart"]
    assert len(pictures) == 1 and pictures[0].shape_type == 13   # MSO_SHAPE_TYPE.PICTURE


def test_every_slide_has_the_footer_and_nothing_overflows(folders, config):
    _, presentation = deck(folders, config)
    for slide in presentation.slides:
        footer = next(item for item in slide.shapes if item.name == "Footer").text_frame.text
        assert footer.startswith("Fictional data | Portfolio rollup of 3 workbooks | 2026-07-15 | ")
        assert footer.endswith(" | computed metrics only, no AI text")
    check_no_overflow(presentation, "rollup")


def test_a_big_portfolio_splits_the_ranking_over_several_slides_and_still_fits(folders, config):
    data_dir, _ = folders
    for number in range(9):   # 12 companies in all, with the longest names the web page allows
        name = f"Company {number} " + "x" * 30
        shutil.copy(PROJECT_DIR / "data" / f"{COMPANIES[number % 3]}.xlsx", data_dir / f"{name}.xlsx")
    _, presentation = deck(folders, config)
    titles = [slide.shapes.title.text for slide in presentation.slides]
    assert titles[:2] == ["Portfolio ranked by flags tripped, Q2 2026 (1 of 2)",
                          "Portfolio ranked by flags tripped, Q2 2026 (2 of 2)"]
    ranks = [row[0] for slide in list(presentation.slides)[:2] for row in table_rows(slide, "Ranking table")[0][1:]]
    assert ranks == [str(number) for number in range(1, 13)]
    check_no_overflow(presentation, "big rollup")


def test_companies_on_different_quarters_are_titled_by_each_latest_quarter(folders, config):
    ranked = entries(folders, config)
    ranked[0]["data"] = {**ranked[0]["data"], "latest": "Q1 2026"}
    assert rollup.quarter_words(ranked) == "each company's latest quarter"
    assert rollup.quarter_words(entries(folders, config)) == "Q2 2026"


def test_when_no_workbook_can_be_read_the_titles_name_no_quarter(tmp_path, config):
    (tmp_path / "data").mkdir()
    book = Workbook()
    book.active.append(["Quarter", "Starting ARR"])
    book.active.append(["Q1 2025", 100])
    book.save(tmp_path / "data" / "broken.xlsx")
    paths = rollup.save_rollup(config, tmp_path / "data", tmp_path / "output", run_date=RUN_DATE)
    assert [slide.shapes.title.text for slide in Presentation(paths["deck"]).slides] == [
        "Portfolio ranked by flags tripped", "Companies by status", "Runway at current burn by company"]


def test_nothing_in_the_rollup_deck_has_an_em_dash(folders, config):
    _, presentation = deck(folders, config)
    for slide in presentation.slides:
        for item in slide.shapes:
            texts = [item.text_frame.text] if item.has_text_frame else []
            if item.has_table:
                texts = [cell.text for row in item.table.rows for cell in row.cells]
            assert not any(EM_DASH in text for text in texts), item.name


# ---------------------------------------------------------------------------
# The workbook
# ---------------------------------------------------------------------------

def workbook(folders, config):
    paths = rollup.save_rollup(config, *folders, run_date=RUN_DATE)
    return load_workbook(paths["excel"])


def sheet_rows(sheet):
    return [[cell.value for cell in row] for row in sheet.iter_rows()]


def test_the_rollup_workbook_has_three_sheets(folders, config):
    assert workbook(folders, config).sheetnames == ["Ranking", "By status", "Runway"]


def test_the_ranking_sheet_holds_real_numbers_in_the_metrics_formats(folders, config):
    sheet = workbook(folders, config)["Ranking"]
    rows = sheet_rows(sheet)
    assert rows[0] == rollup.RANKING_HEADERS
    assert [row[:6] for row in rows[1:]] == [[1, "Fernhollow", "Q2 2026", 7, 1, 9],
                                             [2, "Northwind", "Q2 2026", 6, 0, 9],
                                             [3, "Alderpeak", "Q2 2026", 0, 0, 9]]
    header = rows[0]
    worst, value, threshold, runway = (header.index(name) for name in
                                       ("Worst flag", "Worst flag value", "Threshold", "Runway at current burn"))
    assert rows[1][worst] == "Runway at current burn"
    assert rows[1][value] == pytest.approx(3300 / (1650 / 3))        # Fernhollow's answer key
    assert rows[1][threshold] == 12
    assert rows[3][worst] == rollup.NONE_TRIPPED and rows[3][value] is None
    assert rows[3][runway] == pytest.approx(7200 / (200 / 3))       # Alderpeak's answer key
    assert sheet.cell(row=2, column=value + 1).number_format == '0.0" mo"'
    assert sheet.cell(row=4, column=runway + 1).number_format == '0.0" mo"'


def test_the_ranking_sheet_colors_the_flag_status_like_the_metrics_workbook(folders, config):
    sheet = workbook(folders, config)["Ranking"]
    column = rollup.RANKING_HEADERS.index("Flag status") + 1
    fills = [sheet.cell(row=row, column=column).fill.fgColor.rgb[-6:] for row in (2, 3, 4)]
    assert fills == [EXCEL_STATUS_COLORS[TRIP][0], EXCEL_STATUS_COLORS[TRIP][0], EXCEL_STATUS_COLORS[PASS][0]]


def test_the_by_status_sheet_has_the_same_counts_as_the_deck(folders, config):
    rows = sheet_rows(workbook(folders, config)["By status"])
    assert rows[0] == ["Group", "Status", "Companies", "Which"]
    assert ["Flag status", "Flags tripped", 2, "Fernhollow, Northwind"] in rows
    assert ["Review status", "Not generated", 3, "Alderpeak, Fernhollow, Northwind"] in rows


def test_the_runway_sheet_is_in_chart_order_with_the_flag_status(folders, config):
    rows = sheet_rows(workbook(folders, config)["Runway"])
    assert [row[0] for row in rows[1:]] == ["Fernhollow", "Northwind", "Alderpeak"]
    assert [row[-1] for row in rows[1:]] == ["Tripped", "Tripped", "Passed"]


def test_an_unreadable_workbook_is_on_every_output_with_its_reason_and_the_rest_still_builds(folders, config):
    book = Workbook()
    book.active.append(["Quarter", "Starting ARR"])
    book.active.append(["Q1 2025", 100])
    book.save(folders[0] / "broken.xlsx")
    paths = rollup.save_rollup(config, *folders, run_date=RUN_DATE)
    presentation = Presentation(paths["deck"])
    rows, _ = table_rows(presentation.slides[0], "Ranking table")
    assert rows[4][:5] == ["-", "Broken", "-", "-", rollup.UNREADABLE_TEXT]
    flag_rows, _ = table_rows(presentation.slides[1], "Flag status table")
    assert flag_rows[4] == ["Workbook can't be read", "1", "Broken"]
    check_no_overflow(presentation, "rollup with an unreadable workbook")
    ranking = sheet_rows(load_workbook(paths["excel"])["Ranking"])
    assert ranking[4][1] == "Broken" and ranking[4][-3] == "Workbook can't be read" and ranking[4][-1]
    assert sheet_rows(load_workbook(paths["excel"])["Runway"])[-1][-1] == "Workbook can't be read"


def test_the_command_line_writes_both_files_into_output(folders, config, monkeypatch, capsys):
    data_dir, output_dir = folders
    monkeypatch.setattr(rollup, "DATA_DIR", data_dir)
    monkeypatch.setattr(rollup, "OUTPUT_DIR", output_dir)
    rollup.main([])
    assert rollup.rollup_paths(output_dir)["deck"].exists() and rollup.rollup_paths(output_dir)["excel"].exists()
    assert "Fernhollow" in capsys.readouterr().out


def test_an_empty_portfolio_stops_with_plain_words(tmp_path, config):
    (tmp_path / "data").mkdir()
    with pytest.raises(ValueError, match="nothing to roll up"):
        rollup.save_rollup(config, tmp_path / "data", tmp_path / "output")


# ---------------------------------------------------------------------------
# Review status follows approvals; the web page's download
# ---------------------------------------------------------------------------

def test_an_approved_company_counts_as_approved(folders, config, monkeypatch):
    # A current manifest with an approval: faked through run_state, the page's own judge.
    monkeypatch.setattr(rollup, "run_state", lambda path, output_dir: {
        "status": "approved by Tyler Ho on 2026-07-15T10:00:00" if Path(path).stem == "northwind" else NOT_REVIEWED})
    counts = dict((label, count) for label, count, _ in rollup.review_counts(entries(folders, config)))
    assert counts == {"Approved": 1, "Not reviewed": 2, "Out of date": 0, "Not generated": 0}


def test_a_download_builds_in_a_temporary_folder_and_writes_nothing_to_output(folders, config):
    data_dir, output_dir = folders
    output_dir.mkdir()
    save_manifest(output_dir / "keep.json", {"x": 1})
    before = sorted(path.name for path in output_dir.rglob("*"))
    name, content = rollup.rollup_download("deck", config, data_dir, output_dir)
    assert name == "portfolio_rollup.pptx" and content[:2] == b"PK"     # a zip, as every .pptx is
    name, content = rollup.rollup_download("excel", config, data_dir, output_dir)
    assert name == "portfolio_rollup.xlsx" and content[:2] == b"PK"
    assert sorted(path.name for path in output_dir.rglob("*")) == before
