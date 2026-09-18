"""Unit tests for main.py: the batch outputs (summary CSV, quarter warning) and the AI step (Task 3).

The summary tests use results written out by hand, so no workbook is needed. The AI-step tests run
data/northwind.xlsx through run_batch into a temporary folder with a fake Claude client, so the
real output/ folder is untouched and no API call can happen (a guard makes creating a real client fail).
Run from the project folder:  pytest
"""

import csv
import json
from pathlib import Path
from types import SimpleNamespace

import anthropic
import pytest
from pptx import Presentation

import analyze
import main
from analyze import BoardSummary
from build_deck import PLACEHOLDER_TEXT
from main import quarter_mismatch_warning, write_summary_csv
from provenance import NOT_REVIEWED, file_sha256, git_commit, manifest_path, read_manifest, save_manifest
from text_fit import TextDoesNotFitError

PROJECT_DIR = Path(__file__).parent.parent
NORTHWIND = PROJECT_DIR / "data" / "northwind.xlsx"


def ok(company, quarter, tripped=(), cannot_evaluate=(), gap_count=0, blank_quarters=(), ai="skipped"):
    """A successful company's result."""
    return {"company": company, "quarter": quarter, "flags_total": 9, "tripped": list(tripped),
            "cannot_evaluate": list(cannot_evaluate), "gap_count": gap_count,
            "blank_quarters": list(blank_quarters), "ai": ai, "error": None}


def failed(company, error):
    return {"company": company, "error": error}


# ---------------------------------------------------------------------------
# Summary table and CSV
# ---------------------------------------------------------------------------

def test_summary_csv_has_one_row_per_company(tmp_path):
    results = [
        ok("Northwind", "Q2 2026", tripped=["a"] * 6, gap_count=19, blank_quarters=["Q1 2025"], ai="ok"),
        ok("Fernhollow", "Q2 2026", tripped=["a"] * 7, cannot_evaluate=["Rule of 40"], gap_count=20,
           blank_quarters=["Q2 2025"], ai="failed"),
        ok("Alderpeak", "Q2 2026", ai="skipped"),
        failed("Broken", "ValueError: row 3 (header) is missing columns: pipeline"),
    ]
    path = write_summary_csv(results, tmp_path / "batch_summary.csv")
    with open(path, newline="") as file:
        rows = list(csv.reader(file))
    assert rows == [
        ["Company", "Latest quarter", "Flags tripped", "Flags total", "Cannot evaluate", "Data gaps",
         "Blank quarters", "Result"],
        ["Northwind", "Q2 2026", "6", "9", "0", "19", "Q1 2025", "OK"],
        ["Fernhollow", "Q2 2026", "7", "9", "1", "20", "Q2 2025", "OK (AI failed)"],
        ["Alderpeak", "Q2 2026", "0", "9", "0", "0", "", "OK (AI skipped)"],
        ["Broken", "", "", "", "", "", "", "FAILED: ValueError: row 3 (header) is missing columns: pipeline"],
    ]


def test_no_warning_when_every_company_ends_on_the_same_quarter():
    results = [ok("Alderpeak", "Q2 2026"), ok("Northwind", "Q2 2026"), failed("Broken", "ValueError: x")]
    assert quarter_mismatch_warning(results) is None


def test_no_warning_for_a_single_company():
    assert quarter_mismatch_warning([ok("Northwind", "Q2 2026")]) is None


def test_warning_names_each_quarter_and_its_companies():
    results = [ok("Alderpeak", "Q2 2026"), ok("Northwind", "Q1 2026"), ok("Fernhollow", "Q2 2026"),
               failed("Broken", "ValueError: x")]  # a failed company has no quarter and is left out
    assert quarter_mismatch_warning(results) == (
        "⚠ Companies end on different quarters (Q2 2026: Alderpeak, Fernhollow; Q1 2026: Northwind) "
        "- compare them with care")


def test_result_text_says_what_happened_to_the_ai_summary():
    assert main.result_text(ok("Northwind", "Q2 2026", ai="ok")) == "OK"
    assert main.result_text(ok("Northwind", "Q2 2026", ai="skipped")) == "OK (AI skipped)"
    assert main.result_text(ok("Northwind", "Q2 2026", ai="failed")) == "OK (AI failed)"
    assert main.result_text(failed("Broken", "ValueError: x")) == "FAILED: ValueError: x"


def test_ai_failed_warning_names_the_companies():
    results = [ok("Northwind", "Q2 2026", ai="failed"), ok("Alderpeak", "Q2 2026", ai="ok"),
               ok("Fernhollow", "Q2 2026", ai="failed"), failed("Broken", "ValueError: x")]
    assert main.ai_failed_warning(results) == (
        "⚠ AI summary unavailable for 2 companies (Northwind, Fernhollow): their decks show the placeholder; "
        "the reasons are in output/<company>_analysis.json")
    assert main.ai_failed_warning([ok("Northwind", "Q2 2026", ai="ok")]) is None


# ---------------------------------------------------------------------------
# The AI step, with a fake Claude client
# ---------------------------------------------------------------------------

@pytest.fixture(autouse=True)
def no_real_client(monkeypatch):
    """Creating a real Anthropic client fails the test, so no test here can reach the API."""
    def refuse(*args, **kwargs):
        raise AssertionError("a test tried to create a real Anthropic client")
    monkeypatch.setattr(analyze.anthropic, "Anthropic", refuse)


class FakeClient:
    """Stands in for anthropic.Anthropic(): every call returns the same answer (or raises), and calls are counted."""

    def __init__(self, summary=None, error=None):
        self.summary, self.error, self.calls = summary, error, 0
        self.messages = self  # analyze.py calls client.messages.parse(...)

    def parse(self, **kwargs):
        self.calls += 1
        if self.error:
            raise self.error
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=self.summary.model_dump_json())],
                               parsed_output=self.summary, stop_reason="end_turn",
                               usage=SimpleNamespace(input_tokens=100, output_tokens=50))


def summary(headline="Retention is the main question for the board."):
    """A valid answer with no numbers in it, so the number check has nothing to reject."""
    point = {"title": "Steady base", "detail": "Customers stayed."}
    return BoardSummary.model_validate({
        "headline": headline, "wins": [point] * 3, "risks": [point] * 3,
        "questions": ["What drives churn?", "Where is pipeline coming from?", "How is hiring going?"]})


def run_northwind(tmp_path, skip_ai=False, client=None, draft=False):
    """run_batch on Northwind, saving into tmp_path. Returns Northwind's result."""
    return main.run_batch([NORTHWIND], main.load_config(), skip_ai, client=client, output_dir=tmp_path,
                          draft=draft)[0]


def headline_on_deck(tmp_path):
    """The Headline box on the last slide, AI commentary (slide 4)."""
    slide = Presentation(tmp_path / "northwind_board_pack.pptx").slides[-1]
    return next(shape.text_frame.text for shape in slide.shapes if shape.name == "Headline")


def saved_analysis(tmp_path):
    return json.loads((tmp_path / "northwind_analysis.json").read_text())


def test_ai_that_passes_goes_on_the_deck_and_is_saved(tmp_path):
    client = FakeClient(summary())
    result = run_northwind(tmp_path, client=client)
    assert result["error"] is None and main.result_text(result) == "OK"
    assert client.calls == 1
    assert headline_on_deck(tmp_path) == "Retention is the main question for the board."
    saved = saved_analysis(tmp_path)
    assert saved["summary"]["headline"] == "Retention is the main question for the board."
    assert saved["run_info"]["passed"] is True and saved["error"] is None
    assert saved["payload"]["company"] == "Northwind"


def test_ai_that_fails_validation_twice_still_builds_the_deck(tmp_path):
    # An old, good analysis is already there: it must be replaced, never put on this run's deck.
    run_northwind(tmp_path, client=FakeClient(summary()))
    assert saved_analysis(tmp_path)["summary"] is not None

    client = FakeClient(summary(headline="ARR grew 555.5% this quarter."))  # a number not in the data
    result = run_northwind(tmp_path, client=client)
    assert result["error"] is None and main.result_text(result) == "OK (AI failed)"
    assert client.calls == analyze.MAX_ATTEMPTS  # first try + one retry
    assert headline_on_deck(tmp_path) == PLACEHOLDER_TEXT
    saved = saved_analysis(tmp_path)
    assert saved["summary"] is None and saved["run_info"]["attempts"] == analyze.MAX_ATTEMPTS
    assert "555.5" in saved["error"]


def test_api_error_is_ai_failed_not_a_failed_company(tmp_path):
    client = FakeClient(error=anthropic.AnthropicError("simulated outage"))
    result = run_northwind(tmp_path, client=client)
    assert result["error"] is None and main.result_text(result) == "OK (AI failed)"
    assert headline_on_deck(tmp_path) == PLACEHOLDER_TEXT
    saved = saved_analysis(tmp_path)
    assert saved["summary"] is None and saved["error"] == "AnthropicError: simulated outage"


def test_unexpected_error_in_the_ai_step_fails_the_company(tmp_path):
    # A bug (not a validation failure or an API error) must not hide behind "AI failed".
    run_northwind(tmp_path, client=FakeClient(summary()))  # an old, good analysis from an earlier run
    result = run_northwind(tmp_path, client=FakeClient(error=KeyError("simulated bug")))
    assert result["error"] == "KeyError: 'simulated bug'"
    assert not (tmp_path / "northwind_analysis.json").exists()  # the old analysis can't pass for this run's


def test_ai_text_too_long_for_the_slides_is_ai_failed_not_a_failed_company(tmp_path):
    # Review finding 1a, the last line of defence: an analysis that passed when it was made but
    # doesn't fit the slides still gets a deck, with the placeholder, and the run reads OK (AI failed).
    long_point = {"title": "Steady base", "detail": "customers " * 45}
    saved = {"summary": {"headline": "Retention is the main question for the board.",
                         "wins": [long_point] * 3, "risks": [long_point] * 3,
                         "questions": ["What drives churn?", "Where is pipeline from?", "How is hiring?"]},
             "payload": {"company": "Northwind", "latest_quarter": "Q2 2026"}}
    path = tmp_path / "northwind_analysis.json"
    path.write_text(json.dumps(saved))

    why_unavailable = main.deck_step(NORTHWIND, main.load_config(), path, tmp_path)
    assert "does not fit slide 4" in why_unavailable
    assert headline_on_deck(tmp_path) == PLACEHOLDER_TEXT
    assert main.RESULT_TEXTS[main.ai_status(False, why_unavailable)] == "OK (AI failed)"


def test_a_text_that_does_not_fit_prints_one_clear_line_not_a_traceback(capsys):
    # Review finding 2: the message already names the slide and the box, so a traceback (which means
    # "this is a bug in the code") only makes the real problem harder to see.
    main.describe_error(TextDoesNotFitError("Slide 3, Risks and flags: text doesn't fit"))
    printed = capsys.readouterr()
    assert "Slide 3, Risks and flags" in printed.out
    assert "Traceback" not in printed.out + printed.err


def test_skip_ai_builds_the_placeholder_deck_and_never_calls_claude(tmp_path, monkeypatch):
    def refuse(*args, **kwargs):
        raise AssertionError("--skip-ai called analyze()")
    monkeypatch.setattr(main, "analyze", refuse)
    old = tmp_path / "northwind_analysis.json"
    old.write_text("an analysis from an earlier run")

    result = run_northwind(tmp_path, skip_ai=True)
    assert result["error"] is None and main.result_text(result) == "OK (AI skipped)"
    assert headline_on_deck(tmp_path) == PLACEHOLDER_TEXT
    assert old.read_text() == "an analysis from an earlier run"  # --skip-ai leaves an earlier analysis alone


# ---------------------------------------------------------------------------
# The manifest: where each deck came from (provenance)
# ---------------------------------------------------------------------------

CONFIG_FILE = PROJECT_DIR / "config.yaml"


def northwind_manifest(tmp_path):
    return read_manifest(manifest_path(NORTHWIND, tmp_path))


def test_a_run_writes_a_manifest_for_the_company(tmp_path):
    client = FakeClient(summary())
    run_northwind(tmp_path, client=client)
    saved = northwind_manifest(tmp_path)

    assert saved["company"] == "Northwind"
    assert saved["input"] == {"file": "northwind.xlsx", "sha256": file_sha256(NORTHWIND)}
    assert saved["config"] == {"file": "config.yaml", "sha256": file_sha256(CONFIG_FILE)}
    assert saved["code"] == git_commit()
    assert saved["ai"]["model"] == analyze.DEFAULT_MODEL
    assert saved["ai"]["prompt_version"] == analyze.PROMPT_VERSION
    assert saved["ai"]["validation"] == "passed" and saved["ai"]["attempts"] == 1
    assert saved["ai"]["input_tokens"] == 100 and saved["ai"]["output_tokens"] == 50
    assert saved["ai"]["cost_usd"] > 0
    assert saved["deck"] == {"file": "northwind_board_pack.pptx", "ai_text": True, "status": NOT_REVIEWED}
    assert saved["approval"] is None  # nobody has reviewed it yet
    assert saved["run_at"]


def test_the_manifest_records_a_skipped_ai_step(tmp_path):
    run_northwind(tmp_path, skip_ai=True)
    saved = northwind_manifest(tmp_path)
    assert saved["ai"]["validation"] == "skipped"
    assert saved["ai"]["model"] is None and saved["ai"]["cost_usd"] is None
    assert saved["deck"]["ai_text"] is False


def test_the_manifest_records_an_ai_failure_and_that_the_deck_has_the_placeholder(tmp_path):
    run_northwind(tmp_path, client=FakeClient(summary(headline="ARR grew 555.5% this quarter.")))
    saved = northwind_manifest(tmp_path)
    assert saved["ai"]["validation"].startswith("failed") and "555.5" in saved["ai"]["validation"]
    assert saved["ai"]["attempts"] == analyze.MAX_ATTEMPTS and saved["ai"]["cost_usd"] > 0
    assert saved["deck"]["ai_text"] is False


def test_the_manifest_hash_follows_the_workbook(tmp_path, monkeypatch):
    # The hash is of the file that was read, so a different workbook can't pass for this one.
    run_northwind(tmp_path, skip_ai=True)
    first = northwind_manifest(tmp_path)["input"]["sha256"]

    copy = tmp_path / "northwind.xlsx"
    copy.write_bytes(NORTHWIND.read_bytes() + b"a change")
    main.run_batch([copy], main.load_config(), True, output_dir=tmp_path)
    assert northwind_manifest(tmp_path)["input"]["sha256"] != first


def test_a_re_run_keeps_an_approval_that_is_still_valid(tmp_path):
    run_northwind(tmp_path, skip_ai=True)
    saved = northwind_manifest(tmp_path)
    saved["approval"] = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00",
                         "input_sha256": file_sha256(NORTHWIND), "config_sha256": file_sha256(CONFIG_FILE)}
    save_manifest(manifest_path(NORTHWIND, tmp_path), saved)

    run_northwind(tmp_path, skip_ai=True)  # same workbook, same thresholds
    after = northwind_manifest(tmp_path)
    assert after["approval"]["reviewer"] == "Tyler Ho"
    assert after["deck"]["status"] == "approved by Tyler Ho on 2026-09-17T15:00:00"


def test_a_run_after_the_workbook_changed_drops_back_to_draft(tmp_path):
    run_northwind(tmp_path, skip_ai=True)
    saved = northwind_manifest(tmp_path)
    saved["approval"] = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00",
                         "input_sha256": "the-hash-of-an-older-workbook", "config_sha256": file_sha256(CONFIG_FILE)}
    save_manifest(manifest_path(NORTHWIND, tmp_path), saved)

    run_northwind(tmp_path, skip_ai=True)
    assert northwind_manifest(tmp_path)["deck"]["status"] == NOT_REVIEWED


# ---------------------------------------------------------------------------
# The DRAFT watermark is opt-in (--draft); the footer always says whether the deck was reviewed
# ---------------------------------------------------------------------------

def deck_watermarks(tmp_path):
    deck = Presentation(tmp_path / "northwind_board_pack.pptx")
    return [item.name for slide in deck.slides for item in slide.shapes if item.name == "Watermark"]


def deck_footer(tmp_path):
    deck = Presentation(tmp_path / "northwind_board_pack.pptx")
    return [item for item in deck.slides[0].shapes if item.name == "Footer"][0].text_frame.text


def test_a_run_has_no_watermark_by_default_and_the_footer_says_not_reviewed(tmp_path):
    run_northwind(tmp_path, skip_ai=True)
    assert deck_watermarks(tmp_path) == []
    assert deck_footer(tmp_path).endswith(" | AI-drafted | not reviewed")


def test_draft_watermarks_every_slide_of_an_unreviewed_deck(tmp_path):
    run_northwind(tmp_path, skip_ai=True, draft=True)
    assert len(deck_watermarks(tmp_path)) == 4


def test_the_printout_names_the_slide_the_ai_text_is_on(tmp_path, capsys):
    # The deck has 4 slides and the AI text is only on slide 4 (AI commentary).
    run_northwind(tmp_path, client=FakeClient(summary()))
    assert "(AI text on slide 4)" in capsys.readouterr().out
    run_northwind(tmp_path, skip_ai=True)
    assert f"({PLACEHOLDER_TEXT} on slide 4)" in capsys.readouterr().out


def test_the_footer_of_a_re_run_names_a_reviewer_whose_approval_still_counts(tmp_path):
    run_northwind(tmp_path, skip_ai=True)
    saved = northwind_manifest(tmp_path)
    saved["approval"] = {"reviewer": "Tyler Ho", "approved_at": "2026-09-17T15:00:00",
                         "input_sha256": file_sha256(NORTHWIND), "config_sha256": file_sha256(CONFIG_FILE)}
    save_manifest(manifest_path(NORTHWIND, tmp_path), saved)

    run_northwind(tmp_path, skip_ai=True, draft=True)
    assert deck_footer(tmp_path).endswith(" | AI-drafted | reviewed by Tyler Ho on 2026-09-17")
    assert deck_watermarks(tmp_path) == []  # --draft never stamps an approved deck


def test_draft_is_passed_from_the_command_line_to_the_batch(monkeypatch):
    seen = {}
    monkeypatch.setattr(main, "run_batch", lambda *args, **kwargs: seen.update(kwargs) or [])
    monkeypatch.setattr(main, "write_summary_csv", lambda results: PROJECT_DIR / "output" / "batch_summary.csv")
    main.main(["--all", "--skip-ai", "--draft"])
    assert seen["draft"] is True
    main.main(["--all", "--skip-ai"])
    assert seen["draft"] is False


# ---------------------------------------------------------------------------
# The API key is checked before any company runs
# ---------------------------------------------------------------------------

def test_api_key_problem(monkeypatch):
    monkeypatch.setattr(main, "load_dotenv", lambda: None)  # don't let .env fill the key back in
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    assert main.api_key_problem() == main.NO_KEY_MESSAGE
    monkeypatch.setenv("ANTHROPIC_API_KEY", "   ")
    assert main.api_key_problem() == main.NO_KEY_MESSAGE
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    assert main.api_key_problem() is None


def test_no_api_key_stops_before_any_company(monkeypatch, capsys):
    monkeypatch.setattr(main, "load_dotenv", lambda: None)
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.setattr(main, "run_batch", lambda *args, **kwargs: pytest.fail("the batch ran without a key"))
    assert main.main(["--all"]) == 1
    assert main.NO_KEY_MESSAGE in capsys.readouterr().out


def test_skip_ai_never_needs_a_key(monkeypatch):
    monkeypatch.setattr(main, "api_key_problem", lambda: pytest.fail("--skip-ai looked for an API key"))
    monkeypatch.setattr(main, "run_batch", lambda *args, **kwargs: [])
    monkeypatch.setattr(main, "write_summary_csv", lambda results: PROJECT_DIR / "output" / "batch_summary.csv")
    assert main.main(["--all", "--skip-ai"]) == 0
