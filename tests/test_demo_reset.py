"""Tests for demo_reset.py (Task 12): output/ back to a known good state before a demo, and DEMO.md's script.

Every test works on temporary data/ and output/ folders holding copies of the three workbooks, so
the real ones are never touched. The saved analyses are made here (a fixed answer with no numbers
in it), not copied from output/, so the tests run on a fresh clone. Creating a real Anthropic client
fails the test, so nothing here can reach the API.
Run from the project folder:  pytest
"""

import hashlib
import json
import re
import shutil
from pathlib import Path

import pytest
from pptx import Presentation
from streamlit.testing.v1 import AppTest

import analyze
import demo_reset
import mapping
import portfolio
from analyze import BoardSummary, build_payload, save_analysis
from build_deck import flag_count_text
from clean import clean_workbook
from diff_runs import HEADING
from export import export_paths
from main import AI_REUSED
from metrics import load_config
from provenance import NOT_REVIEWED, manifest_path, read_manifest
from rollup import rollup_paths
from test_app import render
from test_docs import code_strings
from valid_answer import summary_dict

PROJECT_DIR = Path(__file__).parent.parent
DEMO_MD = PROJECT_DIR / "DEMO.md"
COMPANIES = ["alderpeak", "fernhollow", "northwind"]
LEFTOVERS = ["notes.txt", "task_logs/run.log"]   # files a person put in output/: the reset must leave them


def summary():
    """A valid answer whose only numbers are thresholds every company's payload holds (valid_answer.py)."""
    return BoardSummary.model_validate(summary_dict("A steady quarter."))


def save_matching_analysis(data_dir, output_dir, stem, stale=False):
    """A saved analysis of today's numbers, where main.py puts it. stale=True: Claude saw different
    words for one fact, as Fernhollow's real one did before the em dash change (FINAL_REPORT Task 4)."""
    actuals, next_budget = clean_workbook(data_dir / f"{stem}.xlsx")
    payload = build_payload(stem.title(), actuals, next_budget, load_config())
    if stale:
        payload["company_note"] = "facts worded the old way"
    save_analysis(output_dir / f"{stem}_analysis.json", payload, summary(), {"model": "claude-sonnet-5"})


def hashes(output_dir):
    return {path.name: hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted(Path(output_dir).glob("*_analysis.json"))}


def every_file(output_dir):
    return sorted(str(path.relative_to(output_dir)) for path in Path(output_dir).rglob("*") if path.is_file())


def refuse(*args, **kwargs):
    raise AssertionError("a test tried to create a real Anthropic client")


def leave_practice_mess(data_dir, output_dir):
    """What a practice run leaves: an approved deck, exports, a rollup, a half-finished batch, other files."""
    config = load_config()
    portfolio.generate_company(data_dir / "northwind.xlsx", config, False, output_dir)
    assert list((output_dir / "logs").glob("run_*.jsonl"))   # the practice Generate was logged (Task 15)
    assert portfolio.approve_company("northwind", "Practice Reviewer", data_dir, output_dir)["ok"]
    for path in list(export_paths(data_dir / "northwind.xlsx", output_dir).values()) + list(rollup_paths(output_dir).values()):
        path.write_text("left from practice")
    (output_dir / ".staging" / "northwind_x").mkdir(parents=True)
    (output_dir / "batch_manifest.json").write_text("{}")
    for name in LEFTOVERS:
        (output_dir / name).parent.mkdir(parents=True, exist_ok=True)
        (output_dir / name).write_text("mine")


@pytest.fixture(scope="module")
def reset_twice(tmp_path_factory):
    """Three companies after a practice run, reset twice. Returns what each reset said and left behind."""
    root = tmp_path_factory.mktemp("demo")
    data_dir, output_dir = root / "data", root / "output"
    data_dir.mkdir()
    output_dir.mkdir()
    for stem in COMPANIES:
        shutil.copy(PROJECT_DIR / "data" / f"{stem}.xlsx", data_dir)
    with pytest.MonkeyPatch.context() as patch:
        patch.setattr(analyze.anthropic, "Anthropic", refuse)
        patch.setattr(mapping, "MAPPINGS_DIR", root / "mappings")
        for stem in COMPANIES:
            save_matching_analysis(data_dir, output_dir, stem, stale=stem == "fernhollow")
        leave_practice_mess(data_dir, output_dir)
        before = hashes(output_dir)
        first = demo_reset.reset(data_dir, output_dir)
        files_after_first = every_file(output_dir)
        second = demo_reset.reset(data_dir, output_dir)
        yield {"data_dir": data_dir, "output_dir": output_dir, "before": before, "first": first,
               "second": second, "files_after_first": files_after_first}


def manifest(state, stem):
    return read_manifest(manifest_path(state["data_dir"] / f"{stem}.xlsx", state["output_dir"]))


# ---------------------------------------------------------------------------
# Which files the reset removes, and which it never touches
# ---------------------------------------------------------------------------

def test_built_files_are_what_the_tool_writes_and_never_the_saved_analysis(tmp_path):
    names = [path.name for path in demo_reset.built_files("northwind", tmp_path)]
    assert names == ["northwind_board_pack.pptx", "northwind_board_memo.docx", "northwind_board_memo.pdf",
                     "northwind_metrics.xlsx", "northwind_metrics.csv", "northwind_flags.csv",
                     "northwind_export.json", "northwind_email.html", "northwind_manifest.json"]
    assert "northwind_analysis.json" not in names


def test_shared_files_are_the_batch_rollup_charts_staging_and_run_logs(tmp_path):
    names = [path.name for path in demo_reset.shared_built(tmp_path)]
    assert names == ["batch_summary.csv", "batch_manifest.json", "portfolio_rollup.pptx", "portfolio_rollup.xlsx",
                     "charts", ".staging", "logs"]


def test_the_reset_clears_the_run_logs_and_logs_nothing_itself(reset_twice):
    # Task 15: the practice run's Generate wrote a log; a demo starts with "No runs yet".
    assert not (reset_twice["output_dir"] / "logs").exists()


def test_company_stems_include_a_company_whose_workbook_was_removed(tmp_path):
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()
    shutil.copy(PROJECT_DIR / "data" / "northwind.xlsx", data_dir)
    (output_dir / "blue river_manifest.json").write_text("{}")   # added in practice, workbook since deleted
    assert demo_reset.company_stems(data_dir, output_dir) == ["blue river", "northwind"]


def test_the_reset_keeps_every_saved_analysis_byte_for_byte(reset_twice):
    assert set(reset_twice["before"]) == {f"{stem}_analysis.json" for stem in COMPANIES}
    assert hashes(reset_twice["output_dir"]) == reset_twice["before"]


def test_the_reset_leaves_files_it_did_not_build_and_says_so(reset_twice):
    for name in LEFTOVERS:
        assert (reset_twice["output_dir"] / name).read_text() == "mine"
    assert reset_twice["first"]["left_alone"] == sorted(LEFTOVERS)


def test_practice_leftovers_are_gone(reset_twice):
    output_dir = reset_twice["output_dir"]
    for path in list(export_paths("northwind.xlsx", output_dir).values()) + list(rollup_paths(output_dir).values()):
        assert not path.exists(), path.name
    assert not (output_dir / ".staging").exists()
    assert not (output_dir / "batch_manifest.json").exists()


def test_a_cleared_approval_is_named(reset_twice):
    assert reset_twice["first"]["cleared_approvals"] == ["Northwind: approved by Practice Reviewer"]
    assert reset_twice["second"]["cleared_approvals"] == []
    assert manifest(reset_twice, "northwind").get("approval") is None


# ---------------------------------------------------------------------------
# The state it leaves: every company built from today's workbook, nobody's approval
# ---------------------------------------------------------------------------

def test_every_company_is_built_from_todays_workbook_and_not_reviewed(reset_twice):
    for stem in COMPANIES:
        state = portfolio.run_state(reset_twice["data_dir"] / f"{stem}.xlsx", reset_twice["output_dir"])
        assert state["current"] and state["status"] == NOT_REVIEWED, stem
        assert manifest(reset_twice, stem)["deck"]["draft"] is False
        assert all(path.exists() for path in portfolio.output_files(reset_twice["data_dir"] / f"{stem}.xlsx",
                                                                    reset_twice["output_dir"]).values())
    assert (reset_twice["output_dir"] / "batch_summary.csv").exists()


def test_the_demo_company_has_ai_text_and_a_comparison_with_last_quarter(reset_twice):
    saved = manifest(reset_twice, "northwind")
    assert saved["deck"]["ai_text"] is True and saved["memo"]["ai_text"] is True
    assert saved["ai"]["validation"] == AI_REUSED
    assert saved["previous_run"]["results"]["quarter"] == "Q1 2026"
    assert "batch_events" not in saved   # the practice runs' history is gone


def test_every_company_with_an_answer_key_is_compared_with_last_quarter(reset_twice):
    for stem in COMPANIES:
        assert manifest(reset_twice, stem)["previous_run"]["results"]["quarter"] == "Q1 2026", stem


def test_a_company_without_a_matching_analysis_is_reported_not_hidden(reset_twice):
    assert manifest(reset_twice, "fernhollow")["deck"]["ai_text"] is False
    notes = [line for line in reset_twice["first"]["notes"] if "Fernhollow" in line]
    assert notes and "AI summary unavailable" in notes[0]


def test_the_reset_says_ready_with_no_problems(reset_twice):
    assert reset_twice["first"]["problems"] == [] and reset_twice["second"]["problems"] == []


def test_resetting_twice_leaves_the_same_files(reset_twice):
    assert every_file(reset_twice["output_dir"]) == reset_twice["files_after_first"]


# ---------------------------------------------------------------------------
# When the demo isn't ready, it says so and exits 1
# ---------------------------------------------------------------------------

@pytest.fixture
def broken(tmp_path, monkeypatch):
    """Northwind with no saved analysis, a workbook clean.py can't read, and a company added in practice."""
    monkeypatch.setattr(analyze.anthropic, "Anthropic", refuse)
    monkeypatch.setattr(mapping, "MAPPINGS_DIR", tmp_path / "mappings")
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()
    shutil.copy(PROJECT_DIR / "data" / "northwind.xlsx", data_dir)
    shutil.copy(PROJECT_DIR / "data" / "alderpeak.xlsx", data_dir / "blue river.xlsx")
    (data_dir / "bad.xlsx").write_text("not a workbook")
    return data_dir, output_dir


def test_a_missing_demo_analysis_is_a_problem_that_names_the_fix(broken):
    answer = demo_reset.reset(*broken)
    northwind = [line for line in answer["problems"] if "Northwind" in line]
    assert northwind and "northwind_analysis.json" in northwind[0]
    assert "python main.py data/northwind.xlsx" in northwind[0]


def test_a_workbook_that_cannot_be_read_is_a_problem(broken):
    answer = demo_reset.reset(*broken)
    assert any(line.startswith("Bad:") for line in answer["problems"])


def test_a_company_added_in_practice_is_noted_and_built_without_history(broken):
    answer = demo_reset.reset(*broken)
    # Its own words: the no-AI-text note also names data/blue river.xlsx (found by a planted bug).
    assert any("Blue River" in line and "isn't a demo company" in line for line in answer["notes"])
    data_dir, output_dir = broken
    assert read_manifest(manifest_path(data_dir / "blue river.xlsx", output_dir))["previous_run"] is None


@pytest.fixture
def reset_copy(reset_twice, tmp_path):
    """(data folder, output folder): copies of a finished reset's, so a test can spoil one thing."""
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    shutil.copytree(reset_twice["data_dir"], data_dir)
    shutil.copytree(reset_twice["output_dir"], output_dir)
    return data_dir, output_dir


def test_a_ready_company_has_no_problems(reset_copy):
    data_dir, output_dir = reset_copy
    assert demo_reset.company_readiness(data_dir / "northwind.xlsx", output_dir)[0] == []


def test_files_not_built_from_todays_workbook_are_a_problem(reset_copy):
    # A correct reset never leaves this, so only a direct check reaches it (found by a planted bug).
    data_dir, output_dir = reset_copy
    shutil.copy(PROJECT_DIR / "data" / "alderpeak.xlsx", data_dir / "northwind.xlsx")   # today's workbook changed
    problems, _ = demo_reset.company_readiness(data_dir / "northwind.xlsx", output_dir)
    assert any("aren't a fresh, unreviewed build" in line for line in problems)


def test_no_earlier_run_to_compare_with_is_a_problem(reset_copy):
    data_dir, output_dir = reset_copy
    path = manifest_path(data_dir / "northwind.xlsx", output_dir)
    saved = read_manifest(path)
    del saved["previous_run"]
    path.write_text(json.dumps(saved))
    problems, _ = demo_reset.company_readiness(data_dir / "northwind.xlsx", output_dir)
    assert any("no earlier run to compare with" in line for line in problems)


def test_main_exits_1_when_not_ready_and_prints_why(broken, capsys):
    assert demo_reset.main(*broken) == 1
    printed = capsys.readouterr().out
    assert demo_reset.NOT_READY in printed and "northwind_analysis.json" in printed


def test_main_exits_0_and_says_ready(reset_twice, capsys, monkeypatch):
    monkeypatch.setattr(analyze.anthropic, "Anthropic", refuse)
    monkeypatch.setattr(mapping, "MAPPINGS_DIR", reset_twice["data_dir"].parent / "mappings")
    assert demo_reset.main(reset_twice["data_dir"], reset_twice["output_dir"]) == 0
    assert demo_reset.READY in capsys.readouterr().out


def test_the_refusing_client_fails_any_use():
    with pytest.raises(RuntimeError, match="never calls the API"):
        demo_reset.NoApiClient().messages.parse()


def test_every_build_is_handed_the_refusing_client(monkeypatch):
    # With the AI box unticked no client is used today, so only this test notices if the refusing
    # client is dropped (found by a Task 12 planted bug re-run in Task 13).
    calls = []
    record = lambda *args, **kwargs: calls.append((args, kwargs))   # noqa: E731 - a one-line stand-in
    monkeypatch.setattr(demo_reset, "generate_all", record)
    monkeypatch.setattr(demo_reset, "run_company", record)
    demo_reset.rebuild({}, Path("data"), Path("output"))
    demo_reset.last_quarter_run("northwind", {}, Path("output"))
    assert len(calls) == 2
    assert all(isinstance(kwargs.get("client"), demo_reset.NoApiClient) for _, kwargs in calls)
    assert calls[0][0][1] is False   # generate_all's ask_claude: the AI box unticked


def test_a_saved_analysis_the_rebuild_deletes_or_changes_is_put_back_and_reported(tmp_path, monkeypatch):
    # The rebuild never touches a saved analysis; if a future change made it, the reset puts the copy
    # it took first back, and says the demo isn't ready.
    monkeypatch.setattr(analyze.anthropic, "Anthropic", refuse)
    monkeypatch.setattr(mapping, "MAPPINGS_DIR", tmp_path / "mappings")
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    data_dir.mkdir()
    output_dir.mkdir()
    for stem in ("northwind", "alderpeak"):
        shutil.copy(PROJECT_DIR / "data" / f"{stem}.xlsx", data_dir)
        save_matching_analysis(data_dir, output_dir, stem)
    before = hashes(output_dir)

    def careless_rebuild(config, data_dir, output_dir):
        (output_dir / "northwind_analysis.json").unlink()
        (output_dir / "alderpeak_analysis.json").write_text("{}")
        return []
    monkeypatch.setattr(demo_reset, "rebuild", careless_rebuild)
    answer = demo_reset.reset(data_dir, output_dir)
    assert hashes(output_dir) == before
    assert [line for line in answer["problems"] if "analysis" in line] == [
        "alderpeak_analysis.json changed during the reset: the copy taken before it was put back",
        "northwind_analysis.json changed during the reset: the copy taken before it was put back"]


# ---------------------------------------------------------------------------
# DEMO.md: the clicks it names are on the page, and the numbers it quotes are today's
# ---------------------------------------------------------------------------

def demo_text():
    return DEMO_MD.read_text()


def test_demo_md_names_the_reset_and_the_launcher():
    for name in ("python demo_reset.py", "run_app.command"):
        assert name in demo_text()


def test_every_button_demo_md_clicks_is_on_the_page():
    # Clicks are written "Click **Label**"; a label must be a string in app.py, or a company name.
    labels = set(re.findall(r"[Cc]lick \*\*(.+?)\*\*", demo_text()))
    page = {text for path in (PROJECT_DIR / "app.py", PROJECT_DIR / "portfolio.py") for text in code_strings(path)}
    assert labels, "DEMO.md names no clicks"
    missing = [label for label in labels if label not in page and label.lower() not in COMPANIES]
    assert missing == []


def test_demo_md_quotes_each_companys_flag_count_as_the_page_shows_it():
    config = load_config()
    words = " ".join(demo_text().split())   # a line break reads as a space, as Markdown shows it
    todays = set()
    for stem in COMPANIES:
        data, problem = portfolio.load_company(PROJECT_DIR / "data" / f"{stem}.xlsx", config)
        assert problem is None
        assert flag_count_text(data["flags"]) in words, stem
        todays.add(re.match(r"\d+ of \d+", flag_count_text(data["flags"])).group())
    # And every count it quotes is one of today's: a wrong one beside a right one is still wrong
    # (found by a planted bug re-run in Task 13).
    assert set(re.findall(r"\d+ of \d+(?= flags)", words)) <= todays


def test_demo_md_timings_add_up_to_five_minutes():
    # Each step's heading says when it starts and ends ("(1:30 to 2:30)"); they follow on and end at 5:00.
    spans = re.findall(r"^## .*\((\d:\d\d) to (\d:\d\d)\)", demo_text(), re.MULTILINE)
    assert spans and spans[0][0] == "0:00" and spans[-1][1] == "5:00"
    assert all(end == start for (_, end), (start, _) in zip(spans, spans[1:]))


def test_demo_md_walkthrough_runs_on_the_page_after_a_reset(reset_twice, tmp_path):
    # DEMO.md's steps 2 to 6 with the same clicks, on a copy of the reset folders.
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    shutil.copytree(reset_twice["data_dir"], data_dir)
    shutil.copytree(reset_twice["output_dir"], output_dir)
    test = AppTest.from_function(render, args=(str(data_dir), str(output_dir)), default_timeout=120).run()
    shown = [element.value for element in test.text]
    for count in ("0 of 9 flags tripped", "6 of 9 flags tripped", "7 of 9 flags tripped, 1 cannot evaluate"):
        assert count in shown                                            # step 2
    assert shown.count(NOT_REVIEWED) == 3

    test.button(key="open_northwind").click().run()                      # step 3
    assert [header.value for header in test.subheader][:2] == ["6 of 9 flags tripped", HEADING]
    assert any("Tripped (was Passed)" in text.value for text in test.markdown)
    assert "**A steady quarter.**" in [text.value for text in test.markdown]   # the AI commentary
    deck_button = next(button for button in test.get("download_button") if button.proto.label == "Download deck")
    assert not deck_button.proto.disabled                                # step 4

    test.text_input(key="reviewer_northwind").set_value("Demo Viewer").run()   # step 5
    # Clicked by key, so check the label DEMO.md names too ("Approve" is also the card's heading, so
    # a renamed button would otherwise pass: found by a planted bug re-run in Task 13).
    assert test.button(key="approve").label == "Approve"
    assert test.button(key="generate_page").label == "Generate"
    test.button(key="approve").click().run()
    test.button(key="generate_page").click().run()
    assert not test.exception
    assert "Deck status: approved by Demo Viewer" in test.caption[0].value
    slide = Presentation(output_dir / "northwind_board_pack.pptx").slides[0]
    assert any("AI-drafted | reviewed by Demo Viewer" in shape.text_frame.text
               for shape in slide.shapes if shape.has_text_frame)

    test.button(key="back").click().run()                                # step 6
    assert test.title[0].value == "Board Pack Generator"


def test_uncommitted_code_is_noted_because_every_footer_shows_it():
    assert demo_reset.code_note({"commit": "abc1234", "uncommitted_changes": False}) is None
    note = demo_reset.code_note({"commit": "abc1234", "uncommitted_changes": True})
    assert "abc1234*" in note and "commit" in note
