"""Tests for the evaluation set: eval/make_eval_data.py (12 companies with answer keys) and eval/run_eval.py.

The main test runs the whole evaluation and requires every company to match its answer key, so
`python -m pytest` includes it. The others prove the scorecard names a mismatch instead of passing it:
each one breaks a copy of an answer key (never the real one) and looks for the company, the check and
the detail in what is printed.
"""

import copy
import math

import pytest
from openpyxl import load_workbook

import make_eval_data
import run_eval
from make_data_common import check_true_data
from metrics import load_config

LATEST = "Q2 2026"


def company(name):
    """A deep copy of one eval company, so a test can break its answer key without touching the real one."""
    found = [c for c in make_eval_data.COMPANIES if c["name"] == name]
    assert len(found) == 1, f"no eval company called {name}"
    return copy.deepcopy(found[0])


def run(companies, capsys):
    """Run the scorecard on these companies: (exit code, printed text)."""
    code = run_eval.main(companies)
    return code, capsys.readouterr().out


# ---------------------------------------------------------------------------
# The set itself
# ---------------------------------------------------------------------------

def test_the_set_has_twelve_companies_with_unique_names_and_files():
    companies = make_eval_data.COMPANIES
    assert len(companies) == 12
    assert len({c["name"] for c in companies}) == 12
    assert len({c["output_path"] for c in companies}) == 12
    assert all(c["output_path"].parent == make_eval_data.DATA_DIR for c in companies)


def test_the_set_covers_every_case_asked_for():
    stories = {c["name"]: c["story"] for c in make_eval_data.COMPANIES}
    text = " ".join(stories.values()).lower()
    for case in ("healthy", "distressed", "exactly at every threshold", "zero revenue", "negative budget",
                 "pipeline falling", "short history", "pasted twice", "missing"):
        assert case in text, f"no eval company for {case!r}"


def test_the_blank_quarter_sits_first_second_to_last_and_last():
    positions = {c["quarters"].index(c["blank_quarter"]) for c in make_eval_data.COMPANIES
                 if c["blank_quarter"] and not c.get("stop")}
    assert {0, 6, 7} <= positions   # Northwind (2) and Fernhollow (3) cover the middle already


def test_two_companies_must_stop_one_repeated_row_and_one_missing_row():
    edits = sorted(c["row_edit"][0] for c in make_eval_data.COMPANIES if c.get("stop"))
    assert edits == ["drop", "repeat"]


@pytest.mark.parametrize("eval_company", make_eval_data.COMPANIES, ids=lambda c: c["name"])
def test_every_answer_key_ties_out(eval_company):
    check_true_data(eval_company)   # ARR and cash roll forward, money in 10s, text cells on real quarters


def test_the_threshold_company_sits_exactly_on_every_threshold():
    # Its hand formulas must give each config.yaml threshold exactly, so "exactly at the threshold passes"
    # is tested on every flag at once.
    config = load_config()
    expected = company("Tidewell")["expected_latest"]
    for _, column, key, _ in run_eval.FLAG_RULES:
        assert math.isclose(expected[column], config[key], rel_tol=1e-12), column


def test_the_saved_workbooks_are_what_make_eval_data_writes(tmp_path):
    # A stale eval/data file would test old numbers, so regenerate and compare every cell.
    make_eval_data.write_all(tmp_path)
    for eval_company in make_eval_data.COMPANIES:
        saved = load_workbook(eval_company["output_path"])
        fresh = load_workbook(tmp_path / eval_company["output_path"].name)
        assert saved.sheetnames == fresh.sheetnames, eval_company["name"]
        for name in saved.sheetnames:
            saved_rows = list(saved[name].iter_rows(values_only=True))
            assert saved_rows == list(fresh[name].iter_rows(values_only=True)), (eval_company["name"], name)


# ---------------------------------------------------------------------------
# The evaluation passes
# ---------------------------------------------------------------------------

def test_every_company_matches_its_answer_key(capsys):
    code, out = run(make_eval_data.COMPANIES, capsys)
    assert code == 0, out
    assert "12 of 12 companies match their answer keys" in out
    assert "MISMATCH" not in out


def test_the_stop_companies_stop_with_the_row_to_fix(capsys):
    _, out = run([company("Kestrelwood"), company("Lanternreach")], capsys)
    assert out.count("stopped as expected") == 2


# ---------------------------------------------------------------------------
# A mismatch is named, never passed
# ---------------------------------------------------------------------------

def test_a_wrong_flag_is_named(capsys):
    broken = company("Larkspur")
    broken["expected_flags"]["NRR (annualized)"] = "trip"
    code, out = run([broken], capsys)
    assert code == 1
    assert "Larkspur, Flags: NRR (annualized): expected trip, got pass" in out
    assert "0 of 1 companies match their answer keys" in out


def test_a_wrong_cannot_evaluate_reason_is_named(capsys):
    broken = company("Ivywick")
    broken["expected_flags"]["Rule of 40"] = "cannot evaluate: missing input"
    _, out = run([broken], capsys)
    assert "Ivywick, Flags: Rule of 40: expected cannot evaluate: missing input, " \
           "got cannot evaluate: no prior period" in out


def test_a_wrong_latest_value_is_named(capsys):
    broken = company("Quillmoor")
    broken["expected_latest"]["nrr"] += 0.001
    _, out = run([broken], capsys)
    assert "Quillmoor, Metrics: nrr in Q2 2026: expected" in out


def test_a_not_meaningful_cell_the_answer_key_does_not_list_is_named(capsys):
    broken = company("Emberfall")
    del broken["not_meaningful"]["gross_margin"]
    _, out = run([broken], capsys)
    assert "Emberfall, Metrics: gross_margin in Q3 2024: expected a number, got not meaningful" in out


def test_a_wrong_budget_row_runway_is_named(capsys):
    broken = company("Glenmarsh")
    broken["expected_runway_at_budget"] = 50.0
    _, out = run([broken], capsys)
    assert "Glenmarsh, Metrics: runway at next quarter's budgeted burn: expected 50.0, got inf" in out


def test_a_wrong_cleaned_value_is_named(capsys):
    broken = company("Tidewell")
    broken["true_data"]["revenue"][-1] = 22010
    _, out = run([broken], capsys)
    assert "Tidewell, Clean: Q2 2026 revenue: expected 22010, got 22000.0" in out


def test_a_gap_the_rules_do_not_predict_is_named(capsys):
    broken = company("Copperlane")
    broken["blank_quarter"] = "Q4 2025"   # the workbook's blank is Q1 2026
    _, out = run([broken], capsys)
    assert "Copperlane, Gaps:" in out


def test_a_trip_in_any_quarter_fails_a_company_that_must_never_trip(capsys):
    broken = company("Hollowmere")   # trips NRR and Rule of 40 in the latest quarter
    broken["never_trips"] = True
    _, out = run([broken], capsys)
    assert "Hollowmere, Flags: Q2 2026 tripped NRR (annualized), Rule of 40" in out


def test_a_workbook_that_should_stop_but_reads_is_named(capsys):
    broken = company("Kestrelwood")
    broken["output_path"] = company("Larkspur")["output_path"]   # a workbook with nothing wrong
    code, out = run([broken], capsys)
    assert code == 1
    assert "Kestrelwood, Clean: expected clean.py to stop, but the workbook was read" in out


def test_a_stop_with_the_wrong_message_is_named(capsys):
    broken = company("Lanternreach")
    broken["stop"] = ["appears twice"]
    _, out = run([broken], capsys)
    assert "Lanternreach, Clean: stopped, but the message lacks 'appears twice':" in out


def test_a_missing_workbook_says_how_to_make_it(capsys, tmp_path):
    broken = company("Larkspur")
    broken["output_path"] = tmp_path / "larkspur.xlsx"
    code, out = run([broken], capsys)
    assert code == 1
    assert "python eval/make_eval_data.py" in out


# ---------------------------------------------------------------------------
# The gap rules, by blank position (worked out by hand from CLAUDE.md)
# ---------------------------------------------------------------------------

def test_a_blank_first_quarter_misses_the_next_quarter_and_a_year_later():
    gaps = run_eval.predicted_gaps(company("Brackenfield"), ["arr_qoq", "arr_yoy", "nrr"])
    assert gaps == {"arr_qoq": ["Q3 2024", "Q4 2024"], "arr_yoy": ["Q3 2024", "Q3 2025"], "nrr": ["Q3 2024"]}


def test_a_blank_last_quarter_misses_only_itself_and_every_flag():
    gaps = run_eval.predicted_gaps(company("Duskhaven"), ["arr_qoq", "arr_yoy"])
    assert gaps["arr_qoq"] == [LATEST] and gaps["arr_yoy"] == [LATEST]
    assert len([name for name in gaps if name.startswith("flag: ")]) == 9


def test_no_prior_period_is_never_a_gap():
    assert run_eval.predicted_gaps(company("Ivywick"), ["arr_qoq", "arr_yoy", "rule_of_40"]) == {}
