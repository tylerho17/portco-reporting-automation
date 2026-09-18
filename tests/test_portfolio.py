"""Smoke tests for portfolio.py: the web page's rows, search, Generate, Add a company and Approve, without a browser.

Every test copies the three workbooks into a temporary data/ folder and builds into a temporary
output/ folder, so the real data/ and output/ are never touched. A guard makes creating a real
Anthropic client fail, so no test can reach the API; the AI step uses a fake client.
Run from the project folder:  pytest
"""

import csv
import json
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from pptx import Presentation

import analyze
import portfolio
from analyze import BoardSummary, build_payload, save_analysis
from build_deck import PLACEHOLDER_TEXT
from clean import clean_workbook
from main import AI_REUSED
from metrics import load_config
from provenance import NOT_REVIEWED, manifest_path, read_manifest

PROJECT_DIR = Path(__file__).parent.parent
COMPANIES = ["alderpeak", "fernhollow", "northwind"]


class FakeClient:
    """Stands in for anthropic.Anthropic() (as in tests/test_main.py): one fixed answer, calls counted."""

    def __init__(self, summary=None, error=None):
        self.summary, self.error, self.calls = summary, error, 0
        self.messages = self  # analyze.py calls client.messages.parse(...)

    def parse(self, **kwargs):
        self.calls += 1
        if self.error:
            raise self.error
        return SimpleNamespace(content=[SimpleNamespace(type="text", text=self.summary.model_dump_json())],
                               parsed_output=self.summary, stop_reason="end_turn",
                               usage=SimpleNamespace(input_tokens=100, output_tokens=50))


def summary(headline):
    """A valid answer with no numbers in it, so the number check has nothing to reject."""
    point = {"title": "Steady base", "detail": "Customers stayed."}
    return BoardSummary.model_validate({
        "headline": headline, "wins": [point] * 3, "risks": [point] * 3,
        "questions": ["What drives churn?", "Where is pipeline coming from?", "How is hiring going?"]})


@pytest.fixture(autouse=True)
def no_real_client(monkeypatch):
    """Creating a real Anthropic client fails the test, so no test here can reach the API."""
    def refuse(*args, **kwargs):
        raise AssertionError("a test tried to create a real Anthropic client")
    monkeypatch.setattr(analyze.anthropic, "Anthropic", refuse)


@pytest.fixture
def folders(tmp_path):
    """(data folder with copies of the three workbooks, empty output folder)."""
    data_dir, output_dir = tmp_path / "data", tmp_path / "output"
    data_dir.mkdir()
    for name in COMPANIES:
        shutil.copy(PROJECT_DIR / "data" / f"{name}.xlsx", data_dir)
    return data_dir, output_dir


def unreadable_workbook_bytes(tmp_path):
    """A real .xlsx that clean.py stops on (no budget columns), and clean.py's own message for it."""
    book = Workbook()
    book.active.append(["Quarter", "ARR"])
    book.active.append(["Q1 2025", 100])
    path = tmp_path / "bad.xlsx"
    book.save(path)
    with pytest.raises(ValueError) as expected:
        clean_workbook(path)
    return path.read_bytes(), str(expected.value)


def rows_by_company(folders):
    data_dir, output_dir = folders
    return {row["company"]: row for row in portfolio.portfolio_rows(load_config(), data_dir, output_dir)}


def generate(folders, name="northwind", ask_claude=False, client=None):
    data_dir, output_dir = folders
    return portfolio.generate_company(data_dir / f"{name}.xlsx", load_config(), ask_claude, output_dir, client)


def save_northwind_analysis(folders, headline):
    """A saved analysis of today's Northwind numbers, where main.py would have put it."""
    data_dir, output_dir = folders
    actuals, next_budget = clean_workbook(data_dir / "northwind.xlsx")
    payload = build_payload("Northwind", actuals, next_budget, load_config())
    save_analysis(output_dir / "northwind_analysis.json", payload, summary(headline), {"model": "claude-sonnet-5"})


def headline_on_deck(folders):
    slide = Presentation(folders[1] / "northwind_board_pack.pptx").slides[-1]
    return next(shape.text_frame.text for shape in slide.shapes if shape.name == "Headline")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def test_plain_replaces_every_em_dash():
    em_dash = chr(0x2014)   # by its Unicode number, so this file shows none
    assert portfolio.plain(f"Cannot evaluate {em_dash} missing input") == "Cannot evaluate: missing input"
    assert em_dash not in portfolio.plain(f"a{em_dash}b")


def test_last_run_is_shown_to_the_minute():
    assert portfolio.last_run_text("2026-09-17T14:03:11") == "2026-09-17 14:03"
    assert portfolio.last_run_text(None) == portfolio.NEVER_RUN


def test_a_company_name_becomes_a_safe_file_name():
    assert portfolio.company_stem("  Blue   River ") == "blue river"
    assert portfolio.company_stem("Acme-2") == "acme-2"
    for bad in ("", "   ", "../northwind", "a/b", "2cool", "x" * 41):
        with pytest.raises(ValueError, match="company name"):
            portfolio.company_stem(bad)


def test_the_suggested_name_comes_from_the_file_name():
    assert portfolio.suggested_name("bluefin.xlsx") == "bluefin"
    assert portfolio.suggested_name("Bluefin - Q2 2026 (final).xlsx") == "Bluefin - Q2 2026 final"


# ---------------------------------------------------------------------------
# The portfolio table
# ---------------------------------------------------------------------------

def test_one_row_per_workbook_with_each_companys_story(folders):
    rows = rows_by_company(folders)
    assert list(rows) == ["Alderpeak", "Fernhollow", "Northwind"]
    assert rows["Northwind"]["latest"] == "Q2 2026"
    assert rows["Northwind"]["flags"] == "6 of 9 flags tripped"
    assert rows["Northwind"]["gaps"] == "19 metrics/flags (blank: Q1 2025)"
    assert rows["Fernhollow"]["flags"] == "7 of 9 flags tripped, 1 cannot evaluate"
    assert rows["Alderpeak"]["flags"] == "0 of 9 flags tripped" and rows["Alderpeak"]["gaps"] == "none"


def test_before_any_run_the_row_says_never_and_not_generated(folders):
    row = rows_by_company(folders)["Northwind"]
    assert row["last_run"] == portfolio.NEVER_RUN and row["status"] == portfolio.NOT_GENERATED
    assert row["current"] is False and row["problem"] is None


def test_an_unreadable_workbook_shows_clean_py_s_message_and_no_numbers(folders, tmp_path):
    data, message = unreadable_workbook_bytes(tmp_path)
    (folders[0] / "broken.xlsx").write_bytes(data)
    rows = rows_by_company(folders)
    assert rows["Broken"]["problem"] == portfolio.plain(message)
    assert rows["Broken"]["latest"] == rows["Broken"]["flags"] == rows["Broken"]["gaps"] == portfolio.NO_VALUE
    assert rows["Northwind"]["flags"] == "6 of 9 flags tripped"        # the others are unaffected


def test_a_file_that_is_not_a_workbook_gets_the_plain_message(folders):
    (folders[0] / "notes.xlsx").write_text("just some text")
    assert rows_by_company(folders)["Notes"]["problem"] == portfolio.NOT_A_WORKBOOK


def test_an_unexpected_error_is_named_without_a_traceback(folders, monkeypatch):
    def broken(*args, **kwargs):
        raise KeyError("oops")
    monkeypatch.setattr(portfolio, "collect_deck_data", broken)
    problem = rows_by_company(folders)["Northwind"]["problem"]
    assert problem.startswith(portfolio.UNEXPECTED_ERROR_START) and "KeyError" in problem
    assert "Traceback" not in problem


# ---------------------------------------------------------------------------
# Search
# ---------------------------------------------------------------------------

def test_search_matches_part_of_a_name_in_any_case(folders):
    rows = list(rows_by_company(folders).values())
    assert [row["company"] for row in portfolio.search_rows(rows, "NORTH")] == ["Northwind"]
    assert len(portfolio.search_rows(rows, "  ")) == 3
    assert portfolio.search_message(rows, "north") is None and portfolio.search_message(rows, "") is None


def test_a_searched_name_with_no_workbook_says_upload_one(folders):
    rows = list(rows_by_company(folders).values())
    assert portfolio.search_rows(rows, "Bluefin") == []
    message = portfolio.search_message(rows, " Bluefin ")
    assert message == 'No KPI workbook found for "Bluefin". Upload one under "Add a company".'


def test_find_workbook_by_its_file_name(folders):
    assert portfolio.find_workbook("northwind", folders[0]) == folders[0] / "northwind.xlsx"
    assert portfolio.find_workbook("bluefin", folders[0]) is None
    assert portfolio.find_workbook("../data/northwind", folders[0]) is None   # never outside data/


# ---------------------------------------------------------------------------
# Generate
# ---------------------------------------------------------------------------

def test_generate_builds_every_file_and_the_row_reads_the_manifest(folders):
    outcome = generate(folders)
    assert outcome["ok"] is True and outcome["company"] == "Northwind"
    files = portfolio.output_files(folders[0] / "northwind.xlsx", folders[1])
    assert all(path.exists() for path in files.values())
    row = rows_by_company(folders)["Northwind"]
    manifest = read_manifest(manifest_path(folders[0] / "northwind.xlsx", folders[1]))
    assert row["last_run"] == portfolio.last_run_text(manifest["run_at"])
    assert row["status"] == NOT_REVIEWED and row["current"] is True


def test_generate_without_ai_and_no_saved_analysis_says_so(folders):
    outcome = generate(folders)
    assert headline_on_deck(folders) == PLACEHOLDER_TEXT
    assert outcome["message"] == f"Northwind: generated. {portfolio.AI_NOT_ASKED_NOTE}"


def test_generate_reuses_a_saved_analysis_of_the_same_numbers_even_with_the_box_unticked(folders):
    save_northwind_analysis(folders, "Saved headline.")
    client = FakeClient(summary("Fresh headline."))
    outcome = generate(folders, ask_claude=True, client=client)
    assert client.calls == 0 and headline_on_deck(folders) == "Saved headline."
    assert outcome["message"].endswith(portfolio.AI_REUSED_NOTE)
    assert read_manifest(manifest_path(folders[0] / "northwind.xlsx", folders[1]))["ai"]["validation"] == AI_REUSED
    generate(folders)                                                       # box unticked: still reused
    assert headline_on_deck(folders) == "Saved headline."


def test_generate_with_the_box_ticked_asks_claude_once(folders):
    client = FakeClient(summary("Fresh headline."))
    outcome = generate(folders, ask_claude=True, client=client)
    assert client.calls == 1 and headline_on_deck(folders) == "Fresh headline."
    assert outcome["message"].endswith(portfolio.AI_NEW_NOTE)


def test_generate_with_no_api_key_still_builds_and_says_why(folders, monkeypatch):
    monkeypatch.setattr(portfolio, "api_key_problem", lambda: "no key here")
    outcome = generate(folders, ask_claude=True)
    assert outcome["ok"] is True and headline_on_deck(folders) == PLACEHOLDER_TEXT
    assert outcome["message"].endswith(portfolio.NO_KEY_NOTE)


def test_an_ai_failure_still_builds_the_deck_and_says_why(folders):
    client = FakeClient(error=analyze.anthropic.AnthropicError("simulated outage"))
    outcome = generate(folders, ask_claude=True, client=client)
    assert outcome["ok"] is True and headline_on_deck(folders) == PLACEHOLDER_TEXT
    assert PLACEHOLDER_TEXT in outcome["message"] and "simulated outage" in outcome["message"]


def test_generate_on_an_unreadable_workbook_gives_clean_py_s_message(folders, tmp_path):
    data, message = unreadable_workbook_bytes(tmp_path)
    (folders[0] / "broken.xlsx").write_bytes(data)
    outcome = generate(folders, name="broken")
    assert outcome["ok"] is False and outcome["message"] == f"Broken: {portfolio.plain(message)}"


def test_generate_all_carries_on_past_a_failure_and_reports_progress(folders, tmp_path):
    data, _ = unreadable_workbook_bytes(tmp_path)
    (folders[0] / "broken.xlsx").write_bytes(data)
    data_dir, output_dir = folders
    progress = []
    outcomes = portfolio.generate_all(load_config(), False, data_dir, output_dir,
                                      on_progress=lambda done, total, name: progress.append((done, total, name)))
    assert [(o["company"], o["ok"]) for o in outcomes] == [
        ("Alderpeak", True), ("Broken", False), ("Fernhollow", True), ("Northwind", True)]
    assert progress == [(0, 4, "Alderpeak"), (1, 4, "Broken"), (2, 4, "Fernhollow"), (3, 4, "Northwind"),
                        (4, 4, None)]
    with open(output_dir / "batch_summary.csv", newline="") as file:
        results = [row[-1] for row in csv.reader(file)][1:]
    assert results[0] == "OK (AI skipped)" and results[1].startswith("FAILED: ValueError")


def test_the_row_goes_out_of_date_when_the_workbook_changes(folders):
    generate(folders)
    shutil.copy(PROJECT_DIR / "data" / "alderpeak.xlsx", folders[0] / "northwind.xlsx")
    row = rows_by_company(folders)["Northwind"]
    assert row["current"] is False
    assert row["status"] == "Out of date: the workbook has changed since the last run. Generate again."


def test_downloads_are_offered_only_for_current_files(folders):
    workbook = folders[0] / "northwind.xlsx"
    assert portfolio.download(workbook, folders[1], "deck", current=False) is None
    generate(folders)
    name, data = portfolio.download(workbook, folders[1], "deck", current=True)
    assert name == "northwind_board_pack.pptx" and data == (folders[1] / name).read_bytes()
    assert portfolio.download(workbook, folders[1], "memo", current=True)[0] == "northwind_board_memo.pdf"
    assert portfolio.download(workbook, folders[1], "excel", current=True)[0] == "northwind_metrics.xlsx"
    assert portfolio.download(workbook, folders[1], "deck", current=False) is None   # there, but out of date
    (folders[1] / "northwind_metrics.xlsx").unlink()
    assert portfolio.download(workbook, folders[1], "excel", current=True) is None


def test_a_changed_workbook_takes_its_old_files_off_the_download_buttons(folders):
    generate(folders)
    shutil.copy(PROJECT_DIR / "data" / "alderpeak.xlsx", folders[0] / "northwind.xlsx")
    row = rows_by_company(folders)["Northwind"]
    for kind in ("deck", "memo", "memo_docx", "excel"):
        assert portfolio.download(row["workbook"], folders[1], kind, row["current"]) is None


# ---------------------------------------------------------------------------
# Add a company
# ---------------------------------------------------------------------------

def add(folders, data, name, replace=False, file_name="upload.xlsx"):
    return portfolio.add_company(file_name, data, name, load_config(), folders[0], replace)


def test_a_readable_workbook_is_added_to_data(folders):
    data = (PROJECT_DIR / "data" / "northwind.xlsx").read_bytes()
    outcome = add(folders, data, "Blue River")
    assert outcome["ok"] is True
    assert (folders[0] / "blue river.xlsx").read_bytes() == data
    assert outcome["message"] == "Added Blue River (latest quarter Q2 2026). Click Generate on its row to build its files."
    assert "Blue River" in rows_by_company(folders)


def test_an_unreadable_workbook_is_not_added_and_clean_py_s_message_is_shown(folders, tmp_path):
    data, message = unreadable_workbook_bytes(tmp_path)
    outcome = add(folders, data, "Bluefin")
    assert outcome == {"ok": False, "message": portfolio.plain(message)}
    assert not (folders[0] / "bluefin.xlsx").exists()


def test_a_file_that_is_not_a_workbook_is_not_added(folders):
    deck = (PROJECT_DIR / "templates" / "base.pptx").read_bytes()   # a zip, but not Excel
    assert add(folders, deck, "Bluefin") == {"ok": False, "message": portfolio.NOT_A_WORKBOOK}
    assert add(folders, b"text", "Bluefin") == {"ok": False, "message": portfolio.NOT_A_WORKBOOK}
    assert not (folders[0] / "bluefin.xlsx").exists()


def test_an_existing_company_is_replaced_only_when_asked(folders):
    data = (PROJECT_DIR / "data" / "alderpeak.xlsx").read_bytes()
    before = (folders[0] / "northwind.xlsx").read_bytes()
    outcome = add(folders, data, "Northwind")
    assert outcome["ok"] is False and "already" in outcome["message"]
    assert (folders[0] / "northwind.xlsx").read_bytes() == before
    assert add(folders, data, "northwind", replace=True)["ok"] is True
    assert (folders[0] / "northwind.xlsx").read_bytes() == data


def test_a_bad_name_is_refused_before_anything_is_saved(folders):
    data = (PROJECT_DIR / "data" / "northwind.xlsx").read_bytes()
    outcome = add(folders, data, "../../elsewhere")
    assert outcome == {"ok": False, "message": portfolio.NAME_RULE}
    assert sorted(path.name for path in folders[0].iterdir()) == [f"{name}.xlsx" for name in COMPANIES]


# ---------------------------------------------------------------------------
# Approve
# ---------------------------------------------------------------------------

def approve(folders, reviewer):
    return portfolio.approve_company("northwind", reviewer, folders[0], folders[1], now="2026-09-18T10:00:00")


def test_approve_records_the_reviewer_as_approve_py_does(folders):
    generate(folders)
    outcome = approve(folders, "  Dana Reviewer ")
    assert outcome["ok"] is True
    approval = read_manifest(manifest_path(folders[0] / "northwind.xlsx", folders[1]))["approval"]
    assert approval["reviewer"] == "Dana Reviewer" and approval["documents"] == ["deck", "memo"]
    assert rows_by_company(folders)["Northwind"]["status"] == "approved by Dana Reviewer on 2026-09-18T10:00:00"
    assert "Generate" in outcome["message"]          # the footer changes only when the deck is rebuilt


def test_approve_needs_a_typed_name(folders):
    generate(folders)
    assert approve(folders, "  ") == {"ok": False, "message": portfolio.REVIEWER_NEEDED}
    assert read_manifest(manifest_path(folders[0] / "northwind.xlsx", folders[1]))["approval"] is None


def test_approve_before_generate_or_after_a_change_is_refused_in_the_pages_words(folders):
    # approve.py would refuse too, but its words are for the command line ("run main.py ... again").
    assert approve(folders, "Dana") == {"ok": False, "message": portfolio.NOT_APPROVABLE.format(
        status=portfolio.NOT_GENERATED)}
    generate(folders)
    shutil.copy(PROJECT_DIR / "data" / "alderpeak.xlsx", folders[0] / "northwind.xlsx")
    outcome = approve(folders, "Dana")
    assert outcome["ok"] is False and "changed" in outcome["message"] and "main.py" not in outcome["message"]
    assert read_manifest(manifest_path(folders[0] / "northwind.xlsx", folders[1]))["approval"] is None


# ---------------------------------------------------------------------------
# The company page's AI commentary
# ---------------------------------------------------------------------------

def test_saved_commentary_only_for_exactly_these_numbers(folders):
    workbook = folders[0] / "northwind.xlsx"
    assert portfolio.saved_commentary(workbook, load_config(), folders[1]) is None
    save_northwind_analysis(folders, "Saved headline.")
    assert portfolio.saved_commentary(workbook, load_config(), folders[1]).headline == "Saved headline."
    path = folders[1] / "northwind_analysis.json"
    saved = json.loads(path.read_text())
    saved["payload"]["flags"] = []
    path.write_text(json.dumps(saved))
    assert portfolio.saved_commentary(workbook, load_config(), folders[1]) is None
