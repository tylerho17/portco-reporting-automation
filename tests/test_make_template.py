"""Unit tests for make_template.py: the brand template every deck is built on.

Run from the project folder:  pytest
"""

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

from make_template import CONTENT_LAYOUT, TEMPLATE_PATH, TITLE_LAYOUT, build_template
from theme import FONT, LINE, NAVY, SLATE, SURFACE

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def test_slides_are_16_by_9():
    presentation = build_template()
    # 13.333 in x 7.5 in: 16 / 9 = 1.777...
    assert (presentation.slide_width, presentation.slide_height) == (Emu(12192000), Emu(6858000))


def test_exactly_a_title_layout_and_a_title_and_content_layout():
    names = [layout.name for layout in build_template().slide_layouts]
    assert names == [TITLE_LAYOUT, CONTENT_LAYOUT]


def test_template_has_no_slides():
    assert len(build_template().slides) == 0


def placeholder_types(layout):
    return {placeholder.placeholder_format.type for placeholder in layout.placeholders}


def test_content_layout_has_title_body_and_footer_only():
    layout = build_template().slide_layouts.get_by_name(CONTENT_LAYOUT)
    # The run date goes in the footer text, so date and slide number placeholders are removed.
    assert placeholder_types(layout) == {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.FOOTER}


def test_every_placeholder_sits_inside_the_slide():
    presentation = build_template()
    for layout in presentation.slide_layouts:
        for shape in layout.placeholders:
            assert shape.left >= 0 and shape.top >= 0, f"{layout.name}: {shape.name}"
            assert shape.left + shape.width <= presentation.slide_width, f"{layout.name}: {shape.name}"
            assert shape.top + shape.height <= presentation.slide_height, f"{layout.name}: {shape.name}"


def test_body_sits_between_title_and_footer():
    layout = build_template().slide_layouts.get_by_name(CONTENT_LAYOUT)
    shapes = {shape.placeholder_format.type: shape for shape in layout.placeholders}
    title, body, footer = shapes[PP_PLACEHOLDER.TITLE], shapes[PP_PLACEHOLDER.OBJECT], shapes[PP_PLACEHOLDER.FOOTER]
    assert title.top + title.height <= body.top
    assert body.top + body.height <= footer.top


def test_brand_name_fits_one_line_beside_the_footer():
    # The footer box was widened to hold the review status, so the brand box is only as wide as
    # the brand name needs. The two must not overlap, and the name must not wrap.
    from text_fit import text_width_pt
    presentation = build_template()
    brand = [shape for shape in presentation.slide_master.shapes if shape.name == "Brand name"][0]
    layout = presentation.slide_layouts.get_by_name(CONTENT_LAYOUT)
    footer = [shape for shape in layout.placeholders if shape.placeholder_format.type == PP_PLACEHOLDER.FOOTER][0]
    assert footer.left + footer.width <= brand.left
    frame = brand.text_frame
    room = Emu(brand.width - frame.margin_left - frame.margin_right).pt
    assert text_width_pt(frame.text, 12, bold=True) <= room


def theme(presentation):
    """The theme XML (colors and fonts) behind the slide master."""
    return etree.fromstring(presentation.slide_master.part.part_related_by(RT.THEME).blob)


def test_theme_colors_are_navy_and_gray_and_font_is_arial():
    root = theme(build_template())
    assert root.find(f".//{A}dk2/{A}srgbClr").get("val") == NAVY
    assert root.find(f".//{A}accent1/{A}srgbClr").get("val") == NAVY
    assert root.find(f".//{A}majorFont/{A}latin").get("typeface") == "Arial"
    assert root.find(f".//{A}minorFont/{A}latin").get("typeface") == "Arial"


def test_saved_template_opens_again(tmp_path):
    from pptx import Presentation

    from make_template import save_template
    path = save_template(tmp_path / "base.pptx")
    reopened = Presentation(path)
    assert [layout.name for layout in reopened.slide_layouts] == [TITLE_LAYOUT, CONTENT_LAYOUT]


# ---------------------------------------------------------------------------
# Task 3: the palette and sizes from theme.py
# ---------------------------------------------------------------------------

def text_style(presentation, kind, level=1):
    """(size in pt, color) of the master's title or body text at one level."""
    styles = presentation.slide_master.element.find(f"{P}txStyles")
    defaults = styles.find(f"{P}{kind}/{A}lvl{level}pPr/{A}defRPr")
    return int(defaults.get("sz")) / 100, defaults.find(f"{A}solidFill/{A}srgbClr").get("val")


P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def test_master_text_is_navy_titles_at_28_and_slate_body_at_15():
    presentation = build_template()
    assert text_style(presentation, "titleStyle") == (28, "0B2545")
    assert text_style(presentation, "bodyStyle") == (15, "334155")
    assert text_style(presentation, "bodyStyle", level=2) == (13, "334155")


def test_theme_carries_the_new_palette():
    root = theme(build_template())
    assert root.find(f".//{A}dk2/{A}srgbClr").get("val") == "0B2545"
    assert root.find(f".//{A}lt2/{A}srgbClr").get("val") == SURFACE
    assert root.find(f".//{A}accent2/{A}srgbClr").get("val") == "64748B"


def master_shape(presentation, name):
    return [shape for shape in presentation.slide_master.shapes if shape.name == name][0]


def test_top_bar_is_navy_and_footer_rule_is_the_line_color():
    presentation = build_template()
    assert str(master_shape(presentation, "Top bar").fill.fore_color.rgb) == NAVY
    assert str(master_shape(presentation, "Footer rule").fill.fore_color.rgb) == LINE


def test_brand_name_is_example_capital_in_navy_arial_at_the_floor():
    run = master_shape(build_template(), "Brand name").text_frame.paragraphs[0].runs[0]
    assert run.text == "Example Capital"
    assert (run.font.size.pt, str(run.font.color.rgb)) == (12, NAVY)


def test_cover_layout_is_navy_with_a_28_pt_white_title():
    layout = build_template().slide_layouts.get_by_name(TITLE_LAYOUT)
    assert str(layout.background.fill.fore_color.rgb) == NAVY
    title = [shape for shape in layout.placeholders if shape.placeholder_format.type == PP_PLACEHOLDER.CENTER_TITLE][0]
    defaults = title.text_frame._txBody.find(f"{A}lstStyle/{A}lvl1pPr/{A}defRPr")
    assert int(defaults.get("sz")) / 100 == 28
    assert defaults.find(f"{A}solidFill/{A}srgbClr").get("val") == "FFFFFF"


def test_the_saved_template_was_rebuilt_with_the_palette():
    # templates/base.pptx is committed: this fails until python make_template.py has been run again.
    from pptx import Presentation
    saved = Presentation(TEMPLATE_PATH)
    assert theme(saved).find(f".//{A}dk2/{A}srgbClr").get("val") == NAVY
    assert text_style(saved, "bodyStyle") == (15, SLATE)
    assert theme(saved).find(f".//{A}minorFont/{A}latin").get("typeface") == FONT
